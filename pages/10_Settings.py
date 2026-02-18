# Summary: update mapping UI labels and template filenames to NEXT branding.
import io
import uuid
import time
from datetime import date, datetime
import re
from typing import Optional, List, Dict, Tuple, Any

import pandas as pd
import streamlit as st
from core import bulkload_v2
from core.ado_recon import load_explorer_feature_rows
from core.ado_profile import (
    get_active_ado_profile,
    normalize_profile_config,
    build_workitems_select_expand,
    extract_program_team,
    extract_pi_label,
    extract_app_name,
    extract_points,
)
from core.cache_utils import cache_data_portfolio
from core.freshness import post_write_refresh
from core.portfolio_context import get_active_profile_key, get_active_pat_env_key
from core.ado_http import get_ado_auth_context, ado_odata_get, fetch_ado_odata_metadata, parse_ado_metadata_fields
from utils.ado_pat import resolve_ado_pat_source
from core.data import dominant_iteration_root, list_iteration_roots
from core.nwf_taxonomy import PLANNED_NWF_SUBCOMPONENTS, normalize_nwf_subcomponent
from utils.theme import use_theme
from utils.reconciliation import render_reconciliation_tab
import base64
import json
import math
import sys
import os
from pathlib import Path
from urllib.parse import urlparse, quote, parse_qs
from utils.ado import (
    fetch_ado_odata,
    fetch_ado_iterations_odata,
    transform_ado_odata_to_expected,
    normalize_to_canonical,
    repair_leaf_teams,
)
from db import (
    ensure_ado_profiles_table,
    set_profile_sync_mode,
    set_profile_last_sync_at,
    get_sql_trace_debug,
)


def _settings_post_write_refresh(
    context: str,
    *,
    rerun: bool = False,
    bump_version: Optional[bool] = None,
) -> None:
    post_write_refresh(context, ensure_views=False, rerun=rerun, bump_version=bump_version)

try:
    import requests
except Exception:  # handled at runtime if missing
    requests = None  # type: ignore

from db import (
    execute,
    fetch_df,
    ensure_ado_minimal_tables,
    ensure_ado_iteration_calendar_table,
    ensure_ado_features_enriched_view,
    ensure_ado_workitem_lookup_table,
    ensure_tables,
    ensure_apptio_actuals_lines_table,
    ensure_map_apptio_to_cost_type_table,
    ensure_tco_feature_demand_view,
    recompute_feature_iteration_mapping,
    repair_ado_features_leaf_teams,
    fetch_ado_profiles_full,
    set_active_ado_profile,
    upsert_ado_profile,
    delete_ado_profile,
    load_ado_portfolio_settings,
    save_ado_portfolio_settings,
    delete_ado_portfolio_settings,
    ensure_analytics_views_ok,

    # Lookups + upserts (align with your schema)
    list_programs,
    list_teams,
    list_vendors,
    list_application_groups,
    list_applications,
    upsert_program,
    upsert_team,
    upsert_vendor,
    upsert_application_group,
    upsert_application_instance,
    upsert_invoice,
    upsert_ado_features,
    upsert_ado_iteration_calendar,
    ensure_contracts_table,
    upsert_contract,
    list_program_apptio_workids,
    upsert_apptio_actuals,
    upsert_apptio_actuals_lines,
    upsert_apptio_to_cost_type_mappings,
    fetch_apptio_actuals_by_program_breakdown,
)
from db import fetch_df_active as fetch_df, execute_active as execute
from db.control_db import (
    control_db_available,
    ensure_control_schema,
    ensure_control_automation_tables,
    list_automation_settings,
    upsert_automation_setting,
    queue_automation_run,
    list_automation_runs,
)
from core.db_session import db_scalar
from utils.toast import toast_error, toast_success
from core.debug import is_debug_enabled

# Master-data utilities (no cost math; helps mapping and default instance behavior)
from core.ado_candidates import refresh_ado_app_candidates
from core.app_instances import ensure_default_instances

show_debug = bool(is_debug_enabled(label="Debug"))

# Auth guard: contributors only
try:
    from utils.auth import require_role, ensure_sso
except Exception:
    st.error("Authentication is not available. Contact an admin.")
    st.stop()


def _readonly_preview_active() -> bool:
    try:
        from utils.auth import is_readonly_preview_active

        return bool(is_readonly_preview_active())
    except Exception:
        return False


def _preview_blocked_notice(action_label: str) -> bool:
    if not _readonly_preview_active():
        return False
    st.warning(
        f"{action_label} is blocked while Scope preview is active. "
        "Set Scope preview to 'Actual scope (me)' in the sidebar first."
    )
    return True

# -------------------------
# Page setup
# -------------------------
from core.init import init_page

from utils.app_shell import bootstrap_page
bootstrap_page()
page_theme = init_page("Settings", page_path=__file__)
_page_theme = page_theme
is_admin = True
st.title("Settings")
st.caption(
    "Maintain mappings, rates, and reference data"
)
if is_admin:
    try:
        db_name = db_scalar("SELECT DB_NAME()")
        if db_name:
            st.caption(f"Active DB: {db_name}")
    except Exception:
        pass

# Ensure minimal schema is ready (non-fatal if creation fails so UI still loads)
st.session_state.setdefault("_settings_schema_checked", False)
_SETTINGS_SCHEMA_VERSION = 2
if st.session_state.get("_settings_schema_version") != _SETTINGS_SCHEMA_VERSION:
    st.session_state["_settings_schema_checked"] = False
    st.session_state["_settings_schema_version"] = _SETTINGS_SCHEMA_VERSION
if st.button("Re-run schema checks", type="secondary", key="settings_schema_rerun_btn"):
    st.session_state["_settings_schema_checked"] = False
    st.rerun()

if not st.session_state.get("_settings_schema_checked"):
    with st.spinner("Ensuring minimal ADO schema..."):
        try:
            ensure_ado_minimal_tables()
            st.success("ADO base tables: OK")
        except Exception as e:
            st.warning(f"ADO base tables check failed: {e}")
        try:
            if callable(ensure_tco_feature_demand_view):
                ensure_tco_feature_demand_view()
                st.success("VW_TCO_FEATURE_DEMAND: OK")
        except Exception as e:
            st.warning(f"View build (feature demand) failed: {e}")
        try:
            if callable(ensure_ado_features_enriched_view):
                ensure_ado_features_enriched_view()
                st.success("VW_ADO_FEATURES_ENRICHED: OK")
        except Exception as e:
            st.error(f"VW_ADO_FEATURES_ENRICHED failed: {e}")
        try:
            ensure_ado_workitem_lookup_table()
            st.success("ADO_WORKITEM_LOOKUP: OK")
        except Exception as e:
            st.error(f"ADO_WORKITEM_LOOKUP failed: {e}")
    st.session_state["_settings_schema_checked"] = True

# -------------------------
# Session state
# -------------------------
for state_key, default in [
    ("ado_parsed_raw", None),
    ("ado_parsed_norm", None),
    ("ado_query_mode", "Advanced OData Builder"),
    ("_ado_active_portfolio_key", None),
    ("ado_portfolio_settings_loaded", False),
    ("ado_data_cache_bust", 0),
    ("one_sheet_df", None),
    ("colmap", {}),
    ("previews", {}),
    ("profiles_cache_bust", 0),
]:
    if state_key not in st.session_state:
        st.session_state[state_key] = default

if not st.session_state.get("ado_portfolio_settings_loaded"):
    try:
        st.session_state["ado_portfolio_settings"] = load_ado_portfolio_settings()
    except Exception:
        st.session_state["ado_portfolio_settings"] = {}
    st.session_state["ado_portfolio_settings_loaded"] = True

# -------------------------
# Column expectations (ADO)
# -------------------------
EXPECTED = {
    "Team": ["Team", "System.Team", "Area Team"],
    "Custom_ApplicationName": ["Custom_ApplicationName", "Application", "App Name"],
    "Custom_InvestmentDimension": ["Custom_InvestmentDimension", "Investment Dimension", "INVESTMENT_DIMENSION"],  # NEW
    "BusinessValue": ["BusinessValue", "Business Value", "Microsoft.VSTS.Common.BusinessValue", "Microsoft_VSTS_Common_BusinessValue"],
    "Iteration": ["Iteration", "Iteration Path", "System.IterationPath", "Iteration.IterationLevel3.2"],
    "Title": ["Title", "System.Title"],
    "State": ["State", "System.State"],
    "ID": ["ID", "Work Item ID", "System.Id", "WorkItemId", "Work Item Id"],
    "CreatedDate": ["Created Date", "System.CreatedDate", "CreatedDate"],
    "ChangedDate": ["Changed Date", "System.ChangedDate", "ChangedDate"],
    "Year": ["Year", "ADO Year", "ADO_YEAR"],
    "AreaPath": ["AreaPath", "Area.Path", "Area.AreaPath"],
    "AreaLevel1": ["AreaLevel1", "Area.Level1", "Area.AreaLevel1"],
    "AreaLevel2": ["AreaLevel2", "Area.Level2", "Area.AreaLevel2", "Team Area Level2"],
    "AreaLevel3": ["AreaLevel3", "Area.Level3", "Area.AreaLevel3"],
    "AreaLevel4": ["AreaLevel4", "Area.Level4", "Area.AreaLevel4"],
}

# -------------------------
# Utilities
# -------------------------
def _table_count(table: str) -> int:
    try:
        df = fetch_df(f"SELECT COUNT(*) AS N FROM {table}")
        return int(df.iloc[0]["N"]) if df is not None and not df.empty else 0
    except Exception:
        return 0

def _get_preview(name: str) -> pd.DataFrame:
    obj = st.session_state.get("previews", {}).get(name)
    return obj if isinstance(obj, pd.DataFrame) else pd.DataFrame()

def _blank_or_nan(s) -> bool:
    """True for None, '', 'nan' (string), and values that stringify to blank."""
    if s is None:
        return True
    try:
        txt = str(s).strip()
    except Exception:
        return True
    return (txt == "") or (txt.lower() == "nan")

@cache_data_portfolio(ttl=180, show_spinner=False)
def _ado_feature_columns() -> set:
    df = fetch_df(
        """
        SELECT COLUMN_NAME
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_NAME = 'ADO_FEATURES'
        """
    )
    if df is None or df.empty:
        return set()
    return {str(c).strip().upper() for c in df["COLUMN_NAME"].tolist()}


@cache_data_portfolio(ttl=60, show_spinner=False)
def _story_points_coverage_summary(cache_bust: int = 0) -> Dict[str, Any]:
    _ = cache_bust
    cols = _ado_feature_columns()
    where_parts: List[str] = []
    if "STATE" in cols:
        where_parts.append("(STATE IS NULL OR STATE <> 'Removed')")
    where_sql = " WHERE " + " AND ".join(where_parts) if where_parts else ""
    total = 0
    populated = 0
    try:
        df_total = fetch_df(f"SELECT COUNT(*) AS N FROM ADO_FEATURES{where_sql}")
        total = int(df_total.iloc[0]["N"]) if df_total is not None and not df_total.empty else 0
        sp_parts = list(where_parts)
        sp_parts.append("TRY_CONVERT(FLOAT, STORY_POINTS) IS NOT NULL")
        sp_parts.append("TRY_CONVERT(FLOAT, STORY_POINTS) > 0")
        sp_where = " WHERE " + " AND ".join(sp_parts) if sp_parts else ""
        df_pop = fetch_df(f"SELECT COUNT(*) AS N FROM ADO_FEATURES{sp_where}")
        populated = int(df_pop.iloc[0]["N"]) if df_pop is not None and not df_pop.empty else 0
    except Exception:
        total = 0
        populated = 0
    coverage_pct = (float(populated) / float(total) * 100.0) if total else 0.0
    return {"total": total, "populated": populated, "coverage_pct": coverage_pct}


def _render_story_points_coverage(coverage: Dict[str, Any]) -> None:
    total = int(coverage.get("total") or 0)
    populated = int(coverage.get("populated") or 0)
    pct = float(coverage.get("coverage_pct") or 0.0)
    if total <= 0:
        st.caption("Coverage: — (no synced features)")
        return
    msg = f"Coverage: {pct:.1f}% ({populated:,}/{total:,} Features)"
    if pct >= 80.0:
        st.success(msg)
    elif pct >= 30.0:
        st.info(msg)
    else:
        st.warning(f"{msg}. Low coverage — consider selecting a different Story Points field and re-sync.")
    st.caption("Coverage based on last synced/restored ADO_FEATURES data.")

def _ado_org_proj_defaults() -> Tuple[str, str]:
    """Best-effort org/project resolution from session or secrets for ADO links."""
    try:
        _ado_cfg = getattr(st, "secrets", {}).get("ado", {})  # type: ignore
    except Exception:
        _ado_cfg = {}
    org = (st.session_state.get("ado_org") or _ado_cfg.get("org") or "").strip()
    proj = (st.session_state.get("ado_proj") or _ado_cfg.get("project") or "").strip()
    return org, proj

def _ado_query_url(org: str, proj: str, feature_ids: List[Any]) -> str:
    """Build a WIQL query URL to list the given feature IDs in ADO."""
    if not org or not proj or not feature_ids:
        return ""
    ids = [str(int(float(i))) for i in feature_ids if pd.notna(i)]
    if not ids:
        return ""
    ids = ids[:200]  # keep URL reasonable
    wiql = f"Select [System.Id], [System.Title] From WorkItems Where [System.Id] In ({','.join(ids)})"
    return f"https://dev.azure.com/{org}/{proj}/_workitems/query?wiql={quote(wiql)}"

# -------------------------
# ADO OData helpers (direct sync)
# -------------------------
def _build_ado_odata_url(
    org: str,
    project: str,
    entity: str = "WorkItems",
    select: Optional[str] = None,
    expand: Optional[str] = None,
    filter_: Optional[str] = None,
    orderby: Optional[str] = None,
) -> str:
    base = f"https://analytics.dev.azure.com/{org}/{project}/_odata/v3.0-preview/{entity}"
    qs: List[str] = []
    if select:
        qs.append(f"$select={select}")
    if expand:
        qs.append(f"$expand={expand}")
    if filter_:
        qs.append(f"$filter={filter_}")
    if orderby:
        qs.append(f"$orderby={orderby}")
    return base + ("?" + "&".join(qs) if qs else "")


## fetch_ado_odata provided by utils.ado


## transform_ado_odata_to_expected provided by utils.ado

# -------------------------
# File parsing helpers
# -------------------------
def _auto_header_index(df_no_header: pd.DataFrame, expected_samples: List[str], max_scan: int = 20) -> Optional[int]:
    for i in range(min(max_scan, len(df_no_header))):
        row_vals = df_no_header.iloc[i].astype(str).str.strip().str.lower().tolist()
        hits = sum(1 for e in expected_samples if e.lower() in row_vals)
        if hits >= 2:
            return i
    return None

def _list_excel_sheets(data: bytes) -> List[str]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        return list(wb.sheetnames)
    except Exception:
        pass
    try:
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        return wb.sheet_names()
    except Exception:
        pass
    try:
        from pyxlsb import open_workbook
        with open_workbook(fileobj=io.BytesIO(data)) as wb:
            return [s.name for s in wb.sheets]
    except Exception:
        pass
    return []

def _read_excel_any(data: bytes, sheet_name: Optional[str], diag: Dict[str, Any]) -> Optional[pd.DataFrame]:
    errors: List[str] = []
    for eng in ("openpyxl", "xlrd", "pyxlsb"):
        try:
            __import__(eng)
            df = pd.read_excel(io.BytesIO(data), sheet_name=(sheet_name or 0), engine=eng)
            diag.setdefault("excel_engines_used", []).append(eng)
            return df
        except ModuleNotFoundError as e:
            errors.append(f"{eng} not installed: {e}")
        except Exception as e:
            errors.append(f"{eng} failed: {e}")
    try:
        df = pd.read_excel(io.BytesIO(data), sheet_name=(sheet_name or 0))
        diag.setdefault("excel_engines_used", []).append("auto")
        return df
    except Exception as e:
        errors.append(f"pandas auto engine failed: {e}")
    try:
        df_raw = pd.read_excel(io.BytesIO(data), sheet_name=(sheet_name or 0), header=None)
        hi = _auto_header_index(df_raw, expected_samples=["Title", "ID", "Team", "Iteration", "State", "Business Value", "BusinessValue"])
        if hi is not None:
            df = pd.read_excel(io.BytesIO(data), sheet_name=(sheet_name or 0), header=hi)
            diag.setdefault("header_autodetected", True)
            return df
        errors.append("Header auto-detect failed.")
    except Exception as e:
        errors.append(f"header=None strategy failed: {e}")

    diag["excel_errors"] = errors
    return None

def _read_csv_any(data: bytes, diag: Dict[str, Any]) -> Optional[pd.DataFrame]:
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin-1", "utf-16", "utf-16le", "utf-16be"]
    seps: List[Optional[str]] = [",", ";", "\t", None]
    errors: List[str] = []
    for enc in encodings:
        for sep in seps:
            try:
                df = pd.read_csv(io.BytesIO(data), encoding=enc, sep=sep, engine="python")
                if df.shape[1] >= 2:
                    diag.setdefault("csv_attempts", []).append({"encoding": enc, "sep": sep or "auto"})
                    return df
            except Exception as e:
                errors.append(f"csv {enc}/{sep or 'auto'} failed: {e}")
    diag["csv_errors"] = errors
    return None

def _read_file_any(upl, sheet_name: Optional[str], diag: Dict[str, Any]) -> pd.DataFrame:
    """
    IMPORTANT: never use `A or B` with DataFrames. Use explicit None/empty checks.
    """
    if upl is None:
        raise ValueError("No file uploaded.")
    upl.seek(0)
    raw = upl.read()
    upl.seek(0)
    name_lower = (upl.name or "").lower()
    looks_like_excel = any(ext in name_lower for ext in (".xlsx", ".xlsm", ".xls", ".xlsb")) or (b"\x00" in raw)

    df: Optional[pd.DataFrame] = None
    if looks_like_excel:
        df = _read_excel_any(raw, sheet_name, diag)
        if df is None or df.empty:
            df = _read_csv_any(raw, diag)
    else:
        df = _read_csv_any(raw, diag)
        if df is None or df.empty:
            df = _read_excel_any(raw, sheet_name, diag)

    if df is None or df.empty:
        raise ValueError("Could not parse the file. Try a clean XLSX (preferred) or CSV UTF‑8.")
    # Normalize column header BOM/whitespace
    df.columns = [str(c).strip().replace("\ufeff", "") for c in df.columns]
    return df


_MONTH_ABBR = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12,
}


def _norm_work_id(val) -> str:
    s = str(val or "").strip()
    s = re.sub(r"[^0-9A-Za-z]", "", s)
    s = s.lstrip("0")
    return s if s else "0"


def _normalize_apptio_actuals(df: pd.DataFrame) -> pd.DataFrame:
    """Normalize Apptio actuals wide table to tidy rows.

    Returns at least: WORK_ID, FISCAL_YEAR, MONTH, AMOUNT.
    When available, also returns dimensions for taxonomy mapping:
      - LEDGER_ACCOUNT_L3_DESC
      - PRODUCT_ID
      - PRODUCT_NAME
    """
    if df is None or df.empty:
        raise ValueError("Apptio sheet is empty.")
    cols_lower = {str(c).strip().lower(): c for c in df.columns}
    work_col = None
    for key in ("work id", "work_id", "workid"):
        if key in cols_lower:
            work_col = cols_lower[key]
            break
    if not work_col:
        raise ValueError("Work ID column not found. Add a 'Work ID' column.")

    month_cols: Dict[str, Tuple[int, int]] = {}
    for col in df.columns:
        m = re.match(r"\s*([A-Za-z]{3,5})\s*FY\s*(\d{4})", str(col), re.IGNORECASE)
        if not m:
            continue
        mon_abbr = m.group(1).strip().lower()
        fy = int(m.group(2))
        if mon_abbr in _MONTH_ABBR:
            month_cols[col] = (_MONTH_ABBR[mon_abbr], fy)

    if not month_cols:
        raise ValueError("No month columns found (expected headers like 'Jan FY 2025').")

    def _canon_header(s: Any) -> str:
        # Robust against Excel oddities: line breaks, multiple spaces, punctuation, etc.
        return re.sub(r"[^0-9a-z]+", " ", str(s or "").lower()).strip()

    cols_canon_in_order: list[tuple[str, str]] = [(_canon_header(c), str(c)) for c in df.columns]

    def _pick_col(candidates: list[str]) -> Optional[str]:
        candidates_canon = [_canon_header(c) for c in candidates if str(c or "").strip()]
        # Pass 1: exact canonical match (prefer candidate priority, then left-most column).
        for cand in candidates_canon:
            for ccanon, orig in cols_canon_in_order:
                if ccanon == cand:
                    return orig
        # Pass 2: substring canonical match (prefer candidate priority, then left-most column).
        for cand in candidates_canon:
            for ccanon, orig in cols_canon_in_order:
                if cand and cand in ccanon:
                    return orig
        return None

    ledger_col = _pick_col(
        [
            "ledger accounts : account l3 medium text description",
            "account l3 medium text description",
            "ledger account l3",
            "ledger l3",
        ]
    )
    product_id_col = _pick_col(
        [
            "nwf / products : nwf id",
            "nwf id",
            "product id",
            "nwf_product_id",
        ]
    )
    product_name_col = _pick_col(
        [
            # Primary: column G in the standard Apptio export
            "nwf / products : nwf id name",
            "nwf / products : nwf name",
            "nwf name",
            "product name",
            "nwf_product_name",
        ]
    )

    extra_cols: list[str] = []
    if ledger_col:
        extra_cols.append(ledger_col)
    if product_id_col:
        extra_cols.append(product_id_col)
    if product_name_col:
        extra_cols.append(product_name_col)

    df_work = df[[work_col] + extra_cols + list(month_cols.keys())].copy()
    df_work = df_work.rename(columns={work_col: "WORK_ID"})
    df_work["WORK_ID"] = df_work["WORK_ID"].apply(_norm_work_id)
    rename_map: dict[str, str] = {}
    if ledger_col:
        rename_map[ledger_col] = "LEDGER_ACCOUNT_L3_DESC"
    if product_id_col:
        rename_map[product_id_col] = "PRODUCT_ID"
    if product_name_col:
        rename_map[product_name_col] = "PRODUCT_NAME"
    if rename_map:
        df_work = df_work.rename(columns=rename_map)

    id_vars = ["WORK_ID"] + [c for c in ["LEDGER_ACCOUNT_L3_DESC", "PRODUCT_ID", "PRODUCT_NAME"] if c in df_work.columns]
    long_df = df_work.melt(id_vars=id_vars, value_vars=list(month_cols.keys()), var_name="PERIOD", value_name="AMOUNT")
    long_df["AMOUNT"] = pd.to_numeric(long_df["AMOUNT"], errors="coerce")
    long_df = long_df.dropna(subset=["AMOUNT"])
    long_df["AMOUNT"] = long_df["AMOUNT"].astype(float)
    long_df["MONTH"] = long_df["PERIOD"].map({k: v[0] for k, v in month_cols.items()})
    long_df["FISCAL_YEAR"] = long_df["PERIOD"].map({k: v[1] for k, v in month_cols.items()})
    long_df = long_df.dropna(subset=["MONTH", "FISCAL_YEAR"])
    long_df["MONTH"] = long_df["MONTH"].astype(int)
    long_df["FISCAL_YEAR"] = long_df["FISCAL_YEAR"].astype(int)
    if "LEDGER_ACCOUNT_L3_DESC" not in long_df.columns:
        long_df["LEDGER_ACCOUNT_L3_DESC"] = None
    if "PRODUCT_ID" not in long_df.columns:
        long_df["PRODUCT_ID"] = None
    if "PRODUCT_NAME" not in long_df.columns:
        long_df["PRODUCT_NAME"] = None

    long_df["LEDGER_ACCOUNT_L3_DESC"] = long_df.get("LEDGER_ACCOUNT_L3_DESC").fillna("").astype(str).str.strip()
    long_df.loc[long_df["LEDGER_ACCOUNT_L3_DESC"].eq(""), "LEDGER_ACCOUNT_L3_DESC"] = None
    long_df["PRODUCT_ID"] = long_df.get("PRODUCT_ID").fillna("").astype(str).str.strip()
    long_df.loc[long_df["PRODUCT_ID"].eq(""), "PRODUCT_ID"] = None
    long_df["PRODUCT_NAME"] = long_df.get("PRODUCT_NAME").fillna("").astype(str).str.strip()
    long_df.loc[long_df["PRODUCT_NAME"].eq(""), "PRODUCT_NAME"] = None

    keys = ["WORK_ID", "FISCAL_YEAR", "MONTH", "PRODUCT_ID", "LEDGER_ACCOUNT_L3_DESC"]

    def _first_nonempty(s: pd.Series) -> Optional[str]:
        for v in s.dropna().astype(str).tolist():
            vv = str(v).strip()
            if vv:
                return vv
        return None

    agg = (
        long_df[keys + ["AMOUNT", "PRODUCT_NAME"]]
        .groupby(keys, as_index=False)
        .agg({"AMOUNT": "sum", "PRODUCT_NAME": _first_nonempty})
    )
    return agg

# -------------------------
# Template generator (Bulk Load)
# -------------------------
def generate_bulk_template() -> bytes:
    """Build an Excel template for the Bulk Load tab with required columns.

    If tables already have values, include a Lists sheet and add Excel
    data‑validation dropdowns for entity names on the Data sheet.
    """
    buf = io.BytesIO()
    # Recommended column order (covers required + optional fields used by import)
    cols = [
        # Core entities
        "PROGRAMNAME", "TEAMNAME", "GROUPNAME", "APPNAME", "VENDORNAME",
        # Contract (required for Recurring invoices)
        "CONTRACT_START_FY", "CONTRACT_END_FY",
        "CONTRACT_ANNUAL_AMOUNT", "CONTRACT_ESCALATION_PCT", "CONTRACT_STATUS", "CONTRACT_AGREEMENT_NUMBER",
        "CONTRACT_COMPANY_CODE", "CONTRACT_COST_CENTER", "CONTRACT_SERVICE_TYPE", "CONTRACT_RENEWAL_DATE",
        "INVOICE_RENEWAL_DATE", "CONTRACT_TOTAL_COST",
        # Invoice requireds
        "AMOUNT", "FISCAL_YEAR", "RENEWAL_MONTH",
        # Invoice optional
        "INVOICE_STATUS", "CONTRACT_ACTIVE", "SERIAL_NUMBER", "WORK_ORDER", "COMPANY_CODE",
        "COST_CENTER", "PRODUCT_OWNER", "NOTES",
    ]
    df = pd.DataFrame(columns=cols)

    # Read current values from DB (best effort)
    try:
        progs_df = list_programs()
    except Exception:
        progs_df = pd.DataFrame()
    try:
        teams_df = list_teams()
    except Exception:
        teams_df = pd.DataFrame()
    try:
        groups_df = list_application_groups()
    except Exception:
        groups_df = pd.DataFrame()
    try:
        apps_df = list_applications()
    except Exception:
        apps_df = pd.DataFrame()
    try:
        vendors_df = list_vendors()
    except Exception:
        vendors_df = pd.DataFrame()

    def uniq_col(df0: pd.DataFrame, name: str) -> List[str]:
        if df0 is None or df0.empty or name not in df0.columns:
            return []
        return sorted(df0[name].dropna().astype(str).str.strip().unique().tolist())

    # Force a controlled list for Programs to avoid unnecessary entities
    prog_names = ALLOWED_PROGRAMS
    team_names = uniq_col(teams_df, "TEAMNAME")
    group_names = uniq_col(groups_df, "GROUPNAME")
    app_names = uniq_col(apps_df, "APPLICATIONNAME")
    vendor_names = uniq_col(vendors_df, "VENDORNAME")

    # Instructions sheet content
    instr: List[Dict[str, Any]] = [
        {"Field": "PROGRAMNAME",      "Required": "Yes", "Description": "Program (pick from allowed list)", "Example": "R&M"},
        {"Field": "TEAMNAME",         "Required": "Yes", "Description": "Team name (linked to Program if provided)", "Example": "ERNE WEST"},
        {"Field": "GROUPNAME",        "Required": "Yes", "Description": "This is an application group; multiple application instances can be aggregated here. If you don't have multiple application instances, the group can be named after your unique application instance/site.", "Example": "Application - MHM"},
        {"Field": "APPNAME",          "Required": "Yes", "Description": "Application instance name", "Example": "MHM Baytown"},
        {"Field": "VENDORNAME",       "Required": "Yes", "Description": "Vendor name", "Example": "Emerson"},
        {"Field": "CONTRACT_START_FY","Required": "Yes", "Description": "Contract start fiscal year", "Example": "2024"},
        {"Field": "CONTRACT_END_FY",  "Required": "Yes", "Description": "Contract end fiscal year (Contract Due)", "Example": "2027"},
        # CONTRACT_RENEWAL_MONTH intentionally omitted from template (use CONTRACT_RENEWAL_DATE or INVOICE_RENEWAL_DATE)
        {"Field": "CONTRACT_ANNUAL_AMOUNT", "Required": "Yes", "Description": "Annual amount at START_FY (numeric)", "Example": "12500.00"},
        {"Field": "CONTRACT_ESCALATION_PCT", "Required": "No", "Description": "% escalation per year (0..100)", "Example": "3.0"},
        {"Field": "CONTRACT_STATUS", "Required": "No", "Description": "Active/Terminated", "Example": "Active"},
        {"Field": "CONTRACT_AGREEMENT_NUMBER", "Required": "No", "Description": "Agreement/reference number", "Example": "AGR-123"},
        {"Field": "CONTRACT_COMPANY_CODE", "Required": "No", "Description": "Default Company Code for invoices", "Example": "1000"},
        {"Field": "CONTRACT_COST_CENTER", "Required": "No", "Description": "Default Cost Center for invoices", "Example": "CC-789"},
        {"Field": "CONTRACT_SERVICE_TYPE", "Required": "No", "Description": "Default Service Type", "Example": "SaaS"},
        {"Field": "CONTRACT_RENEWAL_DATE", "Required": "No", "Description": "Contract renewal date (MM/DD/YYYY)", "Example": "01/15/2025"},
        {"Field": "AMOUNT",           "Required": "No",  "Description": "Invoice amount (numeric). Optional — if omitted and covered by a Contract, planned amount is seeded from Contract.", "Example": "12500.00"},
        {"Field": "FISCAL_YEAR",      "Required": "Yes", "Description": "Year (integer)", "Example": "2025"},
        {"Field": "RENEWAL_MONTH",    "Required": "Yes", "Description": "Renewal month 1–12", "Example": "4"},
        {"Field": "INVOICE_STATUS",   "Required": "No",  "Description": "Recurring invoice status (Planned/Completed)", "Example": "Planned"},
        {"Field": "CONTRACT_ACTIVE",  "Required": "No",  "Description": "Derived from Contract Status (Active → TRUE, Terminated → FALSE)", "Example": "(derived)"},
        {"Field": "SERIAL_NUMBER",    "Required": "No",  "Description": "Serial / reference", "Example": "SN-123"},
        {"Field": "WORK_ORDER",       "Required": "No",  "Description": "Work order", "Example": "WO-456"},
        {"Field": "COMPANY_CODE",     "Required": "No",  "Description": "Company code", "Example": "1000"},
        {"Field": "COST_CENTER",      "Required": "No",  "Description": "Cost center", "Example": "CC-789"},
        {"Field": "PRODUCT_OWNER",    "Required": "No",  "Description": "Owner", "Example": "Alice"},
        {"Field": "NOTES",            "Required": "No",  "Description": "Free text notes", "Example": "Renewal changed to Apr"},
    ]

    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill, Font
    from openpyxl.formatting.rule import FormulaRule

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data")
        pd.DataFrame(instr).to_excel(writer, index=False, sheet_name="Instructions")

        wb = writer.book
        ws_data = writer.sheets["Data"]
        # Freeze header row for convenience in Excel (Mac/Windows)
        try:
            ws_data.freeze_panes = "A2"
        except Exception:
            pass
        # Map header name -> column index for styling and validation
        header_to_idx = {h: i+1 for i, h in enumerate(cols)}

        # Create Lists sheet with current values (if any)
        any_lists = any([prog_names, team_names, group_names, app_names, vendor_names])
        if any_lists:
            ws_lists = wb.create_sheet("Lists")
            ws_lists["A1"] = "PROGRAMNAME";    ws_lists["B1"] = "TEAMNAME"
            ws_lists["C1"] = "GROUPNAME";      ws_lists["D1"] = "APPLICATIONNAME"
            ws_lists["E1"] = "VENDORNAME";     ws_lists["F1"] = "CONTRACT_STATUS"
            ws_lists["G1"] = "SERVICE_TYPE";   ws_lists["H1"] = "INVOICE_STATUS"

            for i, v in enumerate(prog_names, start=2):
                ws_lists.cell(row=i, column=1, value=v)
            for i, v in enumerate(team_names, start=2):
                ws_lists.cell(row=i, column=2, value=v)
            for i, v in enumerate(group_names, start=2):
                ws_lists.cell(row=i, column=3, value=v)
            for i, v in enumerate(app_names, start=2):
                ws_lists.cell(row=i, column=4, value=v)
            for i, v in enumerate(vendor_names, start=2):
                ws_lists.cell(row=i, column=5, value=v)
            # Status + Service type choices
            status_opts = ["Active", "Terminated"]
            svc_opts = ["OnPrem", "SaaS", "IaaS", "Other"]
            inv_status_opts = ["Planned", "Completed"]
            for i, v in enumerate(status_opts, start=2):
                ws_lists.cell(row=i, column=6, value=v)
            for i, v in enumerate(svc_opts, start=2):
                ws_lists.cell(row=i, column=7, value=v)
            for i, v in enumerate(inv_status_opts, start=2):
                ws_lists.cell(row=i, column=8, value=v)

            # Helper to add DV to an entire column in Data sheet
            def add_list_validation(header: str, lists_col: int, last_row: int):
                if last_row < 2:
                    return
                try:
                    idx = cols.index(header) + 1
                except ValueError:
                    return
                col_letter = get_column_letter(idx)
                first = 2; dest_last = 2000
                src_range = f"Lists!${get_column_letter(lists_col)}$2:${get_column_letter(lists_col)}${last_row}"
                dv = DataValidation(type="list", formula1=f"={src_range}", allow_blank=True)
                ws_data.add_data_validation(dv)
                dv.add(f"{col_letter}{first}:{col_letter}{dest_last}")

            add_list_validation("PROGRAMNAME", 1, 1 + len(prog_names))
            add_list_validation("TEAMNAME",    2, 1 + len(team_names))
            add_list_validation("GROUPNAME",   3, 1 + len(group_names))
            add_list_validation("APPNAME",     4, 1 + len(app_names))
            add_list_validation("VENDORNAME",  5, 1 + len(vendor_names))
            add_list_validation("CONTRACT_STATUS", 6, 1 + len(status_opts))
            add_list_validation("CONTRACT_SERVICE_TYPE", 7, 1 + len(svc_opts))
            add_list_validation("INVOICE_STATUS", 8, 1 + len(inv_status_opts))

        # Numeric validations
        def add_numeric_validation(header: str, dv: DataValidation):
            try:
                idx = cols.index(header) + 1
            except ValueError:
                return
            col_letter = get_column_letter(idx)
            first = 2; dest_last = 2000
            ws_data.add_data_validation(dv)
            dv.add(f"{col_letter}{first}:{col_letter}{dest_last}")

        # Force text formatting for code fields that may contain leading zeros
        def enforce_text_column(header: str):
            try:
                idx = cols.index(header) + 1
            except ValueError:
                return
            for r in range(2, 2001):
                ws_data.cell(row=r, column=idx).number_format = '@'

        # Months 1-12
        for header in ("RENEWAL_MONTH",):
            dv = DataValidation(type="whole", operator="between", formula1="1", formula2="12", allow_blank=True)
            add_numeric_validation(header, dv)
        # Fiscal years
        for header in ("FISCAL_YEAR", "CONTRACT_START_FY", "CONTRACT_END_FY"):
            dv = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100", allow_blank=False)
            add_numeric_validation(header, dv)
        # Non-negative amounts and escalation percent
        for header in ("AMOUNT", "CONTRACT_ANNUAL_AMOUNT", "CONTRACT_TOTAL_COST"):
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
            add_numeric_validation(header, dv)
        dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True)
        add_numeric_validation("CONTRACT_ESCALATION_PCT", dv)

        # Ensure codes are treated as text (not numeric)
        enforce_text_column("CONTRACT_COMPANY_CODE")
        enforce_text_column("COMPANY_CODE")

        # Header comments (tooltips)
        try:
            field_help = {row["Field"]: f"{row['Description']}" + (f" (Example: {row['Example']})" if row.get('Example') else "") for row in instr}
            for h in cols:
                if h in field_help:
                    i = header_to_idx.get(h)
                    if i:
                        cell = ws_data.cell(row=1, column=i)
                        cell.comment = Comment(field_help[h], "NEXT")
        except Exception:
            pass

        # Header colors: mark required fields (obligated) for quick visual guidance
        try:
            # Use a bright Excel-friendly yellow for visibility (Mac/Windows)
            req_header_fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
            for h in req_fields:
                i = header_to_idx.get(h)
                if i:
                    c = ws_data.cell(row=1, column=i)
                    c.fill = req_header_fill
                    try:
                        c.font = Font(bold=True)
                    except Exception:
                        pass
        except Exception:
            pass

        # Also highlight required rows in Instructions sheet for clarity
        try:
            ws_instr = writer.sheets.get("Instructions")
            if ws_instr is not None:
                # Find header positions in Instructions
                hdr_map = {}
                for col in range(1, ws_instr.max_column + 1):
                    val = ws_instr.cell(row=1, column=col).value
                    if isinstance(val, str):
                        hdr_map[val.strip().upper()] = col
                col_required = hdr_map.get("REQUIRED")
                col_field = hdr_map.get("FIELD")
                if col_required and col_field:
                    for r in range(2, ws_instr.max_row + 1):
                        if str(ws_instr.cell(row=r, column=col_required).value).strip().lower() == "yes":
                            cell = ws_instr.cell(row=r, column=col_field)
                            cell.fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
                            try:
                                cell.font = Font(bold=True)
                            except Exception:
                                pass
        except Exception:
            pass

        # Conditional formatting: missing requireds
        req_fields = [
            "PROGRAMNAME","TEAMNAME","GROUPNAME","APPNAME","VENDORNAME",
            "CONTRACT_START_FY","CONTRACT_END_FY","CONTRACT_ANNUAL_AMOUNT",
            "FISCAL_YEAR","RENEWAL_MONTH",
        ]
        try:
            red_fill = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
            for h in req_fields:
                if h not in header_to_idx:
                    continue
                col_letter = get_column_letter(header_to_idx[h])
                rng = f"{col_letter}2:{col_letter}2000"
                ws_data.conditional_formatting.add(rng, FormulaRule(formula=[f"LEN(${col_letter}2)=0"], fill=red_fill))
        except Exception:
            pass

        # Conditional formatting: duplicate detection (Contracts: TEAMNAME+APPNAME; Invoices: TEAMNAME+APPNAME+FISCAL_YEAR)
        try:
            yellow_fill = PatternFill(start_color="FFFFEB9C", end_color="FFFFEB9C", fill_type="solid")
            # Contracts dupes on TEAMNAME + APPNAME
            if all(k in header_to_idx for k in ("TEAMNAME","APPNAME")):
                t_col = get_column_letter(header_to_idx["TEAMNAME"]) ; a_col = get_column_letter(header_to_idx["APPNAME"]) 
                rng = f"{t_col}2:{t_col}2000"
                formula = f"COUNTIFS($${t_col}$2:$${t_col}$2000,$${t_col}2,$${a_col}$2:$${a_col}$2000,$${a_col}2)>1".replace('$$', '$')
                ws_data.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=yellow_fill))
                rng2 = f"{a_col}2:{a_col}2000"
                ws_data.conditional_formatting.add(rng2, FormulaRule(formula=[formula], fill=yellow_fill))
            # Invoice dupes on TEAMNAME + APPNAME + FISCAL_YEAR
            if all(k in header_to_idx for k in ("TEAMNAME","APPNAME","FISCAL_YEAR")):
                t_col = get_column_letter(header_to_idx["TEAMNAME"]) ; a_col = get_column_letter(header_to_idx["APPNAME"]) ; y_col = get_column_letter(header_to_idx["FISCAL_YEAR"])
                rng = f"{y_col}2:{y_col}2000"
                formula = (
                    f"COUNTIFS($${t_col}$2:$${t_col}$2000,$${t_col}2,$${a_col}$2:$${a_col}$2000,$${a_col}2,$${y_col}$2:$${y_col}$2000,$${y_col}2)>1"
                ).replace('$$', '$')
                ws_data.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=yellow_fill))
        except Exception:
            pass

        # Pre-fill first data row with examples from actual values when available
        def set_cell(name: str, value):
            i = header_to_idx.get(name)
            if i:
                ws_data.cell(row=2, column=i, value=value)

        try:
            import datetime as _dt
            # Names: best-effort pick first available
            if prog_names:
                set_cell("PROGRAMNAME", "R&M")
            if team_names:
                set_cell("TEAMNAME", team_names[0])
            if group_names:
                set_cell("GROUPNAME", group_names[0])
            if app_names:
                set_cell("APPNAME", app_names[0])
            if vendor_names:
                set_cell("VENDORNAME", vendor_names[0])
            # Contracts
            year_now = _dt.date.today().year
            set_cell("CONTRACT_START_FY", year_now)
            set_cell("CONTRACT_END_FY", year_now + 2)
            set_cell("CONTRACT_ANNUAL_AMOUNT", 10000.00)
            set_cell("CONTRACT_ESCALATION_PCT", 3.0)
            set_cell("CONTRACT_STATUS", "Active")
            set_cell("CONTRACT_AGREEMENT_NUMBER", "AGR-001")
            set_cell("CONTRACT_COMPANY_CODE", "1000")
            set_cell("CONTRACT_COST_CENTER", "CC-001")
            set_cell("CONTRACT_SERVICE_TYPE", "SaaS")
            set_cell("CONTRACT_TOTAL_COST", 30000.00)
            set_cell("CONTRACT_RENEWAL_DATE", "01/15/2025")
            # Invoices
            set_cell("AMOUNT", 10000.00)
            set_cell("FISCAL_YEAR", year_now)
            set_cell("RENEWAL_MONTH", 1)
            set_cell("INVOICE_STATUS", "Planned")
            set_cell("CONTRACT_ACTIVE", True)
            set_cell("PRODUCT_OWNER", "Alice")
            set_cell("NOTES", "Example row; replace with your data")
        except Exception:
            pass

        # Auto-fit simple: based on header/examples
        try:
            from openpyxl.utils import get_column_letter as _gcl
            for h in cols:
                i = header_to_idx.get(h)
                if not i:
                    continue
                values = [str(h), str(ws_data.cell(row=2, column=i).value or "")]
                width = max(len(v) for v in values) + 4
                ws_data.column_dimensions[_gcl(i)].width = max(12, min(width, 60))
        except Exception:
            pass

        # Add end-of-row red warning text on example row
        try:
            from openpyxl.styles import Font as _Font
            max_idx = max(header_to_idx.values()) if header_to_idx else 1
            note_col = get_column_letter(max_idx + 1)
            ws_data.cell(row=2, column=max_idx + 1, value="<- THIS IS JUST AN EXAMPLE, DELETE THIS ROW BEFORE THE IMPORT").font = _Font(color="FF0000", bold=True)
        except Exception:
            pass

    buf.seek(0)
    return buf.getvalue()

ALLOWED_PROGRAMS = ["R&M", "Process Ops", "QM", "WFE", "Central Services"]

def generate_contracts_template() -> bytes:
    """Build an Excel template focused on Contracts (and core entities), without invoice columns.

    Includes dropdowns, numeric validations, header tooltips/colors, and a first example row.
    """
    buf = io.BytesIO()
    cols = [
        # Core entities
        "PROGRAMNAME", "TEAMNAME", "GROUPNAME", "APPNAME", "VENDORNAME",
        # Contracts only
        "CONTRACT_START_FY", "CONTRACT_END_FY",
        "CONTRACT_ANNUAL_AMOUNT", "CONTRACT_ESCALATION_PCT", "CONTRACT_STATUS", "CONTRACT_AGREEMENT_NUMBER",
        "CONTRACT_COMPANY_CODE", "CONTRACT_COST_CENTER", "CONTRACT_SERVICE_TYPE", "CONTRACT_RENEWAL_DATE",
        "INVOICE_RENEWAL_DATE",
        "CONTRACT_TOTAL_COST",
    ]
    df = pd.DataFrame(columns=cols)

    # Current values from DB for dropdowns
    try:
        progs_df = list_programs()
    except Exception:
        progs_df = pd.DataFrame()
    try:
        teams_df = list_teams()
    except Exception:
        teams_df = pd.DataFrame()
    try:
        groups_df = list_application_groups()
    except Exception:
        groups_df = pd.DataFrame()
    try:
        apps_df = list_applications()
    except Exception:
        apps_df = pd.DataFrame()
    try:
        vendors_df = list_vendors()
    except Exception:
        vendors_df = pd.DataFrame()

    def uniq_col(df0: pd.DataFrame, name: str) -> List[str]:
        if df0 is None or df0.empty or name not in df0.columns:
            return []
        return sorted(df0[name].dropna().astype(str).str.strip().unique().tolist())

    # Use controlled list for Programs to avoid unnecessary entities
    prog_names = ALLOWED_PROGRAMS
    team_names = uniq_col(teams_df, "TEAMNAME")
    group_names = uniq_col(groups_df, "GROUPNAME")
    app_names = uniq_col(apps_df, "APPLICATIONNAME")
    vendor_names = uniq_col(vendors_df, "VENDORNAME")

    instr: List[Dict[str, Any]] = [
        {"Field": "PROGRAMNAME",      "Required": "Yes", "Description": "Program (pick from allowed list)", "Example": "R&M"},
        {"Field": "TEAMNAME",         "Required": "Yes", "Description": "Team name (linked to Program if provided)", "Example": "ERNE WEST"},
        {"Field": "GROUPNAME",        "Required": "Yes", "Description": "Application. Multiple application instances can be aggregated here. If you only have a single instance/site, you can name the group after the instance/site.", "Example": "Application - MHM"},
        {"Field": "APPNAME",          "Required": "Yes", "Description": "Application instance name", "Example": "MHM Baytown"},
        {"Field": "VENDORNAME",       "Required": "Yes", "Description": "Vendor name", "Example": "Emerson"},
        {"Field": "CONTRACT_START_FY","Required": "Yes", "Description": "Contract start fiscal year", "Example": "2024"},
        {"Field": "CONTRACT_END_FY",  "Required": "Yes", "Description": "Contract end fiscal year (Contract Due)", "Example": "2027"},
        # CONTRACT_RENEWAL_MONTH intentionally omitted from template (use CONTRACT_RENEWAL_DATE or INVOICE_RENEWAL_DATE)
        {"Field": "CONTRACT_ANNUAL_AMOUNT", "Required": "Yes", "Description": "Annual amount at START_FY (numeric)", "Example": "12500.00"},
        {"Field": "CONTRACT_ESCALATION_PCT", "Required": "No", "Description": "% escalation per year (0..100)", "Example": "3.0"},
        {"Field": "CONTRACT_STATUS", "Required": "No", "Description": "Active/Terminated", "Example": "Active"},
        {"Field": "CONTRACT_AGREEMENT_NUMBER", "Required": "No", "Description": "Agreement/reference number", "Example": "AGR-123"},
        {"Field": "CONTRACT_COMPANY_CODE", "Required": "No", "Description": "Default Company Code for invoices", "Example": "1000"},
        {"Field": "CONTRACT_COST_CENTER", "Required": "No", "Description": "Default Cost Center for invoices", "Example": "CC-789"},
        {"Field": "CONTRACT_SERVICE_TYPE", "Required": "No", "Description": "Default Service Type", "Example": "SaaS"},
        {"Field": "CONTRACT_RENEWAL_DATE", "Required": "No", "Description": "Contract renewal date (MM/DD/YYYY)", "Example": "01/15/2025"},
        {"Field": "INVOICE_RENEWAL_DATE", "Required": "No", "Description": "Default invoice renewal date (MM/DD)", "Example": "01/10"},
        {"Field": "CONTRACT_TOTAL_COST", "Required": "No", "Description": "Total contract cost for entire period (sum of annuals)", "Example": "30000.00"},
    ]

    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill, Font
    from openpyxl.formatting.rule import FormulaRule

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data")
        pd.DataFrame(instr).to_excel(writer, index=False, sheet_name="Instructions")
        wb = writer.book
        ws_data = writer.sheets["Data"]
        try:
            ws_data.freeze_panes = "A2"
        except Exception:
            pass
        header_to_idx = {h: i+1 for i, h in enumerate(cols)}

        # Lists sheet
        any_lists = any([prog_names, team_names, group_names, app_names, vendor_names])
        if any_lists:
            ws_lists = wb.create_sheet("Lists")
            ws_lists["A1"] = "PROGRAMNAME";    ws_lists["B1"] = "TEAMNAME"
            ws_lists["C1"] = "GROUPNAME";      ws_lists["D1"] = "APPLICATIONNAME"
            ws_lists["E1"] = "VENDORNAME";     ws_lists["F1"] = "CONTRACT_STATUS"
            ws_lists["G1"] = "SERVICE_TYPE"
            for i, v in enumerate(prog_names, start=2): ws_lists.cell(row=i, column=1, value=v)
            for i, v in enumerate(team_names, start=2): ws_lists.cell(row=i, column=2, value=v)
            for i, v in enumerate(group_names, start=2): ws_lists.cell(row=i, column=3, value=v)
            for i, v in enumerate(app_names, start=2): ws_lists.cell(row=i, column=4, value=v)
            for i, v in enumerate(vendor_names, start=2): ws_lists.cell(row=i, column=5, value=v)
            status_opts = ["Active", "Terminated"]
            svc_opts = ["OnPrem", "SaaS", "IaaS", "Other"]
            for i, v in enumerate(status_opts, start=2): ws_lists.cell(row=i, column=6, value=v)
            for i, v in enumerate(svc_opts, start=2): ws_lists.cell(row=i, column=7, value=v)

            def add_list_validation(header: str, lists_col: int, last_row: int):
                if last_row < 2:
                    return
                try:
                    idx = cols.index(header) + 1
                except ValueError:
                    return
                col_letter = get_column_letter(idx)
                first = 2; dest_last = 2000
                src_range = f"Lists!${get_column_letter(lists_col)}$2:${get_column_letter(lists_col)}${last_row}"
                dv = DataValidation(type="list", formula1=f"={src_range}", allow_blank=True)
                ws_data.add_data_validation(dv)
                dv.add(f"{col_letter}{first}:{col_letter}{dest_last}")

            add_list_validation("PROGRAMNAME", 1, 1 + len(prog_names))
            add_list_validation("TEAMNAME",    2, 1 + len(team_names))
            add_list_validation("GROUPNAME",   3, 1 + len(group_names))
            add_list_validation("APPNAME",     4, 1 + len(app_names))
            add_list_validation("VENDORNAME",  5, 1 + len(vendor_names))
            add_list_validation("CONTRACT_STATUS", 6, 1 + len(status_opts))
            add_list_validation("CONTRACT_SERVICE_TYPE", 7, 1 + len(svc_opts))

        # Numeric validations
        def add_numeric_validation(header: str, dv: DataValidation):
            try:
                idx = cols.index(header) + 1
            except ValueError:
                return
            col_letter = get_column_letter(idx)
            first = 2; dest_last = 2000
            ws_data.add_data_validation(dv)
            dv.add(f"{col_letter}{first}:{col_letter}{dest_last}")

        # Force text formatting for code fields that may contain leading zeros
        def enforce_text_column(header: str):
            try:
                idx = cols.index(header) + 1
            except ValueError:
                return
            for r in range(2, 2001):
                ws_data.cell(row=r, column=idx).number_format = '@'

        # CONTRACT_RENEWAL_MONTH intentionally omitted from template
        for header in ("CONTRACT_START_FY", "CONTRACT_END_FY"):
            dv = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100", allow_blank=False)
            add_numeric_validation(header, dv)
        for header in ("CONTRACT_ANNUAL_AMOUNT", "CONTRACT_TOTAL_COST"):
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
            add_numeric_validation(header, dv)
        dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True)
        add_numeric_validation("CONTRACT_ESCALATION_PCT", dv)

        # Ensure CONTRACT_COMPANY_CODE is treated as text (not numeric)
        enforce_text_column("CONTRACT_COMPANY_CODE")

        # Header tooltips and colors + banner note
        try:
            field_help = {row["Field"]: f"{row['Description']}" + (f" (Example: {row['Example']})" if row.get('Example') else "") for row in instr}
            for h in cols:
                i = header_to_idx.get(h)
                if i and h in field_help:
                    cell = ws_data.cell(row=1, column=i)
                    cell.comment = Comment(field_help[h], "NEXT")
        except Exception:
            pass
        # Add a banner row in Instructions
        try:
            ws_instr = writer.sheets.get("Instructions")
            if ws_instr is not None:
                ws_instr.insert_rows(1)
                ws_instr.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
                cell = ws_instr.cell(row=1, column=1)
                cell.value = (
                    "Importing Contracts will create planned Recurring Invoices automatically for each fiscal year between Start and End FY. "
                    "INVOICE_RENEWAL_DATE (optional) sets the default invoice renewal date."
                )
            
        except Exception:
            pass
        try:
            req_fields = [
                "PROGRAMNAME","TEAMNAME","GROUPNAME","APPNAME","VENDORNAME",
                "CONTRACT_START_FY","CONTRACT_END_FY","CONTRACT_ANNUAL_AMOUNT",
            ]
            req_header_fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
            for h in req_fields:
                i = header_to_idx.get(h)
                if i:
                    c = ws_data.cell(row=1, column=i)
                    c.fill = req_header_fill
                    try:
                        c.font = Font(bold=True)
                    except Exception:
                        pass
        except Exception:
            pass

        # Prefill example row
        def set_cell(name: str, value):
            i = header_to_idx.get(name)
            if i:
                ws_data.cell(row=2, column=i, value=value)
        try:
            import datetime as _dt
            if prog_names: set_cell("PROGRAMNAME", "R&M")
            if team_names: set_cell("TEAMNAME", team_names[0])
            if group_names: set_cell("GROUPNAME", group_names[0])
            if app_names:  set_cell("APPNAME", "MHM Baytown")
            if vendor_names: set_cell("VENDORNAME", "Emerson")
            year_now = _dt.date.today().year
            set_cell("CONTRACT_START_FY", year_now)
            set_cell("CONTRACT_END_FY", year_now + 2)
            set_cell("CONTRACT_ANNUAL_AMOUNT", 10000.00)
            set_cell("CONTRACT_ESCALATION_PCT", 3.0)
            set_cell("CONTRACT_STATUS", "Active")
            set_cell("CONTRACT_AGREEMENT_NUMBER", "AGR-001")
            set_cell("CONTRACT_COMPANY_CODE", "1000")
            set_cell("CONTRACT_COST_CENTER", "CC-001")
            set_cell("CONTRACT_SERVICE_TYPE", "SaaS")
            set_cell("CONTRACT_TOTAL_COST", 30000.00)
            set_cell("INVOICE_RENEWAL_DATE", "01/10")
        except Exception:
            pass

        # Auto-fit simple: based on header/examples
        try:
            from openpyxl.utils import get_column_letter as _gcl
            for h in cols:
                i = header_to_idx.get(h)
                if not i:
                    continue
                values = [str(h), str(ws_data.cell(row=2, column=i).value or "")]
                width = max(len(v) for v in values) + 4
                ws_data.column_dimensions[_gcl(i)].width = max(12, min(width, 60))
        except Exception:
            pass

        # Add end-of-row red warning text on example row
        try:
            from openpyxl.styles import Font as _Font
            max_idx = max(header_to_idx.values()) if header_to_idx else 1
            note_col = get_column_letter(max_idx + 1)
            ws_data.cell(row=2, column=max_idx + 1, value="<- THIS IS JUST AN EXAMPLE, DELETE THIS ROW BEFORE THE IMPORT").font = _Font(color="FF0000", bold=True)
        except Exception:
            pass

    buf.seek(0)
    return buf.getvalue()

def generate_core_entities_template() -> bytes:
    buf = io.BytesIO()
    cols = [
        "PROGRAMNAME", "TEAMNAME", "PRODUCT_OWNER", "GROUPNAME", "APPNAME", "VENDORNAME",
    ]
    df = pd.DataFrame(columns=cols)
    try:
        progs_df = list_programs()
    except Exception:
        progs_df = pd.DataFrame()
    try:
        teams_df = list_teams()
    except Exception:
        teams_df = pd.DataFrame()
    try:
        groups_df = list_application_groups()
    except Exception:
        groups_df = pd.DataFrame()
    try:
        apps_df = list_applications()
    except Exception:
        apps_df = pd.DataFrame()
    try:
        vendors_df = list_vendors()
    except Exception:
        vendors_df = pd.DataFrame()

    team_names = sorted(teams_df["TEAMNAME"].dropna().astype(str).str.strip().unique().tolist()) if not teams_df.empty else []
    group_names = sorted(groups_df["GROUPNAME"].dropna().astype(str).str.strip().unique().tolist()) if not groups_df.empty else []
    app_names  = sorted(apps_df["APPLICATIONNAME"].dropna().astype(str).str.strip().unique().tolist()) if not apps_df.empty else []
    vendor_names = sorted(vendors_df["VENDORNAME"].dropna().astype(str).str.strip().unique().tolist()) if not vendors_df.empty else []

    instr = [
        {"Field": "PROGRAMNAME", "Required": "Yes", "Description": "Program (pick from allowed list)", "Example": "R&M"},
        {"Field": "TEAMNAME", "Required": "Yes", "Description": "Team name", "Example": "ERNE WEST"},
        {"Field": "PRODUCT_OWNER", "Required": "No", "Description": "Default Product Owner for Team", "Example": "Alice"},
        {"Field": "GROUPNAME", "Required": "Yes", "Description": "Application. Multiple application instances can be aggregated here.", "Example": "Application - MHM"},
        {"Field": "APPNAME", "Required": "Yes", "Description": "Application instance name", "Example": "MHM Baytown"},
        {"Field": "VENDORNAME", "Required": "Yes", "Description": "Vendor name", "Example": "Emerson"},
    ]

    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill, Font

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data")
        pd.DataFrame(instr).to_excel(writer, index=False, sheet_name="Instructions")
        wb = writer.book
        ws_data = writer.sheets["Data"]
        ws_data.freeze_panes = "A2"
        header_to_idx = {h: i+1 for i, h in enumerate(cols)}

        ws_lists = wb.create_sheet("Lists")
        ws_lists["A1"] = "PROGRAMNAME"; ws_lists["B1"] = "TEAMNAME"; ws_lists["C1"] = "GROUPNAME"; ws_lists["D1"] = "APPLICATIONNAME"; ws_lists["E1"] = "VENDORNAME"
        for i, v in enumerate(ALLOWED_PROGRAMS, start=2): ws_lists.cell(row=i, column=1, value=v)
        for i, v in enumerate(team_names, start=2): ws_lists.cell(row=i, column=2, value=v)
        for i, v in enumerate(group_names, start=2): ws_lists.cell(row=i, column=3, value=v)
        for i, v in enumerate(app_names, start=2): ws_lists.cell(row=i, column=4, value=v)
        for i, v in enumerate(vendor_names, start=2): ws_lists.cell(row=i, column=5, value=v)

        def add_list_validation(header: str, lists_col: int, last_row: int):
            if last_row < 2:
                return
            try:
                idx = cols.index(header) + 1
            except ValueError:
                return
            col_letter = get_column_letter(idx)
            first = 2; dest_last = 2000
            src_range = f"Lists!${get_column_letter(lists_col)}$2:${get_column_letter(lists_col)}${last_row}"
            dv = DataValidation(type="list", formula1=f"={src_range}", allow_blank=True)
            ws_data.add_data_validation(dv)
            dv.add(f"{col_letter}{first}:{col_letter}{dest_last}")

        add_list_validation("PROGRAMNAME", 1, 1 + len(ALLOWED_PROGRAMS))
        add_list_validation("TEAMNAME", 2, 1 + len(team_names))
        add_list_validation("GROUPNAME", 3, 1 + len(group_names))
        add_list_validation("APPNAME", 4, 1 + len(app_names))
        add_list_validation("VENDORNAME", 5, 1 + len(vendor_names))

        # Header tooltips/colors
        field_help = {row["Field"]: f"{row['Description']}" + (f" (Example: {row['Example']})" if row.get('Example') else "") for row in instr}
        for h in cols:
            i = header_to_idx.get(h)
            if i and h in field_help:
                cell = ws_data.cell(row=1, column=i)
                cell.comment = Comment(field_help[h], "NEXT")
        req_header_fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
        for h in ["PROGRAMNAME","TEAMNAME","GROUPNAME","APPNAME","VENDORNAME"]:
            i = header_to_idx.get(h)
            if i:
                c = ws_data.cell(row=1, column=i)
                c.fill = req_header_fill
                c.font = Font(bold=True)

        # Prefill row 2 example
        def set_cell(name: str, value):
            i = header_to_idx.get(name)
            if i:
                ws_data.cell(row=2, column=i, value=value)
        set_cell("PROGRAMNAME", "R&M")
        set_cell("TEAMNAME", "ERNE WEST")
        set_cell("GROUPNAME", "Application - MHM")
        set_cell("APPNAME", "MHM Baytown")
        set_cell("VENDORNAME", "Emerson")
        set_cell("PRODUCT_OWNER", "Alice")

        # Auto-fit simple: based on header/examples
        try:
            from openpyxl.utils import get_column_letter as _gcl
            for h in cols:
                i = header_to_idx.get(h)
                if not i:
                    continue
                values = [str(h), str(ws_data.cell(row=2, column=i).value or "")]
                width = max(len(v) for v in values) + 4
                ws_data.column_dimensions[_gcl(i)].width = max(12, min(width, 60))
        except Exception:
            pass

    buf.seek(0)
    return buf.getvalue()

# -------------------------
# ADO parsing (core)
# -------------------------
def _auto_rename_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip().replace("\ufeff", "") for c in df.columns]
    rename: Dict[str, str] = {}
    for canonical, candidates in EXPECTED.items():
        for cand in candidates:
            match = [col for col in df.columns if isinstance(col, str) and col.lower() == cand.lower()]
            if match:
                rename[match[0]] = canonical
                break
        else:
            match2 = [col for col in df.columns if isinstance(col, str) and canonical.lower() in col.lower()]
            if match2:
                rename[match2[0]] = canonical
    return df.rename(columns=rename)

def read_ado_upload_any(upl, sheet_name: Optional[str], diag: Dict[str, Any]) -> pd.DataFrame:
    if upl is None:
        raise ValueError("No file uploaded.")
    df = _read_file_any(upl, sheet_name, diag)
    df = _auto_rename_columns(df)

    for col in (
        "Team",
        "Custom_ApplicationName",
        "Custom_InvestmentDimension",
        "Iteration",
        "Title",
        "State",
        "AreaPath",
        "AreaLevel1",
        "AreaLevel2",
        "AreaLevel3",
        "AreaLevel4",
    ):
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip()

    for num_col in ("BusinessValue",):
        if num_col in df.columns:
            df[num_col] = pd.to_numeric(df[num_col], errors="coerce")

    if "ID" in df.columns:
        df["ID"] = df["ID"].astype(str).str.strip()

    for dcol in ("CreatedDate", "ChangedDate"):
        if dcol in df.columns:
            df[dcol] = pd.to_datetime(df[dcol], errors="coerce")

    if "Year" in df.columns:
        df["Year"] = df["Year"].astype(str).str.replace(",", ".", regex=False)
        df["Year"] = pd.to_numeric(df["Year"], errors="coerce").astype("Int64")

    return df

# --- helper: pick first existing column by name(s) ---
## Normalization implemented in utils.ado.normalize_to_canonical

def _to_py(v: Any):
    import pandas as _pd, numpy as _np
    if v is None:
        return None
    if isinstance(v, _pd._libs.tslibs.nattype.NaTType):
        return None
    if isinstance(v, _pd.Timestamp):
        return v.to_pydatetime()
    if _pd.isna(v):
        return None
    if isinstance(v, _np.floating):
        return None if _np.isnan(v) else float(v)
    if isinstance(v, _np.integer):
        return int(v)
    if isinstance(v, str):
        s = v.strip()
        return s if s != "" else None
    return v

# Note: upsert_ado_features now comes from db facade (centralized merge)

def load_ado_distincts() -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    teams = fetch_df("""
        WITH base AS (
          SELECT COALESCE(t.TEAMNAME, af.TEAM_RAW) AS TEAMNAME
          FROM ADO_FEATURES af
          LEFT JOIN MAP_ADO_TEAM_TO_TCO_TEAM m
            ON UPPER(m.ADO_TEAM_KEY) = COALESCE(
                 af.TEAM_VARIANT_KEY,
                 UPPER(NULLIF(LTRIM(RTRIM(af.TEAM_RAW)), ''))
               )
          LEFT JOIN TEAMS t ON t.TEAMID = m.TEAMID
          WHERE af.TEAM_RAW IS NOT NULL AND LTRIM(RTRIM(af.TEAM_RAW)) <> ''
        )
        SELECT DISTINCT TEAMNAME
        FROM base
        WHERE TEAMNAME IS NOT NULL AND LTRIM(RTRIM(TEAMNAME)) <> ''
        ORDER BY TEAMNAME
    """)
    apps  = fetch_df("""
        SELECT DISTINCT APP_NAME_RAW
        FROM ADO_FEATURES
        WHERE APP_NAME_RAW IS NOT NULL AND TRIM(APP_NAME_RAW) <> ''
        ORDER BY APP_NAME_RAW
    """)
    iters = fetch_df("""
        SELECT DISTINCT ITERATION_PATH
        FROM ADO_FEATURES
        WHERE ITERATION_PATH IS NOT NULL AND TRIM(ITERATION_PATH) <> ''
        ORDER BY ITERATION_PATH
    """)
    return teams, apps, iters

def ado_features_base_query(
    where_sql: str = "",
    params: Optional[tuple] = None,
    *,
    post_where_sql: str = "",
    post_params: Optional[tuple] = None,
) -> pd.DataFrame:
    ado_cols = _ado_feature_columns()
    business_value_expr = "f.BUSINESS_VALUE" if "BUSINESS_VALUE" in ado_cols else "CAST(NULL AS FLOAT)"
    ado_year_expr = "TRY_CONVERT(INT, f.ADO_YEAR)" if "ADO_YEAR" in ado_cols else "CAST(NULL AS INT)"
    business_value_sql = f"{business_value_expr} AS BUSINESS_VALUE"
    iteration_level3_sql = (
        "f.ITERATION_LEVEL3 AS ITERATION_LEVEL3"
        if "ITERATION_LEVEL3" in ado_cols
        else "CAST(NULL AS NVARCHAR(400)) AS ITERATION_LEVEL3"
    )
    iteration_sk_sel_sql = (
        "NULLIF(LTRIM(RTRIM(f.ITERATION_SK)), '') AS ITERATION_SK"
        if "ITERATION_SK" in ado_cols
        else "CAST(NULL AS NVARCHAR(100)) AS ITERATION_SK"
    )
    iteration_sk_join_sql = "NULLIF(LTRIM(RTRIM(f.ITERATION_SK)), '')" if "ITERATION_SK" in ado_cols else "CAST(NULL AS NVARCHAR(100))"
    # Avoid implicit nvarchar→int conversions (some environments have bad/mixed ADO_YEAR values like "2025 I1").
    cal_year_expr = """
        COALESCE(
          TRY_CONVERT(INT, LEFT(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3), 4)),
          TRY_CONVERT(INT, ic_sk.YEAR),
          TRY_CONVERT(INT, ic_path.YEAR),
          TRY_CONVERT(INT, f.ADO_YEAR)
        )
    """.strip()
    sql_with_demand = f"""
      WITH filtered AS (
        SELECT *
        FROM ADO_FEATURES
        {where_sql}
      )
      SELECT
        f.FEATURE_ID,
        f.TITLE,
        f.STATE,
        f.TEAM_RAW,
        COALESCE(t.TEAMNAME, f.TEAM_RAW) AS TEAMNAME,
        COALESCE(p.PROGRAMNAME, CAST(NULL AS NVARCHAR(255))) AS PROGRAMNAME,
        f.APP_NAME_RAW,
        COALESCE(ag.GROUPNAME, CAST(NULL AS NVARCHAR(255))) AS GROUPNAME,
        f.INVESTMENT_DIMENSION,
        {ado_year_expr} AS ADO_YEAR,
        {business_value_sql},
        f.ITERATION_PATH,
        {iteration_level3_sql},
        {iteration_sk_sel_sql},
        /* CAL_* prefix = fields from ADO_ITERATION_CALENDAR (debug visibility for iteration mapping) */
        COALESCE(ic_sk.ITERATION_SK, ic_path.ITERATION_SK) AS CAL_ITERATION_SK,
        COALESCE(ic_sk.ITERATION_PATH, ic_path.ITERATION_PATH) AS CAL_ITERATION_PATH,
        COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3) AS CAL_ITERATION_LEVEL3,
        COALESCE(ic_sk.ITERATION_GRAIN, ic_path.ITERATION_GRAIN) AS CAL_ITERATION_GRAIN,
        COALESCE(ic_sk.START_DATE, ic_path.START_DATE) AS CAL_START_DATE,
        COALESCE(ic_sk.END_DATE, ic_path.END_DATE) AS CAL_END_DATE,
        {cal_year_expr} AS CAL_YEAR,
        CASE
          WHEN COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3) IS NULL
               OR CHARINDEX('I', UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3))) = 0
          THEN NULL
          ELSE TRY_CONVERT(
                 INT,
                 SUBSTRING(
                   SUBSTRING(UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3)), CHARINDEX('I', UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3))), 10),
                   NULLIF(PATINDEX('%[1-4]%', SUBSTRING(UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3)), CHARINDEX('I', UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3))), 10)), 0),
                   1
                 )
               )
        END AS CAL_PI_NUM,
        fp.FEATURE_IS_SPRINT_LEVEL,
        CASE
          WHEN fp.FEATURE_IS_SPRINT_LEVEL = 1
               AND COALESCE(ic_sk.ITERATION_PATH, ic_path.ITERATION_PATH) IS NOT NULL
          THEN 1
          ELSE 0
        END AS SPRINT_TO_PI_RESOLVED,
        CASE
          WHEN COALESCE(ic_sk.ITERATION_PATH, ic_path.ITERATION_PATH) IS NULL THEN 1 ELSE 0
        END AS MISSING_PI_CALENDAR,
        TRY_CONVERT(FLOAT, dem.DERIVED_FTE_FEATURE) AS DERIVED_FTE,
        CAST(NULL AS FLOAT) AS FTE_DIFF,
        f.CHANGED_AT
      FROM filtered f
      OUTER APPLY (
        /* Keep this mapping logic aligned with the canonical feature→PI mapping (sprint-level Features map to PI parent path). */
        SELECT
          CASE
            WHEN f.ITERATION_PATH IS NULL OR LTRIM(RTRIM(f.ITERATION_PATH)) = '' THEN 0
            WHEN PATINDEX('%\\I[0-9] S[0-9]%', UPPER(f.ITERATION_PATH)) > 0 THEN 1
            WHEN PATINDEX('%\\S[0-9]%', UPPER(f.ITERATION_PATH)) > 0 THEN 1
            ELSE 0
          END AS FEATURE_IS_SPRINT_LEVEL,
          CASE
            WHEN f.ITERATION_PATH IS NULL OR LTRIM(RTRIM(f.ITERATION_PATH)) = '' THEN NULL
            WHEN (
              PATINDEX('%\\I[0-9] S[0-9]%', UPPER(f.ITERATION_PATH)) > 0
              OR PATINDEX('%\\S[0-9]%', UPPER(f.ITERATION_PATH)) > 0
            )
            AND CHARINDEX('\\', f.ITERATION_PATH) > 0
            THEN LEFT(f.ITERATION_PATH, LEN(f.ITERATION_PATH) - CHARINDEX('\\', REVERSE(f.ITERATION_PATH)))
            ELSE f.ITERATION_PATH
          END AS FEATURE_PI_PARENT_ITERATION_PATH
      ) fp
      LEFT JOIN MAP_ADO_TEAM_TO_TCO_TEAM m
        ON UPPER(m.ADO_TEAM_KEY) = COALESCE(
             f.TEAM_VARIANT_KEY,
             UPPER(NULLIF(LTRIM(RTRIM(f.TEAM_RAW)), ''))
           )
      LEFT JOIN TEAMS t ON t.TEAMID = m.TEAMID
      LEFT JOIN PROGRAMS p ON p.PROGRAMID = t.PROGRAMID
      LEFT JOIN MAP_ADO_APP_TO_TCO_GROUP mag
        ON UPPER(mag.ADO_APP) = UPPER(LTRIM(RTRIM(f.APP_NAME_RAW)))
      LEFT JOIN APPLICATION_GROUPS ag ON ag.GROUPID = mag.APP_GROUP
      LEFT JOIN ADO_ITERATION_CALENDAR ic_sk
        ON UPPER(LTRIM(RTRIM(ic_sk.ITERATION_SK))) = UPPER(LTRIM(RTRIM({iteration_sk_join_sql})))
       AND UPPER(LTRIM(RTRIM(ic_sk.ITERATION_GRAIN))) = 'PI'
      LEFT JOIN ADO_ITERATION_CALENDAR ic_path
        ON UPPER(LTRIM(RTRIM(ic_path.ITERATION_PATH))) = UPPER(LTRIM(RTRIM(fp.FEATURE_PI_PARENT_ITERATION_PATH)))
       AND UPPER(LTRIM(RTRIM(ic_path.ITERATION_GRAIN))) = 'PI'
      LEFT JOIN VW_TCO_FEATURE_DEMAND dem
        ON dem.FEATURE_ID = f.FEATURE_ID
       AND TRY_CONVERT(INT, dem.YEAR) = {cal_year_expr}
       AND UPPER(LTRIM(RTRIM(COALESCE(dem.PI_LABEL, '')))) = UPPER(LTRIM(RTRIM(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3, ''))))
      {post_where_sql}
      ORDER BY COALESCE(f.CHANGED_AT, TRY_CONVERT(DATETIME2,'1900-01-01')) DESC, f.FEATURE_ID
    """
    sql_fallback = f"""
      WITH filtered AS (
        SELECT *
        FROM ADO_FEATURES
        {where_sql}
      )
      SELECT
        f.FEATURE_ID,
        f.TITLE,
        f.STATE,
        f.TEAM_RAW,
        COALESCE(t.TEAMNAME, f.TEAM_RAW) AS TEAMNAME,
        COALESCE(p.PROGRAMNAME, CAST(NULL AS NVARCHAR(255))) AS PROGRAMNAME,
        f.APP_NAME_RAW,
        COALESCE(ag.GROUPNAME, CAST(NULL AS NVARCHAR(255))) AS GROUPNAME,
        f.INVESTMENT_DIMENSION,
        {ado_year_expr} AS ADO_YEAR,
        {business_value_sql},
        f.ITERATION_PATH,
        {iteration_level3_sql},
        {iteration_sk_sel_sql},
        /* CAL_* prefix = fields from ADO_ITERATION_CALENDAR (debug visibility for iteration mapping) */
        COALESCE(ic_sk.ITERATION_SK, ic_path.ITERATION_SK) AS CAL_ITERATION_SK,
        COALESCE(ic_sk.ITERATION_PATH, ic_path.ITERATION_PATH) AS CAL_ITERATION_PATH,
        COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3) AS CAL_ITERATION_LEVEL3,
        COALESCE(ic_sk.ITERATION_GRAIN, ic_path.ITERATION_GRAIN) AS CAL_ITERATION_GRAIN,
        COALESCE(ic_sk.START_DATE, ic_path.START_DATE) AS CAL_START_DATE,
        COALESCE(ic_sk.END_DATE, ic_path.END_DATE) AS CAL_END_DATE,
        {cal_year_expr} AS CAL_YEAR,
        CASE
          WHEN COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3) IS NULL
               OR CHARINDEX('I', UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3))) = 0
          THEN NULL
          ELSE TRY_CONVERT(
                 INT,
                 SUBSTRING(
                   SUBSTRING(UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3)), CHARINDEX('I', UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3))), 10),
                   NULLIF(PATINDEX('%[1-4]%', SUBSTRING(UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3)), CHARINDEX('I', UPPER(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3))), 10)), 0),
                   1
                 )
               )
        END AS CAL_PI_NUM,
        fp.FEATURE_IS_SPRINT_LEVEL,
        CASE
          WHEN fp.FEATURE_IS_SPRINT_LEVEL = 1
               AND COALESCE(ic_sk.ITERATION_PATH, ic_path.ITERATION_PATH) IS NOT NULL
          THEN 1
          ELSE 0
        END AS SPRINT_TO_PI_RESOLVED,
        CASE
          WHEN COALESCE(ic_sk.ITERATION_PATH, ic_path.ITERATION_PATH) IS NULL THEN 1 ELSE 0
        END AS MISSING_PI_CALENDAR,
        TRY_CONVERT(FLOAT, dem.DERIVED_FTE_FEATURE) AS DERIVED_FTE,
        CAST(NULL AS FLOAT) AS FTE_DIFF,
        f.CHANGED_AT
      FROM filtered f
      OUTER APPLY (
        SELECT
          CASE
            WHEN f.ITERATION_PATH IS NULL OR LTRIM(RTRIM(f.ITERATION_PATH)) = '' THEN 0
            WHEN PATINDEX('%\\I[0-9] S[0-9]%', UPPER(f.ITERATION_PATH)) > 0 THEN 1
            WHEN PATINDEX('%\\S[0-9]%', UPPER(f.ITERATION_PATH)) > 0 THEN 1
            ELSE 0
          END AS FEATURE_IS_SPRINT_LEVEL,
          CASE
            WHEN f.ITERATION_PATH IS NULL OR LTRIM(RTRIM(f.ITERATION_PATH)) = '' THEN NULL
            WHEN (
              PATINDEX('%\\I[0-9] S[0-9]%', UPPER(f.ITERATION_PATH)) > 0
              OR PATINDEX('%\\S[0-9]%', UPPER(f.ITERATION_PATH)) > 0
            )
            AND CHARINDEX('\\', f.ITERATION_PATH) > 0
            THEN LEFT(f.ITERATION_PATH, LEN(f.ITERATION_PATH) - CHARINDEX('\\', REVERSE(f.ITERATION_PATH)))
            ELSE f.ITERATION_PATH
          END AS FEATURE_PI_PARENT_ITERATION_PATH
      ) fp
      LEFT JOIN MAP_ADO_TEAM_TO_TCO_TEAM m
        ON UPPER(m.ADO_TEAM_KEY) = COALESCE(
             f.TEAM_VARIANT_KEY,
             UPPER(NULLIF(LTRIM(RTRIM(f.TEAM_RAW)), ''))
           )
      LEFT JOIN TEAMS t ON t.TEAMID = m.TEAMID
      LEFT JOIN PROGRAMS p ON p.PROGRAMID = t.PROGRAMID
      LEFT JOIN MAP_ADO_APP_TO_TCO_GROUP mag
        ON UPPER(mag.ADO_APP) = UPPER(LTRIM(RTRIM(f.APP_NAME_RAW)))
      LEFT JOIN APPLICATION_GROUPS ag ON ag.GROUPID = mag.APP_GROUP
      LEFT JOIN ADO_ITERATION_CALENDAR ic_sk
        ON UPPER(LTRIM(RTRIM(ic_sk.ITERATION_SK))) = UPPER(LTRIM(RTRIM({iteration_sk_join_sql})))
       AND UPPER(LTRIM(RTRIM(ic_sk.ITERATION_GRAIN))) = 'PI'
      LEFT JOIN ADO_ITERATION_CALENDAR ic_path
        ON UPPER(LTRIM(RTRIM(ic_path.ITERATION_PATH))) = UPPER(LTRIM(RTRIM(fp.FEATURE_PI_PARENT_ITERATION_PATH)))
       AND UPPER(LTRIM(RTRIM(ic_path.ITERATION_GRAIN))) = 'PI'
      LEFT JOIN VW_TCO_FEATURE_DEMAND dem
        ON dem.FEATURE_ID = f.FEATURE_ID
       AND TRY_CONVERT(INT, dem.YEAR) = {cal_year_expr}
       AND UPPER(LTRIM(RTRIM(COALESCE(dem.PI_LABEL, '')))) = UPPER(LTRIM(RTRIM(COALESCE(ic_sk.ITERATION_LEVEL3, ic_path.ITERATION_LEVEL3, ''))))
      {post_where_sql}
      ORDER BY COALESCE(f.CHANGED_AT, TRY_CONVERT(DATETIME2,'1900-01-01')) DESC, f.FEATURE_ID
    """

    combined_params: Optional[tuple]
    if params and post_params:
        combined_params = tuple(params) + tuple(post_params)
    elif params:
        combined_params = params
    elif post_params:
        combined_params = post_params
    else:
        combined_params = None

    try:
        df = fetch_df(sql_with_demand, combined_params)
    except Exception:
        df = fetch_df(sql_fallback, combined_params)
    if df is not None and not df.empty:
        for num_col in ("BUSINESS_VALUE", "CAL_PI_NUM", "CAL_YEAR", "ADO_YEAR"):
            if num_col in df.columns:
                df[num_col] = pd.to_numeric(df[num_col], errors="coerce")
    return df if df is not None else pd.DataFrame()

def render_ado_summary_header() -> None:
    with st.container(border=True):
        st.markdown("### ADO Configuration Summary")
        org = str(st.session_state.get("ado_org") or "").strip()
        proj = str(st.session_state.get("ado_proj") or "").strip()
        ident = f"{org}/{proj}".strip("/")
        query_mode = st.session_state.get("ado_query_mode", "Advanced OData Builder")
        filter_val = str(st.session_state.get("ado_filter") or "").strip()
        features_only = st.session_state.get("ado_features_only")
        c1, c2 = st.columns([2, 3])
        with c1:
            st.write(f"Org/Project: {ident or '-'}")
            st.write(f"Query mode: {query_mode}")
        with c2:
            if filter_val:
                st.write(f"Filter: {filter_val}")
            elif features_only:
                st.write("Filter: WorkItemType eq 'Feature' (default)")
            else:
                st.write("Filter: —")
        st.caption("Next steps: Validate → Sync → Mapping")

def render_ado_profiles_tab() -> None:
    st.subheader("ADO Profiles")
    _ = st.session_state.get("profiles_cache_bust", 0)
    try:
        ensure_ado_profiles_table()
    except Exception as e:
        st.error(f"Could not ensure ADO_PROFILES table: {e}")
        return
    try:
        profiles_full = fetch_ado_profiles_full()
    except Exception as e:
        st.error(f"Could not load ADO profiles: {e}")
        return

    if profiles_full is None or profiles_full.empty:
        st.info("No ADO profiles found yet. Run schema checks to seed the default profile.")
        return

    profiles_full = profiles_full.copy()
    profiles_full["PROFILE_NAME"] = profiles_full["PROFILE_NAME"].astype(str)
    profiles_full["PROFILE_ID"] = profiles_full["PROFILE_ID"].astype(str)
    active_mask = profiles_full["IS_ACTIVE"].astype(int) == 1
    active_id = profiles_full.loc[active_mask, "PROFILE_ID"].iloc[0] if active_mask.any() else ""
    active_name = profiles_full.loc[active_mask, "PROFILE_NAME"].iloc[0] if active_mask.any() else ""

    id_to_name = {r.PROFILE_ID: r.PROFILE_NAME for _, r in profiles_full.iterrows()}
    profile_ids = profiles_full["PROFILE_ID"].tolist()
    selected_id = st.selectbox(
        "Profile",
        profile_ids,
        format_func=lambda pid: f"{id_to_name.get(pid, pid)}{' (active)' if pid == active_id else ''}",
        key="ado_profile_selected_id",
    )
    selected_row = profiles_full.loc[profiles_full["PROFILE_ID"] == selected_id].iloc[0]
    selected_name = str(selected_row["PROFILE_NAME"])
    is_active = bool(int(selected_row.get("IS_ACTIVE", 0)))
    def _coerce_dt(val: Any) -> Optional[datetime]:
        if val is None or (isinstance(val, float) and math.isnan(val)):
            return None
        try:
            if isinstance(val, datetime):
                return val
            return pd.to_datetime(val, errors="coerce").to_pydatetime()  # type: ignore[return-value]
        except Exception:
            return None
    profile_updated_at = _coerce_dt(selected_row.get("UPDATED_AT"))
    profile_last_sync_at = _coerce_dt(selected_row.get("LAST_SYNC_AT"))
    profile_out_of_sync = False
    if profile_last_sync_at is None:
        profile_out_of_sync = True
    elif profile_updated_at and profile_updated_at > profile_last_sync_at:
        profile_out_of_sync = True

    with st.container(border=True):
        st.markdown("### Profile Actions")
        pat_env_key = get_active_pat_env_key()
        pat_val, pat_source = resolve_ado_pat_source(selected_name, pat_env_key=pat_env_key)
        st.write(f"PAT configured for this profile: {'Yes' if pat_val else 'No'}")
        st.caption(f"Configured via env/secrets: {'Yes' if pat_source in ('env', 'secrets') else 'No'}")
        st.caption(f"PAT source: {pat_source if pat_source != 'missing' else 'missing'}")
        col1, col2, col3, col4 = st.columns(4)
        with col1:
            if st.button("Set Active", disabled=is_active, key="ado_profile_set_active_btn"):
                try:
                    set_active_ado_profile(selected_id)
                    st.success("Active ADO profile updated.")
                    try:
                        _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                    except Exception:
                        pass
                    st.session_state["profiles_cache_bust"] = st.session_state.get("profiles_cache_bust", 0) + 1
                    st.rerun()
                except Exception as e:
                    st.error(f"Failed to set active profile: {e}")
        with col2:
            if st.button("Duplicate", key="ado_profile_duplicate_btn"):
                try:
                    new_name = f"{selected_name} (copy)"
                    config_json = selected_row.get("CONFIG_JSON") or "{}"
                    updated_by = st.session_state.get("auth_user", {}).get("email") or "ADMIN"
                    upsert_ado_profile(
                        profile_id=None,
                        profile_name=new_name,
                        config_json=config_json,
                        is_active=False,
                        updated_by=updated_by,
                    )
                    st.success("Profile duplicated.")
                    try:
                        _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                    except Exception:
                        pass
                    st.session_state["profiles_cache_bust"] = st.session_state.get("profiles_cache_bust", 0) + 1
                    st.rerun()
                except Exception as e:
                    st.error(f"Duplicate failed: {e}")
        with col3:
            config_json = selected_row.get("CONFIG_JSON") or "{}"
            try:
                parsed = json.loads(config_json) if isinstance(config_json, str) else {}
            except Exception:
                parsed = {}
            export_payload = json.dumps(parsed, indent=2, sort_keys=True)
            export_name = selected_name.strip().replace(" ", "_") or "ado_profile"
            st.download_button(
                "Export JSON",
                data=export_payload,
                file_name=f"{export_name}.json",
                mime="application/json",
                key="ado_profile_export_btn",
            )
        with col4:
            import_clicked = st.button("Import JSON", key="ado_profile_import_btn")

        upl = st.file_uploader("Import JSON file", type=["json"], key="ado_profile_import_file")
        import_name = st.text_input(
            "Profile name",
            value=(upl.name.rsplit(".", 1)[0] if upl else ""),
            key="ado_profile_import_name",
        ).strip()
        if import_clicked:
            if upl is None:
                st.error("Choose a JSON file to import.")
            elif not import_name:
                st.error("Enter a profile name.")
            else:
                try:
                    payload_text = upl.getvalue().decode("utf-8")
                    json.loads(payload_text)
                    updated_by = st.session_state.get("auth_user", {}).get("email") or "ADMIN"
                    upsert_ado_profile(
                        profile_id=None,
                        profile_name=import_name,
                        config_json=payload_text,
                        is_active=False,
                        updated_by=updated_by,
                    )
                    st.success("Profile imported.")
                    try:
                        _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                    except Exception:
                        pass
                    st.session_state["profiles_cache_bust"] = st.session_state.get("profiles_cache_bust", 0) + 1
                    st.rerun()
                except Exception as e:
                    st.error(f"Import failed: {e}")

    st.markdown("### Profile Editor")
    overrides = st.session_state.setdefault("ado_profile_editor_overrides", {})
    raw_config = selected_row.get("CONFIG_JSON") or "{}"
    if selected_id in overrides:
        config = overrides[selected_id]
    else:
        try:
            config = json.loads(raw_config) if isinstance(raw_config, str) else {}
        except Exception:
            config = {}
            st.error("Profile JSON is invalid; use Advanced JSON editor to fix.")
    config = normalize_profile_config(config if isinstance(config, dict) else {})

    defaults = {
        "mode": "legacy",
        "work_item_types": {"feature": "Feature", "epic": "Epic"},
        "area_mapping": {"program_level": "AreaLevel2", "team_level": "Leaf"},
        "iteration_mapping": {"pi_field": "IterationLevel3", "path_field": "IterationPath"},
        "fields": {
            "app_name": "Custom_ApplicationName",
            "investment_dimension": "Custom_InvestmentDimension",
            "story_points": "Effort",
            "business_value": "BusinessValue",
        },
        "swag": {"points_per_fte": 65.0},
        "forecast": {"derived_fte_driver": "SWAG", "velocity_row_fallback": True},
        "filters": {"exclude_states": ["Removed"], "years_prefix": ["2025", "2026", "2027"]},
        "epic_resolution": {"max_depth": 6},
    }

    def _get_cfg(path: str, default_val: Any) -> Any:
        cur = config if isinstance(config, dict) else {}
        for key in path.split("."):
            if not isinstance(cur, dict) or key not in cur:
                return default_val
            cur = cur.get(key)
        return cur if cur is not None else default_val

    with st.form(key=f"ado_profile_form_{selected_id}"):
        profile_name_input = st.text_input(
            "Profile name",
            value=selected_name,
            key=f"ado_profile_name_{selected_id}",
        ).strip()
        save_as_new = st.checkbox(
            "Save as new profile",
            value=False,
            key=f"ado_profile_save_as_new_{selected_id}",
        )
        set_active = st.checkbox(
            "Set Active",
            value=is_active,
            key=f"ado_profile_set_active_{selected_id}",
        )
        st.markdown("#### Connection")
        odata_base_url = st.text_input(
            "OData base URL",
            value=str(_get_cfg("odata_base_url", defaults.get("odata_base_url", ""))),
            key=f"ado_profile_odata_{selected_id}",
        ).strip()

        st.markdown("#### Work item types")
        wi_feature = st.text_input(
            "Feature type",
            value=str(_get_cfg("work_item_types.feature", defaults["work_item_types"]["feature"])),
            key=f"ado_profile_wi_feature_{selected_id}",
        ).strip()
        wi_epic = st.text_input(
            "Epic type",
            value=str(_get_cfg("work_item_types.epic", defaults["work_item_types"]["epic"])),
            key=f"ado_profile_wi_epic_{selected_id}",
        ).strip()

        st.markdown("#### Area mapping")
        area_opts = ["AreaLevel1", "AreaLevel2", "AreaLevel3", "AreaLevel4", "Leaf"]
        program_level = st.selectbox(
            "Program level",
            options=area_opts,
            index=area_opts.index(
                str(_get_cfg("area_mapping.program_level", defaults["area_mapping"]["program_level"]))
                if str(_get_cfg("area_mapping.program_level", defaults["area_mapping"]["program_level"])) in area_opts
                else "AreaLevel2"
            ),
            key=f"ado_profile_program_level_{selected_id}",
        )
        team_level = st.selectbox(
            "Team level",
            options=area_opts,
            index=area_opts.index(
                str(_get_cfg("area_mapping.team_level", defaults["area_mapping"]["team_level"]))
                if str(_get_cfg("area_mapping.team_level", defaults["area_mapping"]["team_level"])) in area_opts
                else "Leaf"
            ),
            key=f"ado_profile_team_level_{selected_id}",
        )

        st.markdown("#### Iteration mapping")
        iter_opts = ["IterationLevel1", "IterationLevel2", "IterationLevel3", "IterationPath"]
        pi_field = st.selectbox(
            "PI field",
            options=iter_opts,
            index=iter_opts.index(
                str(_get_cfg("iteration_mapping.pi_field", defaults["iteration_mapping"]["pi_field"]))
                if str(_get_cfg("iteration_mapping.pi_field", defaults["iteration_mapping"]["pi_field"])) in iter_opts
                else "IterationLevel3"
            ),
            key=f"ado_profile_pi_field_{selected_id}",
        )
        path_field = st.selectbox(
            "Path field",
            options=["IterationPath", "IterationLevel3"],
            index=0
            if str(_get_cfg("iteration_mapping.path_field", defaults["iteration_mapping"]["path_field"])) == "IterationPath"
            else 1,
            key=f"ado_profile_path_field_{selected_id}",
        )

        st.markdown("#### Story Points & Effort Model")
        discovered = st.session_state.get("ado_discovered_fields", {})
        discovered_top = discovered.get("metadata_fields") or discovered.get("top") or []
        points_options = ["(None - disable demand/FTE derivation)"]
        if discovered_top:
            points_options += sorted({p for p in discovered_top if str(p).strip() != "Custom_FTE"})
        else:
            points_options += ["Effort", "StoryPoints", "BusinessValue"]
        points_value = str(_get_cfg("fields.story_points", defaults["fields"]["story_points"]))
        if points_value and points_value not in points_options:
            points_options.append(points_value)
        points_index = points_options.index(points_value) if points_value in points_options else 0
        story_points_field = st.selectbox(
            "Story Points field (ADO)",
            options=points_options,
            index=points_index,
            key=f"ado_profile_field_story_points_{selected_id}",
            help="FTE is not ingested from ADO. Demand FTE is always derived from Story Points using SWAG.",
        )
        coverage = _story_points_coverage_summary(int(st.session_state.get("ado_data_cache_bust", 0)))
        _render_story_points_coverage(coverage)
        st.caption("Coverage is computed from last synced data. Changing the field requires a new Sync.")
        if not discovered_top:
            st.caption("Run Discover Fields in Validate tab to load available field names.")
        if profile_out_of_sync:
            st.warning("Profile changes require a new ADO Sync to recompute demand and expected costs.")
        else:
            st.caption("Profile is in sync with latest ADO data.")
        with st.expander("Preview sample values", expanded=False):
            st.caption("Story Points must be defined on Feature work items to drive demand. Derived FTE = Story Points / SWAG.")
            if story_points_field.startswith("(None") or not story_points_field.strip():
                st.info("Select a Story Points field to preview values.")
            else:
                sample = get_field_sample(story_points_field, config, selected_id, limit=20)
                rows = sample.get("rows") or []
                source = sample.get("source") or "none"
                if not rows:
                    st.info("No sample data available from ADO or DB for this field.")
                else:
                    df_s = pd.DataFrame(rows)
                    val_col = "FIELD_VALUE" if "FIELD_VALUE" in df_s.columns else story_points_field
                    total = len(df_s.index)
                    non_null = int(pd.to_numeric(df_s.get(val_col), errors="coerce").notna().sum()) if val_col in df_s.columns else 0
                    coverage = (non_null / total * 100.0) if total else 0.0
                    numeric_vals = pd.to_numeric(df_s.get(val_col), errors="coerce") if val_col in df_s.columns else pd.Series(dtype=float)
                    stats = ""
                    if numeric_vals.notna().any():
                        stats = f" | Min: {numeric_vals.min():.2f} | Avg: {numeric_vals.mean():.2f} | Max: {numeric_vals.max():.2f}"
                    st.write(f"Coverage: {coverage:.1f}% ({non_null} / {total} rows){stats}")
                    st.caption(f"Source: {source}")
                    if coverage < 30.0:
                        st.warning("This field is sparsely populated for Features.")
                    preview_cols = ["FEATURE_ID", "TITLE", val_col]
                    preview_cols = [c for c in preview_cols if c in df_s.columns]
                    st.dataframe(df_s[preview_cols].head(10), use_container_width=True, height=260)

        st.markdown("#### Fields")
        app_name_field = st.text_input(
            "Application field",
            value=str(_get_cfg("fields.app_name", defaults["fields"]["app_name"])),
            key=f"ado_profile_field_app_{selected_id}",
        ).strip()
        bv_options = ["(None)"]
        if discovered_top:
            bv_options += sorted(set(discovered_top + ["BusinessValue"]))
        else:
            bv_options += ["BusinessValue", "Custom_BusinessValue"]
        bv_value = str(_get_cfg("fields.business_value", defaults["fields"]["business_value"]))
        if bv_value and bv_value not in bv_options:
            bv_options.append(bv_value)
        bv_index = bv_options.index(bv_value) if bv_value in bv_options else 0
        business_value_field = st.selectbox(
            "Business Value field (optional)",
            options=bv_options,
            index=bv_index,
            key=f"ado_profile_field_bv_{selected_id}",
            help="Used for prioritization/insights only. Does not affect cost/capacity.",
        )
        if profile_out_of_sync:
            st.warning("Profile changes require a new ADO Sync to recompute demand and expected costs.")
        else:
            st.caption("Profile is in sync with latest ADO data.")
        with st.expander("Preview sample values", expanded=False):
            if business_value_field == "(None)" or not business_value_field.strip():
                st.info("Select a Business Value field to preview values.")
            else:
                sample = get_field_sample(business_value_field, config, selected_id, limit=20)
                rows = sample.get("rows") or []
                source = sample.get("source") or "none"
                if not rows:
                    st.info("No sample data available from ADO or DB for this field.")
                else:
                    df_s = pd.DataFrame(rows)
                    val_col = "FIELD_VALUE" if "FIELD_VALUE" in df_s.columns else business_value_field
                    total = len(df_s.index)
                    non_null = int(pd.to_numeric(df_s.get(val_col), errors="coerce").notna().sum()) if val_col in df_s.columns else 0
                    coverage = (non_null / total * 100.0) if total else 0.0
                    numeric_vals = pd.to_numeric(df_s.get(val_col), errors="coerce") if val_col in df_s.columns else pd.Series(dtype=float)
                    stats = ""
                    if numeric_vals.notna().any():
                        stats = f" | Min: {numeric_vals.min():.2f} | Avg: {numeric_vals.mean():.2f} | Max: {numeric_vals.max():.2f}"
                    st.write(f"Coverage: {coverage:.1f}% ({non_null} / {total} rows){stats}")
                    st.caption(f"Source: {source}")
                    if coverage < 30.0:
                        st.warning("This field is sparsely populated for Features.")
                    preview_cols = ["FEATURE_ID", "TITLE", val_col]
                    preview_cols = [c for c in preview_cols if c in df_s.columns]
                    st.dataframe(df_s[preview_cols].head(10), use_container_width=True, height=260)
        inv_field = st.text_input(
            "Investment dimension field",
            value=str(_get_cfg("fields.investment_dimension", defaults["fields"]["investment_dimension"])),
            key=f"ado_profile_field_inv_{selected_id}",
        ).strip()

        st.markdown("#### SWAG Conversion")
        points_per_fte = st.number_input(
            "Story Points per 1 FTE",
            value=float(_get_cfg("swag.points_per_fte", defaults["swag"]["points_per_fte"])),
            min_value=0.0,
            step=1.0,
            key=f"ado_profile_points_per_fte_{selected_id}",
        )
        st.caption("Derived FTE = Story Points ÷ Story Points per 1 FTE")

        st.markdown("#### Forecast demand driver")
        raw_driver = str(_get_cfg("forecast.derived_fte_driver", defaults["forecast"]["derived_fte_driver"])).strip().upper()
        if raw_driver in {"VELOCITY", "SNAPSHOT", "SNAPSHOT_VELOCITY"}:
            forecast_driver = "SNAPSHOT_VELOCITY"
        else:
            forecast_driver = "SWAG"
        driver_options = ["SWAG", "SNAPSHOT_VELOCITY"]
        driver_labels = {
            "SWAG": "SWAG (Derived FTE)",
            "SNAPSHOT_VELOCITY": "Velocity snapshot (Derived FTE)",
        }
        forecast_driver = st.selectbox(
            "Expected/Forecast demand driver",
            options=driver_options,
            index=driver_options.index(forecast_driver),
            format_func=lambda v: driver_labels.get(v, v),
            key=f"ado_profile_forecast_driver_{selected_id}",
            help=(
                "Controls how Expected/Forecast demand is derived in canonical costing. "
                "SWAG uses Derived FTE from Story Points. Velocity snapshot converts SWAG points "
                "using team baseline points from TEAM_VELOCITY_SNAPSHOT."
            ),
        )
        velocity_row_fallback = st.checkbox(
            "Velocity mode: fallback missing rows to global SWAG baseline",
            value=bool(_get_cfg("forecast.velocity_row_fallback", defaults["forecast"]["velocity_row_fallback"])),
            key=f"ado_profile_velocity_row_fallback_{selected_id}",
            help=(
                "When enabled, features without a team snapshot row still use the global baseline "
                "instead of zero demand."
            ),
            disabled=(forecast_driver != "SNAPSHOT_VELOCITY"),
        )
        if forecast_driver == "SNAPSHOT_VELOCITY":
            st.caption("Velocity mode selected. Ensure TEAM_VELOCITY_SNAPSHOT is refreshed after ADO sync.")

        st.markdown("#### Filters")
        exclude_states = st.multiselect(
            "Exclude states",
            options=["Removed", "Closed", "Done", "Resolved"],
            default=_get_cfg("filters.exclude_states", defaults["filters"]["exclude_states"]),
            key=f"ado_profile_exclude_states_{selected_id}",
        )
        years_prefix = st.text_input(
            "Years prefix (comma-separated)",
            value=",".join(_get_cfg("filters.years_prefix", defaults["filters"]["years_prefix"])),
            key=f"ado_profile_years_prefix_{selected_id}",
        ).strip()

        st.markdown("#### Epic resolution")
        max_depth = st.number_input(
            "Max depth",
            value=int(_get_cfg("epic_resolution.max_depth", defaults["epic_resolution"]["max_depth"])),
            min_value=1,
            max_value=20,
            step=1,
            key=f"ado_profile_epic_depth_{selected_id}",
        )

        save_clicked = st.form_submit_button("Save Profile")

    if save_clicked:
        if not odata_base_url:
            st.error("OData base URL is required.")
        elif not profile_name_input:
            st.error("Profile name is required.")
        elif not wi_feature:
            st.error("Feature type is required.")
        else:
            years_list = [y.strip() for y in years_prefix.split(",") if y.strip()]
            payload = {
                "odata_base_url": odata_base_url,
                "work_item_types": {"feature": wi_feature, "epic": wi_epic or "Epic"},
                "area_mapping": {"program_level": program_level, "team_level": team_level},
                "iteration_mapping": {"pi_field": pi_field, "path_field": path_field},
                "fields": {
                    "app_name": app_name_field,
                    "investment_dimension": inv_field,
                    "story_points": "" if story_points_field.startswith("(None") else story_points_field,
                    "business_value": "" if business_value_field == "(None)" else business_value_field,
                },
                "swag": {"points_per_fte": float(points_per_fte or 0.0)},
                "forecast": {
                    "derived_fte_driver": str(forecast_driver or "SWAG"),
                    "velocity_row_fallback": bool(velocity_row_fallback),
                },
                "filters": {"exclude_states": list(exclude_states), "years_prefix": years_list},
                "epic_resolution": {"max_depth": int(max_depth)},
            }
            payload = normalize_profile_config(payload)
            overrides[selected_id] = payload
            updated_by = st.session_state.get("auth_user", {}).get("email") or st.session_state.get("user_email") or "ADMIN"
            try:
                target_id = None if save_as_new else selected_id
                upsert_ado_profile(
                    profile_id=target_id,
                    profile_name=profile_name_input,
                    config_json=json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False),
                    is_active=set_active,
                    updated_by=updated_by,
                )
                if set_active:
                    active_id = target_id
                    if active_id is None:
                        try:
                            df_active = fetch_df(
                                "SELECT TOP 1 PROFILE_ID FROM ADO_PROFILES WHERE PROFILE_NAME = %s ORDER BY UPDATED_AT DESC",
                                (profile_name_input,),
                            )
                            if df_active is not None and not df_active.empty:
                                active_id = str(df_active.iloc[0].get("PROFILE_ID") or "").strip() or None
                        except Exception:
                            active_id = None
                    if active_id:
                        set_active_ado_profile(active_id)
                st.success("Profile saved.")
                try:
                    _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                except Exception:
                    pass
                st.session_state["profiles_cache_bust"] = st.session_state.get("profiles_cache_bust", 0) + 1
                st.rerun()
            except Exception as e:
                st.error(f"Save failed: {e}")

    with st.expander("Advanced JSON editor", expanded=False):
        json_key = f"ado_profile_json_text_{selected_id}"
        if json_key not in st.session_state:
            st.session_state[json_key] = json.dumps(config if isinstance(config, dict) else {}, indent=2, sort_keys=True)
        json_text = st.text_area("Profile JSON", value=st.session_state[json_key], height=240, key=json_key)
        if st.button("Apply JSON", key=f"ado_profile_apply_json_{selected_id}"):
            try:
                parsed = json.loads(json_text)
            except Exception as e:
                st.error(f"Invalid JSON: {e}")
            else:
                overrides[selected_id] = parsed
                st.success("JSON applied to editor.")
                st.rerun()

    if is_active:
        st.info("Active profile cannot be deleted.")
    else:
        with st.expander("Delete profile", expanded=False):
            confirm = st.checkbox("I understand this will delete the selected profile.", key="ado_profile_delete_confirm")
            if st.button("Delete profile", disabled=not confirm, key="ado_profile_delete_btn"):
                try:
                    delete_ado_profile(selected_id)
                    st.success("Profile deleted.")
                    try:
                        _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                    except Exception:
                        pass
                    st.session_state["profiles_cache_bust"] = st.session_state.get("profiles_cache_bust", 0) + 1
                    st.rerun()
                except Exception as e:
                    st.error(f"Delete failed: {e}")

    if show_debug:
        with st.expander("Profiles debug", expanded=False):
            try:
                df_cnt = fetch_df("SELECT COUNT(*) AS N FROM ADO_PROFILES")
                count_val = int(df_cnt.iloc[0]["N"]) if df_cnt is not None and not df_cnt.empty else 0
                df_active = fetch_df(
                    "SELECT TOP 1 PROFILE_NAME FROM ADO_PROFILES WHERE IS_ACTIVE = 1 ORDER BY UPDATED_AT DESC"
                )
                active_name_dbg = (
                    str(df_active.iloc[0]["PROFILE_NAME"])
                    if df_active is not None and not df_active.empty
                    else "—"
                )
                df_latest = fetch_df(
                    "SELECT TOP 1 PROFILE_NAME, UPDATED_AT FROM ADO_PROFILES ORDER BY UPDATED_AT DESC"
                )
                latest_name = (
                    str(df_latest.iloc[0]["PROFILE_NAME"])
                    if df_latest is not None and not df_latest.empty
                    else "—"
                )
                latest_updated = (
                    str(df_latest.iloc[0]["UPDATED_AT"])
                    if df_latest is not None and not df_latest.empty
                    else "—"
                )
                st.write(f"row_count: {count_val}")
                st.write(f"active_profile: {active_name_dbg}")
                st.write(f"latest_profile: {latest_name} ({latest_updated})")
                st.write(f"selected_profile_id: {selected_id}")
            except Exception as e:
                st.error(f"Debug query failed: {e}")

def _ensure_iteration_expand_for_preview(expand_str: str) -> str:
    exp = (expand_str or "").strip()
    if not exp:
        return "Iteration($select=IterationLevel3,IterationPath,IterationSK)"
    m = re.search(r"Iteration\(\s*\$select\s*=\s*([^\)]*)\)", exp, flags=re.IGNORECASE)
    if not m:
        return exp + ("," if exp and not exp.endswith(",") else "") + "Iteration($select=IterationLevel3,IterationPath,IterationSK)"
    raw_fields = m.group(1)
    fields = [f.strip() for f in raw_fields.split(",") if f.strip()]
    fields_l = {f.lower() for f in fields}
    required = ["IterationLevel3", "IterationPath", "IterationSK"]
    for req in required:
        if req.lower() not in fields_l:
            fields.append(req)
    fixed = ",".join(fields)
    return exp[: m.start(1)] + fixed + exp[m.end(1) :]


def _ado_workitem_url(org: str, project: str, work_item_id: Any) -> str:
    return f"https://dev.azure.com/{org}/{project}/_workitems/edit/{work_item_id}"


def _parse_org_project_from_odata_base(base_url: str) -> Tuple[Optional[str], Optional[str]]:
    base = str(base_url or "").strip()
    if not base:
        return None, None
    try:
        parsed = urlparse(base)
        parts = [p for p in parsed.path.split("/") if p]
        if len(parts) >= 2:
            return parts[0], parts[1]
    except Exception:
        return None, None
    return None, None


def _field_role_from_profile(field_name: str, profile: Dict[str, Any]) -> str:
    name = str(field_name or "").strip().lower()
    if not name:
        return ""
    fields = profile.get("fields") if isinstance(profile.get("fields"), dict) else {}
    for role in ("story_points", "business_value", "app_name", "investment_dimension"):
        cfg_val = str(fields.get(role) or "").strip().lower()
        if cfg_val and cfg_val == name:
            return role
    return ""


def _resolve_ado_features_column(
    field_name: str,
    cols: List[str],
    *,
    field_role: str = "",
) -> Optional[str]:
    name = str(field_name or "").strip()
    if not name:
        return None
    role = str(field_role or "").strip().lower()
    if role == "story_points":
        return "STORY_POINTS" if "STORY_POINTS" in cols else ("EFFORT_POINTS" if "EFFORT_POINTS" in cols else None)
    if role == "business_value":
        return "BUSINESS_VALUE" if "BUSINESS_VALUE" in cols else None
    if role == "app_name":
        return "APP_NAME_RAW" if "APP_NAME_RAW" in cols else None
    if role == "investment_dimension":
        return "INVESTMENT_DIMENSION" if "INVESTMENT_DIMENSION" in cols else None
    if name in ("Effort", "StoryPoints", "STORY_POINTS"):
        return "STORY_POINTS" if "STORY_POINTS" in cols else ("EFFORT_POINTS" if "EFFORT_POINTS" in cols else None)
    if name in ("BusinessValue", "BUSINESS_VALUE"):
        return "BUSINESS_VALUE" if "BUSINESS_VALUE" in cols else None
    if name in ("Custom_ApplicationName", "APP_NAME_RAW"):
        return "APP_NAME_RAW" if "APP_NAME_RAW" in cols else None
    if name in ("Custom_InvestmentDimension", "INVESTMENT_DIMENSION"):
        return "INVESTMENT_DIMENSION" if "INVESTMENT_DIMENSION" in cols else None
    return name if name in cols else None


def _candidate_odata_field_names(field_name: str) -> List[str]:
    raw = str(field_name or "").strip()
    if not raw:
        return []
    candidates: List[str] = [raw]
    if raw.startswith("Fields."):
        candidates.append(raw[len("Fields."):])
    else:
        candidates.append(f"Fields.{raw}")
    if "." in raw:
        candidates.append(raw.replace(".", "_"))
    if "_" in raw:
        candidates.append(raw.replace("_", "."))
    out: List[str] = []
    seen = set()
    for name in candidates:
        n = str(name or "").strip()
        if not n:
            continue
        low = n.lower()
        if low in seen:
            continue
        seen.add(low)
        out.append(n)
    return out


def _resolve_df_column_name(df: pd.DataFrame, field_name: str) -> Optional[str]:
    if df is None or df.empty:
        return None
    candidates = _candidate_odata_field_names(field_name)
    if not candidates:
        return None
    cols = [str(c) for c in df.columns]
    by_lower = {c.lower(): c for c in cols}
    for cand in candidates:
        if cand in df.columns:
            return cand
        low = cand.lower()
        if low in by_lower:
            return by_lower[low]
    return None


def _apply_profile_field_mappings(df_expected: pd.DataFrame, profile_fields: Dict[str, Any]) -> Dict[str, Optional[str]]:
    resolved: Dict[str, Optional[str]] = {
        "story_points": None,
        "app_name": None,
        "investment_dimension": None,
        "business_value": None,
    }
    if df_expected is None or df_expected.empty or not isinstance(profile_fields, dict):
        return resolved

    sp_field = str(profile_fields.get("story_points") or "").strip()
    app_field = str(profile_fields.get("app_name") or "").strip()
    inv_field = str(profile_fields.get("investment_dimension") or "").strip()
    bv_field = str(profile_fields.get("business_value") or "").strip()

    resolved["story_points"] = _resolve_df_column_name(df_expected, sp_field) if sp_field else None
    resolved["app_name"] = _resolve_df_column_name(df_expected, app_field) if app_field else None
    resolved["investment_dimension"] = _resolve_df_column_name(df_expected, inv_field) if inv_field else None
    resolved["business_value"] = _resolve_df_column_name(df_expected, bv_field) if bv_field else None

    app_col = resolved.get("app_name")
    if app_col:
        df_expected["Custom_ApplicationName"] = df_expected.get(app_col)

    inv_col = resolved.get("investment_dimension")
    if inv_col:
        df_expected["Custom_InvestmentDimension"] = df_expected.get(inv_col)

    bv_col = resolved.get("business_value")
    if bv_col:
        df_expected["BusinessValue"] = pd.to_numeric(df_expected.get(bv_col), errors="coerce")

    return resolved


@cache_data_portfolio(ttl=300, show_spinner=False)
def _get_field_sample_cached(
    profile_id: str,
    base_url: str,
    profile_name: str,
    feature_type: str,
    field_role: str,
    field_name: str,
    limit: int,
    cache_bust: int,
) -> Dict[str, Any]:
    profile_cfg = get_active_ado_profile()
    auth_ctx = get_ado_auth_context(
        {"_profile_id": profile_id, "_profile_name": profile_name}, profile_cfg
    )
    rows: List[Dict[str, Any]] = []
    source = "none"
    if base_url and auth_ctx.get("has_pat"):
        try:
            base = base_url.rstrip("/")
            if not base.lower().endswith("/workitems"):
                base = base + "/WorkItems"
            select_fields = ["WorkItemId", "Title", field_name]
            feature_type_safe = str(feature_type or "Feature").replace("'", "''")
            filter_clause = f"WorkItemType eq '{feature_type_safe}'"
            url = base + "?" + "&".join(
                [f"$select={','.join(select_fields)}", f"$filter={filter_clause}", f"$top={limit}"]
            )
            payload = ado_odata_get(url, auth_ctx, {}, max_pages=1)
            rows = payload.get("value") or []
            source = "odata"
        except Exception:
            rows = []
            source = "none"
    if not rows:
        try:
            cols_df = fetch_df(
                """
                SELECT COLUMN_NAME
                FROM INFORMATION_SCHEMA.COLUMNS
                WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s
                """,
                ("dbo", "ADO_FEATURES"),
            )
            cols = cols_df["COLUMN_NAME"].tolist() if cols_df is not None and not cols_df.empty else []
            col = _resolve_ado_features_column(field_name, cols, field_role=field_role)
            if col:
                df = fetch_df(
                    f"""
                    SELECT TOP ({int(limit)})
                      FEATURE_ID,
                      TITLE,
                      {col} AS FIELD_VALUE
                    FROM ADO_FEATURES
                    WHERE {col} IS NOT NULL
                    ORDER BY FEATURE_ID DESC
                    """
                )
                if df is not None and not df.empty:
                    rows = df.to_dict(orient="records")
                    source = "db"
        except Exception:
            rows = []
            source = source if source != "none" else "none"
    return {"rows": rows, "source": source}


def get_field_sample(field_name: str, profile: Dict[str, Any], profile_id: str, limit: int = 20) -> Dict[str, Any]:
    base_url = str(profile.get("odata_base_url") or "").strip()
    cache_bust = int(st.session_state.get("profiles_cache_bust", 0))
    profile_name = str(profile.get("profile_name") or profile.get("PROFILE_NAME") or profile_id or "").strip()
    feature_type = str((profile.get("work_item_types") or {}).get("feature") or "Feature").strip() or "Feature"
    field_role = _field_role_from_profile(field_name, profile)
    return _get_field_sample_cached(
        profile_id,
        base_url,
        profile_name,
        feature_type,
        field_role,
        field_name,
        int(limit),
        cache_bust,
    )

def _ado_preview_url_from_state() -> str:
    query_mode = st.session_state.get("ado_query_mode", "Advanced OData Builder")
    full_url = str(st.session_state.get("ado_full_url_input") or "").strip()
    if query_mode == "Advanced OData Builder":
        org = str(st.session_state.get("ado_org") or "").strip()
        proj = str(st.session_state.get("ado_proj") or "").strip()
        entity = str(st.session_state.get("ado_entity") or "WorkItems").strip() or "WorkItems"
        select = str(st.session_state.get("ado_select") or "").strip() or None
        expand = str(st.session_state.get("ado_expand") or "").strip() or None
        filter_val = str(st.session_state.get("ado_filter") or "").strip() or None
        features_only = bool(st.session_state.get("ado_features_only", True))
        if features_only and not filter_val:
            filter_val = "WorkItemType eq 'Feature'"
        if expand is not None:
            expand = _ensure_iteration_expand_for_preview(expand)
        if org and proj:
            return _build_ado_odata_url(org, proj, entity, select or None, expand or None, filter_val or None, None)
        return ""
    return full_url


def resolve_ado_pat(
    profile_key: str = "",
    *,
    pat_env_key: str = "",
    fallback_pat: Optional[str] = None,
) -> Tuple[Optional[str], Dict[str, Any]]:
    key_raw = str(profile_key or "").strip()
    key_norm = re.sub(r"[^A-Za-z0-9]+", "_", key_raw).strip("_").upper() if key_raw else ""
    env_candidates: List[str] = []
    if pat_env_key:
        env_candidates.append(str(pat_env_key).strip())
    if key_raw:
        env_candidates.append(f"ADO_PAT_{key_raw}")
    if key_norm:
        env_candidates.append(f"ADO_PAT_{key_norm}")

    profile_pat: Optional[str] = None
    env_key_used: Optional[str] = None
    for env_key in env_candidates:
        val = str(os.getenv(env_key, "") or "").strip()
        if val:
            profile_pat = val
            env_key_used = env_key
            break

    legacy_pat: Optional[str] = None
    try:
        ado_secret_cfg = getattr(st, "secrets", {}).get("ado", {})  # type: ignore[attr-defined]
        if isinstance(ado_secret_cfg, dict):
            legacy_raw = str(ado_secret_cfg.get("pat") or "").strip()
            legacy_pat = legacy_raw or None
    except Exception:
        legacy_pat = None

    fb_pat = str(fallback_pat or "").strip() or None
    if profile_pat:
        source = "env_key"
        resolved = profile_pat
    elif legacy_pat:
        source = "secrets"
        resolved = legacy_pat
    elif fb_pat:
        source = "fallback"
        resolved = fb_pat
    else:
        source = "missing"
        resolved = None

    diag: Dict[str, Any] = {
        "source": source,
        "env_key_used": env_key_used,
        "profile_key": key_raw,
        "pat_env_key": str(pat_env_key or "").strip(),
        "has_profile_pat": bool(profile_pat),
        "has_legacy_pat": bool(legacy_pat),
        "has_fallback_pat": bool(fb_pat),
        "profile_pat": profile_pat,
        "legacy_pat": legacy_pat,
    }
    return resolved, diag


def render_ado_validate_tab() -> None:
    st.subheader("Validation / Preview")
    profile = get_active_ado_profile()
    active_profile_id = None
    active_profile_name = ""
    try:
        profiles_full = fetch_ado_profiles_full()
        if profiles_full is not None and not profiles_full.empty:
            active_mask = profiles_full["IS_ACTIVE"].astype(int) == 1
            if active_mask.any():
                active_profile_id = str(profiles_full.loc[active_mask, "PROFILE_ID"].iloc[0])
                active_profile_name = str(profiles_full.loc[active_mask, "PROFILE_NAME"].iloc[0])
    except Exception:
        active_profile_id = None
        active_profile_name = ""
    pat_key = str(get_active_profile_key() or active_profile_name or "").strip()
    pat_env_key = str(get_active_pat_env_key() or "").strip()
    pat_val, pat_diag = resolve_ado_pat(pat_key, pat_env_key=pat_env_key)
    has_pat = bool(pat_val)
    pat_source = str(pat_diag.get("source") or "missing")
    auth_ctx = {
        "pat": pat_val or "",
        "has_pat": has_pat,
        "source": pat_source,
        "auth_header": {},
    }
    def _org_project_from_url(url: str) -> str:
        base_url = str(url or "").strip()
        if not base_url:
            return "Unknown"
        try:
            parsed = urlparse(base_url)
            parts = [p for p in parsed.path.split("/") if p]
            if len(parts) >= 2:
                return f"{parts[0]} / {parts[1]}"
        except Exception:
            return "Unknown"
        return "Unknown"

    def _status_from_error(msg: str) -> str:
        m = re.search(r"\\((\\d{3})\\)", msg or "")
        return m.group(1) if m else "Unknown"

    with st.container(border=True):
        st.markdown("#### Authentication status")
        st.write(f"PAT configured: {'Yes' if has_pat else 'No'}")
        st.write(f"Source: {pat_source}")
        st.caption("PAT must have access to Analytics for this project.")
    if not has_pat:
        st.warning("No PAT configured. Provide ADO_PAT_<PROFILE> via env/secrets or paste a session PAT.")

    _select_fields_ignored, expand_clause = build_workitems_select_expand(profile)
    sp_field = str((profile.get("fields") or {}).get("story_points") or "").strip()
    bv_field_sel = str((profile.get("fields") or {}).get("business_value") or "").strip()
    app_field_sel = str((profile.get("fields") or {}).get("app_name") or "").strip()
    inv_field_sel = str((profile.get("fields") or {}).get("investment_dimension") or "").strip()
    profile_fields = {f for f in [sp_field, bv_field_sel, app_field_sel, inv_field_sel] if f}
    extra_fields = []
    for key, val in (profile.get("fields") or {}).items():
        val_str = str(val or "").strip()
        if val_str and val_str not in profile_fields:
            extra_fields.append(val_str)
    if extra_fields:
        profile_fields.update(extra_fields)

    base_fields = ["WorkItemId", "WorkItemType", "Title", "State", "ChangedDate", "ParentWorkItemId"]
    prioritized = []
    for f in [sp_field, bv_field_sel, app_field_sel, inv_field_sel]:
        if f and f not in prioritized:
            prioritized.append(f)
    remaining = sorted(profile_fields - set(prioritized))
    max_profile_fields = 10
    if len(prioritized) + len(remaining) > max_profile_fields:
        remaining = remaining[: max(0, max_profile_fields - len(prioritized))]
    final_fields = base_fields + prioritized + remaining
    select_fields = []
    seen = set()
    for f in final_fields:
        if f and f not in seen:
            select_fields.append(f)
            seen.add(f)
    feature_type = str(
        (profile.get("work_item_types") or {}).get("feature") or "Feature"
    ).strip() or "Feature"
    filter_clause = f"WorkItemType eq '{feature_type}'"
    select_str = ",".join(select_fields)

    base_url = str(profile.get("odata_base_url") or "").strip()
    if base_url:
        base = base_url.rstrip("/")
        if not base.lower().endswith("/workitems"):
            base = base + "/WorkItems"
        parts = [f"$select={select_str}", f"$expand={expand_clause}", f"$filter={filter_clause}", "$top=20"]
        preview_url = base + "?" + "&".join(parts)
    else:
        org, proj = _ado_org_proj_defaults()
        if org and proj:
            preview_url = _build_ado_odata_url(org, proj, "WorkItems", select_str, expand_clause, filter_clause, None)
            preview_url = preview_url + ("&" if "?" in preview_url else "?") + "$top=20"
        else:
            preview_url = ""

    if preview_url:
        st.code(preview_url, language="text")
    else:
        st.info("Configure a profile base URL or legacy org/project in Sync to generate a preview URL.")

    if st.button("Run Preview", key="ado_validate_run_preview", disabled=not preview_url or not has_pat):
        diag_api: Dict[str, Any] = {}
        try:
            payload = ado_odata_get(preview_url, auth_ctx, diag_api, max_pages=1)
        except Exception as e:
            msg = str(e)
            if "401" in msg or "403" in msg or "not authorized" in msg.lower():
                st.error("Not authorized. Confirm PAT has access to this org/project and Analytics is enabled.")
                st.caption(f"Org/Project: {_org_project_from_url(base_url or preview_url)}")
                st.caption(f"HTTP status: {_status_from_error(msg)}")
                st.caption("Request URL:")
                st.code(preview_url, language="text")
                return
            st.error(f"Preview failed: {e}")
            with st.expander("Debug - diagnostics", expanded=True):
                st.json(diag_api)
            return

        rows = payload.get("value") or []
        if not rows:
            st.warning("Preview returned no rows.")
            return

        st.session_state["ado_validate_preview_rows"] = rows
        with st.expander("Preview request details", expanded=False):
            st.caption("Preview select fields:")
            st.write(", ".join(select_fields))

        def _build_obj(row: Dict[str, Any], prefix: str, fields: List[str]) -> Dict[str, Any]:
            if isinstance(row.get(prefix), dict):
                return row.get(prefix) or {}
            out: Dict[str, Any] = {}
            for field in fields:
                out[field] = row.get(f"{prefix}.{field}") or row.get(field)
            return out

        iter_fields = ["IterationLevel3", "IterationPath", "IterationSK"]
        area_fields = ["AreaLevel1", "AreaLevel2", "AreaLevel3", "AreaLevel4"]
        points_key = sp_field

        def _has_val(row: Dict[str, Any], key: str) -> bool:
            if key in row:
                return pd.notna(row.get(key))
            return False

        iter_present = []
        area_present = []
        parent_present = []
        points_present = []
        app_present = []
        bv_present = []
        program_team = ("", "")
        pi_label = ""
        app_name = ""

        for idx, row in enumerate(rows):
            iter_obj = _build_obj(row, "Iteration", iter_fields)
            area_obj = _build_obj(row, "Area", area_fields)
            if idx == 0:
                program_team = extract_program_team(profile, area_obj)
                pi_label = extract_pi_label(profile, iter_obj)
                app_name = extract_app_name(profile, row)
            iter_present.append(any(pd.notna(iter_obj.get(f)) for f in iter_fields))
            area_present.append(any(pd.notna(area_obj.get(f)) for f in area_fields))
            parent_present.append(_has_val(row, "ParentWorkItemId"))
            points_present.append(_has_val(row, points_key) if points_key else False)
            app_present.append(bool(app_name) if idx == 0 else bool(extract_app_name(profile, row)))
            bv_key = bv_field_sel
            bv_present.append(_has_val(row, bv_key) if bv_key else False)

        diagnostics = {
            "Has Iteration": f"{sum(iter_present)}/{len(rows)}",
            "Has Area": f"{sum(area_present)}/{len(rows)}",
            "Has ParentWorkItemId": f"{sum(parent_present)}/{len(rows)}",
            "Has story points field": f"{sum(points_present)}/{len(rows)}",
            "Parsed Program": program_team[0],
            "Parsed Team": program_team[1],
            "Parsed PI label": pi_label,
            "Parsed App name": app_name,
        }
        st.dataframe(pd.DataFrame([diagnostics]), use_container_width=True)

        parent_pct = (sum(parent_present) / len(rows) * 100.0) if rows else 0.0
        st.caption(f"ParentWorkItemId present: {parent_pct:.1f}% of preview rows. Epic resolution will walk the parent chain later.")

        points_field = sp_field
        app_field = str((profile.get("fields") or {}).get("app_name") or "").strip()
        bv_field = bv_field_sel
        points_rate = (sum(points_present) / len(rows) * 100.0) if rows else 0.0
        app_rate = (sum(app_present) / len(rows) * 100.0) if rows else 0.0
        bv_rate = (sum(bv_present) / len(rows) * 100.0) if rows else 0.0
        parent_rate = (sum(parent_present) / len(rows) * 100.0) if rows else 0.0
        with st.container(border=True):
            st.markdown("#### Profile Suggestions")
            st.write(f"Points field `{points_field or '—'}` present: {points_rate:.1f}%")
            st.write(f"App field `{app_field or '—'}` present: {app_rate:.1f}%")
            st.write(f"Business Value field `{bv_field or '—'}` present: {bv_rate:.1f}%")
            st.write(f"ParentWorkItemId present: {parent_rate:.1f}%")
            if points_rate < 30.0:
                st.warning("Story points field seems mostly null; verify profile.fields.story_points.")
            if app_rate < 30.0:
                st.warning("App name field seems missing; verify profile.fields.app_name.")
            if bv_rate < 30.0 and bv_field:
                st.warning("Business Value field seems mostly null; verify profile.fields.business_value.")
            if parent_rate < 30.0:
                st.warning("ParentWorkItemId is often missing; epic grouping may be limited.")

        df_preview = pd.json_normalize(rows)
        display_cols: List[str] = []
        for col in ["WorkItemId", "Title", "State", "ParentWorkItemId"]:
            if col not in df_preview.columns:
                df_preview[col] = pd.NA
            display_cols.append(col)
        epic_cols = [c for c in ["EpicId", "EpicTitle", "EpicState", "EPIC_ID", "EPIC_TITLE", "EPIC_STATE"] if c in df_preview.columns]
        display_cols.extend(epic_cols)
        col_labels: Dict[str, str] = {}
        profile_order = [f for f in (prioritized + remaining) if f]
        label_map = {
            sp_field: "Story Points",
            bv_field_sel: "Business Value",
            app_field_sel: "App Name",
            inv_field_sel: "Investment Dimension",
        }
        for field_name in profile_order:
            label = label_map.get(field_name, field_name)
            if field_name in label_map:
                col_labels[field_name] = f"{label} (field: {field_name})"
            else:
                col_labels[field_name] = label
            if field_name not in df_preview.columns:
                df_preview[field_name] = pd.NA
                st.warning(f"Selected field '{field_name}' not present in preview response. Check Discover Fields or your OData permissions.")
            display_cols.append(field_name)

        for field_name, label in [
            (sp_field, "Story Points"),
            (bv_field_sel, "Business Value"),
            (app_field_sel, "App Name"),
            (inv_field_sel, "Investment Dimension"),
        ]:
            if field_name:
                col_labels[field_name] = f"{label} (field: {field_name})"

        preview_view = df_preview[display_cols] if display_cols else df_preview
        st.dataframe(
            preview_view.rename(columns=col_labels),
            use_container_width=True,
            height=280,
        )

    preview_rows = st.session_state.get("ado_validate_preview_rows") or []
    if preview_rows:
        parent_present = [
            pd.notna(r.get("ParentWorkItemId")) for r in preview_rows
        ]
        parent_pct = (sum(parent_present) / len(preview_rows) * 100.0) if preview_rows else 0.0
        with st.container(border=True):
            st.markdown("#### Epic Readiness")
            st.write(f"ParentWorkItemId present: {parent_pct:.1f}%")
            st.write("Parent types: unknown (parents not fetched in preview)")
            if parent_pct < 50.0:
                st.warning("Most preview rows lack ParentWorkItemId; epic grouping will be limited.")
            else:
                st.info("ParentWorkItemId present on many rows; epic grouping should be feasible.")

        if st.button("Resolve Epic for preview rows", key="ado_validate_resolve_epics_btn"):
            try:
                from scripts.ado_epic_resolver import (
                    fetch_workitems_by_ids,
                    fetch_workitems_by_filter,
                    resolve_epics_for_features,
                )
            except Exception as e:
                st.error(f"Epic resolver not available: {e}")
                return
            if not has_pat:
                st.warning("No PAT configured. Provide ADO_PAT_<PROFILE> via env/secrets or paste a session PAT.")
                return
            base_url = str(profile.get("odata_base_url") or "").strip()
            if not base_url:
                st.error("Active profile is missing odata_base_url. Configure it in Profiles.")
                return
            select_fields = [
                "WorkItemId",
                "WorkItemType",
                "Title",
                "State",
                "ParentWorkItemId",
            ]

            def _fetch_ids(ids: List[int]) -> List[Dict[str, Any]]:
                return fetch_workitems_by_ids(base_url, ids, auth_ctx, select_fields)

            _fetch_ids.fetch_by_filter = lambda filter_clause: fetch_workitems_by_filter(
                base_url, filter_clause, auth_ctx, select_fields
            )

            epic_map = resolve_epics_for_features(preview_rows, profile, _fetch_ids, max_rounds=2)
            table_rows = []
            for row in preview_rows:
                fid = str(row.get("WorkItemId") or "")
                mapped = epic_map.get(str(fid), {})
                table_rows.append(
                    {
                        "FeatureId": fid,
                        "FeatureTitle": row.get("Title"),
                        "ParentId": mapped.get("parent_id"),
                        "EpicId": mapped.get("epic_id"),
                        "EpicTitle": mapped.get("epic_title"),
                        "EpicState": mapped.get("epic_state"),
                    }
                )
            st.dataframe(pd.DataFrame(table_rows), use_container_width=True)

    with st.container(border=True):
        st.markdown("#### Discover ADO Fields")
        use_metadata = st.toggle(
            "Use $metadata (recommended)",
            value=True,
            key="ado_validate_use_metadata",
        )
        if st.button("Discover fields ($top=25)", key="ado_validate_discover_fields", disabled=not preview_url or not has_pat):
            base_url = str(profile.get("odata_base_url") or "").strip()
            feature_type = str((profile.get("work_item_types") or {}).get("feature") or "Feature").strip() or "Feature"
            exclude_states = list((profile.get("filters") or {}).get("exclude_states") or [])
            years_prefix = list((profile.get("filters") or {}).get("years_prefix") or [])
            pi_field = str((profile.get("iteration_mapping") or {}).get("pi_field") or "IterationLevel3").strip() or "IterationLevel3"
            if not base_url:
                st.error("Active profile is missing odata_base_url. Configure it in Profiles.")
                return
            if not has_pat:
                st.warning("No PAT configured. Provide ADO_PAT_<PROFILE> via env/secrets or paste a session PAT.")
                return
            metadata_fields: List[str] = []
            if use_metadata:
                try:
                    @cache_data_portfolio(ttl=3600, show_spinner=False)
                    def _metadata_fields_cached(url: str, profile_id: str) -> List[str]:
                        xml_text = fetch_ado_odata_metadata(url, auth_ctx)
                        return parse_ado_metadata_fields(xml_text)
                    metadata_fields = _metadata_fields_cached(base_url, active_profile_id or "")
                except Exception as e:
                    st.warning(f"$metadata discovery failed; falling back to sample keys. {e}")
                    metadata_fields = []
            base = base_url.rstrip("/")
            if not base.lower().endswith("/workitems"):
                base = base + "/WorkItems"
            select_fields = [
                "WorkItemId",
                "WorkItemType",
                "Title",
                "State",
                "ChangedDate",
                "ParentWorkItemId",
                "Effort",
                "BusinessValue",
            ]
            if sp_field and sp_field not in select_fields:
                select_fields.append(sp_field)
            if bv_field_sel and bv_field_sel not in select_fields:
                select_fields.append(bv_field_sel)
            custom_candidates = []
            if metadata_fields:
                patterns = ("Custom_Value", "Custom_Points", "Custom_Effort", "Custom_BV", "Custom_Business")
                for f in metadata_fields:
                    if f.startswith("Custom_") and any(p.lower() in f.lower() for p in patterns):
                        custom_candidates.append(f)
            custom_candidates = sorted(set(custom_candidates))[:20]
            for f in custom_candidates:
                if f not in select_fields:
                    select_fields.append(f)
            expand_clause = "Iteration($select=IterationLevel3,IterationPath,IterationSK),Area($select=AreaLevel1,AreaLevel2,AreaLevel3,AreaLevel4,AreaPath)"
            filter_parts = [f"WorkItemType eq '{feature_type}'"]
            for state in exclude_states:
                state_val = str(state or "").strip()
                if state_val:
                    filter_parts.append(f"State ne '{state_val}'")
            if years_prefix:
                ors = [f"startswith(Iteration/{pi_field},'{str(y).strip()}')" for y in years_prefix if str(y).strip()]
                if ors:
                    filter_parts.append("(" + " or ".join(ors) + ")")
            filter_clause = " and ".join(filter_parts)
            discover_url = base + "?" + "&".join(
                [f"$select={','.join(select_fields)}", f"$expand={expand_clause}", f"$filter={filter_clause}", "$top=25"]
            )
            diag_discover: Dict[str, Any] = {}
            try:
                payload = ado_odata_get(discover_url, auth_ctx, diag_discover, max_pages=1)
            except Exception as e:
                msg = str(e)
                if "401" in msg or "403" in msg or "not authorized" in msg.lower():
                    st.error("Not authorized. Confirm PAT has access to this org/project and Analytics is enabled.")
                    st.caption(f"Org/Project: {_org_project_from_url(base_url)}")
                    st.caption(f"HTTP status: {_status_from_error(msg)}")
                    st.caption("Request URL:")
                    st.code(discover_url, language="text")
                    return
                st.error(f"Discover fields failed: {e}")
                with st.expander("Debug - diagnostics", expanded=True):
                    st.json(diag_discover)
            else:
                rows = payload.get("value") or []
                if not rows:
                    st.warning("Discover fields returned no rows.")
                else:
                    row = rows[0]
                    top_keys = sorted(row.keys())
                    if metadata_fields:
                        st.write(f"$metadata fields: {len(metadata_fields)}")
                        custom_meta = sorted([k for k in metadata_fields if k.startswith("Custom_")])
                        if custom_meta:
                            st.write("Custom_* fields (metadata):")
                            st.code(", ".join(custom_meta))
                    st.write("Top-level keys:")
                    st.code(", ".join(top_keys) if top_keys else "—")
                    iter_obj = row.get("Iteration")
                    if isinstance(iter_obj, dict):
                        st.write("Iteration keys:")
                        st.code(", ".join(sorted(iter_obj.keys())) if iter_obj else "—")
                    else:
                        iter_keys = sorted({k.split("Iteration.", 1)[1] for k in row.keys() if k.startswith("Iteration.")})
                        st.write("Iteration keys:")
                        st.code(", ".join(iter_keys) if iter_keys else "—")
                    area_obj = row.get("Area")
                    if isinstance(area_obj, dict):
                        st.write("Area keys:")
                        st.code(", ".join(sorted(area_obj.keys())) if area_obj else "—")
                    else:
                        area_keys = sorted({k.split("Area.", 1)[1] for k in row.keys() if k.startswith("Area.")})
                        st.write("Area keys:")
                        st.code(", ".join(area_keys) if area_keys else "—")
                    coverage_rows = []
                    for field in select_fields:
                        non_null = sum(pd.notna(r.get(field)) for r in rows)
                        coverage = (non_null / len(rows) * 100.0) if rows else 0.0
                        note = "numeric candidate" if field in {"Effort", "StoryPoints", "BusinessValue"} or field.startswith("Custom_") else ""
                        coverage_rows.append({"Field": field, "Coverage%": round(coverage, 1), "Notes": note})
                    st.dataframe(pd.DataFrame(coverage_rows), use_container_width=True, height=260)
                    st.session_state["ado_discovered_fields"] = {
                        "top": top_keys,
                        "iteration": iter_obj if isinstance(iter_obj, dict) else iter_keys,
                        "area": area_obj if isinstance(area_obj, dict) else area_keys,
                        "metadata_fields": metadata_fields,
                    }
                    custom_keys = sorted([k for k in row.keys() if str(k).startswith("Custom_")])
                    if custom_keys:
                        st.write("Custom_* keys:")
                        st.code(", ".join(custom_keys))
                    candidate_points = [
                        k
                        for k in top_keys
                        if (k in {"Effort", "StoryPoints"} or str(k).startswith("Custom_"))
                        and str(k).strip() != "Custom_FTE"
                    ]
                    if candidate_points:
                        st.write("Candidate Story Points fields:")
                        st.code(", ".join(candidate_points))

def render_ado_sync_tab() -> None:
    try:
        info = fetch_df("SELECT COUNT(*) AS N FROM ADO_FEATURES")
        n_rows = int(info.iloc[0]["N"]) if info is not None and not info.empty else 0
        st.info(f"Current rows in ADO_FEATURES: **{n_rows}**")
    except Exception:
        st.warning("Could not query ADO_FEATURES row count (check connection/secrets).")

    if "ado_portfolio_select_pending" in st.session_state:
        st.session_state["ado_portfolio_select"] = st.session_state.pop("ado_portfolio_select_pending")
    if "ado_profile_pat_input_pending" in st.session_state:
        st.session_state["ado_profile_pat_input"] = st.session_state.pop("ado_profile_pat_input_pending")

    profile_cfg = normalize_profile_config(get_active_ado_profile())
    active_profile_name = "Unknown"
    active_profile_id = None
    active_profile_has_pat = False
    active_profile_pat_source = "missing"
    active_profile_ctx_key = get_active_profile_key()
    active_profile_sync_mode = None
    active_profile_updated_at = None
    active_profile_last_sync_at = None
    active_profile_out_of_sync = False

    def _coerce_dt(val: Any) -> Optional[datetime]:
        if val is None or (isinstance(val, float) and math.isnan(val)):
            return None
        try:
            if isinstance(val, datetime):
                return val
            return pd.to_datetime(val, errors="coerce").to_pydatetime()  # type: ignore[return-value]
        except Exception:
            return None
    try:
        profiles_full = fetch_ado_profiles_full()
        if profiles_full is not None and not profiles_full.empty:
            active_mask = profiles_full["IS_ACTIVE"].astype(int) == 1
            if active_mask.any():
                row = profiles_full.loc[active_mask].iloc[0]
                active_profile_name = str(row.get("PROFILE_NAME") or "Unknown")
                active_profile_id = str(row.get("PROFILE_ID") or "").strip() or None
                pat_key = active_profile_ctx_key or active_profile_name
                pat_env_key = get_active_pat_env_key()
                pat_val, pat_diag = resolve_ado_pat(pat_key, pat_env_key=pat_env_key)
                pat_source = str(pat_diag.get("source") or "missing")
                active_profile_has_pat = bool(pat_val)
                active_profile_pat_source = pat_source
                active_profile_sync_mode = str(row.get("SYNC_MODE") or "").strip() or None
                active_profile_updated_at = _coerce_dt(row.get("UPDATED_AT"))
                active_profile_last_sync_at = _coerce_dt(row.get("LAST_SYNC_AT"))
                if active_profile_last_sync_at is None:
                    active_profile_out_of_sync = True
                elif active_profile_updated_at and active_profile_updated_at > active_profile_last_sync_at:
                    active_profile_out_of_sync = True
    except Exception:
        active_profile_name = "Unknown"

    prev_profile = st.session_state.get("ado_active_profile_id_prev")
    if "ado_sync_mode" not in st.session_state or active_profile_id != prev_profile:
        if active_profile_sync_mode == "profile":
            st.session_state["ado_sync_mode"] = "Profile-driven (ADO Profiles)"
        else:
            st.session_state["ado_sync_mode"] = "Legacy (portfolio settings)"
        st.session_state["ado_active_profile_id_prev"] = active_profile_id
    sync_mode_value = st.session_state.get("ado_sync_mode", "Legacy (portfolio settings)")
    pat_diag_sync: Dict[str, Any] = {}
    pat_key_sync = str(active_profile_ctx_key or active_profile_name or "").strip()
    pat_env_key_sync = str(get_active_pat_env_key() or "").strip()
    _resolved_sync_pat, pat_diag_sync = resolve_ado_pat(
        pat_key_sync,
        pat_env_key=pat_env_key_sync,
    )
    profile_pat: Optional[str] = pat_diag_sync.get("profile_pat")
    legacy_pat: Optional[str] = pat_diag_sync.get("legacy_pat")
    with st.container(border=True):
        st.markdown("### Active Configuration Summary")
        story_points_field = str((profile_cfg.get("fields") or {}).get("story_points") or "").strip() or "None"
        bv_field = str((profile_cfg.get("fields") or {}).get("business_value") or "").strip() or "Not configured"
        app_field = str((profile_cfg.get("fields") or {}).get("app_name") or "").strip() or "Not configured"
        swag_ppf = float((profile_cfg.get("swag") or {}).get("points_per_fte") or 65.0)
        base_url = str(profile_cfg.get("odata_base_url") or "").strip()
        org_project = "Unknown"
        if base_url:
            try:
                parsed = urlparse(base_url)
                parts = [p for p in parsed.path.split("/") if p]
                if len(parts) >= 2:
                    org_project = f"{parts[0]} / {parts[1]}"
            except Exception:
                org_project = "Unknown"
        config_source = "ADO Profiles" if sync_mode_value == "Profile-driven (ADO Profiles)" else "Portfolio settings"
        st.write(f"Active profile: {active_profile_name}")
        if active_profile_ctx_key:
            st.caption(f"Profile key: {active_profile_ctx_key}")
        st.write(f"PAT configured: {'Yes' if active_profile_has_pat else 'No'}")
        st.caption(f"PAT source: {active_profile_pat_source if active_profile_pat_source != 'missing' else 'missing'}")
        st.caption("Mode: Online Sync available" if active_profile_has_pat else "Mode: Offline restore")
        st.write(f"Org/Project: {org_project}")
        st.write(f"Config source: {config_source}")
        st.write(f"Sync mode: {sync_mode_value}")
        st.write(f"Profile sync status: {'Requires sync' if active_profile_out_of_sync else 'Up to date'}")
        coverage = _story_points_coverage_summary(int(st.session_state.get("ado_data_cache_bust", 0)))
        c_sp1, c_sp2 = st.columns([2, 2])
        with c_sp1:
            st.write(f"Story Points field: {story_points_field}")
        with c_sp2:
            _render_story_points_coverage(coverage)
        st.write(f"SWAG points per FTE: {swag_ppf:g}")
        st.write(f"Derived FTE: Story Points ÷ {swag_ppf:g}")
        st.write("Derived FTE is always computed from Story Points using SWAG. No FTE field exists in ADO.")
        st.write(f"Business Value field: {bv_field}")
        st.write(f"App Name field: {app_field}")

    def _on_sync_mode_change() -> None:
        if not active_profile_id:
            return
        mode = st.session_state.get("ado_sync_mode", "Legacy (portfolio settings)")
        stored = "profile" if mode == "Profile-driven (ADO Profiles)" else "legacy"
        updated_by = st.session_state.get("auth_user", {}).get("email") or st.session_state.get("user_email") or "ADMIN"
        try:
            set_profile_sync_mode(active_profile_id, stored, updated_by=updated_by)
        except Exception as e:
            st.error(f"Failed to persist sync mode: {e}")

    tab_run, tab_connection, tab_advanced = st.tabs(["Run Sync", "Connection", "Advanced"])

    with tab_run:
        if active_profile_out_of_sync:
            st.warning("Profile changes detected. Run Sync to apply updated demand logic.")
        sync_mode_options = ["Legacy (portfolio settings)", "Profile-driven (ADO Profiles)"]
        sync_index = sync_mode_options.index(sync_mode_value) if sync_mode_value in sync_mode_options else 0
        sync_mode = st.selectbox(
            "Sync mode",
            sync_mode_options,
            index=sync_index,
            key="ado_sync_mode",
            on_change=_on_sync_mode_change,
        )
        sync_mode_value = sync_mode
        st.caption("Saved with Active Profile.")
        st.subheader("Run Sync")
        st.caption("Run sync using the active configuration.")
        pat_source_sync = str(pat_diag_sync.get("source") or "missing")
        if pat_source_sync == "env_key":
            st.caption("PAT source: env var  ✅")
        elif pat_source_sync in {"secrets", "fallback"}:
            st.caption("PAT source: secrets ✅")
        else:
            st.caption("PAT missing ❌")

    # Try to read defaults from secrets if you added them
    try:
        _ado_cfg = getattr(st, "secrets", {}).get("ado", {})  # type: ignore
    except Exception:
        _ado_cfg = {}
    _ado_defaults = {
        "pat": (_ado_cfg.get("pat") if isinstance(_ado_cfg, dict) else None) or "",
        "org": (_ado_cfg.get("org") if isinstance(_ado_cfg, dict) else None) or "",
        "project": (_ado_cfg.get("project") if isinstance(_ado_cfg, dict) else None) or "",
        "entity": (_ado_cfg.get("entity") if isinstance(_ado_cfg, dict) else None) or "WorkItems",
        "select": (_ado_cfg.get("select") if isinstance(_ado_cfg, dict) else None)
        or "WorkItemId,Title,State,CreatedDate,ChangedDate,BusinessValue,Custom_ApplicationName,Custom_InvestmentDimension",
        "expand": (_ado_cfg.get("expand") if isinstance(_ado_cfg, dict) else None)
        or "Team($select=TeamName),Iteration($select=IterationLevel3,IterationPath,IterationSK),Area($select=AreaPath,AreaLevel1,AreaLevel2,AreaLevel3,AreaLevel4)",
        "filter": (_ado_cfg.get("filter") if isinstance(_ado_cfg, dict) else None) or "WorkItemType eq 'Feature'",
        "url": (_ado_cfg.get("url") if isinstance(_ado_cfg, dict) else None) or "",
    }
    _resolved_sync_pat, pat_diag_sync = resolve_ado_pat(
        pat_key_sync,
        pat_env_key=pat_env_key_sync,
        fallback_pat=_ado_defaults.get("pat"),
    )
    profile_pat = pat_diag_sync.get("profile_pat")
    legacy_pat = pat_diag_sync.get("legacy_pat")

    def _ensure_iteration_expand(expand_str: str) -> str:
        exp = (expand_str or "").strip()
        if not exp:
            return "Iteration($select=IterationLevel3,IterationPath,IterationSK)"
        # Ensure Iteration($select=...) contains join keys needed for iteration calendar mapping.
        m = re.search(r"Iteration\\(\\s*\\$select\\s*=\\s*([^\\)]*)\\)", exp, flags=re.IGNORECASE)
        if not m:
            # No Iteration expand at all: append it.
            return exp + ("," if exp and not exp.endswith(",") else "") + "Iteration($select=IterationLevel3,IterationPath,IterationSK)"
        raw_fields = m.group(1)
        fields = [f.strip() for f in raw_fields.split(",") if f.strip()]
        fields_l = {f.lower() for f in fields}
        required = ["IterationLevel3", "IterationPath", "IterationSK"]
        for req in required:
            if req.lower() not in fields_l:
                fields.append(req)
        fixed = ",".join(fields)
        return exp[: m.start(1)] + fixed + exp[m.end(1) :]

    def _validate_full_url(url: str) -> Optional[str]:
        u = (url or "").strip()
        if not u:
            return "Missing OData URL."
        bad_tokens = ["$metadata", "@odata.context", "odata.context"]
        if any(tok.lower() in u.lower() for tok in bad_tokens):
            return "This looks like metadata/context, not an OData query URL. Paste the WorkItems query URL (with $select/$expand/$filter)."
        # Require Iteration join keys for iteration calendar mapping.
        decoded = u.replace("%28", "(").replace("%29", ")").replace("%2C", ",").replace("%24", "$")
        if "Iteration(" not in decoded and "iteration(" not in decoded:
            return "Missing Iteration expand. Add: $expand=Iteration($select=IterationLevel3,IterationPath,IterationSK)"
        for req in ("IterationLevel3", "IterationPath", "IterationSK"):
            if req.lower() not in decoded.lower():
                return "Iteration expand must include IterationLevel3, IterationPath, and IterationSK for calendar mapping."
        return None

    with tab_connection:
        st.markdown("### Connection & portfolio")
        require_pat = True
        profile_key = str(get_active_profile_key() or active_profile_name or "").strip()
        pat_env_key = get_active_pat_env_key()
        if profile_key:
            pat_val, pat_diag_conn = resolve_ado_pat(
                profile_key,
                pat_env_key=pat_env_key,
                fallback_pat=_ado_defaults.get("pat"),
            )
            pat_source = str(pat_diag_conn.get("source") or "missing")
        else:
            pat_val, pat_source = None, "missing"
        pat_configured = bool(pat_val)
        built_url = ""
        sel_portfolio = "(unsaved)"
        portfolio_name = ""
        current_settings: Dict[str, Any] = {}
        stored_pat = ""
        store_pat = False
        max_pages = int(st.session_state.get("ado_max_pages", 5))
        query_mode = st.session_state.get("ado_query_mode", "Advanced OData Builder")
        full_url = str(st.session_state.get("ado_full_url_input") or _ado_defaults["url"] or "")
        org = st.session_state.get("ado_org", _ado_defaults["org"])
        proj = st.session_state.get("ado_proj", _ado_defaults["project"])
        entity = st.session_state.get("ado_entity", _ado_defaults["entity"])
        select = st.session_state.get("ado_select", _ado_defaults["select"])
        expand = st.session_state.get("ado_expand", _ado_defaults["expand"])
        filter_ = st.session_state.get("ado_filter", _ado_defaults["filter"])
        features_only = st.session_state.get("ado_features_only", True)
        pat = str(st.session_state.get("ado_pat") or "")
        with st.container(border=True):
            st.markdown("#### Authentication")
            st.write(f"Active profile: {active_profile_name}")
            st.write(f"PAT configured for active profile: {'Yes' if pat_configured else 'No'}")
            st.caption(f"Configured via env/secrets: {'Yes' if pat_source in ('env', 'secrets') else 'No'}")
            st.caption(f"PAT source: {pat_source if pat_source != 'missing' else 'missing'}")
            pat_input = st.text_input("Paste PAT for this session", value="", type="password", key="ado_profile_pat_input")
            col_pat_a, col_pat_b = st.columns([1, 1])
            with col_pat_a:
                if st.button("Save PAT to Session", key="ado_profile_pat_save"):
                    if not profile_key:
                        st.error("No active profile selected.")
                    elif not pat_input:
                        st.error("Enter a PAT to store for this session.")
                    else:
                        st.session_state.setdefault("ado_pat_overrides", {})
                        st.session_state["ado_pat_overrides"][profile_key] = pat_input
                        st.success("PAT saved for this session.")
                        st.rerun()
            with col_pat_b:
                if st.button("Clear Session PAT", key="ado_profile_pat_clear"):
                    st.session_state.setdefault("ado_pat_overrides", {})
                    if profile_key in st.session_state["ado_pat_overrides"]:
                        st.session_state["ado_pat_overrides"].pop(profile_key, None)
                        st.success("Session PAT cleared.")
                        st.rerun()
            if is_admin:
                with st.expander("Migration: remove stored PATs from DB", expanded=False):
                    st.caption("This clears PAT_ENC for all ADO profiles. PATs will need env/secrets or session paste.")
                    confirm_clear = st.checkbox("I understand this will remove all stored PATs.", key="ado_pat_clear_confirm")
                    if st.button("Remove stored PATs from DB", disabled=not confirm_clear, key="ado_pat_clear_btn"):
                        try:
                            execute("UPDATE ADO_PROFILES SET PAT_ENC = NULL", None)
                            st.success("Stored PATs cleared from DB.")
                        except Exception as e:
                            st.error(f"Failed to clear PATs: {e}")

        st.caption("Legacy mode settings are optional. Profile-driven uses the Active Profile base URL and field mapping.")
        with st.expander("Legacy OData Builder (optional)", expanded=False):
            col_conn, col_port = st.columns([3, 2])
            with col_port:
                with st.container(border=True):
                    st.markdown("#### Portfolio")
                    # Portfolio-level settings persisted in session
                    st.session_state.setdefault("ado_portfolio_settings", {})
                    portfolio_options = ["(unsaved)"] + sorted(st.session_state["ado_portfolio_settings"].keys())
                    sel_portfolio = st.selectbox("Portfolio", portfolio_options, index=0, key="ado_portfolio_select")
                    portfolio_name = st.text_input("Portfolio name", value=("" if sel_portfolio == "(unsaved)" else sel_portfolio), key="ado_portfolio_name").strip()
            current_settings = st.session_state["ado_portfolio_settings"].get(sel_portfolio, {}) if sel_portfolio != "(unsaved)" else {}
            stored_pat = ""
        
            # Apply stored defaults only when switching portfolios (avoid overwriting user input on reruns).
            active_key = sel_portfolio if sel_portfolio != "(unsaved)" else "(unsaved)"
            if st.session_state.get("_ado_active_portfolio_key") != active_key:
                st.session_state["_ado_active_portfolio_key"] = active_key
                st.session_state["ado_query_mode"] = current_settings.get("query_mode", st.session_state.get("ado_query_mode", "Advanced OData Builder"))
                for key, fallback in [
                    ("ado_full_url_input", current_settings.get("url", _ado_defaults["url"])),
                    ("ado_pat", _ado_defaults["pat"]),
                    ("ado_max_pages", int(current_settings.get("max_pages", 5))),
                    ("ado_org", current_settings.get("org", _ado_defaults["org"])),
                    ("ado_proj", current_settings.get("project", _ado_defaults["project"])),
                    ("ado_entity", current_settings.get("entity", _ado_defaults["entity"])),
                    ("ado_select", current_settings.get("select", _ado_defaults["select"])),
                    ("ado_expand", current_settings.get("expand", _ado_defaults["expand"])),
                    ("ado_filter", current_settings.get("filter", _ado_defaults["filter"])),
                    ("ado_features_only", current_settings.get("features_only", True)),
                ]:
                    st.session_state[key] = fallback
        
            with col_conn:
                with st.container(border=True):
                    st.markdown("#### Authentication (legacy)")
                    pat = st.text_input("Personal Access Token (PAT)", value=_ado_defaults["pat"], type="password", key="ado_pat")
                    store_pat = False
                    st.caption("PATs are not stored in the database. Use env/secrets or session paste.")
                    max_pages = st.number_input("Max pages to fetch", value=int(current_settings.get("max_pages") or 5), min_value=1, max_value=100, step=1, key="ado_max_pages")
                    st.info("Source mode: Cloud (PAT required)")

                with st.container(border=True):
                    st.markdown("#### Query mode")
                    query_mode = st.radio(
                        "OData query mode",
                        options=["Advanced OData Builder", "Full OData URL (Expert)"],
                        horizontal=True,
                        key="ado_query_mode",
                    )
                full_url = st.session_state.get("ado_full_url_input") or current_settings.get("url") or st.session_state.get("ado_prefill_full_url") or _ado_defaults["url"]
                full_url = str(full_url or "").strip()

                org = st.session_state.get("ado_org", current_settings.get("org", _ado_defaults["org"]))
                proj = st.session_state.get("ado_proj", current_settings.get("project", _ado_defaults["project"]))
                entity = st.session_state.get("ado_entity", current_settings.get("entity", _ado_defaults["entity"]))
                select = st.session_state.get("ado_select", current_settings.get("select", _ado_defaults["select"]))
                expand = st.session_state.get("ado_expand", current_settings.get("expand", _ado_defaults["expand"]))
                filter_ = st.session_state.get("ado_filter", current_settings.get("filter", _ado_defaults["filter"]))
                features_only = st.session_state.get("ado_features_only", current_settings.get("features_only", True))

                if query_mode == "Full OData URL (Expert)":
                    with st.container(border=True):
                        st.markdown("#### Full OData URL (Expert)")
                        full_url = st.text_input(
                            "Full OData URL",
                            value=full_url,
                            key="ado_full_url_input",
                            disabled=(query_mode == "Advanced OData Builder"),
                        ).strip()

            built_url = ""
            if query_mode == "Advanced OData Builder":
                with st.container(border=True):
                    st.markdown("#### Legacy OData Builder")
                    st.caption("Set org/project/select/expand/filter, then apply or run directly in Advanced Builder mode.")
                    colb1, colb2 = st.columns(2)
                    with colb1:
                        org = st.text_input("Organization", value=org, placeholder="your-org", key="ado_org").strip()
                        proj = st.text_input("Project", value=proj, placeholder="your-project", key="ado_proj").strip()
                        entity = (st.text_input("Entity", value=entity, key="ado_entity").strip() or "WorkItems")
                        features_only = st.checkbox("Only WorkItemType == 'Feature'", value=bool(features_only), key="ado_features_only")
                    with colb2:
                        select = st.text_area("$select", value=select, height=70, key="ado_select").strip()
                        expand = st.text_input("$expand", value=expand, key="ado_expand").strip()
                        filter_ = st.text_input("$filter (optional)", value=filter_, placeholder="WorkItemType eq 'Feature'", key="ado_filter").strip()

                    expand_fixed = _ensure_iteration_expand(expand)
                    if expand_fixed != expand:
                        st.warning("Builder $expand is missing required Iteration join fields; URL generation will add them for execution.")
                        if st.button("Fix $expand to include Iteration join keys", key="ado_fix_expand_keys_btn"):
                            st.session_state["ado_expand"] = expand_fixed
                            st.rerun()
                    filter_for_build = filter_ or ("WorkItemType eq 'Feature'" if features_only else None)
                    built_url = _build_ado_odata_url(org, proj, entity, select or None, expand_fixed or None, filter_for_build, None) if (org and proj) else ""
                    st.code(built_url or "(set org/project to build URL)")
                    if st.button("Apply built URL", key="ado_apply_built", disabled=not built_url):
                        st.session_state["ado_full_url_input"] = built_url
                        st.success("Copied builder URL into Full OData URL field.")

            with st.container(border=True):
                st.markdown("#### Save portfolio settings")
                if st.button("Save portfolio settings", icon=":material/save:", disabled=not portfolio_name, key="btn_save_portfolio"):
                    save_name = (portfolio_name or sel_portfolio or "(unsaved)").strip()
                    settings_payload = {
                        "query_mode": query_mode,
                        "url": full_url,
                        "max_pages": max_pages,
                        "org": org,
                        "project": proj,
                        "entity": entity,
                        "select": select,
                        "expand": _ensure_iteration_expand(expand),
                        "filter": filter_,
                        "features_only": features_only,
                        "sync_mode": st.session_state.get("ado_sync_mode", "Legacy (portfolio settings)"),
                        "store_pat": False,
                    }
                    user_email = st.session_state.get("auth_user", {}).get("email")
                    try:
                        save_ado_portfolio_settings(save_name, settings_payload, updated_by=user_email)
                    except Exception as e:
                        st.error(f"Failed to save portfolio settings: {e}")
                    else:
                        st.session_state["ado_portfolio_settings"][save_name] = settings_payload
                        toast_success("Portfolio settings saved.")
                        if sel_portfolio == "(unsaved)" and save_name:
                            st.session_state["ado_portfolio_select_pending"] = save_name
                            st.rerun()

            with st.expander("Danger zone", expanded=False):
                if sel_portfolio == "(unsaved)":
                    st.caption("Select a saved portfolio to delete or reset its stored settings.")
                else:
                    confirm = st.checkbox(
                        "I understand this will remove saved ADO settings for this portfolio",
                        value=False,
                        key="ado_delete_confirm",
                    )
                    if st.button(
                        "Delete portfolio ADO settings",
                        type="secondary",
                        disabled=not confirm,
                        key="ado_delete_portfolio_btn",
                    ):
                        try:
                            delete_ado_portfolio_settings(sel_portfolio)
                        except Exception as e:
                            st.error(f"Failed to delete portfolio settings: {e}")
                            st.stop()
                        st.session_state["ado_portfolio_settings"].pop(sel_portfolio, None)
                        for k in [
                            "ado_full_url_input",
                            "ado_pat",
                            "ado_store_pat",
                            "ado_max_pages",
                            "ado_org",
                            "ado_proj",
                            "ado_entity",
                            "ado_select",
                            "ado_expand",
                            "ado_filter",
                            "ado_features_only",
                            "ado_query_mode",
                            "ado_auto_fetched",
                            "ado_parsed_raw",
                            "ado_parsed_norm",
                        ]:
                            st.session_state.pop(k, None)
                        st.session_state["_ado_active_portfolio_key"] = None
                        st.session_state["ado_portfolio_select_pending"] = "(unsaved)"
                        st.success("Portfolio ADO settings deleted. Defaults will be used.")
                        st.rerun()

                    if st.button(
                        "Reset ADO settings to defaults (keep record)",
                        type="secondary",
                        key="ado_reset_portfolio_btn",
                    ):
                        reset_payload = {
                            "query_mode": "Advanced OData Builder",
                            "url": _ado_defaults["url"],
                            "max_pages": 5,
                            "org": _ado_defaults["org"],
                            "project": _ado_defaults["project"],
                            "entity": _ado_defaults["entity"],
                            "select": _ado_defaults["select"],
                            "expand": _ado_defaults["expand"],
                            "filter": _ado_defaults["filter"],
                            "features_only": True,
                            "store_pat": False,
                        }
                        user_email = st.session_state.get("auth_user", {}).get("email")
                        try:
                            save_ado_portfolio_settings(sel_portfolio, reset_payload, updated_by=user_email)
                        except Exception as e:
                            st.error(f"Failed to reset portfolio settings: {e}")
                            st.stop()
                        st.session_state["ado_portfolio_settings"][sel_portfolio] = reset_payload
                        st.session_state["_ado_active_portfolio_key"] = None
                        st.success("Portfolio ADO settings reset to defaults.")
                        st.rerun()

        legacy_pat = str(legacy_pat or "").strip()
        if not pat_configured and legacy_pat:
            st.warning("No PAT saved on profile. Falling back to legacy portfolio PAT (if configured).")
        pat = str(profile_pat or legacy_pat or (_ado_defaults.get("pat") or "")).strip()
    
    
        profile_url = ""
        profile_debug: Dict[str, Any] = {}
        progress_feature_type = "Feature"
        progress_epic_type = "Epic"
        progress_story_types_csv = "User Story,Product Backlog Item"
        progress_story_points_field = "StoryPoints"
        progress_exclude_states_cfg: List[str] = []
        progress_scope_filter = ""
        if sync_mode == "Profile-driven (ADO Profiles)":
            profile = get_active_ado_profile()
            base_url = str(profile.get("odata_base_url") or "").strip()
            if not base_url:
                st.error("Active profile is missing odata_base_url. Configure it in Profiles.")
            else:
                _select_fields_ignored, expand_clause = build_workitems_select_expand(profile)
                core_fields = [
                    "WorkItemId",
                    "WorkItemType",
                    "Title",
                    "State",
                    "ChangedDate",
                    "ParentWorkItemId",
                ]
                sp_field = str((profile.get("fields") or {}).get("story_points") or "").strip()
                app_field = str((profile.get("fields") or {}).get("app_name") or "").strip()
                inv_field = str((profile.get("fields") or {}).get("investment_dimension") or "").strip()
                bv_field = str((profile.get("fields") or {}).get("business_value") or "").strip()
                mapped_fields = [sp_field, app_field, inv_field, bv_field]
                select_fields = []
                seen = set()
                for field in core_fields + [f for f in mapped_fields if f]:
                    if field and field not in seen:
                        select_fields.append(field)
                        seen.add(field)
                feature_type = str((profile.get("work_item_types") or {}).get("feature") or "Feature").strip() or "Feature"
                epic_type = str((profile.get("work_item_types") or {}).get("epic") or "Epic").strip() or "Epic"
                exclude_states = list((profile.get("filters") or {}).get("exclude_states") or [])
                years_prefix = list((profile.get("filters") or {}).get("years_prefix") or [])
                pi_field = str((profile.get("iteration_mapping") or {}).get("pi_field") or "IterationLevel3").strip() or "IterationLevel3"
                progress_feature_type = feature_type
                progress_epic_type = epic_type
                progress_story_points_field = sp_field or "StoryPoints"
                progress_exclude_states_cfg = [str(s).strip() for s in exclude_states if str(s).strip()]
    
                filter_parts = [f"WorkItemType eq '{feature_type}'"]
                for state in exclude_states:
                    state_val = str(state or "").strip()
                    if state_val:
                        filter_parts.append(f"State ne '{state_val}'")
                if years_prefix:
                    ors = [f"startswith(Iteration/{pi_field},'{str(y).strip()}')" for y in years_prefix if str(y).strip()]
                    if ors:
                        filter_parts.append("(" + " or ".join(ors) + ")")
                progress_filter_parts: List[str] = []
                for state in exclude_states:
                    state_val = str(state or "").strip()
                    if state_val:
                        progress_filter_parts.append(f"State ne '{state_val}'")
                if years_prefix:
                    ors = [f"startswith(Iteration/{pi_field},'{str(y).strip()}')" for y in years_prefix if str(y).strip()]
                    if ors:
                        progress_filter_parts.append("(" + " or ".join(ors) + ")")
                progress_scope_filter = " and ".join(progress_filter_parts)
    
                filter_clause = " and ".join(filter_parts)
                select_str = ",".join(select_fields)
    
                base = base_url.rstrip("/")
                if not base.lower().endswith("/workitems"):
                    base = base + "/WorkItems"
                profile_url = base + "?" + "&".join(
                    [f"$select={select_str}", f"$expand={expand_clause}", f"$filter={filter_clause}"]
                )
                profile_debug = {
                    "base_url": base,
                    "select_fields": select_fields,
                    "expand_clause": expand_clause,
                    "filter_clause": filter_clause,
                }
    
        # Final URL used for fetch (strict mode precedence).
        final_url = built_url if query_mode == "Advanced OData Builder" else full_url
        with tab_advanced:
            with st.expander("Final OData URL sent to Azure DevOps", expanded=False):
                st.code(final_url or "(no URL)")
            if sync_mode == "Profile-driven (ADO Profiles)" and show_debug:
                with st.expander("Profile-driven URL (debug)", expanded=False):
                    st.code(profile_url or "(no URL)")
                    if profile_debug:
                        st.json(profile_debug)
    
    with tab_run:
        st.markdown("### Demand driver (Expected/Forecast)")
        try:
            _active_cfg = normalize_profile_config(get_active_ado_profile())
            _driver = str((_active_cfg.get("forecast") or {}).get("derived_fte_driver") or "SWAG").strip().upper()
            _fallback = bool((_active_cfg.get("forecast") or {}).get("velocity_row_fallback", True))
        except Exception:
            _driver = "SWAG"
            _fallback = True
        if _driver in {"VELOCITY", "SNAPSHOT", "SNAPSHOT_VELOCITY"}:
            st.info(
                "Active profile demand driver: Velocity snapshot (Derived FTE). "
                f"Row fallback to global baseline: {'ON' if _fallback else 'OFF'}."
            )
        else:
            st.info(
                "Active profile demand driver: SWAG (Derived FTE from Story Points). "
                "Features without SWAG contribute zero demand."
            )

        effective_url = profile_url if sync_mode == "Profile-driven (ADO Profiles)" else final_url
        effective_features_only = features_only if sync_mode == "Legacy (portfolio settings)" else False

        diag_api: Dict[str, Any] = {}
        pat_key = active_profile_ctx_key or str(active_profile_name or "")
        pat_env_key = get_active_pat_env_key()
        pat_val, pat_diag_run = resolve_ado_pat(
            pat_key,
            pat_env_key=pat_env_key,
            fallback_pat=_ado_defaults.get("pat"),
        )
        pat_source = str(pat_diag_run.get("source") or "missing")
        pat = pat_val or ""
        can_fetch = bool(effective_url) and bool(pat)
        if not pat:
            st.warning("No PAT configured. Provide ADO_PAT_<PROFILE> via env/secrets or paste a session PAT.")

        def _resolve_ado_org_project() -> Tuple[str, str]:
            o = (org or "").strip()
            p = (proj or "").strip()
            if o and p:
                return o, p
            try:
                parsed = urlparse(full_url or "")
                parts = [x for x in (parsed.path or "").strip("/").split("/") if x]
                if len(parts) >= 2:
                    return parts[0], parts[1]
            except Exception:
                pass
            return "", ""

        # Helper to execute the fetch + normalize and render diagnostics
        def _do_odata_fetch():
            try:
                url_to_use = effective_url
                if sync_mode == "Legacy (portfolio settings)" and query_mode == "Full OData URL (Expert)":
                    err = _validate_full_url(url_to_use)
                    if err:
                        st.error(err)
                        return False
                raw_df = fetch_ado_odata(url_to_use, pat if require_pat else "", diag_api, max_pages=int(max_pages))
                if raw_df is None or raw_df.empty:
                    st.warning("OData call returned no rows.")
                    return False
                if effective_features_only:
                    try:
                        cols = [c for c in raw_df.columns if c.lower().endswith("workitemtype") or c.lower() == "workitemtype"]
                        if cols:
                            c0 = cols[0]
                            raw_df = raw_df[raw_df[c0].astype(str).str.strip().str.lower() == "feature"].copy()
                    except Exception:
                        pass

                epic_map: Dict[str, Dict[str, Any]] = {}
                epic_summary: Dict[str, Any] = {}
                if sync_mode == "Profile-driven (ADO Profiles)":
                    try:
                        from scripts.ado_epic_resolver import (
                            fetch_workitems_by_ids,
                            fetch_workitems_by_filter,
                            resolve_epics_for_features,
                            LAST_EPIC_SUMMARY,
                        )

                        profile = get_active_ado_profile()
                        base_url = str(profile.get("odata_base_url") or "").strip()
                        if base_url:
                            select_fields = [
                                "WorkItemId",
                                "WorkItemType",
                                "Title",
                                "State",
                                "ParentWorkItemId",
                            ]
                            auth_ctx = {"pat": pat}

                            def _fetch_ids(ids: List[int]) -> List[Dict[str, Any]]:
                                return fetch_workitems_by_ids(base_url, ids, auth_ctx, select_fields)

                            _fetch_ids.fetch_by_filter = lambda filter_clause: fetch_workitems_by_filter(
                                base_url, filter_clause, auth_ctx, select_fields
                            )

                            epic_map = resolve_epics_for_features(raw_df.to_dict(orient="records"), profile, _fetch_ids)
                            epic_summary = dict(LAST_EPIC_SUMMARY)
                    except Exception:
                        epic_map = {}
                        epic_summary = {}

                df_expected = transform_ado_odata_to_expected(raw_df)
                profile_cfg_sync: Dict[str, Any] = {}
                profile_fields_sync: Dict[str, Any] = {}
                profile_field_resolution: Dict[str, Optional[str]] = {}
                if sync_mode == "Profile-driven (ADO Profiles)":
                    profile_cfg_sync = normalize_profile_config(get_active_ado_profile())
                    profile_fields_sync = (
                        profile_cfg_sync.get("fields")
                        if isinstance(profile_cfg_sync.get("fields"), dict)
                        else {}
                    )
                    profile_field_resolution = _apply_profile_field_mappings(df_expected, profile_fields_sync)
                norm_diag: Dict[str, Any] = {}
                df_norm = normalize_to_canonical(_auto_rename_columns(df_expected), diag=norm_diag)
                df_norm = repair_leaf_teams(df_norm)

                # Regression guard: prevent silently dropping key ADO fields during normalization.
                def _guard_field_not_dropped(
                    *,
                    label: str,
                    raw_aliases: List[str],
                    expected_col: str,
                    norm_col: str,
                ) -> None:
                    def _raw_nonnull_count() -> int:
                        raw_cols_by_lower = {str(c).lower(): c for c in raw_df.columns}
                        for a in raw_aliases:
                            col = a if a in raw_df.columns else raw_cols_by_lower.get(str(a).lower())
                            if col in raw_df.columns:
                                try:
                                    n = int(pd.to_numeric(raw_df[col], errors="coerce").notna().sum())
                                except Exception:
                                    n = int(raw_df[col].notna().sum())
                                if n > 0:
                                    return n
                        return 0

                    raw_n = _raw_nonnull_count()
                    if raw_n <= 0:
                        return
                    exp_n = int(pd.to_numeric(df_expected.get(expected_col), errors="coerce").notna().sum()) if expected_col in df_expected.columns else 0
                    norm_n = int(pd.to_numeric(df_norm.get(norm_col), errors="coerce").notna().sum()) if norm_col in df_norm.columns else 0
                    if exp_n <= 0:
                        raise RuntimeError(
                            f"Normalization bug: {label} exists in raw OData but was not mapped into expected columns. "
                            f"Check your $select includes the field and that the field name matches your Analytics schema."
                        )
                    if norm_n <= 0:
                        raise RuntimeError(
                            f"Normalization bug: {label} was dropped during normalization. Aborting sync to avoid data loss."
                        )

                bv_aliases = [
                    "BusinessValue",
                    "Fields.BusinessValue",
                    "Microsoft.VSTS.Common.BusinessValue",
                    "Fields.Microsoft.VSTS.Common.BusinessValue",
                    "Microsoft_VSTS_Common_BusinessValue",
                    "Fields.Microsoft_VSTS_Common_BusinessValue",
                ]
                if sync_mode == "Profile-driven (ADO Profiles)":
                    bv_field_cfg = str((profile_fields_sync or {}).get("business_value") or "").strip()
                    if bv_field_cfg:
                        bv_aliases.extend(_candidate_odata_field_names(bv_field_cfg))
                    bv_field_resolved = str((profile_field_resolution or {}).get("business_value") or "").strip()
                    if bv_field_resolved:
                        bv_aliases.append(bv_field_resolved)
                _guard_field_not_dropped(
                    label="BusinessValue",
                    raw_aliases=list(dict.fromkeys(bv_aliases)),
                    expected_col="BusinessValue",
                    norm_col="BUSINESS_VALUE",
                )
                if sync_mode == "Profile-driven (ADO Profiles)":
                    profile_cfg = profile_cfg_sync if profile_cfg_sync else normalize_profile_config(get_active_ado_profile())
                    sp_field = str((profile_cfg.get("fields") or {}).get("story_points") or "").strip()
                    sp_source_col = str((profile_field_resolution or {}).get("story_points") or "").strip()
                    if not sp_field:
                        st.warning("Profile story points field is not configured; STORY_POINTS set to null.")
                        df_norm["STORY_POINTS"] = pd.NA
                    elif not sp_source_col or sp_source_col not in df_expected.columns:
                        st.warning(
                            f"Story points field '{sp_field}' not present in response; STORY_POINTS set to null."
                        )
                        df_norm["STORY_POINTS"] = pd.NA
                    elif "ID" not in df_expected.columns or "FEATURE_ID" not in df_norm.columns:
                        st.warning("Story points mapping skipped due to missing ID columns.")
                    else:
                        sp_map_df = df_expected[["ID", sp_source_col]].copy()
                        sp_map_df["FEATURE_ID"] = sp_map_df["ID"].astype(str).str.strip()
                        sp_map_df = sp_map_df.dropna(subset=["FEATURE_ID"]).drop_duplicates(subset=["FEATURE_ID"])
                        sp_map = sp_map_df.set_index("FEATURE_ID")[sp_source_col]
                        df_norm["STORY_POINTS"] = pd.to_numeric(
                            df_norm["FEATURE_ID"].astype(str).map(sp_map),
                            errors="coerce",
                        )
                    area_cfg = profile_cfg.get("area_mapping") or {}
                    program_key = str(area_cfg.get("program_level") or "").strip()
                    team_key = str(area_cfg.get("team_level") or "").strip()

                    def _clean_series(series: Any) -> pd.Series:
                        if series is None:
                            return pd.Series([None] * len(df_norm), index=df_norm.index, dtype="object")
                        cleaned = pd.Series(series).astype(str).str.strip()
                        cleaned = cleaned.replace({"": None, "nan": None, "None": None})
                        return cleaned.reindex(df_norm.index, fill_value=None)

                    if "AreaLevel1" in df_expected.columns:
                        df_norm["AREA_LEVEL1_RAW"] = _clean_series(df_expected["AreaLevel1"])
                    if "AreaLevel2" in df_expected.columns:
                        df_norm["AREA_LEVEL2_RAW"] = _clean_series(df_expected["AreaLevel2"])
                    if "AreaLevel3" in df_expected.columns:
                        df_norm["AREA_LEVEL3_RAW"] = _clean_series(df_expected["AreaLevel3"])
                    if "AreaLevel4" in df_expected.columns:
                        df_norm["AREA_LEVEL4_RAW"] = _clean_series(df_expected["AreaLevel4"])
                    if "AreaPath" in df_expected.columns:
                        df_norm["AREA_PATH_RAW"] = _clean_series(df_expected["AreaPath"])
                    if program_key and program_key in df_expected.columns:
                        df_norm["PROGRAM_RAW"] = _clean_series(df_expected[program_key])
                    if team_key and team_key in df_expected.columns:
                        df_norm["TEAM_RAW"] = _clean_series(df_expected[team_key])
                st.session_state["ado_parsed_raw"] = df_expected
                st.session_state["ado_parsed_norm"] = df_norm

                if epic_map:
                    def _norm_id(val: Any) -> Optional[str]:
                        try:
                            if val is None or (isinstance(val, float) and pd.isna(val)):
                                return None
                            return str(int(float(val)))
                        except Exception:
                            txt = str(val).strip()
                            return txt if txt else None

                    story_points_before = None
                    if "STORY_POINTS" in df_norm.columns:
                        try:
                            story_points_before = int(pd.to_numeric(df_norm["STORY_POINTS"], errors="coerce").notna().sum())
                        except Exception:
                            story_points_before = int(df_norm["STORY_POINTS"].notna().sum())

                    if "ParentWorkItemId" in df_expected.columns and "ID" in df_expected.columns:
                        parent_lookup = df_expected[["ID", "ParentWorkItemId"]].copy()
                        parent_lookup["FEATURE_ID"] = parent_lookup["ID"].astype(str).str.strip()
                        parent_lookup["PARENT_ID"] = parent_lookup["ParentWorkItemId"].map(_norm_id)
                        parent_lookup = parent_lookup.drop_duplicates(subset=["FEATURE_ID"])
                        parent_map = parent_lookup.set_index("FEATURE_ID")["PARENT_ID"]
                        df_norm["PARENT_ID"] = df_norm["FEATURE_ID"].astype(str).map(parent_map)
                    else:
                        df_norm["PARENT_ID"] = df_norm.get("PARENT_ID")

                    epic_rows: List[Dict[str, Any]] = []
                    for row in epic_map.values():
                        epic_id = _norm_id(row.get("epic_id"))
                        if not epic_id:
                            continue
                        epic_rows.append(
                            {
                                "EPIC_ID": epic_id,
                                "EPIC_TITLE": row.get("epic_title"),
                                "EPIC_STATE": row.get("epic_state"),
                            }
                        )
                    if epic_rows:
                        epic_lookup = pd.DataFrame(epic_rows)
                        epic_lookup["EPIC_ID"] = epic_lookup["EPIC_ID"].map(_norm_id)
                        epic_lookup = epic_lookup.dropna(subset=["EPIC_ID"]).drop_duplicates(subset=["EPIC_ID"])
                        epic_lookup = epic_lookup[["EPIC_ID", "EPIC_TITLE", "EPIC_STATE"]]
                        df_norm = df_norm.merge(
                            epic_lookup,
                            how="left",
                            left_on="PARENT_ID",
                            right_on="EPIC_ID",
                        )

                    story_points_after = None
                    if "STORY_POINTS" in df_norm.columns:
                        try:
                            story_points_after = int(pd.to_numeric(df_norm["STORY_POINTS"], errors="coerce").notna().sum())
                        except Exception:
                            story_points_after = int(df_norm["STORY_POINTS"].notna().sum())
                    if (
                        story_points_before is not None
                        and story_points_after is not None
                        and story_points_after < story_points_before
                    ):
                        st.warning("Epic enrichment reduced populated Story Points values; check merge collisions.")

                    suffix_cols = [c for c in df_norm.columns if c in {"STORY_POINTS_x", "STORY_POINTS_y"}]
                    if suffix_cols:
                        st.warning("Epic enrichment produced STORY_POINTS_x/y columns; using feature-side values.")
                        if "STORY_POINTS_x" in df_norm.columns:
                            df_norm["STORY_POINTS"] = df_norm["STORY_POINTS_x"]
                        df_norm = df_norm.drop(columns=[c for c in suffix_cols if c in df_norm.columns], errors="ignore")

                for col in ["FEATURE_ID", "PARENT_ID", "EPIC_ID"]:
                    if col in df_norm.columns:
                        df_norm[col] = pd.to_numeric(df_norm[col], errors="coerce").astype("Int64")
                st.session_state["ado_parsed_norm"] = df_norm
                st.success(f"Fetched {len(df_expected)} rows via OData; normalized {len(df_norm)} rows.")
                try:
                    _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=False)
                except Exception:
                    pass
                st.session_state["ado_data_cache_bust"] = st.session_state.get("ado_data_cache_bust", 0) + 1
                st.session_state["ado_last_sync_ts"] = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
                if active_profile_id:
                    try:
                        set_profile_last_sync_at(active_profile_id)
                    except Exception:
                        pass

                def _pct_present(series: pd.Series) -> float:
                    if series is None or series.empty:
                        return 0.0
                    return float(pd.to_numeric(series, errors="coerce").notna().mean() * 100.0)

                d_total = int(len(df_norm))
                bv_pct = _pct_present(df_norm.get("BUSINESS_VALUE", pd.Series(dtype=float)))
                c1, c2 = st.columns(2)
                c1.metric("Rows", f"{d_total:,d}")
                c2.metric("% with Business Value", f"{bv_pct:.1f}%")
                mode = diag_api.get("mode", "?")
                pages = diag_api.get("pages", [])
                st.caption(f"Debug: mode={mode}; pages={len(pages)}")
                with st.expander("Preview (first 200 rows)", expanded=False):
                    st.dataframe(df_expected.head(200), use_container_width=True, height=320)
                with st.expander("Run details", expanded=False):
                    st.json({**diag_api, **norm_diag})
                    sql_debug = get_sql_trace_debug()
                    if sql_debug:
                        st.caption("SQL trace (this rerun)")
                        st.json(sql_debug)
                st.session_state["ado_auto_fetched"] = True
                if sync_mode == "Profile-driven (ADO Profiles)":
                    total = int(len(df_norm)) if df_norm is not None else 0
                    prog_ct = int(df_norm["PROGRAM_RAW"].dropna().nunique()) if "PROGRAM_RAW" in df_norm.columns else 0
                    team_ct = int(df_norm["TEAM_RAW"].dropna().nunique()) if "TEAM_RAW" in df_norm.columns else 0
                    iter_null = int(df_norm["ITERATION_LEVEL3"].isna().sum()) if "ITERATION_LEVEL3" in df_norm.columns else 0
                    area_null = int(df_norm["AREA_LEVEL3_RAW"].isna().sum()) if "AREA_LEVEL3_RAW" in df_norm.columns else 0
                    summary = {
                        "total_features": total,
                        "distinct_programs": prog_ct,
                        "distinct_teams": team_ct,
                        "null_iteration_count": iter_null,
                        "null_area_count": area_null,
                    }
                    with st.expander("Profile-driven sync summary", expanded=False):
                        st.json(summary)
                    if epic_summary:
                        with st.expander("Epic resolution summary", expanded=False):
                            st.json(epic_summary)
                return True
            except Exception as e:
                with st.expander("Run details", expanded=True):
                    st.exception(e)
                    st.json(diag_api)
                st.error("OData fetch failed. Verify URL/PAT and that Analytics is enabled.")
                return False

        def _run_upsert_from_parsed_norm() -> Dict[str, int]:
            df_norm = st.session_state.get("ado_parsed_norm")
            if df_norm is None or (hasattr(df_norm, "empty") and df_norm.empty):
                st.error("No parsed ADO dataset available. Run sync first.")
                return {"rows_upserted": 0, "inserted": 0, "updated": 0, "total_rows": 0}

            # capture pre-insert diff sets for reporting
            try:
                existing = fetch_df("SELECT FEATURE_ID FROM ADO_FEATURES")
                existing_ids = set(existing["FEATURE_ID"].astype(str)) if existing is not None and not existing.empty else set()
            except Exception:
                existing_ids = set()
            upload_ids = set(df_norm["FEATURE_ID"].astype(str))
            to_insert_ids = sorted(list(upload_ids - existing_ids))
            to_update_ids = sorted(list(upload_ids & existing_ids))

            # Full sync always includes iteration calendar refresh.
            sync_iter_cal = True
            if sync_iter_cal:
                org_res = ""
                proj_res = ""
                pat_res = ""
                iter_year_prefixes = ["2025", "2026", "2027"]
                if sync_mode == "Profile-driven (ADO Profiles)":
                    profile_cfg = get_active_ado_profile()
                    org_res, proj_res = _parse_org_project_from_odata_base(
                        str(profile_cfg.get("odata_base_url") or "")
                    )
                    pat_res = str(profile_pat or legacy_pat or _ado_defaults["pat"] or "").strip()
                    try:
                        cfg_years = list((profile_cfg.get("filters") or {}).get("years_prefix") or [])
                    except Exception:
                        cfg_years = []
                    cfg_years = [str(y).strip() for y in cfg_years if str(y).strip()]
                    if cfg_years:
                        iter_year_prefixes = cfg_years
                else:
                    org_res, proj_res = _resolve_ado_org_project()
                    pat_res = str(legacy_pat or pat or _ado_defaults["pat"] or "").strip()
                if not (org_res and proj_res and pat_res):
                    st.error("Iteration calendar sync requires Organization, Project, and PAT.")
                    org_res = ""
                    proj_res = ""
                    pat_res = ""
                diag_iter: Dict[str, Any] = {}
                if org_res and proj_res and pat_res:
                    with st.spinner("Syncing Iteration Calendar (ADO_ITERATION_CALENDAR)…"):
                        try:
                            df_iters = fetch_ado_iterations_odata(
                                org_res,
                                proj_res,
                                pat_res,
                                diag_iter,
                                year_prefixes=iter_year_prefixes,
                                include_sprints=True,
                                max_pages=int(max_pages or 20),
                            )
                            n_iters = upsert_ado_iteration_calendar(df_iters)
                        except Exception as e:
                            with st.expander("Iteration sync diagnostics", expanded=True):
                                st.exception(e)
                                st.json(diag_iter)
                            st.error("Iteration calendar sync failed; continuing feature upsert.")
                        else:
                            try:
                                df_ic_meta2 = fetch_df("SELECT COUNT(*) AS N, MAX(UPDATED_AT) AS LAST_UPDATED FROM ADO_ITERATION_CALENDAR")
                                ic_n2 = int(df_ic_meta2.iloc[0]["N"]) if df_ic_meta2 is not None and not df_ic_meta2.empty else 0
                                ic_last2 = df_ic_meta2.iloc[0].get("LAST_UPDATED") if df_ic_meta2 is not None and not df_ic_meta2.empty else None
                            except Exception:
                                ic_n2, ic_last2 = None, None
                            st.success(f"Iterations synced: **{int(n_iters)}**")
                            st.caption(f"Iteration calendar last updated: {ic_last2 if ic_last2 is not None else '—'} (rows: {ic_n2 if ic_n2 is not None else '—'})")

            total = len(df_norm)
            chunk_size = 1000 if total > 5000 else (500 if total > 2000 else 250)
            prog = st.progress(0, text="Starting upsert…")
            status = st.empty()
            n = 0
            for start in range(0, total, chunk_size):
                stop = min(start + chunk_size, total)
                chunk = df_norm.iloc[start:stop]
                status.write(f"Upserting rows {start+1:,} to {stop:,} of {total:,}…")
                try:
                    n += upsert_ado_features(chunk)
                except Exception as e:
                    status.error(f"Chunk {start+1:,}-{stop:,} failed: {e}")
                    raise
                prog.progress(stop / total, text=f"Upsert progress: {int((stop/total)*100)}%")
            prog.empty()
            status.write("Upsert complete.")
            info2 = fetch_df("SELECT COUNT(*) AS N FROM ADO_FEATURES")
            n_rows2 = int(info2.iloc[0]["N"]) if info2 is not None and not info2.empty else 0
            st.success(
                f"Upserted **{n}** rows into ADO_FEATURES. "
                f"Inserted: **{len(to_insert_ids)}**, Updated: **{len(to_update_ids)}**. "
                f"New total rows: **{n_rows2}**. "
                f"Iteration join keys (IterationPath/IterationSK) updated."
            )
            with st.expander("Show inserted IDs", expanded=False):
                st.write(to_insert_ids[:500])
            with st.expander("Show updated IDs", expanded=False):
                st.write(to_update_ids[:500])
            return {
                "rows_upserted": int(n),
                "inserted": int(len(to_insert_ids)),
                "updated": int(len(to_update_ids)),
                "total_rows": int(n_rows2),
            }

        st.markdown("### Progress Add-on")
        st.caption(
            "Optional sync for Feature/Epic progress only. "
            "This is additive and does not modify ADO_FEATURES. "
            "Removed/Cut stories are excluded from progress %."
        )

        def _resolve_progress_org_project() -> Tuple[str, str]:
            if sync_mode == "Profile-driven (ADO Profiles)":
                profile_cfg_local = get_active_ado_profile()
                o, p = _parse_org_project_from_odata_base(
                    str(profile_cfg_local.get("odata_base_url") or "")
                )
                return str(o or "").strip(), str(p or "").strip()
            o = str(org or "").strip()
            p = str(proj or "").strip()
            if o and p:
                return o, p
            try:
                parsed = urlparse(full_url or "")
                parts = [x for x in (parsed.path or "").strip("/").split("/") if x]
                if len(parts) >= 2:
                    return str(parts[0]).strip(), str(parts[1]).strip()
            except Exception:
                pass
            return "", ""

        def _extract_filter_from_url(url_val: str) -> str:
            try:
                parsed = urlparse(str(url_val or "").strip())
                q = parse_qs(parsed.query or "", keep_blank_values=False)
                vals = q.get("$filter") or q.get("%24filter") or []
                if vals:
                    return str(vals[0] or "").strip()
            except Exception:
                pass
            return ""

        progress_org, progress_project = _resolve_progress_org_project()
        progress_extra_filter = ""
        timeline_extra_filter = ""
        if sync_mode == "Profile-driven (ADO Profiles)":
            progress_extra_filter = str(progress_scope_filter or "").strip()
            timeline_extra_filter = ""
        else:
            progress_extra_filter = _extract_filter_from_url(str(full_url or built_url or ""))
            timeline_extra_filter = progress_extra_filter
        can_sync_progress = bool(pat) and bool(progress_org) and bool(progress_project)
        can_sync_timeline = bool(pat) and bool(progress_org) and bool(progress_project)
        if not can_sync_progress:
            st.caption("Progress sync requires Organization, Project, and PAT.")

        def _run_progress_sync_action() -> bool:
            try:
                from scripts.ado_sync_progress import main as run_ado_progress_sync
                import inspect

                stage_box = st.empty()
                stage_progress = st.progress(0, text="Progress sync: starting...")

                def _on_progress_sync(msg: str) -> None:
                    m = str(msg or "").strip()
                    m_low = m.lower()
                    pct = 0.12
                    if "resolving profile" in m_low:
                        pct = 0.12
                    elif "querying workitems" in m_low:
                        pct = 0.35
                    elif "preparing lookup/progress tables" in m_low:
                        pct = 0.55
                    elif "upserting lookup rows" in m_low:
                        pct = 0.68
                    elif "resolving parent chain" in m_low:
                        pct = 0.76
                    elif "rolling up story points" in m_low:
                        pct = 0.86
                    elif "upserting progress rows" in m_low:
                        pct = 0.94
                    elif "progress sync complete" in m_low:
                        pct = 1.0
                    try:
                        stage_progress.progress(float(max(0.0, min(1.0, pct))), text=m)
                    except Exception:
                        pass
                    try:
                        stage_box.info(m)
                    except Exception:
                        pass

                _on_progress_sync("Progress sync: validating filters + scope...")
                _progress_args = [
                    "--org",
                    progress_org,
                    "--project",
                    progress_project,
                    "--pat",
                    str(pat),
                    "--entity",
                    str(entity or "WorkItems"),
                    "--max-pages",
                    str(int(max_pages or 30)),
                    "--extra-filter",
                    str(progress_extra_filter or ""),
                    "--feature-type",
                    str(progress_feature_type or "Feature"),
                    "--epic-type",
                    str(progress_epic_type or "Epic"),
                    "--story-types",
                    str(progress_story_types_csv or "User Story,Product Backlog Item"),
                    "--story-points-field",
                    str(progress_story_points_field or "StoryPoints"),
                    "--exclude-states",
                    ",".join([str(s).strip() for s in (progress_exclude_states_cfg or []) if str(s).strip()]),
                ]
                _progress_sig = inspect.signature(run_ado_progress_sync)
                if "progress_cb" in _progress_sig.parameters:
                    rc = int(run_ado_progress_sync(_progress_args, progress_cb=_on_progress_sync) or 0)
                else:
                    rc = int(run_ado_progress_sync(_progress_args) or 0)
                if rc == 0:
                    try:
                        stage_progress.progress(1.0, text="Progress sync complete.")
                    except Exception:
                        pass
                    st.success("Progress sync completed successfully.")
                    try:
                        _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                    except Exception:
                        pass
                    st.session_state["ado_data_cache_bust"] = st.session_state.get("ado_data_cache_bust", 0) + 1
                    return True
                st.error(f"Progress sync finished with non-zero status: {rc}")
                return False
            except Exception as e:
                st.error(f"Progress sync failed: {e}")
                return False

        def _run_timeline_sync_action() -> bool:
            try:
                from scripts.ado_sync_links_and_dates import main as run_ado_timeline_sync
                import inspect

                stage_box = st.empty()
                stage_progress = st.progress(0, text="Timeline sync: starting...")

                def _on_timeline_progress(msg: str) -> None:
                    m = str(msg or "").strip()
                    m_low = m.lower()
                    pct = 0.15
                    if "resolving ado profile" in m_low:
                        pct = 0.15
                    elif "querying workitems" in m_low:
                        pct = 0.40
                    elif "upserting workitem dates" in m_low:
                        pct = 0.65
                    elif "querying dependency" in m_low:
                        pct = 0.80
                    elif "upserting links" in m_low:
                        pct = 0.92
                    elif "timeline sync complete" in m_low:
                        pct = 1.0
                    try:
                        stage_progress.progress(float(max(0.0, min(1.0, pct))), text=m)
                    except Exception:
                        pass
                    try:
                        stage_box.info(m)
                    except Exception:
                        pass

                _on_timeline_progress("Timeline sync: validating strict profile year scope...")
                _timeline_args = [
                    "--org",
                    progress_org,
                    "--project",
                    progress_project,
                    "--pat",
                    str(pat),
                    "--entity",
                    str(entity or "WorkItems"),
                    "--max-pages",
                    str(int(max_pages or 30)),
                    "--extra-filter",
                    str(timeline_extra_filter or ""),
                ]
                _timeline_sig = inspect.signature(run_ado_timeline_sync)
                if "progress_cb" in _timeline_sig.parameters:
                    rc = int(run_ado_timeline_sync(_timeline_args, progress_cb=_on_timeline_progress) or 0)
                else:
                    rc = int(run_ado_timeline_sync(_timeline_args) or 0)
                if rc == 0:
                    try:
                        stage_progress.progress(1.0, text="Timeline sync complete.")
                    except Exception:
                        pass
                    st.success("Timeline sync completed successfully.")
                    try:
                        _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                    except Exception:
                        pass
                    st.session_state["ado_data_cache_bust"] = st.session_state.get("ado_data_cache_bust", 0) + 1
                    return True
                st.error(f"Timeline sync finished with non-zero status: {rc}")
                return False
            except Exception as e:
                st.error(f"Timeline sync failed: {e}")
                return False

        st.caption("Use this one-click flow to run fetch + upsert + progress in sequence.")
        st.caption(
            "Timeline sync runs in strict year scope from ADO profile `filters.years_prefix` "
            "(no broad all-years fallback)."
        )
        if st.button(
            "Run Full Sync (Features + Iterations + Progress + Timeline)",
            icon=":material/playlist_add_check:",
            disabled=not (can_fetch and can_sync_progress and can_sync_timeline),
            key="btn_full_sync_ado_plus_progress",
            type="primary",
        ):
            if _preview_blocked_notice("Full ADO sync"):
                st.stop()
            flow_status = st.empty()
            flow_progress = st.progress(0, text="Starting full sync…")
            ok_fetch = _do_odata_fetch()
            if not ok_fetch:
                flow_status.error("Full sync stopped at Step 1 (fetch).")
                flow_progress.empty()
            else:
                flow_status.info("Step 2/4: Upserting ADO_FEATURES + Iterations…")
                flow_progress.progress(0.25, text="Step 2/4: Upserting features + iterations…")
                upsert_summary = _run_upsert_from_parsed_norm()
                flow_status.info("Step 3/4: Syncing progress overlay…")
                flow_progress.progress(0.50, text="Step 3/4: Syncing progress…")
                ok_progress = _run_progress_sync_action()
                flow_status.info("Step 4/4: Syncing timeline dates + links…")
                flow_progress.progress(0.75, text="Step 4/4: Syncing timeline…")
                ok_timeline = _run_timeline_sync_action()
                if ok_progress and ok_timeline:
                    flow_progress.progress(1.0, text="Full sync complete.")
                    flow_status.success(
                        "Full sync completed. "
                        f"Upserted {int(upsert_summary.get('rows_upserted') or 0):,} feature rows; "
                        f"iteration calendar + progress overlay + timeline datasets refreshed."
                    )
                elif ok_progress and not ok_timeline:
                    flow_status.warning("Features/progress sync completed, but timeline sync failed.")
                elif not ok_progress and ok_timeline:
                    flow_status.warning("Features/timeline sync completed, but progress sync failed.")
                else:
                    flow_status.warning("Features upsert completed, but both add-on syncs failed.")
                try:
                    flow_progress.empty()
                except Exception:
                    pass

        if st.button(
            "Sync Timeline Add-on (Dates + Links)",
            key="btn_sync_timeline_addon",
            disabled=not can_sync_timeline,
            use_container_width=False,
        ):
            _run_timeline_sync_action()

        try:
            df_fp = fetch_df("SELECT COUNT(*) AS N FROM ADO_FEATURE_PROGRESS")
            df_ep = fetch_df("SELECT COUNT(*) AS N FROM ADO_EPIC_PROGRESS")
            fp_n = int(df_fp.iloc[0]["N"]) if df_fp is not None and not df_fp.empty else 0
            ep_n = int(df_ep.iloc[0]["N"]) if df_ep is not None and not df_ep.empty else 0
            c_pg1, c_pg2 = st.columns(2)
            c_pg1.metric("Feature progress rows", fp_n)
            c_pg2.metric("Epic progress rows", ep_n)
            wd_n = 0
            wl_n = 0
            try:
                df_wd = fetch_df("SELECT COUNT(*) AS N FROM ADO_WORKITEM_DATES")
                df_wl = fetch_df("SELECT COUNT(*) AS N FROM ADO_WORKITEM_LINKS")
                wd_n = int(df_wd.iloc[0]["N"]) if df_wd is not None and not df_wd.empty else 0
                wl_n = int(df_wl.iloc[0]["N"]) if df_wl is not None and not df_wl.empty else 0
                c_tl1, c_tl2 = st.columns(2)
                c_tl1.metric("Timeline date rows", wd_n)
                c_tl2.metric("Timeline link rows", wl_n)
            except Exception:
                pass

            # Keep sprint readiness merged with the existing progress/timeline diagnostics cards.
            pi_rows = 0
            sprint_rows = 0
            iter_total_rows = 0
            iter_last_updated = None
            try:
                df_iter_readiness = fetch_df(
                    """
                    SELECT
                      COUNT(*) AS TOTAL_ROWS,
                      SUM(CASE WHEN UPPER(COALESCE(ITERATION_GRAIN, 'PI')) = 'PI' THEN 1 ELSE 0 END) AS PI_ROWS,
                      SUM(CASE WHEN UPPER(COALESCE(ITERATION_GRAIN, 'PI')) = 'SPRINT' THEN 1 ELSE 0 END) AS SPRINT_ROWS,
                      MAX(UPDATED_AT) AS LAST_UPDATED
                    FROM ADO_ITERATION_CALENDAR
                    """
                )
                if isinstance(df_iter_readiness, pd.DataFrame) and not df_iter_readiness.empty:
                    r = df_iter_readiness.iloc[0]
                    iter_total_rows = int(pd.to_numeric(r.get("TOTAL_ROWS"), errors="coerce") or 0)
                    pi_rows = int(pd.to_numeric(r.get("PI_ROWS"), errors="coerce") or 0)
                    sprint_rows = int(pd.to_numeric(r.get("SPRINT_ROWS"), errors="coerce") or 0)
                    iter_last_updated = r.get("LAST_UPDATED")
            except Exception:
                pass

            st.markdown("##### Sync readiness")
            r1, r2, r3, r4 = st.columns(4)
            r1.metric("Iteration rows", f"{iter_total_rows:,}")
            r1.caption(f"PI: {pi_rows:,} • Sprint: {sprint_rows:,}")
            r2.metric("Progress overlay", "Ready" if (fp_n > 0 or ep_n > 0) else "Missing")
            r2.caption(f"Feature: {fp_n:,} • Epic: {ep_n:,}")
            r3.metric("Timeline dates", "Ready" if wd_n > 0 else "Missing")
            r3.caption(f"Rows: {wd_n:,}")
            r4.metric("Dependencies", "Ready" if wl_n > 0 else "Optional")
            r4.caption(f"Rows: {wl_n:,}")

            progress_ready = (fp_n > 0 or ep_n > 0)
            timeline_ready = wd_n > 0
            sprint_ready = sprint_rows > 0
            if progress_ready and timeline_ready and sprint_ready:
                st.success("Overall roadmap readiness: Ready (PI + Sprint + Progress + Timeline).")
            elif progress_ready and timeline_ready:
                st.info("Overall roadmap readiness: Partial (PI + Progress + Timeline). Sprint view unavailable.")
            else:
                st.warning("Overall roadmap readiness: Incomplete. Run Full sync and verify diagnostics.")

            if iter_last_updated is not None:
                st.caption(f"Iteration calendar last updated: {iter_last_updated}")
            if sprint_rows <= 0:
                st.warning(
                    "Sprint rows were not found in ADO_ITERATION_CALENDAR. "
                    "Timeline can still run on PI bands, but Sprint mode will be unavailable."
                )

            df_overlap = fetch_df(
                """
                SELECT
                    (SELECT COUNT(*) FROM ADO_FEATURES) AS FEATURE_ROWS,
                    (SELECT COUNT(*) FROM ADO_FEATURE_PROGRESS) AS PROGRESS_ROWS,
                    (
                        SELECT COUNT(*)
                        FROM ADO_FEATURES af
                        JOIN ADO_FEATURE_PROGRESS fp
                          ON TRY_CONVERT(BIGINT, af.FEATURE_ID) = TRY_CONVERT(BIGINT, fp.FEATURE_ID)
                    ) AS JOINED_ROWS,
                    (
                        SELECT COUNT(*)
                        FROM ADO_FEATURES af
                        JOIN ADO_FEATURE_PROGRESS fp
                          ON TRY_CONVERT(BIGINT, af.FEATURE_ID) = TRY_CONVERT(BIGINT, fp.FEATURE_ID)
                        WHERE COALESCE(TRY_CONVERT(FLOAT, fp.TOTAL_SP), 0.0) > 0
                    ) AS JOINED_WITH_SP_ROWS
                """
            )
            if df_overlap is not None and not df_overlap.empty:
                orow = df_overlap.iloc[0]
                joined_rows = int(orow.get("JOINED_ROWS") or 0)
                joined_with_sp = int(orow.get("JOINED_WITH_SP_ROWS") or 0)
                c_ov1, c_ov2 = st.columns(2)
                c_ov1.metric("Progress joined to features", joined_rows)
                c_ov2.metric("Joined rows with SP > 0", joined_with_sp)
                if int(orow.get("PROGRESS_ROWS") or 0) > 0 and joined_rows == 0:
                    st.warning(
                        "Progress rows exist but do not match local ADO_FEATURES IDs. "
                        "Import ADO_FEATURES from the same environment snapshot (or run local ADO sync) "
                        "so Roadmap can display progress."
                    )

            with st.expander("Progress join diagnostics", expanded=False):
                try:
                    df_feature_sample = fetch_df(
                        "SELECT TOP 10 FEATURE_ID FROM ADO_FEATURES ORDER BY TRY_CONVERT(BIGINT, FEATURE_ID)"
                    )
                    df_progress_sample = fetch_df(
                        "SELECT TOP 10 FEATURE_ID, TOTAL_SP, PCT_COMPLETE FROM ADO_FEATURE_PROGRESS ORDER BY TRY_CONVERT(BIGINT, FEATURE_ID)"
                    )
                    df_unmatched = fetch_df(
                        """
                        SELECT TOP 20 fp.FEATURE_ID
                        FROM ADO_FEATURE_PROGRESS fp
                        WHERE NOT EXISTS (
                            SELECT 1
                            FROM ADO_FEATURES af
                            WHERE TRY_CONVERT(BIGINT, af.FEATURE_ID) = TRY_CONVERT(BIGINT, fp.FEATURE_ID)
                        )
                        ORDER BY TRY_CONVERT(BIGINT, fp.FEATURE_ID)
                        """
                    )
                    st.caption("Sample local ADO_FEATURES IDs")
                    st.dataframe(df_feature_sample if df_feature_sample is not None else pd.DataFrame(), use_container_width=True)
                    st.caption("Sample progress IDs")
                    st.dataframe(df_progress_sample if df_progress_sample is not None else pd.DataFrame(), use_container_width=True)
                    st.caption("Progress IDs not found in local ADO_FEATURES")
                    st.dataframe(df_unmatched if df_unmatched is not None else pd.DataFrame(), use_container_width=True)
                except Exception as e_diag:
                    st.info(f"Progress diagnostics unavailable: {e_diag}")

            refresh_note = st.session_state.get("ado_progress_refresh_note")
            if isinstance(refresh_note, dict):
                done_at = str(refresh_note.get("done_at") or "").strip()
                if done_at:
                    st.success(f"Progress overlay refresh completed at {done_at} UTC.")
                else:
                    st.success("Progress overlay refresh completed.")
                if st.button("Dismiss refresh message", key="ado_progress_refresh_dismiss"):
                    st.session_state.pop("ado_progress_refresh_note", None)
                    st.rerun()

            if st.button("Offline: Refresh progress overlay", key="ado_progress_offline_refresh"):
                st.session_state["ado_data_cache_bust"] = st.session_state.get("ado_data_cache_bust", 0) + 1
                try:
                    _settings_post_write_refresh("ado_progress_refresh", bump_version=False)
                except Exception:
                    pass
                st.session_state["ado_progress_refresh_note"] = {
                    "done_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"),
                }
                st.success("Progress caches refreshed. Open Roadmap to validate progress values.")
        except Exception:
            pass
        
        st.markdown("### Iteration Calendar")
        st.caption("Iteration calendar sync is included automatically in Full Sync.")
        try:
            ensure_ado_iteration_calendar_table()
            df_ic_meta = fetch_df("SELECT COUNT(*) AS N, MAX(UPDATED_AT) AS LAST_UPDATED FROM ADO_ITERATION_CALENDAR")
            ic_n = int(df_ic_meta.iloc[0]["N"]) if df_ic_meta is not None and not df_ic_meta.empty else 0
            ic_last = df_ic_meta.iloc[0].get("LAST_UPDATED") if df_ic_meta is not None and not df_ic_meta.empty else None
        except Exception:
            ic_n, ic_last = 0, None
        c_it1, c_it2 = st.columns([1, 2])
        c_it1.metric("Iterations synced", ic_n)
        c_it2.metric("Iteration calendar last updated", (str(ic_last) if ic_last is not None else "—"))
    
        with st.expander("Iteration mapping diagnostics", expanded=False):
            try:
                ensure_ado_iteration_calendar_table()
                df_breakdown = fetch_df(
                    """
                    WITH x AS (
                      SELECT
                        af.FEATURE_ID,
                        af.ITERATION_PATH AS FEATURE_ITERATION_PATH_RAW,
                        CASE
                          WHEN af.ITERATION_PATH IS NULL OR LTRIM(RTRIM(af.ITERATION_PATH)) = '' THEN 0
                          WHEN PATINDEX('%\\I[0-9] S[0-9]%', UPPER(af.ITERATION_PATH)) > 0 THEN 1
                          WHEN PATINDEX('%\\S[0-9]%', UPPER(af.ITERATION_PATH)) > 0 THEN 1
                          ELSE 0
                        END AS FEATURE_IS_SPRINT_LEVEL,
                        CASE
                          WHEN af.ITERATION_PATH IS NULL OR LTRIM(RTRIM(af.ITERATION_PATH)) = '' THEN NULL
                          WHEN (
                            PATINDEX('%\\I[0-9] S[0-9]%', UPPER(af.ITERATION_PATH)) > 0
                            OR PATINDEX('%\\S[0-9]%', UPPER(af.ITERATION_PATH)) > 0
                          )
                          AND CHARINDEX('\\', af.ITERATION_PATH) > 0
                          THEN LEFT(af.ITERATION_PATH, LEN(af.ITERATION_PATH) - CHARINDEX('\\', REVERSE(af.ITERATION_PATH)))
                          ELSE af.ITERATION_PATH
                        END AS FEATURE_PI_PARENT_ITERATION_PATH
                      FROM ADO_FEATURES af
                      WHERE af.ITERATION_PATH IS NOT NULL
                        AND LTRIM(RTRIM(af.ITERATION_PATH)) <> ''
                    ),
                    j AS (
                      SELECT
                        x.*,
                        ic.ITERATION_PATH AS PI_MATCH_PATH
                      FROM x
                      LEFT JOIN ADO_ITERATION_CALENDAR ic
                        ON UPPER(LTRIM(RTRIM(ic.ITERATION_PATH))) = UPPER(LTRIM(RTRIM(x.FEATURE_PI_PARENT_ITERATION_PATH)))
                       AND UPPER(LTRIM(RTRIM(ic.ITERATION_GRAIN))) = 'PI'
                    )
                    SELECT
                      COUNT(*) AS TOTAL_FEATURES,
                      SUM(CASE WHEN FEATURE_IS_SPRINT_LEVEL = 0 AND PI_MATCH_PATH IS NOT NULL THEN 1 ELSE 0 END) AS MAPPED_DIRECT,
                      SUM(CASE WHEN FEATURE_IS_SPRINT_LEVEL = 1 AND PI_MATCH_PATH IS NOT NULL THEN 1 ELSE 0 END) AS SPRINT_MAPPED,
                      SUM(CASE WHEN PI_MATCH_PATH IS NULL THEN 1 ELSE 0 END) AS MISSING_CALENDAR
                    FROM j
                    """
                )
                if df_breakdown is not None and not df_breakdown.empty:
                    # SQL aggregates can return NULL in some edge cases; coerce safely to 0.
                    missing_raw = pd.to_numeric(df_breakdown.iloc[0].get("MISSING_CALENDAR"), errors="coerce")
                    missing_count = 0 if pd.isna(missing_raw) else int(missing_raw)
                else:
                    missing_count = 0
                st.metric("Features missing iteration mapping", missing_count)
                if df_breakdown is not None and not df_breakdown.empty:
                    st.dataframe(df_breakdown, use_container_width=True, height=80)
    
                df_top_missing = fetch_df(
                    """
                    WITH x AS (
                      SELECT
                        af.ITERATION_PATH AS FEATURE_ITERATION_PATH_RAW,
                        CASE
                          WHEN af.ITERATION_PATH IS NULL OR LTRIM(RTRIM(af.ITERATION_PATH)) = '' THEN 0
                          WHEN PATINDEX('%\\I[0-9] S[0-9]%', UPPER(af.ITERATION_PATH)) > 0 THEN 1
                          WHEN PATINDEX('%\\S[0-9]%', UPPER(af.ITERATION_PATH)) > 0 THEN 1
                          ELSE 0
                        END AS FEATURE_IS_SPRINT_LEVEL,
                        CASE
                          WHEN af.ITERATION_PATH IS NULL OR LTRIM(RTRIM(af.ITERATION_PATH)) = '' THEN NULL
                          WHEN (
                            PATINDEX('%\\I[0-9] S[0-9]%', UPPER(af.ITERATION_PATH)) > 0
                            OR PATINDEX('%\\S[0-9]%', UPPER(af.ITERATION_PATH)) > 0
                          )
                          AND CHARINDEX('\\', af.ITERATION_PATH) > 0
                          THEN LEFT(af.ITERATION_PATH, LEN(af.ITERATION_PATH) - CHARINDEX('\\', REVERSE(af.ITERATION_PATH)))
                          ELSE af.ITERATION_PATH
                        END AS FEATURE_PI_PARENT_ITERATION_PATH
                      FROM ADO_FEATURES af
                      WHERE af.ITERATION_PATH IS NOT NULL
                        AND LTRIM(RTRIM(af.ITERATION_PATH)) <> ''
                    )
                    SELECT TOP 20
                      x.FEATURE_PI_PARENT_ITERATION_PATH AS PI_PARENT_ITERATION_PATH,
                      COUNT(*) AS FEATURE_COUNT
                    FROM x
                    LEFT JOIN ADO_ITERATION_CALENDAR ic
                      ON UPPER(LTRIM(RTRIM(ic.ITERATION_PATH))) = UPPER(LTRIM(RTRIM(x.FEATURE_PI_PARENT_ITERATION_PATH)))
                     AND UPPER(LTRIM(RTRIM(ic.ITERATION_GRAIN))) = 'PI'
                    WHERE x.FEATURE_PI_PARENT_ITERATION_PATH IS NOT NULL
                      AND LTRIM(RTRIM(x.FEATURE_PI_PARENT_ITERATION_PATH)) <> ''
                      AND ic.ITERATION_PATH IS NULL
                    GROUP BY x.FEATURE_PI_PARENT_ITERATION_PATH
                    ORDER BY FEATURE_COUNT DESC, x.FEATURE_PI_PARENT_ITERATION_PATH
                    """
                )
                if df_top_missing is None or df_top_missing.empty:
                    st.caption("No missing PI parent iteration paths detected.")
                else:
                    st.caption("Top missing PI parent iteration paths")
                    st.dataframe(df_top_missing, use_container_width=True, height=360)
    
                df_top_sprints = fetch_df(
                    """
                    WITH x AS (
                      SELECT
                        af.ITERATION_PATH AS FEATURE_ITERATION_PATH_RAW,
                        CASE
                          WHEN af.ITERATION_PATH IS NULL OR LTRIM(RTRIM(af.ITERATION_PATH)) = '' THEN 0
                          WHEN PATINDEX('%\\I[0-9] S[0-9]%', UPPER(af.ITERATION_PATH)) > 0 THEN 1
                          WHEN PATINDEX('%\\S[0-9]%', UPPER(af.ITERATION_PATH)) > 0 THEN 1
                          ELSE 0
                        END AS FEATURE_IS_SPRINT_LEVEL,
                        CASE
                          WHEN af.ITERATION_PATH IS NULL OR LTRIM(RTRIM(af.ITERATION_PATH)) = '' THEN NULL
                          WHEN (
                            PATINDEX('%\\I[0-9] S[0-9]%', UPPER(af.ITERATION_PATH)) > 0
                            OR PATINDEX('%\\S[0-9]%', UPPER(af.ITERATION_PATH)) > 0
                          )
                          AND CHARINDEX('\\', af.ITERATION_PATH) > 0
                          THEN LEFT(af.ITERATION_PATH, LEN(af.ITERATION_PATH) - CHARINDEX('\\', REVERSE(af.ITERATION_PATH)))
                          ELSE af.ITERATION_PATH
                        END AS FEATURE_PI_PARENT_ITERATION_PATH
                      FROM ADO_FEATURES af
                      WHERE af.ITERATION_PATH IS NOT NULL
                        AND LTRIM(RTRIM(af.ITERATION_PATH)) <> ''
                    )
                    SELECT TOP 20
                      x.FEATURE_ITERATION_PATH_RAW AS SPRINT_ITERATION_PATH,
                      x.FEATURE_PI_PARENT_ITERATION_PATH AS PI_PARENT_ITERATION_PATH,
                      COUNT(*) AS FEATURE_COUNT
                    FROM x
                    WHERE x.FEATURE_IS_SPRINT_LEVEL = 1
                    GROUP BY x.FEATURE_ITERATION_PATH_RAW, x.FEATURE_PI_PARENT_ITERATION_PATH
                    ORDER BY FEATURE_COUNT DESC, x.FEATURE_ITERATION_PATH_RAW
                    """
                )
                if df_top_sprints is not None and not df_top_sprints.empty:
                    st.caption("Top sprint-level feature iteration paths (mapped to PI parent path)")
                    st.dataframe(df_top_sprints, use_container_width=True, height=360)
            except Exception as e:
                st.exception(e)

        if is_admin:
            st.markdown("### Offline Rebuild (No ADO Connection)")
            st.caption(
                "Use this after restoring ADO_FEATURES / ADO_ITERATION_CALENDAR from production exports. "
                "This does not contact ADO."
            )

            def _object_exists(name: str, obj_type: str = "TABLE") -> bool:
                try:
                    if obj_type.upper() == "VIEW":
                        df_obj = fetch_df(
                            """
                            SELECT COUNT(*) AS N
                            FROM INFORMATION_SCHEMA.VIEWS
                            WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s
                            """,
                            ("dbo", name),
                        )
                    else:
                        df_obj = fetch_df(
                            """
                            SELECT COUNT(*) AS N
                            FROM INFORMATION_SCHEMA.TABLES
                            WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s
                            """,
                            ("dbo", name),
                        )
                    return bool(int(df_obj.iloc[0]["N"])) if df_obj is not None and not df_obj.empty else False
                except Exception:
                    return False

            def _col_exists(table: str, col: str) -> bool:
                try:
                    df_col = fetch_df(
                        """
                        SELECT COUNT(*) AS N
                        FROM INFORMATION_SCHEMA.COLUMNS
                        WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s AND COLUMN_NAME = %s
                        """,
                        ("dbo", table, col),
                    )
                    return bool(int(df_col.iloc[0]["N"])) if df_col is not None and not df_col.empty else False
                except Exception:
                    return False

            def _row_count(table: str, where: str = "") -> int:
                try:
                    where_sql = f" WHERE {where}" if where else ""
                    df_cnt = fetch_df(f"SELECT COUNT(*) AS N FROM {table}{where_sql}")
                    return int(df_cnt.iloc[0]["N"]) if df_cnt is not None and not df_cnt.empty else 0
                except Exception:
                    return 0

            def _normalize_key_expr(expr: str) -> str:
                return (
                    "UPPER(REPLACE(REPLACE(REPLACE(LTRIM(RTRIM({expr})), CHAR(9), ' '), '  ', ' '), '\\\\', '|'))"
                ).format(expr=expr)

            if st.button("Run Offline Rebuild", key="ado_offline_rebuild_btn"):
                profile_cfg = normalize_profile_config(get_active_ado_profile())
                profile_key = str(profile_cfg.get("profile_name") or profile_cfg.get("PROFILE_NAME") or "").strip()
                st.info(f"Active profile: {profile_key or 'Unknown'}")

                if not _object_exists("ADO_FEATURES", "TABLE"):
                    st.error("ADO_FEATURES table not found. Restore it before running Offline Rebuild.")
                    st.stop()
                ado_features_rows = _row_count("ADO_FEATURES")
                st.write(f"ADO_FEATURES rows: {ado_features_rows:,d}")
                if ado_features_rows <= 0:
                    st.error("ADO_FEATURES is empty. Restore data before running Offline Rebuild.")
                    st.stop()

                if _object_exists("ADO_ITERATION_CALENDAR", "TABLE"):
                    ic_rows = _row_count("ADO_ITERATION_CALENDAR")
                    st.write(f"ADO_ITERATION_CALENDAR rows: {ic_rows:,d}")
                else:
                    st.warning("ADO_ITERATION_CALENDAR table not found. Demand views may be incomplete.")

                has_team_variant = _col_exists("ADO_FEATURES", "TEAM_VARIANT_KEY")
                cnt_tvk = (
                    _row_count("ADO_FEATURES", "TEAM_VARIANT_KEY IS NOT NULL AND LTRIM(RTRIM(TEAM_VARIANT_KEY)) <> ''")
                    if has_team_variant
                    else 0
                )

                area_path_col = None
                for cand in ["AREA_PATH_RAW", "AREA_PATH", "AreaPath", "SYSTEM_AREAPATH"]:
                    if _col_exists("ADO_FEATURES", cand):
                        area_path_col = cand
                        break

                if has_team_variant and cnt_tvk > 0:
                    execute(
                        f"""
                        UPDATE ADO_FEATURES
                        SET TEAM_VARIANT_KEY = NULLIF({_normalize_key_expr("TEAM_VARIANT_KEY")}, '')
                        WHERE TEAM_VARIANT_KEY IS NOT NULL
                        """,
                        None,
                    )
                elif area_path_col:
                    execute(
                        f"""
                        UPDATE ADO_FEATURES
                        SET TEAM_VARIANT_KEY = NULLIF({_normalize_key_expr(area_path_col)}, '')
                        WHERE (TEAM_VARIANT_KEY IS NULL OR LTRIM(RTRIM(TEAM_VARIANT_KEY)) = '')
                          AND {area_path_col} IS NOT NULL AND LTRIM(RTRIM({area_path_col})) <> ''
                        """,
                        None,
                    )
                else:
                    st.warning("No AreaPath column found to derive TEAM_VARIANT_KEY.")

                # Ensure offline rebuild also enforces leaf-team canonical keys
                # (prevents ...|TEAM|TEAM splits when running without live ADO sync).
                try:
                    repair_stats = repair_ado_features_leaf_teams()
                    changed = int(repair_stats.get("rows_changed") or 0)
                    before_pct = float(repair_stats.get("dup_tail_pct_before") or 0.0)
                    after_pct = float(repair_stats.get("dup_tail_pct_after") or 0.0)
                    st.caption(
                        f"Leaf team repair: {changed} row(s) updated "
                        f"(dup-tail {before_pct:.3f}% -> {after_pct:.3f}%)."
                    )
                except Exception as e:
                    st.warning(f"Leaf team repair failed: {e}")

                for col in ["ITERATION_PATH", "ITERATION_SK", "PI_LABEL", "ITERATION_LEVEL3"]:
                    if _col_exists("ADO_FEATURES", col):
                        if col in ("PI_LABEL", "ITERATION_LEVEL3"):
                            execute(
                                f"""
                                UPDATE ADO_FEATURES
                                SET {col} = NULLIF(UPPER(LTRIM(RTRIM({col}))), '')
                                WHERE {col} IS NOT NULL
                                """,
                                None,
                            )
                        else:
                            execute(
                                f"""
                                UPDATE ADO_FEATURES
                                SET {col} = NULLIF(LTRIM(RTRIM({col})), '')
                                WHERE {col} IS NOT NULL
                                """,
                                None,
                            )

                if _preview_blocked_notice("Analytics view rebuild"):
                    st.stop()
                try:
                    ensure_analytics_views_ok()
                except Exception as e:
                    st.error(f"Failed to rebuild views: {e}")
                    st.stop()

                try:
                    _settings_post_write_refresh("offline_rebuild", bump_version=True)
                except Exception:
                    pass

                demand_rows = _row_count("VW_TCO_FEATURE_DEMAND") if _object_exists("VW_TCO_FEATURE_DEMAND", "VIEW") else 0
                resolved_rows = _row_count("VW_TCO_FEATURE_DEMAND_RESOLVED") if _object_exists("VW_TCO_FEATURE_DEMAND_RESOLVED", "VIEW") else 0
                st.success(
                    f"Offline rebuild complete. ADO_FEATURES: {ado_features_rows:,d}, "
                    f"VW_TCO_FEATURE_DEMAND: {demand_rows:,d}, "
                    f"VW_TCO_FEATURE_DEMAND_RESOLVED: {resolved_rows:,d}"
                )

            with st.expander("Debug - diagnostics", expanded=False):
                ado_cnt = _row_count("ADO_FEATURES") if _object_exists("ADO_FEATURES", "TABLE") else 0
                v_demand = _row_count("VW_TCO_FEATURE_DEMAND") if _object_exists("VW_TCO_FEATURE_DEMAND", "VIEW") else 0
                v_resolved = _row_count("VW_TCO_FEATURE_DEMAND_RESOLVED") if _object_exists("VW_TCO_FEATURE_DEMAND_RESOLVED", "VIEW") else 0
                st.write(f"ADO_FEATURES rows: {ado_cnt:,d}")
                st.write(f"VW_TCO_FEATURE_DEMAND rows: {v_demand:,d}")
                st.write(f"VW_TCO_FEATURE_DEMAND_RESOLVED rows: {v_resolved:,d}")
                if _col_exists("ADO_FEATURES", "TEAM_VARIANT_KEY"):
                    try:
                        df_top = fetch_df(
                            """
                            SELECT TOP 10 TEAM_VARIANT_KEY, COUNT(*) AS N
                            FROM ADO_FEATURES
                            WHERE TEAM_VARIANT_KEY IS NOT NULL AND LTRIM(RTRIM(TEAM_VARIANT_KEY)) <> ''
                            GROUP BY TEAM_VARIANT_KEY
                            ORDER BY COUNT(*) DESC
                            """
                        )
                        st.dataframe(df_top, use_container_width=True)
                    except Exception as e:
                        st.warning(f"Could not load TEAM_VARIANT_KEY frequency: {e}")

        # Schedule Automation (cron helper)
        with st.expander("Schedule automation (use saved portfolio settings)", expanded=False):
            st.caption("Generate a cron entry to run the CLI weekly. Set MSSQL_* and ADO_PAT in your cron environment.")
            dow_names = ["Sun","Mon","Tue","Wed","Thu","Fri","Sat"]
            c1, c2, c3 = st.columns([1,1,3])
            with c1:
                hour = st.number_input("Hour (24h)", min_value=0, max_value=23, value=2, step=1, key="cron_hour")
            with c2:
                minute = st.number_input("Minute", min_value=0, max_value=59, value=15, step=1, key="cron_minute")
            with c3:
                dow_label = st.selectbox("Day of week", options=dow_names, index=0, key="cron_dow")
            dow = ["0","1","2","3","4","5","6"][dow_names.index(dow_label)]
    
            py = sys.executable
            try:
                script_path = str((Path(__file__).resolve().parents[1] / "scripts" / "ado_sync.py").as_posix())
            except Exception:
                script_path = "scripts/ado_sync.py"
            selected_settings = st.session_state["ado_portfolio_settings"].get(portfolio_name or sel_portfolio or "(unsaved)", {})
            url_for_cmd = selected_settings.get("url") or full_url or "<paste-full-url>"
            max_pages_for_cmd = int(selected_settings.get("max_pages", max_pages))
            org_for_cmd = selected_settings.get("org", org or "YOUR_ORG")
            proj_for_cmd = selected_settings.get("project", proj or "YOUR_PROJECT")
            cron_line = f"{int(minute)} {int(hour)} * * {dow} ADO_PAT='<pat>' MSSQL_SERVER='<server>' MSSQL_DATABASE='<db>' MSSQL_USER='<user>' MSSQL_PASSWORD='<pass>' python {script_path} --url \"{url_for_cmd}\" --max-pages {max_pages_for_cmd} --org \"{org_for_cmd}\" --project \"{proj_for_cmd}\""
            st.code(cron_line)
            st.caption("Save your portfolio first, then use this cron snippet (or Logic App/Scheduler) to keep ADO_FEATURES fresh.")
    
        st.markdown("### Offline note")
        st.caption(
            "For environments without live ADO connection, restore ADO datasets from Backup/Restore and run "
            "'Run Offline Rebuild' above. Manual upload/upsert in this tab has been removed."
        )
    
        # =========================
        # Tab: Map Values
        # =========================
    
def render_ado_mapping_tab() -> None:
    st.subheader("ADO → NEXT Mapping")
    st.caption("These mappings connect raw ADO fields to NEXT Programs, Teams, and Applications.")
    st.info(
        "Why this matters: ADO features drive workforce cost allocation. If mappings are missing, costs land in "
        "unknown buckets or are under-reported, which reduces forecast accuracy."
    )

    cnt = fetch_df("SELECT COUNT(*) AS N FROM ADO_FEATURES")
    current_n = int(cnt.iloc[0]["N"]) if cnt is not None and not cnt.empty else 0
    if current_n == 0:
        st.warning("ADO_FEATURES is empty. Run Full Sync in the first tab.")
    else:
        st.info(f"ADO_FEATURES currently has **{current_n}** rows.")

    ado_programs = fetch_df("""
        SELECT DISTINCT PROGRAM_RAW
        FROM ADO_FEATURES
        WHERE PROGRAM_RAW IS NOT NULL AND LTRIM(RTRIM(PROGRAM_RAW)) <> ''
        ORDER BY PROGRAM_RAW
    """)
    ado_teams = fetch_df("""
        SELECT DISTINCT
               TEAM_VARIANT_KEY,
               TEAM_RAW,
               PROGRAM_RAW,
               AREA_LEVEL3_RAW,
               AREA_LEVEL4_RAW
        FROM ADO_FEATURES
        WHERE TEAM_VARIANT_KEY IS NOT NULL AND LTRIM(RTRIM(TEAM_VARIANT_KEY)) <> ''
        ORDER BY TEAM_VARIANT_KEY
    """)
    ado_apps = fetch_df("""
        SELECT DISTINCT APP_NAME_RAW
        FROM ADO_FEATURES
        WHERE APP_NAME_RAW IS NOT NULL AND APP_NAME_RAW <> ''
        ORDER BY APP_NAME_RAW
    """)

    program_maps = fetch_df("SELECT ADO_PROGRAM, PROGRAMID FROM MAP_ADO_PROGRAM_TO_TCO_PROGRAM ORDER BY ADO_PROGRAM")
    team_maps    = fetch_df("""
        SELECT
          mt.ADO_TEAM_KEY,
          mt.ADO_TEAM,
          mt.PROGRAM_RAW,
          mt.AREA_LEVEL3_RAW,
          mt.AREA_LEVEL4_RAW,
          mt.TEAMID,
          t.PROGRAMID
        FROM MAP_ADO_TEAM_TO_TCO_TEAM mt
        LEFT JOIN TEAMS t ON t.TEAMID = mt.TEAMID
        ORDER BY mt.ADO_TEAM_KEY
    """)
    if team_maps is not None and not team_maps.empty and "ADO_TEAM_KEY" in team_maps.columns:
        team_maps = team_maps.copy()
        team_maps["__mapped"] = team_maps.get("TEAMID").notna()
        team_maps = team_maps.sort_values(by="__mapped", ascending=False)
        team_maps = team_maps.drop_duplicates(subset=["ADO_TEAM_KEY"], keep="first")
        team_maps = team_maps.drop(columns=["__mapped"], errors="ignore")
    app_maps     = fetch_df("SELECT ADO_APP, APP_GROUP FROM MAP_ADO_APP_TO_TCO_GROUP ORDER BY ADO_APP")

    programs_df = fetch_df(
        """
        SELECT
          PROGRAMID,
          COALESCE(NULLIF(LTRIM(RTRIM(PROGRAM_DISPLAY_NAME)), ''), PROGRAMNAME) AS PROGRAMNAME
        FROM PROGRAMS
        ORDER BY COALESCE(NULLIF(LTRIM(RTRIM(PROGRAM_DISPLAY_NAME)), ''), PROGRAMNAME)
        """
    )
    teams_df    = fetch_df(
        """
        SELECT
          TEAMID,
          COALESCE(NULLIF(LTRIM(RTRIM(TEAM_DISPLAY_NAME)), ''), TEAMNAME) AS TEAMNAME
        FROM TEAMS
        ORDER BY COALESCE(NULLIF(LTRIM(RTRIM(TEAM_DISPLAY_NAME)), ''), TEAMNAME)
        """
    )
    groups_df   = fetch_df("SELECT GROUPID, GROUPNAME FROM APPLICATION_GROUPS ORDER BY GROUPNAME")

    def _mapped_stats(values_df: Optional[pd.DataFrame], key_col: str, maps_df: Optional[pd.DataFrame], map_key: str, map_val: str) -> Tuple[int, int]:
        if values_df is None or values_df.empty or key_col not in values_df.columns:
            return 0, 0
        vals = values_df[key_col].dropna().astype(str).str.strip()
        vals = vals[vals != ""]
        total = vals.nunique()
        if maps_df is None or maps_df.empty or map_key not in maps_df.columns or map_val not in maps_df.columns:
            return total, 0
        mapped_keys = maps_df.loc[maps_df[map_val].notna(), map_key].dropna().astype(str).str.strip()
        mapped_keys = mapped_keys[mapped_keys != ""]
        mapped = len(set(vals.str.upper()) & set(mapped_keys.str.upper()))
        return total, mapped

    st.markdown("### Coverage snapshot")
    prog_total, prog_mapped = _mapped_stats(ado_programs, "PROGRAM_RAW", program_maps, "ADO_PROGRAM", "PROGRAMID")
    team_total, team_mapped = _mapped_stats(ado_teams, "TEAM_VARIANT_KEY", team_maps, "ADO_TEAM_KEY", "TEAMID")
    app_total, app_mapped = _mapped_stats(ado_apps, "APP_NAME_RAW", app_maps, "ADO_APP", "APP_GROUP")
    c1, c2, c3 = st.columns(3)
    c1.metric("Programs mapped", f"{prog_mapped}/{prog_total}")
    c2.metric("Teams mapped", f"{team_mapped}/{team_total}")
    c3.metric("Applications mapped", f"{app_mapped}/{app_total}")

    with st.expander("ADO app candidates (discovery)", expanded=False):
        st.caption(
            "Stages distinct ADO app names from ADO_FEATURES into ADO_APP_CANDIDATES with matching suggestions. "
            "This does not write MAP_ADO_APP_TO_TCO_GROUP until you explicitly save mappings in the editor below."
        )
        updated_by_email = str(st.session_state.get("auth_user", {}).get("email") or "").strip() or None

        b1, b2 = st.columns([1, 1])
        with b1:
            if st.button("Refresh candidates", icon=":material/sync:", key="ado_candidates_refresh"):
                try:
                    ensure_ado_minimal_tables()
                    res = refresh_ado_app_candidates(fetch_df, execute)
                    st.success(
                        f"ADO candidates refreshed: {res.new} new, {res.updated} updated, {res.need_review} need review."
                    )
                    st.rerun()
                except Exception as e:
                    st.error(f"Refresh failed: {e}")
        with b2:
            if st.button(
                "Repair default application instances",
                icon=":material/auto_fix_high:",
                key="btn_repair_default_instances",
            ):
                try:
                    r = ensure_default_instances(execute, fetch_df)
                    st.success(
                        "Default instances ensured. "
                        f"Created: {r.groups_created}, Updated: {r.groups_updated}, Repaired multi-default: {r.groups_repaired_multi_default}."
                    )
                    st.rerun()
                except Exception as e:
                    st.error(f"Repair failed: {e}")

        show_all_candidates = st.checkbox("Show all candidate statuses", value=False, key="ado_candidates_show_all")
        try:
            try:
                ensure_ado_minimal_tables()
            except Exception:
                pass
            df_candidates = fetch_df(
                """
                SELECT
                  ADO_APP_RAW,
                  STATUS,
                  FEATURE_COUNT,
                  DERIVED_FTE_SUM,
                  FIRST_SEEN,
                  LAST_SEEN,
                  COALESCE(SUGGESTED_GROUP_NAME, SUGGESTED_TCO_GROUP) AS SUGGESTED_GROUP_NAME,
                  COALESCE(SUGGESTED_GROUP_ID, SUGGESTED_TCO_GROUP_ID) AS SUGGESTED_GROUP_ID,
                  PROGRAMS_SEEN,
                  UPDATED_AT
                FROM ADO_APP_CANDIDATES
                ORDER BY
                  CASE STATUS
                    WHEN 'NEEDS_REVIEW' THEN 0
                    WHEN 'NEW' THEN 1
                    WHEN 'AUTO_SUGGESTED' THEN 2
                    WHEN 'MAPPED' THEN 3
                    WHEN 'IGNORED' THEN 4
                    ELSE 9
                  END,
                  COALESCE(FEATURE_COUNT, 0) DESC,
                  COALESCE(DERIVED_FTE_SUM, 0) DESC
                """,
                None,
            )
        except Exception:
            df_candidates = None

        if df_candidates is None or df_candidates.empty:
            st.caption("No candidate rows found yet. Click “Refresh candidates” to populate ADO_APP_CANDIDATES.")
        else:
            df_candidates = df_candidates.copy()
            if "STATUS" in df_candidates.columns:
                df_candidates["STATUS"] = (
                    df_candidates["STATUS"].astype(str).str.upper().replace({"AUTO_MAPPED": "AUTO_SUGGESTED"})
                )

            try:
                status_counts = df_candidates["STATUS"].fillna("").astype(str).str.upper().value_counts().to_dict()
            except Exception:
                status_counts = {}
            k1, k2, k3, k4, k5 = st.columns(5)
            k1.metric("Total candidates", int(len(df_candidates)))
            k2.metric("Needs review", int(status_counts.get("NEEDS_REVIEW", 0)))
            k3.metric("Auto-suggested", int(status_counts.get("AUTO_SUGGESTED", 0)))
            k4.metric("Mapped", int(status_counts.get("MAPPED", 0)))
            k5.metric("Ignored", int(status_counts.get("IGNORED", 0)))

            tab_queue, tab_manage = st.tabs(["Review queue", "Manage statuses"])

            with tab_queue:
                queue_df = df_candidates.copy()
                if "STATUS" in queue_df.columns:
                    queue_df = queue_df[queue_df["STATUS"].isin(["NEW", "AUTO_SUGGESTED", "NEEDS_REVIEW"])].copy()
                if queue_df.empty:
                    st.caption("No NEW/AUTO_SUGGESTED/NEEDS_REVIEW candidates for the current data set.")
                else:
                    st.dataframe(
                            queue_df[
                                [
                                    c
                                    for c in [
                                        "ADO_APP_RAW",
                                        "STATUS",
                                        "FEATURE_COUNT",
                                        "DERIVED_FTE_SUM",
                                        "LAST_SEEN",
                                        "SUGGESTED_GROUP_NAME",
                                        "PROGRAMS_SEEN",
                                    ]
                                if c in queue_df.columns
                            ]
                        ],
                        use_container_width=True,
                        height=280,
                        hide_index=True,
                    )

                    ado_pick = st.selectbox(
                        "Select candidate ADO app",
                        options=["(Select)"] + queue_df["ADO_APP_RAW"].astype(str).tolist(),
                        index=0,
                        key="ado_candidate_pick",
                    )
                    if ado_pick and ado_pick != "(Select)":
                        row0 = queue_df.loc[queue_df["ADO_APP_RAW"].astype(str) == str(ado_pick)].iloc[0]
                        suggested_name = str(row0.get("SUGGESTED_GROUP_NAME") or "").strip()
                        st.caption(f"Suggested group: {suggested_name or '—'}")

                        groups_df2 = fetch_df("SELECT GROUPID, GROUPNAME FROM APPLICATION_GROUPS ORDER BY GROUPNAME")
                        name_to_id_group2 = {}
                        group_opts = []
                        if groups_df2 is not None and not groups_df2.empty:
                            group_opts = groups_df2["GROUPNAME"].astype(str).tolist()
                            name_to_id_group2 = {str(r.GROUPNAME): str(r.GROUPID) for _, r in groups_df2.iterrows()}

                        col_map, col_create, col_ignore = st.columns([2, 2, 1])
                        with col_map:
                            sel_gname = st.selectbox(
                                "Map to existing app group",
                                options=["(Select)"] + group_opts,
                                index=(1 + group_opts.index(suggested_name) if suggested_name in group_opts else 0),
                                key="ado_candidate_map_group",
                            )
                            if st.button(
                                "Apply mapping",
                                icon=":material/link:",
                                key="ado_candidate_apply_map",
                                disabled=(sel_gname == "(Select)"),
                            ):
                                try:
                                    gid = name_to_id_group2.get(sel_gname)
                                    execute(
                                        """
                                        MERGE INTO MAP_ADO_APP_TO_TCO_GROUP t
                                        USING (SELECT %s AS ADO_APP, %s AS APP_GROUP) s
                                        ON t.ADO_APP = s.ADO_APP
                                        WHEN MATCHED THEN UPDATE SET APP_GROUP = s.APP_GROUP
                                        WHEN NOT MATCHED THEN INSERT (ADO_APP, APP_GROUP) VALUES (s.ADO_APP, s.APP_GROUP);
                                        """,
                                        (str(ado_pick), str(gid)),
                                    )
                                    execute(
                                        """
                                        UPDATE ADO_APP_CANDIDATES
                                        SET STATUS='MAPPED',
                                            SUGGESTED_GROUP_ID=%s,
                                            SUGGESTED_GROUP_NAME=%s,
                                            SUGGESTED_TCO_GROUP_ID=%s,
                                            SUGGESTED_TCO_GROUP=%s,
                                            UPDATED_AT=SYSDATETIME()
                                        WHERE ADO_APP_RAW=%s
                                        """,
                                        (gid, sel_gname, gid, sel_gname, str(ado_pick)),
                                    )
                                    toast_success("Mapping saved and candidate marked MAPPED.")
                                    _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                                    st.rerun()
                                except Exception as e:
                                    st.error(f"Mapping failed: {e}")

                        with col_create:
                            with st.expander("Create new app group", expanded=False):
                                vendors_df = list_vendors()
                                vendor_opts = (
                                    vendors_df["VENDORNAME"].astype(str).tolist()
                                    if vendors_df is not None and not vendors_df.empty
                                    else []
                                )
                                vname = st.selectbox(
                                    "Vendor (required)",
                                    options=["(Select)"] + vendor_opts,
                                    index=0,
                                    key="ado_candidate_vendor",
                                )
                                new_group_name = st.text_input(
                                    "New Application name",
                                    value=str(ado_pick),
                                    key="ado_candidate_new_group_name",
                                )
                                if st.button(
                                    "Create group + map",
                                    icon=":material/add_circle:",
                                    key="ado_candidate_create_group_map",
                                    disabled=(vname == "(Select)") or (not str(new_group_name or "").strip()),
                                ):
                                    try:
                                        vendor_id = None
                                        if vendors_df is not None and not vendors_df.empty and vname != "(Select)":
                                            vendor_id = str(
                                                vendors_df.loc[
                                                    vendors_df["VENDORNAME"].astype(str) == str(vname),
                                                    "VENDORID",
                                                ].iloc[0]
                                            )
                                        group_id = str(uuid.uuid4())
                                        upsert_application_group(
                                            group_id,
                                            str(new_group_name).strip(),
                                            team_id="",
                                            default_vendor_id=vendor_id,
                                            owner=None,
                                            updated_by=updated_by_email,
                                        )
                                        try:
                                            ensure_default_instances(execute, fetch_df)
                                        except Exception:
                                            pass
                                        execute(
                                            """
                                            MERGE INTO MAP_ADO_APP_TO_TCO_GROUP t
                                            USING (SELECT %s AS ADO_APP, %s AS APP_GROUP) s
                                            ON t.ADO_APP = s.ADO_APP
                                            WHEN MATCHED THEN UPDATE SET APP_GROUP = s.APP_GROUP
                                            WHEN NOT MATCHED THEN INSERT (ADO_APP, APP_GROUP) VALUES (s.ADO_APP, s.APP_GROUP);
                                            """,
                                            (str(ado_pick), group_id),
                                        )
                                        execute(
                                            """
                                            UPDATE ADO_APP_CANDIDATES
                                            SET STATUS='MAPPED',
                                                SUGGESTED_GROUP_ID=%s,
                                                SUGGESTED_GROUP_NAME=%s,
                                                SUGGESTED_TCO_GROUP_ID=%s,
                                                SUGGESTED_TCO_GROUP=%s,
                                                UPDATED_AT=SYSDATETIME()
                                            WHERE ADO_APP_RAW=%s
                                            """,
                                            (
                                                group_id,
                                                str(new_group_name).strip(),
                                                group_id,
                                                str(new_group_name).strip(),
                                                str(ado_pick),
                                            ),
                                        )
                                        toast_success("Application created and mapping saved.")
                                        _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                                        st.rerun()
                                    except Exception as e:
                                        st.error(f"Create+map failed: {e}")

                        with col_ignore:
                            if st.button("Ignore", icon=":material/block:", key="ado_candidate_ignore"):
                                try:
                                    execute(
                                        "UPDATE ADO_APP_CANDIDATES SET STATUS='IGNORED', UPDATED_AT=SYSDATETIME() WHERE ADO_APP_RAW=%s",
                                        (str(ado_pick),),
                                    )
                                    toast_success("Candidate marked IGNORED.")
                                    st.rerun()
                                except Exception as e:
                                    st.error(f"Ignore failed: {e}")

                    st.caption(
                        "These candidates are ADO app names that appear in features but are not yet mapped to a NEXT Application. "
                        "Mapping them reduces '(Needs mapping)' / Unassigned buckets in analytics."
                    )

            with tab_manage:
                df_manage = df_candidates.copy()
                if (not show_all_candidates) and "STATUS" in df_manage.columns:
                    df_manage = df_manage[df_manage["STATUS"].isin(["NEW", "AUTO_SUGGESTED", "NEEDS_REVIEW"])].copy()

                if df_manage.empty:
                    st.caption("No candidates to show for the current filter.")
                else:
                    edited_candidates = st.data_editor(
                        df_manage,
                        use_container_width=True,
                        height=320,
                        num_rows="fixed",
                        column_config={
                            "ADO_APP_RAW": st.column_config.TextColumn("ADO App (raw)", disabled=True),
                            "STATUS": st.column_config.SelectboxColumn(
                                "Status",
                                options=["NEW", "AUTO_SUGGESTED", "NEEDS_REVIEW", "IGNORED", "MAPPED"],
                                required=True,
                            ),
                            "SUGGESTED_GROUP_NAME": st.column_config.TextColumn("Suggested Group", disabled=True),
                            "SUGGESTED_GROUP_ID": st.column_config.TextColumn("Suggested Group ID", disabled=True),
                        },
                        key="ado_candidates_editor",
                    )
                    if st.button("Save candidate statuses", icon=":material/save:", key="ado_candidates_save_status"):
                        try:
                            rows = []
                            for _, r in edited_candidates.iterrows():
                                ado_app = str(r.get("ADO_APP_RAW") or "").strip()
                                status = str(r.get("STATUS") or "").strip().upper()
                                if ado_app and status:
                                    rows.append((status, ado_app))
                            if rows:
                                execute(
                                    "UPDATE ADO_APP_CANDIDATES SET STATUS = %s, UPDATED_AT = SYSDATETIME() WHERE ADO_APP_RAW = %s",
                                    rows,
                                    many=True,
                                )
                            toast_success("Candidate statuses saved.")
                            st.rerun()
                        except Exception as e:
                            st.error(f"Save failed: {e}")

    with st.expander("Mapping impact (features missing mapping)", expanded=False):
        try:
            impact_df = fetch_df("""
                SELECT
                  COUNT(*) AS FEATURES_TOTAL,
                  SUM(CASE WHEN mp.PROGRAMID IS NULL THEN 1 ELSE 0 END) AS MISSING_PROGRAM,
                  SUM(CASE WHEN mt.TEAMID IS NULL THEN 1 ELSE 0 END) AS MISSING_TEAM,
                  SUM(CASE WHEN mag.APP_GROUP IS NULL THEN 1 ELSE 0 END) AS MISSING_APP_GROUP
                FROM ADO_FEATURES af
                LEFT JOIN MAP_ADO_PROGRAM_TO_TCO_PROGRAM mp ON UPPER(mp.ADO_PROGRAM) = UPPER(af.PROGRAM_RAW)
                LEFT JOIN MAP_ADO_TEAM_TO_TCO_TEAM mt ON mt.ADO_TEAM_KEY = af.TEAM_VARIANT_KEY
                LEFT JOIN MAP_ADO_APP_TO_TCO_GROUP mag ON mag.ADO_APP = af.APP_NAME_RAW
            """)
            if impact_df is not None and not impact_df.empty:
                row = impact_df.iloc[0]
                c1, c2, c3, c4 = st.columns(4)
                c1.metric("Total features", int(row.get("FEATURES_TOTAL") or 0))
                c2.metric("Missing program map", int(row.get("MISSING_PROGRAM") or 0))
                c3.metric("Missing team map", int(row.get("MISSING_TEAM") or 0))
                c4.metric("Missing app map", int(row.get("MISSING_APP_GROUP") or 0))
        except Exception:
            st.caption("Impact metrics unavailable.")

    guided_mode = st.toggle("Guided mode", value=True, key="ado_map_guided")

    def _run_post_mapping_refresh(reason: str) -> None:
        """Lightweight post-save refresh for ADO mapping edits."""
        if _preview_blocked_notice("Mapping refresh"):
            return
        try:
            # Rebuild demand/iteration mapping chain without a full view bootstrap.
            recompute_feature_iteration_mapping()
            ensure_analytics_views_ok()
        except Exception as e:
            st.warning(f"Mappings saved, but refresh encountered an issue: {e}")
        post_write_refresh(reason, ensure_views=False, rerun=True, bump_version=True)

    st.markdown("### ADO → NEXT Program")
    st.caption("Maps ADO Area Level 2 (Program) to a NEXT Program. Impacts program rollups and cost distribution.")
    if ado_programs is None or ado_programs.empty or programs_df is None or programs_df.empty:
        st.info("Need ADO program values and Programs in NEXT before mapping.")
    else:
        if guided_mode:
            prog_status = "Complete" if prog_total and (prog_mapped >= prog_total) else f"{prog_mapped}/{prog_total} mapped"
            st.caption(f"Step 1 of 3 • Program mapping status: {prog_status}")
            show_unmapped_prog = st.checkbox("Show unmapped programs only", value=True, key="pm_show_unmapped")
        else:
            show_unmapped_prog = st.checkbox("Show unmapped programs only", value=False, key="pm_show_unmapped")
        base_pm = ado_programs.rename(columns={"PROGRAM_RAW": "ADO_PROGRAM"}).copy()
        if program_maps is not None and not program_maps.empty:
            base_pm = base_pm.merge(program_maps, how="left", on="ADO_PROGRAM")
        else:
            base_pm["PROGRAMID"] = None
        if show_unmapped_prog:
            base_pm = base_pm[base_pm["PROGRAMID"].isna()].copy()

        id_to_name_program = {r.PROGRAMID: r.PROGRAMNAME for _, r in programs_df.iterrows()}
        name_to_id_program = {r.PROGRAMNAME: r.PROGRAMID for _, r in programs_df.iterrows()}
        base_pm["TCO_PROGRAMNAME"] = base_pm["PROGRAMID"].map(id_to_name_program)

        edited_prog = st.data_editor(
            base_pm[["ADO_PROGRAM", "TCO_PROGRAMNAME"]],
            use_container_width=True,
            height=240,
            num_rows="fixed",
            column_config={
                "ADO_PROGRAM": st.column_config.TextColumn("ADO Program (Area Level 2)", disabled=True),
                "TCO_PROGRAMNAME": st.column_config.SelectboxColumn(
                    "NEXT Program",
                    options=programs_df["PROGRAMNAME"].tolist(),
                    required=False,
                ),
            },
            key="pm_editor",
        )
        if st.button("Save Program Mappings", icon=":material/save:", key="btn_save_program_mappings"):
            rows_to_upsert_pm: List[Tuple[str, str]] = []
            for _, row in edited_prog.iterrows():
                ado_val = str(row["ADO_PROGRAM"]).strip()
                pname = row.get("TCO_PROGRAMNAME")
                if _blank_or_nan(ado_val):
                    continue
                if pname and pname in name_to_id_program:
                    rows_to_upsert_pm.append((ado_val, name_to_id_program[pname]))
                else:
                    execute("DELETE FROM MAP_ADO_PROGRAM_TO_TCO_PROGRAM WHERE ADO_PROGRAM = %s", (ado_val,))
            if rows_to_upsert_pm:
                merge_sql = """
                MERGE INTO MAP_ADO_PROGRAM_TO_TCO_PROGRAM t
                USING (SELECT %s AS ADO_PROGRAM, %s AS PROGRAMID) s
                ON t.ADO_PROGRAM = s.ADO_PROGRAM
                WHEN MATCHED THEN UPDATE SET PROGRAMID = s.PROGRAMID
                WHEN NOT MATCHED THEN INSERT (ADO_PROGRAM, PROGRAMID) VALUES (s.ADO_PROGRAM, s.PROGRAMID);
                """
                execute(merge_sql, rows_to_upsert_pm, many=True)
            toast_success("Program mappings saved. Refreshing analytics...")
            _run_post_mapping_refresh("ado_program_mapping_save")

    st.markdown("---")

    st.markdown("### ADO → NEXT Team")
    st.caption("Links raw ADO team variants to a NEXT Team. Missing team mapping causes workforce costs to land in unknown teams.")
    if ado_teams is None or ado_teams.empty or teams_df is None or teams_df.empty:
        st.info("Load features (with Area levels) and ensure Teams exist in NEXT before mapping.")
    else:
        if guided_mode:
            team_status = "Complete" if team_total and (team_mapped >= team_total) else f"{team_mapped}/{team_total} mapped"
            st.caption(f"Step 2 of 3 • Team mapping status: {team_status}")
            show_unmapped_team = st.checkbox("Show unmapped teams only", value=True, key="tm_show_unmapped")
        else:
            show_unmapped_team = st.checkbox("Show unmapped teams only", value=False, key="tm_show_unmapped")
        base_tm = ado_teams.copy()
        base_tm["ADO_TEAM_KEY"] = (
            base_tm.get("TEAM_VARIANT_KEY", "").fillna("").astype(str).str.strip()
        )
        if "TEAM_RAW" in base_tm.columns:
            team_raw_norm = base_tm["TEAM_RAW"].fillna("").astype(str).str.strip()
            base_tm["ADO_TEAM_KEY"] = base_tm["ADO_TEAM_KEY"].where(base_tm["ADO_TEAM_KEY"] != "", team_raw_norm)
        # Canonical key parsing for leaf vs area-level nodes
        def _norm_ado_key(key: Any) -> str:
            return str(key or "").strip().upper()

        norm_keys = (
            base_tm["ADO_TEAM_KEY"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
        )
        norm_keys = norm_keys[norm_keys != ""]
        all_norm_keys = sorted(set(norm_keys.tolist()))
        leaf_by_key: Dict[str, bool] = {}
        for key in all_norm_keys:
            prefix = key + "|"
            has_child = any((other != key) and other.startswith(prefix) for other in all_norm_keys)
            leaf_by_key[key] = not has_child

        segs = base_tm["ADO_TEAM_KEY"].astype(str).str.split("|").apply(
            lambda parts: [p.strip() for p in parts if str(p).strip()]
        )
        base_tm["KEY_SEGMENTS"] = segs.apply(len)
        base_tm["KEY_LEAF"] = segs.apply(lambda p: p[-1] if len(p) > 0 else "")
        base_tm["KEY_PREV"] = segs.apply(lambda p: p[-2] if len(p) > 1 else "")
        base_tm["NORM_ADO_TEAM_KEY"] = base_tm["ADO_TEAM_KEY"].map(_norm_ado_key)
        base_tm["IS_LEAF_KEY"] = base_tm["NORM_ADO_TEAM_KEY"].map(lambda k: leaf_by_key.get(k, False))
        # Program-only (single segment) keys are treated as area-level even if they have no children.
        base_tm["IS_AREA_LEVEL"] = (~base_tm["IS_LEAF_KEY"]) | (base_tm["KEY_SEGMENTS"] < 2)
        if not base_tm.empty:
            base_tm = base_tm.sort_values(
                ["ADO_TEAM_KEY", "PROGRAM_RAW", "TEAM_RAW", "AREA_LEVEL3_RAW", "AREA_LEVEL4_RAW"],
                na_position="last"
            ).drop_duplicates(
                subset=["ADO_TEAM_KEY", "PROGRAM_RAW", "TEAM_RAW", "AREA_LEVEL3_RAW", "AREA_LEVEL4_RAW"],
                keep="first"
            ).reset_index(drop=True)
        if team_maps is not None and not team_maps.empty:
            base_tm = base_tm.merge(team_maps, how="left", on="ADO_TEAM_KEY", suffixes=("", "_MAP"))
            for col in ["ADO_TEAM", "PROGRAM_RAW", "AREA_LEVEL3_RAW", "AREA_LEVEL4_RAW", "TEAMID", "PROGRAMID"]:
                map_col = f"{col}_MAP"
                if map_col in base_tm.columns:
                    base_tm[col] = base_tm[col].fillna(base_tm[map_col])
                    base_tm.drop(columns=[map_col], inplace=True)
        else:
            base_tm["TEAMID"] = None
            base_tm["PROGRAMID"] = None

        for col in ["PROGRAM_RAW", "TEAM_RAW", "AREA_LEVEL3_RAW", "AREA_LEVEL4_RAW", "ADO_TEAM", "TEAMID", "PROGRAMID"]:
            if col not in base_tm.columns:
                base_tm[col] = None

        # Fill PROGRAMID from ADO Program mapping when Team is not mapped
        prog_map_lookup = {}
        if program_maps is not None and not program_maps.empty:
            prog_map_lookup = {str(r.ADO_PROGRAM).upper(): r.PROGRAMID for _, r in program_maps.iterrows() if pd.notna(r.ADO_PROGRAM)}
        if prog_map_lookup:
            base_tm["PROGRAMID"] = base_tm["PROGRAMID"].fillna(
                base_tm["PROGRAM_RAW"].apply(lambda x: prog_map_lookup.get(str(x).upper()) if pd.notna(x) else None)
            )

        base_tm["TEAM_DISPLAY"] = (
            base_tm.get("AREA_LEVEL4_RAW").fillna(base_tm.get("TEAM_RAW"))
            .fillna(base_tm.get("AREA_LEVEL3_RAW"))
            .fillna(base_tm.get("PROGRAM_RAW"))
        )
        base_tm["TEAM_DISPLAY"] = base_tm["TEAM_DISPLAY"].fillna("")

        if "ADO_TEAM_KEY" in base_tm.columns:
            base_tm = base_tm.sort_values(by="ADO_TEAM_KEY").drop_duplicates(subset=["ADO_TEAM_KEY"], keep="first")
        meta_lookup = base_tm.set_index("ADO_TEAM_KEY")[
            ["PROGRAM_RAW", "AREA_LEVEL3_RAW", "AREA_LEVEL4_RAW", "TEAM_RAW", "ADO_TEAM", "KEY_LEAF", "KEY_PREV", "IS_AREA_LEVEL", "NORM_ADO_TEAM_KEY", "IS_LEAF_KEY"]
        ].to_dict("index")

        id_to_name_team = {r.TEAMID: r.TEAMNAME for _, r in teams_df.iterrows()}
        name_to_id_team = {r.TEAMNAME: r.TEAMID for _, r in teams_df.iterrows()}
        base_tm["TCO_TEAMNAME"] = base_tm["TEAMID"].map(id_to_name_team)
        base_tm["TCO_PROGRAMNAME"] = base_tm["PROGRAMID"].map(id_to_name_program) if "PROGRAMID" in base_tm.columns else None
        if show_unmapped_team:
            base_tm = base_tm[base_tm["TEAMID"].isna()].copy()
        # Leaf-level only for mapping UI
        area_nodes = base_tm[base_tm["IS_AREA_LEVEL"]].copy()
        base_tm = base_tm[~base_tm["IS_AREA_LEVEL"]].copy()

        with st.expander("Area-level nodes (not mappable)", expanded=False):
            st.caption(f"{len(area_nodes):,} area-level keys detected (read-only).")
            if not area_nodes.empty:
                st.dataframe(
                    area_nodes[["ADO_TEAM_KEY", "PROGRAM_RAW", "TEAM_RAW", "AREA_LEVEL3_RAW", "AREA_LEVEL4_RAW"]].head(200),
                    use_container_width=True,
                    height=240,
                )

        edited = st.data_editor(
            base_tm[["ADO_TEAM_KEY", "PROGRAM_RAW", "TEAM_DISPLAY", "TCO_TEAMNAME", "TCO_PROGRAMNAME"]],
            use_container_width=True,
            height=360,
            num_rows="fixed",
            column_config={
                "ADO_TEAM_KEY": st.column_config.TextColumn("ADO Team Variant Key", disabled=True),
                "PROGRAM_RAW": st.column_config.TextColumn("Area Level 2", disabled=True),
                "TEAM_DISPLAY": st.column_config.TextColumn("ADO Area (Team)", disabled=True),
                "TCO_PROGRAMNAME": st.column_config.TextColumn("NEXT Program (from mapping)", disabled=True),
                "TCO_TEAMNAME": st.column_config.SelectboxColumn(
                    "NEXT Team",
                    options=teams_df["TEAMNAME"].tolist(),
                    required=False,
                ),
            },
            key="tm_editor",
        )
        if st.button("Save Team Mappings", icon=":material/save:", key="btn_save_team_mappings"):
            def _clean(val: Any) -> Optional[str]:
                if val is None:
                    return None
                sval = str(val).strip()
                return sval if sval else None
            def _leaf_from_key(key: str) -> str:
                parts = [p.strip() for p in str(key or "").split("|")]
                return parts[-1] if parts else ""
            def _prev_from_key(key: str) -> str:
                parts = [p.strip() for p in str(key or "").split("|")]
                return parts[-2] if len(parts) > 1 else ""
            def _is_area_level(key: str) -> bool:
                norm_key = _norm_ado_key(key)
                meta_hit = meta_lookup.get(str(key or "").strip(), {})
                if meta_hit:
                    return bool(meta_hit.get("IS_AREA_LEVEL", False))
                parts = [p.strip() for p in str(key or "").split("|") if str(p).strip()]
                if len(parts) < 2:
                    return True
                return not bool(leaf_by_key.get(norm_key, False))

            rows_to_upsert: List[Tuple[str, str, Optional[str], Optional[str], Optional[str], Optional[str]]] = []
            for _, row in edited.iterrows():
                ado_val = str(row["ADO_TEAM_KEY"]).strip()
                tname = row.get("TCO_TEAMNAME")
                if _blank_or_nan(ado_val):
                    continue
                # Canonicalize key to TEAM_VARIANT_KEY when present; fallback to TEAM_RAW only when missing.
                if "TEAM_VARIANT_KEY" in base_tm.columns:
                    tv = base_tm.loc[base_tm["ADO_TEAM_KEY"] == ado_val, "TEAM_VARIANT_KEY"]
                    tv_val = str(tv.iloc[0]).strip() if tv is not None and not tv.empty else ""
                    if tv_val:
                        ado_val = tv_val
                leaf_team = _leaf_from_key(ado_val)
                prev_seg = _prev_from_key(ado_val)
                if _is_area_level(ado_val):
                    st.error(
                        f"ADO_TEAM_KEY is area-level and cannot be mapped: '{ado_val}'. "
                        "Select a leaf team key instead."
                    )
                    st.stop()
                meta = meta_lookup.get(ado_val, {})
                if tname and tname in name_to_id_team:
                    rows_to_upsert.append((
                        ado_val,
                        name_to_id_team[tname],
                        _clean(row.get("PROGRAM_RAW")) or _clean(meta.get("PROGRAM_RAW")),
                        _clean(meta.get("AREA_LEVEL3_RAW")),
                        _clean(meta.get("AREA_LEVEL4_RAW")),
                        leaf_team,
                    ))
                else:
                    execute("DELETE FROM MAP_ADO_TEAM_TO_TCO_TEAM WHERE ADO_TEAM_KEY = %s", (ado_val,))
            if rows_to_upsert:
                merge_sql = """
                MERGE INTO MAP_ADO_TEAM_TO_TCO_TEAM t
                USING (SELECT %s AS ADO_TEAM_KEY, %s AS TEAMID, %s AS PROGRAM_RAW, %s AS AREA_LEVEL3_RAW, %s AS AREA_LEVEL4_RAW, %s AS ADO_TEAM) s
                ON t.ADO_TEAM_KEY = s.ADO_TEAM_KEY
                WHEN MATCHED THEN UPDATE SET TEAMID = s.TEAMID, PROGRAM_RAW = s.PROGRAM_RAW, AREA_LEVEL3_RAW = s.AREA_LEVEL3_RAW, AREA_LEVEL4_RAW = s.AREA_LEVEL4_RAW, ADO_TEAM = s.ADO_TEAM
                WHEN NOT MATCHED THEN INSERT (ADO_TEAM_KEY, TEAMID, PROGRAM_RAW, AREA_LEVEL3_RAW, AREA_LEVEL4_RAW, ADO_TEAM) VALUES (s.ADO_TEAM_KEY, s.TEAMID, s.PROGRAM_RAW, s.AREA_LEVEL3_RAW, s.AREA_LEVEL4_RAW, s.ADO_TEAM);
                """
                execute(merge_sql, rows_to_upsert, many=True)
            toast_success("Team mappings saved. Refreshing analytics...")
            _run_post_mapping_refresh("ado_team_mapping_save")

        if is_admin:
            with st.expander("Admin: Mapping integrity check", expanded=False):
                try:
                    bad = fetch_df(
                        """
                        SELECT
                          ADO_TEAM_KEY,
                          ADO_TEAM
                        FROM MAP_ADO_TEAM_TO_TCO_TEAM
                        WHERE ADO_TEAM IS NOT NULL
                          AND UPPER(LTRIM(RTRIM(ADO_TEAM))) <> UPPER(LTRIM(RTRIM(
                            RIGHT(ADO_TEAM_KEY, CHARINDEX('|', REVERSE(ADO_TEAM_KEY) + '|') - 1)
                          )))
                        """
                    )
                except Exception as e:
                    bad = None
                    st.warning(f"Could not run integrity check: {e}")
                if bad is not None and not bad.empty:
                    st.warning(
                        "Mapping table contains invalid ADO_TEAM values (area-level). "
                        "This can inflate demand/capacity and roadmap features."
                    )
                    st.dataframe(bad.head(200), use_container_width=True)
                else:
                    st.success("No invalid ADO_TEAM values detected.")

    st.markdown("---")

    st.markdown("### ADO → NEXT Application")
    st.caption("Maps ADO Application Name to an Application. Required for application-level cost distribution.")
    st.caption(
        "Multi-instance onboarding note: when you select multiple ADO apps in Applications → Onboard, "
        "this table stores one mapping row per selected ADO app to the same NEXT application group."
    )
    if ado_apps is None or ado_apps.empty or groups_df is None or groups_df.empty:
        st.info("Load features and create Applications first.")
    else:
        if guided_mode:
            app_status = "Complete" if app_total and (app_mapped >= app_total) else f"{app_mapped}/{app_total} mapped"
            st.caption(f"Step 3 of 3 • Application mapping status: {app_status}")
            show_unmapped_app = st.checkbox("Show unmapped app groups only", value=True, key="am_show_unmapped")
        else:
            show_unmapped_app = st.checkbox("Show unmapped app groups only", value=False, key="am_show_unmapped")
        base_am = ado_apps.rename(columns={"APP_NAME_RAW": "ADO_APP"}).copy()
        if app_maps is not None and not app_maps.empty:
            base_am = base_am.merge(app_maps, how="left", on="ADO_APP")
        else:
            base_am["APP_GROUP"] = None
        if show_unmapped_app:
            base_am = base_am[base_am["APP_GROUP"].isna()].copy()

        id_to_name_group = {r.GROUPID: r.GROUPNAME for _, r in groups_df.iterrows()}
        name_to_id_group = {r.GROUPNAME: r.GROUPID for _, r in groups_df.iterrows()}
        base_am["TCO_GROUPNAME"] = base_am["APP_GROUP"].map(id_to_name_group)
        # Pre-fill unmapped apps with high-confidence suggestions from ADO_APP_CANDIDATES (still requires explicit save).
        try:
            cand = fetch_df(
                """
                SELECT
                  ADO_APP_RAW,
                  STATUS,
                  COALESCE(SUGGESTED_GROUP_ID, SUGGESTED_TCO_GROUP_ID) AS SUGGESTED_GROUP_ID
                FROM ADO_APP_CANDIDATES
                """,
                None,
            )
        except Exception:
            cand = None
        if cand is not None and not cand.empty:
            cand = cand.copy()
            cand["ADO_APP_RAW"] = cand["ADO_APP_RAW"].astype(str).str.strip()
            cand["STATUS"] = cand["STATUS"].astype(str).str.upper().replace({"AUTO_MAPPED": "AUTO_SUGGESTED"})
            cand["SUGGESTED_GROUP_ID"] = cand["SUGGESTED_GROUP_ID"].astype(str).str.strip()
            base_am = base_am.merge(cand, how="left", left_on="ADO_APP", right_on="ADO_APP_RAW")
            base_am.drop(columns=["ADO_APP_RAW"], inplace=True, errors="ignore")
            auto_mask = (
                base_am.get("STATUS", "").astype(str).str.upper().eq("AUTO_SUGGESTED")
                & base_am.get("TCO_GROUPNAME").isna()
                & base_am.get("SUGGESTED_GROUP_ID", "").astype(str).str.strip().ne("")
            )
            if auto_mask.any():
                base_am.loc[auto_mask, "TCO_GROUPNAME"] = base_am.loc[auto_mask, "SUGGESTED_GROUP_ID"].map(id_to_name_group)

        edited2 = st.data_editor(
            base_am[["ADO_APP","TCO_GROUPNAME"]],
            use_container_width=True,
            height=360,
            num_rows="fixed",
            column_config={
                "ADO_APP": st.column_config.TextColumn("ADO App Name", disabled=True),
                "TCO_GROUPNAME": st.column_config.SelectboxColumn(
                    "NEXT Application",
                    options=groups_df["GROUPNAME"].tolist() if groups_df is not None and not groups_df.empty else [],
                    required=False,
                    help="This picker shows group names. The mapping will store the group ID."
                ),
            },
            key="am_editor",
        )
        if st.button("Save Application Mappings", icon=":material/save:", key="btn_save_app_group_mappings"):
            rows_to_upsert: List[Tuple[str, str]] = []
            for _, row in edited2.iterrows():
                ado_val = str(row["ADO_APP"]).strip()
                gname = row.get("TCO_GROUPNAME")
                if _blank_or_nan(ado_val):
                    continue
                if gname and gname in name_to_id_group:
                    gid = name_to_id_group[gname]
                    rows_to_upsert.append((ado_val, gid))
                else:
                    execute("DELETE FROM MAP_ADO_APP_TO_TCO_GROUP WHERE ADO_APP = %s", (ado_val,))
            if rows_to_upsert:
                merge_sql = """
                MERGE INTO MAP_ADO_APP_TO_TCO_GROUP t
                USING (SELECT %s AS ADO_APP, %s AS APP_GROUP) s
                ON t.ADO_APP = s.ADO_APP
                WHEN MATCHED THEN UPDATE SET APP_GROUP = s.APP_GROUP
                WHEN NOT MATCHED THEN INSERT (ADO_APP, APP_GROUP) VALUES (s.ADO_APP, s.APP_GROUP);
                """
                execute(merge_sql, rows_to_upsert, many=True)
            toast_success("Application mappings saved. Refreshing analytics...")
            _run_post_mapping_refresh("ado_app_mapping_save")

    # =========================
    # Tab: 🔎 ADO Explorer
    # =========================

def render_ado_advanced_tab() -> None:
    with st.expander("Explorer (Advanced)", expanded=False):
        st.subheader("ADO Explorer")
        st.caption("Browse raw ADO imports and demand diagnostics. Filters apply to all sections.")

        with st.expander("Debug - program/team mapping", expanded=False):
            st.caption("Check raw ADO fields vs enriched view outputs to detect mapping/config issues.")
            diag_cols = st.columns(2)
            profile_cfg = normalize_profile_config(get_active_ado_profile())
            area_cfg = profile_cfg.get("area_mapping") or {}
            program_level = str(area_cfg.get("program_level") or "").strip()
            team_level = str(area_cfg.get("team_level") or "").strip()
            with diag_cols[0]:
                st.markdown("**Active profile mapping**")
                st.write(f"Program level: `{program_level or '—'}`")
                st.write(f"Team level: `{team_level or '—'}`")
            with diag_cols[1]:
                st.markdown("**Legacy settings (if any)**")
                sel_portfolio = st.session_state.get("ado_portfolio_select")
                legacy_settings = st.session_state.get("ado_portfolio_settings", {}).get(sel_portfolio or "", {})
                legacy_program = str(legacy_settings.get("area_program_level") or legacy_settings.get("program_level") or "").strip()
                legacy_team = str(legacy_settings.get("area_team_level") or legacy_settings.get("team_level") or "").strip()
                st.write(f"Program level: `{legacy_program or '—'}`")
                st.write(f"Team level: `{legacy_team or '—'}`")

            st.caption("Quick fix hint: Program level = AreaLevel2, Team level = AreaLevel3 (example).")

            def _safe_count(sql: str) -> int:
                try:
                    df = fetch_df(sql)
                    if df is None or df.empty:
                        return 0
                    return int(df.iloc[0, 0])
                except Exception as e:
                    st.error(f"Debug query failed: {e}")
                    return 0

            col_a, col_b = st.columns(2)
            with col_a:
                st.markdown("**ADO_FEATURES raw distribution**")
                try:
                    df_prog_raw = fetch_df(
                        """
                        SELECT TOP 20 PROGRAM_RAW, COUNT(*) AS FEATURES
                        FROM ADO_FEATURES
                        GROUP BY PROGRAM_RAW
                        ORDER BY COUNT(*) DESC
                        """
                    )
                    st.dataframe(df_prog_raw, use_container_width=True, height=260)
                except Exception as e:
                    st.warning(f"Could not read PROGRAM_RAW distribution: {e}")
                try:
                    df_prog_unmapped = fetch_df(
                        """
                        SELECT COUNT(DISTINCT af.PROGRAM_RAW) AS UNMAPPED_PROGRAMS
                        FROM ADO_FEATURES af
                        LEFT JOIN MAP_ADO_PROGRAM_TO_TCO_PROGRAM mp
                          ON UPPER(mp.ADO_PROGRAM) = UPPER(LTRIM(RTRIM(af.PROGRAM_RAW)))
                        WHERE af.PROGRAM_RAW IS NOT NULL AND LTRIM(RTRIM(af.PROGRAM_RAW)) <> ''
                          AND mp.PROGRAMID IS NULL
                        """
                    )
                    if df_prog_unmapped is not None and not df_prog_unmapped.empty:
                        st.caption(f"Unmapped program candidates: **{int(df_prog_unmapped.iloc[0,0]):,d}**")
                except Exception:
                    pass
                try:
                    df_team_raw = fetch_df(
                        """
                        SELECT TOP 20 TEAM_RAW, COUNT(*) AS FEATURES
                        FROM ADO_FEATURES
                        GROUP BY TEAM_RAW
                        ORDER BY COUNT(*) DESC
                        """
                    )
                    st.dataframe(df_team_raw, use_container_width=True, height=260)
                except Exception as e:
                    st.warning(f"Could not read TEAM_RAW distribution: {e}")
                try:
                    df_team_unmapped = fetch_df(
                        """
                        SELECT COUNT(DISTINCT af.TEAM_RAW) AS UNMAPPED_TEAMS
                        FROM ADO_FEATURES af
                        LEFT JOIN MAP_ADO_TEAM_TO_TCO_TEAM mt
                          ON UPPER(mt.ADO_TEAM) = UPPER(LTRIM(RTRIM(af.TEAM_RAW)))
                        WHERE af.TEAM_RAW IS NOT NULL AND LTRIM(RTRIM(af.TEAM_RAW)) <> ''
                          AND mt.TEAMID IS NULL
                        """
                    )
                    if df_team_unmapped is not None and not df_team_unmapped.empty:
                        st.caption(f"Unmapped team candidates: **{int(df_team_unmapped.iloc[0,0]):,d}**")
                except Exception:
                    pass
            with col_b:
                st.markdown("**Enriched view distribution**")
                try:
                    df_prog_enriched = fetch_df(
                        """
                        SELECT TOP 20 PROGRAMNAME, COUNT(*) AS FEATURES
                        FROM VW_ADO_FEATURES_ENRICHED
                        GROUP BY PROGRAMNAME
                        ORDER BY COUNT(*) DESC
                        """
                    )
                    st.dataframe(df_prog_enriched, use_container_width=True, height=260)
                except Exception as e:
                    st.warning(f"Could not read PROGRAMNAME distribution: {e}")
                try:
                    df_team_enriched = fetch_df(
                        """
                        SELECT TOP 20 TEAMNAME, COUNT(*) AS FEATURES
                        FROM VW_ADO_FEATURES_ENRICHED
                        GROUP BY TEAMNAME
                        ORDER BY COUNT(*) DESC
                        """
                    )
                    st.dataframe(df_team_enriched, use_container_width=True, height=260)
                except Exception as e:
                    st.warning(f"Could not read TEAMNAME distribution: {e}")

            raw_prog_distinct = _safe_count("SELECT COUNT(DISTINCT PROGRAM_RAW) FROM ADO_FEATURES")
            raw_team_distinct = _safe_count("SELECT COUNT(DISTINCT TEAM_RAW) FROM ADO_FEATURES")
            enr_prog_distinct = _safe_count("SELECT COUNT(DISTINCT PROGRAMNAME) FROM VW_ADO_FEATURES_ENRICHED")
            enr_team_distinct = _safe_count("SELECT COUNT(DISTINCT TEAMNAME) FROM VW_ADO_FEATURES_ENRICHED")
            enr_prog_nonnull = _safe_count("SELECT COUNT(*) FROM VW_ADO_FEATURES_ENRICHED WHERE PROGRAMNAME IS NOT NULL AND LTRIM(RTRIM(PROGRAMNAME)) <> ''")
            enr_team_nonnull = _safe_count("SELECT COUNT(*) FROM VW_ADO_FEATURES_ENRICHED WHERE TEAMNAME IS NOT NULL AND LTRIM(RTRIM(TEAMNAME)) <> ''")
            if raw_prog_distinct <= 0:
                st.error("Sync missing AreaLevel fields (enable AreaLevel1..4 in profile + resync).")
            elif raw_prog_distinct > 0 and enr_prog_nonnull <= 0:
                st.warning("Mapping not configured yet (expected during onboarding).")
            if raw_team_distinct <= 0:
                st.error("Sync missing AreaLevel fields (enable AreaLevel1..4 in profile + resync).")
            elif raw_team_distinct > 0 and enr_team_nonnull <= 0:
                st.warning("Mapping not configured yet (expected during onboarding).")

            st.markdown("**PI label normalization check**")
            try:
                df_pi_bad = fetch_df(
                    """
                    SELECT TOP 5 PI_LABEL
                    FROM VW_ADO_FEATURES_ENRICHED
                    WHERE PI_LABEL LIKE '%(%'
                    """
                )
                bad_count = 0 if df_pi_bad is None else int(len(df_pi_bad.index))
                if bad_count == 0:
                    st.success("View PI_LABEL is normalized (no date suffixes detected).")
                else:
                    st.warning("View PI_LABEL still contains raw labels with date suffixes.")
                    st.dataframe(df_pi_bad, use_container_width=True, height=160)
            except Exception as e:
                st.warning(f"PI label check failed: {e}")

            st.markdown("**Reference table checks**")
            ref_cols = st.columns(3)
            programs_ct = _safe_count("SELECT COUNT(*) FROM PROGRAMS")
            teams_ct = _safe_count("SELECT COUNT(*) FROM TEAMS")
            map_team_ct = _safe_count("SELECT COUNT(*) FROM MAP_ADO_TEAM_TO_TCO_TEAM")
            map_app_ct = _safe_count("SELECT COUNT(*) FROM MAP_ADO_APP_TO_TCO_GROUP")
            with ref_cols[0]:
                st.metric("PROGRAMS rows", f"{programs_ct:,d}")
            with ref_cols[1]:
                st.metric("TEAMS rows", f"{teams_ct:,d}")
            with ref_cols[2]:
                st.metric("MAP_ADO_TEAM_TO_TCO_TEAM rows", f"{map_team_ct:,d}")
            st.caption(f"MAP_ADO_APP_TO_TCO_GROUP rows: {map_app_ct:,d}")
            if programs_ct == 0 or teams_ct == 0:
                st.error("Reference tables missing — restore/ensure may have reset them.")

        tab_explore_po, = st.tabs(["Explorer v2 (PO view)"])

        if False:
            with st.expander("Maintenance", expanded=False):
                st.caption("No maintenance actions are exposed for ADO Explorer.")

            df_teams, df_apps, df_iters = load_ado_distincts()
            c1, c2, c3 = st.columns(3)
            c1.metric("Distinct ADO Teams", len(df_teams))
            c2.metric("Distinct ADO Apps", len(df_apps))
            c3.metric("Distinct Iterations", len(df_iters))

            with st.expander("Filters", expanded=True):
                f1, f2, f3 = st.columns([2, 2, 2])
                team_like = f1.text_input("Team contains (raw)", "", key="exp_team_contains")
                app_like = f2.text_input("App contains", "", key="exp_app_contains")
                iter_like = f3.text_input("Iteration contains", "", key="exp_iter_contains")
                st.caption("Filters match raw ADO fields; mapped names may display differently.")

            where: List[str] = []
            params: List[str] = []
            if team_like.strip():
                where.append("UPPER(TEAM_RAW) LIKE UPPER(%s)")
                params.append(f"%{team_like.strip()}%")
            if app_like.strip():
                where.append("UPPER(APP_NAME_RAW) LIKE UPPER(%s)")
                params.append(f"%{app_like.strip()}%")
            if iter_like.strip():
                where.append("UPPER(ITERATION_PATH) LIKE UPPER(%s)")
                params.append(f"%{iter_like.strip()}%")
            where_sql = " WHERE " + " AND ".join(where) if where else ""

            st.markdown("#### Latest features")
            st.caption(
                "Expected demand is driven by the active ADO profile configuration "
                "(SWAG or Velocity snapshot)."
            )
            df_raw = ado_features_base_query(where_sql, tuple(params) if params else None)

            if df_raw is not None and not df_raw.empty:
                fte = pd.to_numeric(df_raw.get("DERIVED_FTE"), errors="coerce").fillna(0.0)
                ready_pct = float((fte > 0).mean()) if len(fte.index) else 0.0
                st.caption(f"SWAG-ready (Derived FTE > 0): {ready_pct*100:.0f}%")

            df_show = df_raw.copy() if isinstance(df_raw, pd.DataFrame) else pd.DataFrame()
            if not df_show.empty:
                def _fmt_fte(v: Any) -> str:
                    try:
                        x = float(v)
                        if pd.isna(x) or not math.isfinite(x):
                            return "N/A"
                        return f"{x:,.2f}"
                    except Exception:
                        return "N/A"

                def _fmt_diff(v: Any) -> str:
                    try:
                        x = float(v)
                        if pd.isna(x) or not math.isfinite(x):
                            return ""
                        return f"{x:,.2f}"
                    except Exception:
                        return ""

                df_show["DERIVED_FTE_DISPLAY"] = df_show.get("DERIVED_FTE", pd.Series([pd.NA] * len(df_show))).apply(_fmt_fte)
                df_show["FTE_DIFF_DISPLAY"] = df_show.get("FTE_DIFF", pd.Series([pd.NA] * len(df_show))).apply(_fmt_diff)

            show_cols = [
                "FEATURE_ID",
                "TITLE",
                "STATE",
                "TEAM_RAW",
                "PROGRAMNAME",
                "ITERATION_LEVEL3",
                "ITERATION_PATH",
                "ITERATION_SK",
                "CAL_ITERATION_LEVEL3",
                "CAL_ITERATION_PATH",
                "CAL_ITERATION_SK",
                "CAL_ITERATION_GRAIN",
                "CAL_START_DATE",
                "CAL_END_DATE",
                "FEATURE_IS_SPRINT_LEVEL",
                "SPRINT_TO_PI_RESOLVED",
                "MISSING_PI_CALENDAR",
                "TEAMNAME",
                "APP_NAME_RAW",
                "GROUPNAME",
                "INVESTMENT_DIMENSION",
                "ADO_YEAR",
                "DERIVED_FTE",
                "BUSINESS_VALUE",
                "CHANGED_AT",
            ]
            show_cols = [c for c in show_cols if c in df_show.columns]
            st.dataframe(
                df_show[show_cols] if show_cols else df_show,
                use_container_width=True,
                height=340,
                column_config={"DERIVED_FTE": st.column_config.NumberColumn("Derived FTE (SWAG)", format="%.2f")},
            )

            st.markdown("#### Summary")
            summary_view = st.selectbox(
                "Summary view",
                ["By iteration", "By year & iteration", "By team & iteration"],
                index=0,
                key="ado_explore_summary_view",
            )

            if summary_view == "By iteration":
                df_iter_sum = fetch_df(f"""
                  SELECT ITERATION_PATH,
                         COUNT(*) AS FEATURES
                  FROM ADO_FEATURES
                  {where_sql}
                  GROUP BY ITERATION_PATH
                  ORDER BY ITERATION_PATH
                """, tuple(params) if params else None)
                st.dataframe(df_iter_sum, use_container_width=True, height=300)
            elif summary_view == "By year & iteration":
                df_year_iter = fetch_df(f"""
                      WITH base AS (
                        SELECT ADO_YEAR, ITERATION_PATH
                        FROM ADO_FEATURES
                        {where_sql}
                      ),
                      labeled AS (
                        SELECT
                          ADO_YEAR,
                      /* Extract the first digit 1-4 after the letter 'I' (case-insensitive) */
                      CASE
                        WHEN ITERATION_PATH IS NULL OR CHARINDEX('I', UPPER(ITERATION_PATH)) = 0 THEN NULL
                        ELSE SUBSTRING(
                               SUBSTRING(UPPER(ITERATION_PATH), CHARINDEX('I', UPPER(ITERATION_PATH)), 10),
                               NULLIF(PATINDEX('%[1-4]%', SUBSTRING(UPPER(ITERATION_PATH), CHARINDEX('I', UPPER(ITERATION_PATH)), 10)), 0),
                               1
                             )
                          END AS ITER_NUM
                        FROM base
                      )
                      SELECT
                        ADO_YEAR AS YEAR,
                        CASE WHEN ITER_NUM IS NOT NULL THEN CONCAT('I', ITER_NUM) ELSE NULL END AS ITERATION,
                        COUNT(*) AS FEATURES
                      FROM labeled
                      GROUP BY ADO_YEAR, ITER_NUM
                      ORDER BY YEAR, TRY_CONVERT(INT, ITER_NUM)
                    """, tuple(params) if params else None)
                st.dataframe(df_year_iter, use_container_width=True, height=300)
            else:
                df_team_iter = fetch_df(f"""
                      WITH base AS (
                        SELECT
                          COALESCE(t.TEAMNAME, af.TEAM_RAW) AS TEAMNAME,
                          af.ITERATION_PATH
                        FROM ADO_FEATURES af
                        LEFT JOIN MAP_ADO_TEAM_TO_TCO_TEAM m
                          ON UPPER(m.ADO_TEAM_KEY) = COALESCE(
                               af.TEAM_VARIANT_KEY,
                           UPPER(NULLIF(LTRIM(RTRIM(af.TEAM_RAW)), ''))
                         )
                    LEFT JOIN TEAMS t ON t.TEAMID = m.TEAMID
                    {where_sql}
                  )
                      SELECT
                        TEAMNAME,
                        ITERATION_PATH,
                        COUNT(*) AS FEATURES
                      FROM base
                      GROUP BY TEAMNAME, ITERATION_PATH
                      ORDER BY TEAMNAME, ITERATION_PATH
                    """, tuple(params) if params else None)
                st.dataframe(df_team_iter, use_container_width=True, height=300)
        with tab_explore_po:
            st.caption(
                "Explorer v2 is the canonical feature-level demand view used by Expected costing."
            )
            profile_cfg = normalize_profile_config(get_active_ado_profile())
            org, project = _parse_org_project_from_odata_base(str(profile_cfg.get("odata_base_url") or ""))
            links_enabled = bool(org and project)
            if not links_enabled:
                st.info("Org/Project could not be derived from the active profile base URL. ADO links are disabled.")

            years_df = None
            try:
                years_df = fetch_df(
                    """
                    SELECT DISTINCT TRY_CONVERT(INT, YEAR) AS Y
                    FROM VW_TCO_FEATURE_DEMAND
                    WHERE TRY_CONVERT(INT, YEAR) IS NOT NULL
                    ORDER BY Y DESC
                    """,
                    None,
                )
            except Exception as e:
                st.warning(f"Could not read VW_TCO_FEATURE_DEMAND years: {e}")

            years = (
                years_df["Y"].dropna().astype(int).tolist()
                if years_df is not None and not years_df.empty and "Y" in years_df.columns
                else []
            )

            if not years:
                st.info("No rows found in VW_TCO_FEATURE_DEMAND yet. Build/refresh ADO demand first.")
                st.caption("Offline mode rebuilds demand using restored ADO_FEATURES + ADO_ITERATION_CALENDAR.")
                if st.button("Offline: Recompute demand mapping", key="ado_offline_recompute_demand"):
                    try:
                        recompute_feature_iteration_mapping()
                    except Exception as e:
                        toast_error(f"Offline recompute failed: {e}")
                    else:
                        _settings_post_write_refresh("offline_recompute_demand", bump_version=True)
                        try:
                            cnt = fetch_df("SELECT COUNT(*) AS N FROM VW_TCO_FEATURE_DEMAND", None)
                            n = int(cnt.iloc[0]["N"]) if cnt is not None and not cnt.empty else 0
                        except Exception:
                            n = 0
                        toast_success(f"Offline recompute complete. VW_TCO_FEATURE_DEMAND rows: {n:,d}.")
            else:
                year = st.selectbox("Year", options=years, index=0, key="ado_po_year")
                feats = load_explorer_feature_rows(
                    years=[int(year)],
                    cache_bust=int(st.session_state.get("ado_data_cache_bust", 0)),
                )

                if feats is None or feats.empty:
                    st.info("No in-scope features found for this year in Explorer v2.")
                else:
                    # SWAG input (points) is now surfaced directly from the canonical Explorer v2 view as SWAG_POINTS.
                    # Keep a safe fallback enrichment from ADO_FEATURES for older deployments.
                    if "SWAG_POINTS" not in feats.columns:
                        try:
                            pts_df = fetch_df(
                                """
                                SELECT
                                  FEATURE_ID,
                                  TRY_CONVERT(FLOAT, STORY_POINTS) AS STORY_POINTS
                                FROM ADO_FEATURES
                                WHERE FEATURE_ID IS NOT NULL
                                """,
                                None,
                            )
                        except Exception:
                            pts_df = None
                        if isinstance(pts_df, pd.DataFrame) and not pts_df.empty and "FEATURE_ID" in pts_df.columns:
                            pts = pts_df.copy()
                            pts["FEATURE_ID"] = pts["FEATURE_ID"].astype(str).str.strip()
                            pts["STORY_POINTS"] = pd.to_numeric(pts.get("STORY_POINTS"), errors="coerce")
                            pts = pts.drop_duplicates(subset=["FEATURE_ID"], keep="first")
                            feats = feats.copy()
                            feats["FEATURE_ID"] = feats["FEATURE_ID"].astype(str).str.strip()
                            feats = feats.merge(pts[["FEATURE_ID", "STORY_POINTS"]], on="FEATURE_ID", how="left")
                            feats["SWAG_POINTS"] = feats["STORY_POINTS"]

                last_sync = st.session_state.get("ado_last_sync_ts")
                st.caption(
                    f"Last sync timestamp: {last_sync or '—'} | Using cache bust: {st.session_state.get('ado_data_cache_bust', 0)}"
                )
                f1, f2, f3, f4 = st.columns([1.0, 1.6, 1.6, 1.6])
                with f1:
                    _ = st.selectbox("Year", options=[int(year)], index=0, disabled=True, key="ado_po_year_locked")
                with f2:
                    prog_opts = ["(All)"] + sorted(feats["PROGRAMNAME"].dropna().astype(str).unique().tolist())
                    prog_sel = st.selectbox("Program", options=prog_opts, index=0, key="ado_po_prog")
                with f3:
                    team_opts = ["(All)"] + sorted(feats["TEAMNAME"].dropna().astype(str).unique().tolist())
                    team_sel = st.selectbox("Team", options=team_opts, index=0, key="ado_po_team")
                with f4:
                    pi_opts = ["(All)"] + sorted(feats["PI_LABEL"].dropna().astype(str).unique().tolist())
                    pi_sel = st.selectbox("PI", options=pi_opts, index=0, key="ado_po_pi")

                filt = feats.copy()
                if prog_sel != "(All)":
                    filt = filt[filt["PROGRAMNAME"] == prog_sel].copy()
                if team_sel != "(All)":
                    filt = filt[filt["TEAMNAME"] == team_sel].copy()
                if pi_sel != "(All)":
                    filt = filt[filt["PI_LABEL"] == pi_sel].copy()

                for c in ["EPIC_TITLE", "EPIC_ID", "EPIC_STATE", "PARENT_ID"]:
                    if c not in filt.columns:
                        filt[c] = None

                if "SWAG_POINTS" in filt.columns:
                    filt["SWAG_POINTS"] = pd.to_numeric(filt.get("SWAG_POINTS"), errors="coerce")
                if "STORY_POINTS" not in filt.columns:
                    filt["STORY_POINTS"] = pd.to_numeric(filt.get("SWAG_POINTS"), errors="coerce")
                else:
                    filt["STORY_POINTS"] = pd.to_numeric(filt.get("STORY_POINTS"), errors="coerce")
                filt["DERIVED_FTE"] = pd.to_numeric(
                    filt.get("DERIVED_FTE_FEATURE", filt.get("DERIVED_FTE")), errors="coerce"
                )
                if "DERIVED_FTE_FEATURE_VELOCITY" in filt.columns:
                    filt["DERIVED_FTE_FEATURE_VELOCITY"] = pd.to_numeric(
                        filt.get("DERIVED_FTE_FEATURE_VELOCITY"), errors="coerce"
                    )
                else:
                    filt["DERIVED_FTE_FEATURE_VELOCITY"] = pd.to_numeric(
                        filt.get("DERIVED_FTE"), errors="coerce"
                    )
                filt["DERIVED_FTE"] = filt["DERIVED_FTE"].fillna(0.0)
                filt["DERIVED_FTE_FEATURE_VELOCITY"] = filt["DERIVED_FTE_FEATURE_VELOCITY"].fillna(filt["DERIVED_FTE"])
                if "SWAG_READY" in filt.columns:
                    filt["SWAG_READY"] = filt.get("SWAG_READY", False).fillna(False).astype(bool)
                else:
                    filt["SWAG_READY"] = filt["STORY_POINTS"].fillna(0.0) > 0
                for col in ["FEATURE_ID", "PARENT_ID", "EPIC_ID"]:
                    if col in filt.columns:
                        filt[col] = pd.to_numeric(filt[col], errors="coerce").astype("Int64")

                bad_cols = [c for c in filt.columns if c.upper().startswith(("EFFORT_", "MANUAL_"))]
                if bad_cols:
                    st.error(f"Explorer v2 dataset unexpectedly contains legacy columns: {', '.join(bad_cols)}")

                total_features = int(filt["FEATURE_ID"].nunique())
                swag_ready_ct = int(filt["SWAG_READY"].fillna(False).astype(bool).sum()) if total_features else 0
                derived_fte_total = float(filt["DERIVED_FTE"].sum()) if total_features else 0.0
                derived_fte_velocity_total = float(filt["DERIVED_FTE_FEATURE_VELOCITY"].sum()) if total_features else 0.0
                derived_delta = derived_fte_velocity_total - derived_fte_total

                m1, m2, m3, m4 = st.columns(4)
                m1.metric("Features (in scope)", f"{total_features:,d}")
                m2.metric("SWAG-ready features", f"{swag_ready_ct:,d}")
                m3.metric("Derived FTE total (SWAG)", f"{derived_fte_total:,.2f}" if derived_fte_total else "0.00")
                m4.metric(
                    "Derived FTE total (Velocity)",
                    f"{derived_fte_velocity_total:,.2f}" if derived_fte_velocity_total else "0.00",
                    delta=f"{derived_delta:+,.2f}",
                )

                not_ready = filt[~filt["SWAG_READY"].fillna(False).astype(bool)].copy()
                if not not_ready.empty:
                    st.warning(
                        f"SWAG not ready for {len(not_ready):,d} feature(s) in this selection (Derived FTE = 0 \u2192 no Expected demand/cost)."
                    )

                tab_flat, tab_epics = st.tabs(["Features (flat)", "Epics \u2192 Features"])

                with tab_flat:
                    flat = filt.copy()
                    flat["EPIC_TITLE"] = flat["EPIC_TITLE"].fillna("No Epic")
                    if "STORY_POINTS" not in flat.columns:
                        flat["STORY_POINTS"] = pd.to_numeric(flat.get("SWAG_POINTS"), errors="coerce")
                    if links_enabled:
                        flat["FEATURE_ID"] = flat["FEATURE_ID"].apply(
                            lambda v: _ado_workitem_url(org, project, v) if pd.notna(v) else None
                        )
                        flat["EPIC_ID"] = flat["EPIC_ID"].apply(
                            lambda v: _ado_workitem_url(org, project, v) if pd.notna(v) else None
                        )
                    show_cols = [
                        "EPIC_ID",
                        "EPIC_TITLE",
                        "FEATURE_ID",
                        "FEATURE_TITLE",
                        "STATE",
                        "PI_LABEL",
                        "PROGRAM_RAW",
                        "TEAMNAME",
                        "PROGRAMNAME",
                        "TEAM_RAW",
                        "AREA_LEVEL1_RAW",
                        "AREA_LEVEL2_RAW",
                        "AREA_LEVEL3_RAW",
                        "AREA_LEVEL4_RAW",
                        "AREA_PATH_RAW",
                        "STORY_POINTS",
                        "DERIVED_FTE",
                        "DERIVED_FTE_FEATURE_VELOCITY",
                    ]
                    show_cols = [c for c in show_cols if c in flat.columns]
                    sort_cols = [c for c in ["EPIC_TITLE", "PI_LABEL", "TEAMNAME", "FEATURE_TITLE"] if c in flat.columns]
                    flat_view = flat[show_cols].sort_values(sort_cols) if show_cols else flat
                    if len(flat_view.index) > 2000:
                        st.warning("Showing first 2,000 rows. Refine filters to narrow results.")
                        flat_view = flat_view.head(2000)
                    column_config = {
                        "STORY_POINTS": st.column_config.NumberColumn("Story Points", format="%.0f"),
                        "DERIVED_FTE": st.column_config.NumberColumn("Derived FTE (SWAG)", format="%.2f"),
                        "DERIVED_FTE_FEATURE_VELOCITY": st.column_config.NumberColumn(
                            "Derived FTE (Velocity)", format="%.2f"
                        ),
                    }
                    if links_enabled:
                        column_config.update(
                            {
                                "EPIC_ID": st.column_config.LinkColumn(
                                    "EPIC_ID", display_text=r".*/([^/]+)$"
                                ),
                                "FEATURE_ID": st.column_config.LinkColumn(
                                    "FEATURE_ID", display_text=r".*/([^/]+)$"
                                ),
                            }
                        )
                        st.data_editor(
                            flat_view,
                            use_container_width=True,
                            height=520,
                            hide_index=True,
                            disabled=True,
                            column_config=column_config,
                        )
                    else:
                        st.dataframe(
                            flat_view,
                            use_container_width=True,
                            height=520,
                            hide_index=True,
                            column_config=column_config,
                        )

                    with tab_epics:
                        epic_df = filt.copy()
                        epic_df["EPIC_TITLE"] = epic_df["EPIC_TITLE"].fillna("No Epic")
                        epic_df["EPIC_ID"] = pd.to_numeric(epic_df["EPIC_ID"], errors="coerce").fillna(-1).astype(int)
                        points_col = "STORY_POINTS" if "STORY_POINTS" in epic_df.columns else ("SWAG_POINTS" if "SWAG_POINTS" in epic_df.columns else None)
                        if points_col:
                            epic_df[points_col] = pd.to_numeric(epic_df[points_col], errors="coerce").fillna(0.0)
                        epic_df["DERIVED_FTE_FALLBACK"] = pd.to_numeric(epic_df.get("DERIVED_FTE"), errors="coerce").fillna(0.0)
                        epic_summary = (
                            epic_df.groupby(["EPIC_ID", "EPIC_TITLE"], dropna=False)
                            .agg(
                                FEATURES_COUNT=("FEATURE_ID", "nunique"),
                                TOTAL_POINTS=(points_col, "sum") if points_col else ("FEATURE_ID", "count"),
                                TOTAL_DERIVED_FTE=("DERIVED_FTE_FALLBACK", "sum"),
                                PROGRAMS=("PROGRAMNAME", "nunique"),
                                TEAMS=("TEAMNAME", "nunique"),
                            )
                            .reset_index()
                        )
                        epic_summary = epic_summary.sort_values(["FEATURES_COUNT", "EPIC_TITLE"], ascending=[False, True])
                        epic_summary["EPIC_ID_VALUE"] = epic_summary["EPIC_ID"]
                        if links_enabled:
                            epic_summary["EPIC_ID"] = epic_summary["EPIC_ID"].apply(
                                lambda v: _ado_workitem_url(org, project, v) if int(v) != -1 else None
                            )
                        summary_cols = [
                            "EPIC_ID",
                            "EPIC_TITLE",
                            "FEATURES_COUNT",
                            "TOTAL_POINTS",
                            "TOTAL_DERIVED_FTE",
                            "PROGRAMS",
                            "TEAMS",
                        ]
                        summary_cols = [c for c in summary_cols if c in epic_summary.columns]
                        summary_config = {
                            "TOTAL_POINTS": st.column_config.NumberColumn("Total Points", format="%.0f"),
                            "TOTAL_DERIVED_FTE": st.column_config.NumberColumn("Total Derived FTE", format="%.2f"),
                        }
                        if links_enabled:
                            summary_config["EPIC_ID"] = st.column_config.LinkColumn(
                                "EPIC_ID", display_text=r".*/([^/]+)$"
                            )
                            st.data_editor(
                                epic_summary[summary_cols] if summary_cols else epic_summary,
                                use_container_width=True,
                                height=320,
                                hide_index=True,
                                disabled=True,
                                column_config=summary_config,
                            )
                        else:
                            st.dataframe(
                                epic_summary[summary_cols] if summary_cols else epic_summary,
                                use_container_width=True,
                                height=320,
                                hide_index=True,
                                column_config=summary_config,
                            )
                        epic_options = []
                        epic_label_to_key: Dict[str, Tuple[int, str]] = {}
                        for _, row in epic_summary.iterrows():
                            eid = int(row.get("EPIC_ID_VALUE", -1))
                            etitle = str(row.get("EPIC_TITLE") or "No Epic")
                            label = f"{etitle} (#{eid})" if eid != -1 else etitle
                            epic_options.append(label)
                            epic_label_to_key[label] = (eid, etitle)
                        selected_epic = st.selectbox("Select Epic", options=epic_options, index=0 if epic_options else None, key="ado_po_epic_select")
                        if selected_epic:
                            eid, etitle = epic_label_to_key.get(selected_epic, (-1, "No Epic"))
                            detail = epic_df[(epic_df["EPIC_ID"] == eid) & (epic_df["EPIC_TITLE"] == etitle)].copy()
                            if "STORY_POINTS" not in detail.columns:
                                detail["STORY_POINTS"] = pd.to_numeric(detail.get("SWAG_POINTS"), errors="coerce")
                            if links_enabled:
                                detail["FEATURE_ID"] = detail["FEATURE_ID"].apply(
                                    lambda v: _ado_workitem_url(org, project, v) if pd.notna(v) else None
                                )
                            detail_cols = [
                                "FEATURE_ID",
                                "FEATURE_TITLE",
                                "STATE",
                                "PI_LABEL",
                                "TEAMNAME",
                                "PROGRAMNAME",
                                "STORY_POINTS",
                                "DERIVED_FTE",
                                "DERIVED_FTE_FEATURE_VELOCITY",
                            ]
                            detail_cols = [c for c in detail_cols if c in detail.columns]
                            if len(detail.index) > 500:
                                st.warning("Showing first 500 features for this epic.")
                                detail = detail.head(500)
                            detail_config = {
                                "STORY_POINTS": st.column_config.NumberColumn("Story Points", format="%.0f"),
                                "DERIVED_FTE": st.column_config.NumberColumn("Derived FTE (SWAG)", format="%.2f"),
                                "DERIVED_FTE_FEATURE_VELOCITY": st.column_config.NumberColumn(
                                    "Derived FTE (Velocity)", format="%.2f"
                                ),
                            }
                            if links_enabled:
                                detail_config["FEATURE_ID"] = st.column_config.LinkColumn(
                                    "FEATURE_ID", display_text=r".*/([^/]+)$"
                                )
                                st.data_editor(
                                    detail[detail_cols] if detail_cols else detail,
                                    use_container_width=True,
                                    height=520,
                                    hide_index=True,
                                    disabled=True,
                                    column_config=detail_config,
                                )
                            else:
                                st.dataframe(
                                    detail[detail_cols] if detail_cols else detail,
                                    use_container_width=True,
                                    height=520,
                                    hide_index=True,
                                    column_config=detail_config,
                                )

# -------------------------
# Tabs (top-level)
# -------------------------
def render_automation_tab() -> None:
    st.subheader("Automation")
    st.caption(
        "Configure background ADO sync/upsert and pipeline refresh schedules. "
        "Jobs are queued in Control DB and executed by your Azure background worker."
    )

    if not is_admin:
        st.warning("Automation settings are available to admins only.")
        return

    if not control_db_available():
        st.error("Control DB is not configured. Set CONTROL_MSSQL_* secrets/env vars first.")
        return

    try:
        ensure_control_schema()
        ensure_control_automation_tables()
    except Exception as e:
        st.error(f"Control DB automation schema check failed: {e}")
        return

    try:
        df_settings = list_automation_settings()
    except Exception as e:
        st.error(f"Could not load automation settings: {e}")
        return

    if df_settings is None or df_settings.empty:
        st.info("No portfolios found in Control DB.")
        return

    rows = []
    for _, row in df_settings.iterrows():
        key = str(row.get("PORTFOLIO_KEY") or "").strip()
        name = str(row.get("DISPLAY_NAME") or "").strip()
        label = f"{key} - {name}" if name and name.upper() != key.upper() else key
        rows.append((key, label))
    options = [key for key, _ in rows]
    labels = [label for _, label in rows]
    default_idx = 0
    current_key = str(st.session_state.get("active_portfolio_key") or "").strip()
    if current_key in options:
        default_idx = options.index(current_key)

    picked_label = st.selectbox("Portfolio", labels, index=default_idx, key="automation_portfolio_pick")
    picked_key = options[labels.index(picked_label)]
    selected = df_settings[df_settings["PORTFOLIO_KEY"].astype(str).str.strip() == picked_key]
    row = selected.iloc[0] if not selected.empty else {}

    def _b(name: str, default: bool) -> bool:
        try:
            return bool(int(row.get(name))) if row.get(name) is not None else bool(default)
        except Exception:
            return bool(default)

    def _i(name: str, default: int, min_v: int, max_v: int) -> int:
        try:
            val = int(row.get(name)) if row.get(name) is not None else int(default)
        except Exception:
            val = int(default)
        return max(min_v, min(max_v, val))

    enable_automation = _b("ENABLE_AUTOMATION", False)
    hourly_enabled = _b("HOURLY_SYNC_ENABLED", False)
    hourly_minute = _i("HOURLY_SYNC_MINUTE", 10, 0, 59)
    nightly_enabled = _b("NIGHTLY_SYNC_ENABLED", True)
    nightly_hour = _i("NIGHTLY_SYNC_HOUR_UTC", 2, 0, 23)
    nightly_minute = _i("NIGHTLY_SYNC_MINUTE", 15, 0, 59)
    refresh_on_change = _b("REFRESH_ON_CHANGE", True)

    st.markdown("### Schedule")
    with st.form("automation_settings_form"):
        c1, c2 = st.columns(2)
        with c1:
            form_enable = st.toggle("Enable automation for this portfolio", value=enable_automation)
            form_hourly = st.toggle("Hourly incremental sync + upsert", value=hourly_enabled)
            form_hourly_min = st.number_input(
                "Hourly minute (UTC)",
                min_value=0,
                max_value=59,
                step=1,
                value=hourly_minute,
                disabled=not form_hourly,
            )
        with c2:
            form_nightly = st.toggle("Nightly full sync + upsert", value=nightly_enabled)
            form_nightly_hour = st.number_input(
                "Nightly hour (UTC)",
                min_value=0,
                max_value=23,
                step=1,
                value=nightly_hour,
                disabled=not form_nightly,
            )
            form_nightly_min = st.number_input(
                "Nightly minute (UTC)",
                min_value=0,
                max_value=59,
                step=1,
                value=nightly_minute,
                disabled=not form_nightly,
            )
        form_refresh_on_change = st.toggle(
            "Run refresh pipeline only when sync/upsert changed data",
            value=refresh_on_change,
        )
        save = st.form_submit_button("Save automation settings", type="primary")

    if save:
        updated_by = str(st.session_state.get("auth_user", {}).get("email") or st.session_state.get("user_email") or "").strip()
        try:
            upsert_automation_setting(
                picked_key,
                enable_automation=bool(form_enable),
                hourly_sync_enabled=bool(form_hourly),
                hourly_sync_minute=int(form_hourly_min),
                nightly_sync_enabled=bool(form_nightly),
                nightly_sync_hour_utc=int(form_nightly_hour),
                nightly_sync_minute=int(form_nightly_min),
                refresh_on_change=bool(form_refresh_on_change),
                updated_by=updated_by,
            )
            st.success("Automation settings saved.")
            st.rerun()
        except Exception as e:
            st.error(f"Save failed: {e}")

    st.markdown("### Run now")
    st.caption("Queues a job for your Azure worker. This does not require manual Azure operations by end users.")
    st.caption(
        "Freshness contract: successful refresh pipeline runs bump DATA_VERSION; "
        "cached reads refresh automatically on next render."
    )
    st.caption(
        "Refresh pipeline is repair/recompute tooling. It is not required after normal app CRUD edits."
    )
    st.caption(
        "Use this when: (1) composition freshness is stale, (2) bulk/backfill data was loaded externally, "
        "(3) sync/upsert changed large volumes and you want immediate recompute."
    )
    rc1, rc2, rc3 = st.columns(3)
    requested_by = str(st.session_state.get("auth_user", {}).get("email") or st.session_state.get("user_email") or "").strip()
    if rc1.button("Queue incremental sync", key="automation_run_now_sync_inc", use_container_width=True):
        try:
            run_id = queue_automation_run(picked_key, "SYNC_INCREMENTAL", requested_by=requested_by, trigger_type="manual")
            st.success(f"Queued incremental sync (run #{run_id}).")
        except Exception as e:
            st.error(f"Queue failed: {e}")
    if rc2.button("Queue full sync", key="automation_run_now_sync_full", use_container_width=True):
        try:
            run_id = queue_automation_run(picked_key, "SYNC_FULL", requested_by=requested_by, trigger_type="manual")
            st.success(f"Queued full sync (run #{run_id}).")
        except Exception as e:
            st.error(f"Queue failed: {e}")
    if rc3.button("Queue refresh pipeline", key="automation_run_now_refresh", use_container_width=True):
        try:
            run_id = queue_automation_run(picked_key, "REFRESH_PIPELINE", requested_by=requested_by, trigger_type="manual")
            st.success(f"Queued refresh pipeline (run #{run_id}).")
        except Exception as e:
            st.error(f"Queue failed: {e}")

    st.markdown("### Current settings")
    try:
        show_cols = [
            "PORTFOLIO_KEY",
            "DISPLAY_NAME",
            "STATUS",
            "ENABLE_AUTOMATION",
            "HOURLY_SYNC_ENABLED",
            "HOURLY_SYNC_MINUTE",
            "NIGHTLY_SYNC_ENABLED",
            "NIGHTLY_SYNC_HOUR_UTC",
            "NIGHTLY_SYNC_MINUTE",
            "REFRESH_ON_CHANGE",
            "UPDATED_AT",
            "UPDATED_BY",
        ]
        st.dataframe(df_settings[show_cols], use_container_width=True, hide_index=True)
    except Exception:
        st.dataframe(df_settings, use_container_width=True, hide_index=True)

    st.markdown("### Recent runs")
    try:
        runs = list_automation_runs(limit=150)
        if runs is None or runs.empty:
            st.info("No runs queued/executed yet.")
        else:
            st.dataframe(runs, use_container_width=True, hide_index=True)
    except Exception as e:
        st.warning(f"Could not load run history: {e}")

    with st.expander("Azure worker setup", expanded=False):
        st.caption(
            "Run a timer-triggered Azure Function/WebJob every 5 minutes. "
            "The worker should read CONTROL_AUTOMATION_SETTINGS, enqueue due jobs, and process queued runs from CONTROL_AUTOMATION_RUNS."
        )
        st.code(
            "python scripts/automation_worker.py --once\n"
            "# or continuously:\n"
            "python scripts/automation_worker.py --loop --interval-sec 300"
        )


_tab_ado, _tab_reconciliation, _tab_automation, _tab_apptio, _tab_loader = st.tabs(
    ["ADO", "Reconciliation", "Automation", "APPTIO", "Data Loader"]
)

with _tab_ado:
    _ado_tab_profiles, _ado_tab_validate, _ado_tab_sync, _ado_tab_mapping, _ado_tab_advanced = st.tabs(
        ["Profiles", "Validate", "Sync", "Mapping", "Advanced"]
    )
    with _ado_tab_profiles:
        render_ado_profiles_tab()
    with _ado_tab_validate:
        render_ado_validate_tab()
    with _ado_tab_sync:
        render_ado_sync_tab()
    with _ado_tab_mapping:
        render_ado_mapping_tab()
    with _ado_tab_advanced:
        render_ado_advanced_tab()

with _tab_reconciliation:
    render_reconciliation_tab(fetch_df=fetch_df, scope=None)

with _tab_automation:
    render_automation_tab()

# =========================
# Tab: 📈 APPTIO
# =========================
with _tab_apptio:
    st.subheader("APPTIO Actuals (Work ID → Program)")
    st.caption("Flow: 1) Upload file  2) Review mapping coverage  3) Preview totals  4) Save actuals")

    with st.expander("PI Calendar (monthly → PI allocation)", expanded=False):
        st.caption("Controls which PI calendar root is used when mapping monthly Apptio program NWF into PIs.")

        # Year selector (used only to populate root options).
        try:
            df_years = fetch_df(
                """
                SELECT DISTINCT TRY_CONVERT(INT, ADO_YEAR) AS YEAR
                FROM dbo.ADO_FEATURES
                WHERE ADO_YEAR IS NOT NULL
                ORDER BY TRY_CONVERT(INT, ADO_YEAR)
                """
            )
            years_opts = (
                sorted(pd.to_numeric(df_years.get("YEAR"), errors="coerce").dropna().astype(int).unique().tolist())
                if df_years is not None and not df_years.empty
                else []
            )
        except Exception:
            years_opts = []

        current_year = date.today().year
        default_year = current_year if current_year in years_opts else (years_opts[-1] if years_opts else current_year)
        pick_year = st.selectbox("Year", years_opts or [default_year], index=0, key="settings_pi_calendar_year")

        roots = list_iteration_roots(int(pick_year)) if pick_year is not None else []
        auto_root = dominant_iteration_root(int(pick_year)) if pick_year is not None else None
        if auto_root:
            st.caption(f"Auto (dominant from ADO) would select: `{auto_root}`")

        options = ["AUTO (dominant from ADO)"] + roots
        override = st.session_state.get("pi_calendar_root_override")
        default_choice = override if override and str(override).strip() in options else "AUTO (dominant from ADO)"
        try:
            default_idx = options.index(default_choice)
        except Exception:
            default_idx = 0
        picked_root = st.selectbox("PI Calendar Root", options, index=default_idx, key="settings_pi_calendar_root")

        if str(picked_root).strip().upper().startswith("AUTO"):
            st.session_state["pi_calendar_root_override"] = None
        else:
            st.session_state["pi_calendar_root_override"] = str(picked_root).strip()

    try:
        ensure_tables()
        ensure_apptio_actuals_lines_table()
        ensure_map_apptio_to_cost_type_table()
    except Exception as e:
        st.warning(f"Could not ensure base tables: {e}")

    mapping_df = list_program_apptio_workids()
    prog_opts = []
    if mapping_df is not None and not mapping_df.empty:
        prog_opts = sorted(mapping_df["PROGRAMNAME"].dropna().astype(str).unique().tolist())
        with st.expander("Current Work ID mappings (edit in Programs page)", expanded=False):
            st.dataframe(mapping_df, use_container_width=True, height=240)
    else:
        st.warning("No Work ID mappings found. Add mappings on the Programs page before importing actuals.")

    with st.expander("Apptio mapping (COST_TYPE / SUBTYPE)", expanded=False):
        st.caption("Map Apptio dimensions to the same COST_TYPE/SUBTYPE taxonomy used in Programs → Additional Costs.")
        try:
            yrs_df = fetch_df("SELECT DISTINCT TRY_CONVERT(INT, FISCAL_YEAR) AS YEAR FROM APPTIO_ACTUALS_LINES WHERE FISCAL_YEAR IS NOT NULL ORDER BY YEAR DESC", None)
            years_map = sorted(pd.to_numeric(yrs_df.get("YEAR"), errors="coerce").dropna().astype(int).unique().tolist()) if yrs_df is not None and not yrs_df.empty else []
        except Exception:
            years_map = []
        if not years_map:
            try:
                yrs_df = fetch_df("SELECT DISTINCT TRY_CONVERT(INT, FISCAL_YEAR) AS YEAR FROM APPTIO_ACTUALS WHERE FISCAL_YEAR IS NOT NULL ORDER BY YEAR DESC", None)
                years_map = sorted(pd.to_numeric(yrs_df.get("YEAR"), errors="coerce").dropna().astype(int).unique().tolist()) if yrs_df is not None and not yrs_df.empty else []
            except Exception:
                years_map = []
        map_year = st.selectbox("Fiscal year (for mapping preview)", years_map or [date.today().year], index=0, key="apptio_map_year")

        planned_opts = [""] + list(PLANNED_NWF_SUBCOMPONENTS)

        def _pair_to_planned(ct: Any, stp: Any) -> str:
            cost_type = str(ct or "").strip()
            subtype = str(stp or "").strip()
            if not cost_type:
                return ""
            if cost_type.lower() == "cloud" and subtype in {"AWS", "Azure"}:
                raw = f"Cloud {subtype}".strip()
            elif cost_type == "NWF" and subtype == "Other":
                raw = "NWF Other"
            elif cost_type in {"Invoices", "Contractor CS", "MSP", "Travel", "Infra"} and subtype == "":
                raw = cost_type
            else:
                raw = f"{cost_type} {subtype}".strip()
            return normalize_nwf_subcomponent(raw)

        def _planned_to_pair(label: Any) -> tuple[str, str]:
            lab = normalize_nwf_subcomponent(str(label or ""))
            if lab == "Cloud AWS":
                return ("Cloud", "AWS")
            if lab == "Cloud Azure":
                return ("Cloud", "Azure")
            if lab == "NWF Other":
                return ("NWF", "Other")
            if lab in {"Invoices", "Contractor CS", "MSP", "Travel", "Infra"}:
                return (lab, "")
            return ("NWF", "Other")

        tab_prod, tab_ledger, tab_rec = st.tabs(["Products", "Ledger fallback", "Reconciliation"])

        with tab_prod:
            st.caption("Tip: map MSP via Product ID (ledger fallback can mix multiple spend types under the same ledger).")
            try:
                prod_df = fetch_df(
                    """
                    SELECT
                      LEDGER_ACCOUNT_L3_DESC,
                      PRODUCT_ID,
                      MAX(PRODUCT_NAME) AS PRODUCT_NAME,
                      COUNT(*) AS ROWS_N,
                      SUM(COALESCE(TRY_CONVERT(FLOAT, AMOUNT), 0.0)) AS AMOUNT
                    FROM APPTIO_ACTUALS_LINES
                    WHERE TRY_CONVERT(INT, FISCAL_YEAR) = ?
                      AND PRODUCT_ID IS NOT NULL
                      AND LTRIM(RTRIM(PRODUCT_ID)) <> ''
                    GROUP BY PRODUCT_ID
                    , LEDGER_ACCOUNT_L3_DESC
                    ORDER BY AMOUNT DESC
                    """,
                    (int(map_year),),
                )
            except Exception:
                prod_df = pd.DataFrame()
            try:
                map_df = fetch_df(
                    """
                    SELECT
                      MAP_ID, FISCAL_YEAR, PRODUCT_ID, LEDGER_ACCOUNT_CONTAINS,
                      COST_TYPE, SUBTYPE, IS_ACTIVE, NOTES
                    FROM MAP_APPTIO_TO_COST_TYPE
                    WHERE PRODUCT_ID IS NOT NULL
                    """,
                    None,
                )
            except Exception:
                map_df = pd.DataFrame()

            try:
                ledger_map_df = fetch_df(
                    """
                    SELECT
                      MAP_ID, FISCAL_YEAR, PRODUCT_ID, LEDGER_ACCOUNT_CONTAINS,
                      COST_TYPE, SUBTYPE, IS_ACTIVE, NOTES
                    FROM MAP_APPTIO_TO_COST_TYPE
                    WHERE PRODUCT_ID IS NULL AND LEDGER_ACCOUNT_CONTAINS IS NOT NULL
                    """,
                    None,
                )
            except Exception:
                ledger_map_df = pd.DataFrame()

            map_id_by_pid: dict[str, Any] = {}
            if prod_df is None or prod_df.empty:
                st.info("No product IDs found in APPTIO_ACTUALS_LINES for this year yet. Upload parses the file; you must click 'Save Apptio actuals' to write APPTIO_ACTUALS_LINES. Then return here.")
                prod_editor_src = pd.DataFrame(
                    columns=[
                        "LEDGER_ACCOUNT_L3_DESC",
                        "PRODUCT_ID",
                        "PRODUCT_NAME",
                        "AMOUNT",
                        "ROWS_N",
                        "PLANNED_NWF_TYPE",
                        "EFFECTIVE_NWF_TYPE",
                        "MAPPED_BY",
                        "IS_ACTIVE",
                        "NOTES",
                    ]
                )
            else:
                p = prod_df.copy()
                p["LEDGER_ACCOUNT_L3_DESC"] = p.get("LEDGER_ACCOUNT_L3_DESC", "").fillna("").astype(str).str.strip()
                p["PRODUCT_ID"] = p.get("PRODUCT_ID", "").fillna("").astype(str).str.strip()
                p["PRODUCT_NAME"] = p.get("PRODUCT_NAME", "").fillna("").astype(str).str.strip()
                p["AMOUNT"] = pd.to_numeric(p.get("AMOUNT"), errors="coerce").fillna(0.0)
                p["ROWS_N"] = pd.to_numeric(p.get("ROWS_N"), errors="coerce").fillna(0).astype(int)
                p = p[p["PRODUCT_ID"].ne("")].copy()

                m = map_df.copy() if map_df is not None and not map_df.empty else pd.DataFrame(columns=["MAP_ID", "FISCAL_YEAR", "PRODUCT_ID", "COST_TYPE", "SUBTYPE", "IS_ACTIVE", "NOTES"])
                if not m.empty:
                    m["PRODUCT_ID"] = m.get("PRODUCT_ID", "").fillna("").astype(str).str.strip()
                    m["FISCAL_YEAR"] = pd.to_numeric(m.get("FISCAL_YEAR"), errors="coerce")
                    m["_pri"] = (m["FISCAL_YEAR"].fillna(int(map_year)) != int(map_year)).astype(int)
                    m["_active_pri"] = (pd.to_numeric(m.get("IS_ACTIVE"), errors="coerce").fillna(1).astype(int) != 1).astype(int)
                    m = m.sort_values(["PRODUCT_ID", "_active_pri", "_pri"]).drop(columns=["_pri", "_active_pri"])
                    m_best = m.groupby("PRODUCT_ID", dropna=False).head(1).copy()
                else:
                    m_best = pd.DataFrame(columns=["MAP_ID", "FISCAL_YEAR", "PRODUCT_ID", "COST_TYPE", "SUBTYPE", "IS_ACTIVE", "NOTES"])

                map_id_by_pid = (
                    m_best.set_index("PRODUCT_ID")["MAP_ID"].to_dict()
                    if m_best is not None and not m_best.empty and "MAP_ID" in m_best.columns and "PRODUCT_ID" in m_best.columns
                    else {}
                )

                prod_editor_src = p.merge(m_best.drop(columns=["MAP_ID"], errors="ignore"), on="PRODUCT_ID", how="left")
                prod_editor_src["IS_ACTIVE"] = prod_editor_src.get("IS_ACTIVE").fillna(1).astype(int)
                prod_editor_src["NOTES"] = prod_editor_src.get("NOTES", "").fillna("").astype(str).str.strip()
                prod_editor_src["PLANNED_NWF_TYPE"] = prod_editor_src.apply(
                    lambda r: _pair_to_planned(r.get("COST_TYPE"), r.get("SUBTYPE")), axis=1
                )

                # Compute effective mapping (product override unless it is blank or NWF Other; otherwise ledger fallback; else default).
                lm = ledger_map_df.copy() if ledger_map_df is not None else pd.DataFrame()
                if lm is not None and not lm.empty:
                    lm["LEDGER_ACCOUNT_CONTAINS"] = lm.get("LEDGER_ACCOUNT_CONTAINS", "").fillna("").astype(str).str.strip()
                    lm = lm[lm["LEDGER_ACCOUNT_CONTAINS"].ne("")].copy()
                    lm["FISCAL_YEAR"] = pd.to_numeric(lm.get("FISCAL_YEAR"), errors="coerce")
                    lm["IS_ACTIVE"] = pd.to_numeric(lm.get("IS_ACTIVE"), errors="coerce").fillna(1).astype(int)
                    lm["PLANNED_NWF_TYPE"] = lm.apply(lambda r: _pair_to_planned(r.get("COST_TYPE"), r.get("SUBTYPE")), axis=1)

                def _best_ledger_planned(desc: str) -> tuple[str, str]:
                    if lm is None or lm.empty:
                        return ("", "")
                    d_up = str(desc or "").strip().upper()
                    if not d_up:
                        return ("", "")
                    cand = lm[lm["LEDGER_ACCOUNT_CONTAINS"].astype(str).str.upper().apply(lambda s: s and s in d_up)].copy()
                    if cand.empty:
                        return ("", "")
                    # Prefer active, year-specific, and the longest match.
                    cand["_apri"] = (cand["IS_ACTIVE"] != 1).astype(int)
                    cand["_ypri"] = (cand["FISCAL_YEAR"].fillna(int(map_year)) != int(map_year)).astype(int)
                    cand["_lpri"] = cand["LEDGER_ACCOUNT_CONTAINS"].astype(str).str.len().mul(-1)
                    cand = cand.sort_values(["_apri", "_ypri", "_lpri"])
                    lab = str(cand.iloc[0].get("PLANNED_NWF_TYPE") or "").strip()
                    return (lab, "Ledger")

                eff_type: list[str] = []
                eff_by: list[str] = []
                for _, rr in prod_editor_src.iterrows():
                    prod_planned = str(rr.get("PLANNED_NWF_TYPE") or "").strip()
                    prod_override = prod_planned if prod_planned and prod_planned != "NWF Other" else ""
                    if prod_override:
                        eff_type.append(prod_override)
                        eff_by.append("Product")
                        continue
                    ledger_planned, by = _best_ledger_planned(str(rr.get("LEDGER_ACCOUNT_L3_DESC") or "").strip())
                    if ledger_planned:
                        eff_type.append(ledger_planned)
                        eff_by.append(by)
                    else:
                        eff_type.append("NWF Other")
                        eff_by.append("Default")
                prod_editor_src["EFFECTIVE_NWF_TYPE"] = eff_type
                prod_editor_src["MAPPED_BY"] = eff_by

                # UX: small filters help quickly map MSP (by Product ID/Name).
                f1, f2 = st.columns([2, 1])
                prod_filter = f1.text_input("Filter (optional)", value="", key="apptio_products_filter")
                show_unmapped_only = f2.checkbox("Show unmapped only", value=False, key="apptio_products_unmapped_only")
                if prod_filter.strip():
                    q = prod_filter.strip().lower()
                    mask = (
                        prod_editor_src["LEDGER_ACCOUNT_L3_DESC"].fillna("").astype(str).str.lower().str.contains(q)
                        | prod_editor_src["PRODUCT_ID"].fillna("").astype(str).str.lower().str.contains(q)
                        | prod_editor_src["PRODUCT_NAME"].fillna("").astype(str).str.lower().str.contains(q)
                    )
                    prod_editor_src = prod_editor_src[mask].copy()
                if show_unmapped_only:
                    prod_editor_src = prod_editor_src[prod_editor_src["PLANNED_NWF_TYPE"].fillna("").astype(str).str.strip().eq("")].copy()

                # Required display order (then editable mapping fields).
                view_cols = [
                    "LEDGER_ACCOUNT_L3_DESC",
                    "PRODUCT_ID",
                    "PRODUCT_NAME",
                    "AMOUNT",
                    "ROWS_N",
                    "PLANNED_NWF_TYPE",
                    "EFFECTIVE_NWF_TYPE",
                    "MAPPED_BY",
                    "IS_ACTIVE",
                    "NOTES",
                ]
                view_cols = [c for c in view_cols if c in prod_editor_src.columns]
                prod_editor_src = prod_editor_src[view_cols].copy()

            edited_prod = st.data_editor(
                prod_editor_src,
                key="apptio_map_products_editor",
                use_container_width=True,
                num_rows="fixed",
                column_config={
                    "LEDGER_ACCOUNT_L3_DESC": st.column_config.TextColumn("LEDGER_ACCOUNT_L3_DESC", disabled=True),
                    "PRODUCT_ID": st.column_config.TextColumn("PRODUCT_ID", disabled=True),
                    "PRODUCT_NAME": st.column_config.TextColumn("PRODUCT_NAME", disabled=True),
                    "AMOUNT": st.column_config.NumberColumn("AMOUNT", format="%.2f", disabled=True),
                    "ROWS_N": st.column_config.NumberColumn("ROWS_N", disabled=True),
                    "PLANNED_NWF_TYPE": st.column_config.SelectboxColumn("Planned NWF Type", options=planned_opts),
                    "EFFECTIVE_NWF_TYPE": st.column_config.TextColumn("Effective NWF Type", disabled=True),
                    "MAPPED_BY": st.column_config.TextColumn("Mapped By", disabled=True),
                    "IS_ACTIVE": st.column_config.CheckboxColumn("IS_ACTIVE"),
                    "NOTES": st.column_config.TextColumn("NOTES"),
                },
                hide_index=True,
            )

            with st.expander(f"Cleanup (FY{int(map_year)})", expanded=False):
                st.warning("This deletes legacy product-level mapping rows that were defaulted to NWF Other and can block ledger fallback.")
                confirm = st.checkbox(f"I understand — delete PRODUCT_ID mappings that are NWF Other for FY{int(map_year)}", value=False, key="apptio_del_prod_nwf_other_confirm")
                if st.button(f"Delete ALL product mappings = NWF Other (FY{int(map_year)})", disabled=not confirm, key="apptio_del_prod_nwf_other_btn"):
                    try:
                        execute(
                            """
                            DELETE FROM MAP_APPTIO_TO_COST_TYPE
                            WHERE PRODUCT_ID IS NOT NULL
                              AND TRY_CONVERT(INT, FISCAL_YEAR) = ?
                              AND UPPER(LTRIM(RTRIM(COST_TYPE))) = 'NWF'
                              AND UPPER(COALESCE(NULLIF(LTRIM(RTRIM(SUBTYPE)), ''), 'Other')) = 'OTHER'
                            """,
                            (int(map_year),),
                        )
                        st.session_state["ver_apptio_mapping"] = int(st.session_state.get("ver_apptio_mapping", 0)) + 1
                        toast_success("Deleted legacy product NWF Other mappings.")
                        try:
                            _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                        except Exception:
                            pass
                    except Exception as e:
                        toast_error(f"Could not delete rows: {e}")

            if st.button("Save product mappings", key="btn_save_apptio_prod_map", type="primary"):
                try:
                    rows_map: list[dict[str, Any]] = []
                    if edited_prod is not None and not edited_prod.empty:
                        # Deduplicate by PRODUCT_ID (mapping is keyed by PRODUCT_ID in SQL; ledger is display-only).
                        ed = edited_prod.copy()
                        ed["PRODUCT_ID"] = ed.get("PRODUCT_ID", "").fillna("").astype(str).str.strip()
                        ed["PLANNED_NWF_TYPE"] = ed.get("PLANNED_NWF_TYPE", "").fillna("").astype(str).str.strip()
                        ed["NOTES"] = ed.get("NOTES", "").fillna("").astype(str).str.strip()
                        ed["IS_ACTIVE"] = ed.get("IS_ACTIVE").fillna(1).astype(int)
                        for pid, grp in ed.groupby("PRODUCT_ID", dropna=False):
                            if not str(pid).strip():
                                continue
                            first = grp.iloc[0]
                            planned = str(first.get("PLANNED_NWF_TYPE") or "").strip()
                            # Do not persist "NWF Other" as a product override; it's treated as default and should not block ledger fallback.
                            if not planned or planned == "NWF Other":
                                continue
                            ct, stp = _planned_to_pair(planned)
                            rows_map.append(
                                {
                                    "MAP_ID": map_id_by_pid.get(str(pid).strip()),
                                    "FISCAL_YEAR": int(map_year),
                                    "PRODUCT_ID": str(pid).strip(),
                                    "LEDGER_ACCOUNT_CONTAINS": None,
                                    "COST_TYPE": ct,
                                    "SUBTYPE": stp,
                                    "IS_ACTIVE": 1 if int(first.get("IS_ACTIVE") or 0) == 1 else 0,
                                    "NOTES": str(first.get("NOTES") or "").strip(),
                                }
                            )
                    if not rows_map:
                        st.info("No mapping rows to save.")
                    else:
                        n = upsert_apptio_to_cost_type_mappings(
                            rows_map, updated_by=st.session_state.get("auth_user", {}).get("email")
                        )
                        st.session_state["ver_apptio_mapping"] = int(st.session_state.get("ver_apptio_mapping", 0)) + 1
                        toast_success(f"Saved {n} mapping rows.")
                except Exception as e:
                    toast_error(f"Could not save mappings: {e}")

        with tab_ledger:
            st.caption("Ledger fallback mappings are only used when no Product mapping matches (or product_id is missing).")
            st.caption("Tip: map MSP via Product ID; ledger fallback is not sufficient for MSP because the same ledger can include Contractor CS.")
            try:
                ledgers_df = fetch_df(
                    """
                    SELECT
                      LEDGER_ACCOUNT_L3_DESC,
                      SUM(COALESCE(TRY_CONVERT(FLOAT, AMOUNT),0)) AS AMOUNT
                    FROM APPTIO_ACTUALS_LINES
                    WHERE TRY_CONVERT(INT,FISCAL_YEAR)=?
                      AND LEDGER_ACCOUNT_L3_DESC IS NOT NULL
                      AND LTRIM(RTRIM(LEDGER_ACCOUNT_L3_DESC))<>''
                    GROUP BY LEDGER_ACCOUNT_L3_DESC
                    ORDER BY AMOUNT DESC
                    """,
                    (int(map_year),),
                )
            except Exception:
                ledgers_df = pd.DataFrame()

            try:
                ledger_map = fetch_df(
                    """
                    SELECT
                      MAP_ID, FISCAL_YEAR, PRODUCT_ID, LEDGER_ACCOUNT_CONTAINS,
                      COST_TYPE, SUBTYPE, IS_ACTIVE, NOTES
                    FROM MAP_APPTIO_TO_COST_TYPE
                    WHERE PRODUCT_ID IS NULL
                    """,
                    None,
                )
            except Exception:
                ledger_map = pd.DataFrame()

            if ledgers_df is None or ledgers_df.empty:
                st.info("No ledger descriptions found in APPTIO_ACTUALS_LINES for this year yet. Upload parses the file; you must click 'Save Apptio actuals' to write APPTIO_ACTUALS_LINES. Then return here.")
                ledger_editor_src = pd.DataFrame(
                    columns=[
                        "MAP_ID",
                        "FISCAL_YEAR",
                        "LEDGER_ACCOUNT_L3_DESC",
                        "AMOUNT",
                        "LEDGER_ACCOUNT_CONTAINS",
                        "PLANNED_NWF_TYPE",
                        "IS_ACTIVE",
                        "NOTES",
                    ]
                )
            else:
                ldf = ledgers_df.copy()
                ldf["LEDGER_ACCOUNT_L3_DESC"] = ldf.get("LEDGER_ACCOUNT_L3_DESC", "").fillna("").astype(str).str.strip()
                ldf["AMOUNT"] = pd.to_numeric(ldf.get("AMOUNT"), errors="coerce").fillna(0.0)
                ldf = ldf[ldf["LEDGER_ACCOUNT_L3_DESC"].ne("")].copy()
                ldf = ldf.head(250).copy()

                m = ledger_map.copy() if ledger_map is not None and not ledger_map.empty else pd.DataFrame(
                    columns=["MAP_ID", "FISCAL_YEAR", "LEDGER_ACCOUNT_CONTAINS", "COST_TYPE", "SUBTYPE", "IS_ACTIVE", "NOTES"]
                )
                if not m.empty:
                    m["LEDGER_ACCOUNT_CONTAINS"] = m.get("LEDGER_ACCOUNT_CONTAINS", "").fillna("").astype(str).str.strip()
                    m = m[m["LEDGER_ACCOUNT_CONTAINS"].ne("")].copy()
                    m["FISCAL_YEAR"] = pd.to_numeric(m.get("FISCAL_YEAR"), errors="coerce")
                    m["IS_ACTIVE"] = pd.to_numeric(m.get("IS_ACTIVE"), errors="coerce").fillna(1).astype(int)

                def _best_match(desc: str) -> pd.Series:
                    if m is None or m.empty:
                        return pd.Series(dtype="object")
                    d_up = str(desc or "").strip().upper()
                    if not d_up:
                        return pd.Series(dtype="object")
                    cand = m[m["LEDGER_ACCOUNT_CONTAINS"].astype(str).str.upper().apply(lambda s: s in d_up)].copy()
                    if cand.empty:
                        return pd.Series(dtype="object")
                    cand["_ypri"] = (cand["FISCAL_YEAR"].fillna(int(map_year)) != int(map_year)).astype(int)
                    cand["_apri"] = (cand["IS_ACTIVE"] != 1).astype(int)
                    cand["_lpri"] = cand["LEDGER_ACCOUNT_CONTAINS"].astype(str).str.len().mul(-1)
                    cand = cand.sort_values(["_apri", "_ypri", "_lpri"])
                    return cand.iloc[0]

                best_rows: list[dict[str, Any]] = []
                for _, rr in ldf.iterrows():
                    desc = str(rr.get("LEDGER_ACCOUNT_L3_DESC") or "").strip()
                    match = _best_match(desc)
                    best_rows.append(
                        {
                            "LEDGER_ACCOUNT_L3_DESC": desc,
                            "AMOUNT": float(rr.get("AMOUNT") or 0.0),
                            "MAP_ID": match.get("MAP_ID") if isinstance(match, pd.Series) and not match.empty else None,
                            "FISCAL_YEAR": (
                                int(match.get("FISCAL_YEAR"))
                                if isinstance(match, pd.Series) and not match.empty and str(match.get("FISCAL_YEAR") or "").strip().isdigit()
                                else int(map_year)
                            ),
                            "LEDGER_ACCOUNT_CONTAINS": (
                                str(match.get("LEDGER_ACCOUNT_CONTAINS") or "").strip()
                                if isinstance(match, pd.Series) and not match.empty and str(match.get("LEDGER_ACCOUNT_CONTAINS") or "").strip()
                                else desc
                            ),
                            "PLANNED_NWF_TYPE": _pair_to_planned(
                                match.get("COST_TYPE") if isinstance(match, pd.Series) and not match.empty else "",
                                match.get("SUBTYPE") if isinstance(match, pd.Series) and not match.empty else "",
                            ),
                            "IS_ACTIVE": (int(match.get("IS_ACTIVE")) if isinstance(match, pd.Series) and not match.empty else 1),
                            "NOTES": str(match.get("NOTES") or "").strip() if isinstance(match, pd.Series) and not match.empty else "",
                        }
                    )
                ledger_editor_src = pd.DataFrame(best_rows)

            edited_ledger = st.data_editor(
                ledger_editor_src,
                key="apptio_map_ledger_editor",
                use_container_width=True,
                num_rows="fixed",
                column_config={
                    "MAP_ID": st.column_config.TextColumn("MAP_ID", disabled=True),
                    "FISCAL_YEAR": st.column_config.NumberColumn("FISCAL_YEAR", help="Optional year filter (blank = all years)."),
                    "LEDGER_ACCOUNT_L3_DESC": st.column_config.TextColumn("LEDGER_ACCOUNT_L3_DESC", disabled=True),
                    "AMOUNT": st.column_config.NumberColumn("AMOUNT", format="%.2f", disabled=True),
                    "LEDGER_ACCOUNT_CONTAINS": st.column_config.TextColumn("LEDGER_ACCOUNT_CONTAINS", help="Substring match; defaults to the full ledger description so you can shorten it."),
                    "PLANNED_NWF_TYPE": st.column_config.SelectboxColumn("Planned NWF Type", options=planned_opts),
                    "IS_ACTIVE": st.column_config.CheckboxColumn("IS_ACTIVE"),
                },
                hide_index=True,
            )
            if st.button("Save ledger mappings", key="btn_save_apptio_ledger_map", type="primary"):
                try:
                    rows_map: list[dict[str, Any]] = []
                    if edited_ledger is not None and not edited_ledger.empty:
                        for _, r in edited_ledger.iterrows():
                            if not str(r.get("LEDGER_ACCOUNT_CONTAINS") or "").strip():
                                continue
                            planned = str(r.get("PLANNED_NWF_TYPE") or "").strip()
                            if not planned:
                                continue
                            ct, stp = _planned_to_pair(planned)
                            rows_map.append(
                                {
                                    "MAP_ID": r.get("MAP_ID"),
                                    "FISCAL_YEAR": int(r.get("FISCAL_YEAR")) if str(r.get("FISCAL_YEAR") or "").strip().isdigit() else None,
                                    "PRODUCT_ID": None,
                                    "LEDGER_ACCOUNT_CONTAINS": str(r.get("LEDGER_ACCOUNT_CONTAINS")).strip(),
                                    "COST_TYPE": ct,
                                    "SUBTYPE": stp,
                                    "IS_ACTIVE": 1 if bool(r.get("IS_ACTIVE")) else 0,
                                    "NOTES": str(r.get("NOTES") or "").strip(),
                                }
                            )
                    if not rows_map:
                        st.info("No mapping rows to save.")
                    else:
                        n = upsert_apptio_to_cost_type_mappings(
                            rows_map, updated_by=st.session_state.get("auth_user", {}).get("email")
                        )
                        st.session_state["ver_apptio_mapping"] = int(st.session_state.get("ver_apptio_mapping", 0)) + 1
                        toast_success(f"Saved {n} mapping rows.")
                except Exception as e:
                    toast_error(f"Could not save mappings: {e}")

        with tab_rec:
            st.caption("Program-level Apptio actuals breakdown (from enriched APPTIO_ACTUALS_LINES + mapping).")
            try:
                df_rec = fetch_apptio_actuals_by_program_breakdown(int(map_year), None)
            except Exception as e:
                st.error(f"Could not load Apptio reconciliation: {e}")
                df_rec = pd.DataFrame()
            if df_rec is None or df_rec.empty:
                st.info("No enriched Apptio rows found for this year. Upload and click 'Save Apptio actuals' to populate APPTIO_ACTUALS_LINES.")
            else:
                wrec = df_rec.copy()
                wrec["PROGRAMNAME"] = wrec.get("PROGRAMNAME", "").fillna("").astype(str).str.strip()
                program_opts = sorted({p for p in wrec["PROGRAMNAME"].dropna().astype(str).tolist() if str(p).strip()})
                sel_prog = st.multiselect("Filter programs (optional)", program_opts, default=program_opts, key="apptio_rec_prog_filter")
                if sel_prog:
                    wrec = wrec[wrec["PROGRAMNAME"].isin(sel_prog)].copy()
                show_cols = ["PROGRAMNAME", "MONTH_KEY", "COST_TYPE", "SUBTYPE", "MAPPING_STATUS", "AMOUNT"]
                show_cols = [c for c in show_cols if c in wrec.columns]
                if "AMOUNT" in wrec.columns:
                    wrec["AMOUNT"] = pd.to_numeric(wrec.get("AMOUNT"), errors="coerce").fillna(0.0)
                st.dataframe(wrec[show_cols].sort_values(["PROGRAMNAME", "MONTH_KEY", "COST_TYPE", "SUBTYPE"]), use_container_width=True, height=520)

apptio_file = None
with _tab_apptio:
    apptio_file = st.file_uploader(
        "Upload Apptio actuals (XLSX/XLSM/XLSB/XLS/CSV) with columns: Work ID, Jan FY 2025, Feb FY 2025, ...",
        type=["xlsx", "xlsm", "xlsb", "xls", "csv"],
        key="apptio_actuals_file",
    )
apptio_sheet: Optional[str] = None
diag_apptio: Dict[str, Any] = {}
parsed_df: Optional[pd.DataFrame] = None

if apptio_file:
    apptio_file.seek(0)
    file_bytes = apptio_file.read()
    apptio_file.seek(0)
    sheets = _list_excel_sheets(file_bytes)
    if sheets:
        apptio_sheet = st.selectbox("Worksheet", sheets, index=0, key="apptio_sheet_select")
    try:
        raw_df = _read_file_any(apptio_file, apptio_sheet, diag_apptio)
        parsed_df = _normalize_apptio_actuals(raw_df)
        st.success(f"Parsed {len(parsed_df)} Work ID-month rows from the file.")
    except Exception as e:
        st.error(f"Could not parse Apptio file: {e}")
        if diag_apptio:
            with st.expander("Debug - diagnostics", expanded=False):
                st.json(diag_apptio)

if parsed_df is not None:
    if mapping_df is None or mapping_df.empty:
        st.info("Upload succeeded but no mappings are available to associate Work IDs to programs.")
    mapped = parsed_df.copy()
    if mapping_df is not None and not mapping_df.empty:
        mapped = mapped.merge(mapping_df[["WORK_ID", "PROGRAMID", "PROGRAMNAME"]], how="left", on="WORK_ID")

    if prog_opts:
        sel_prog = st.multiselect(
            "Programs to include",
            prog_opts,
            default=prog_opts,
            key="apptio_prog_filter",
        )
        if sel_prog:
            mapped = mapped[mapped["PROGRAMNAME"].isin(sel_prog)]

    unmapped = mapped[mapped.get("PROGRAMID").isna()]
    if not unmapped.empty:
        missing_ids = unmapped["WORK_ID"].dropna().unique().tolist()
        st.warning(f"{len(missing_ids)} Work IDs are not mapped to a program and will be skipped.")

    to_upsert = mapped.dropna(subset=["PROGRAMID"]).copy()
    if to_upsert.empty:
        st.info("No rows matched a mapped program.")
    else:
        to_upsert["AMOUNT"] = pd.to_numeric(to_upsert["AMOUNT"], errors="coerce").fillna(0.0)
        summary = (
            to_upsert.groupby(["PROGRAMNAME", "FISCAL_YEAR", "MONTH"], as_index=False)["AMOUNT"]
            .sum()
            .rename(columns={"FISCAL_YEAR": "Year", "MONTH": "Month", "AMOUNT": "Amount"})
        )
        month_names = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
        summary["Month"] = summary["Month"].astype(int).map(lambda m: month_names[m-1] if 1 <= m <= 12 else m)
        st.dataframe(summary, use_container_width=True, height=280)

        # Portfolio/program total across all months shown before save
        total_df = (
            to_upsert.groupby(["PROGRAMNAME"], as_index=False)["AMOUNT"]
            .sum()
            .rename(columns={"AMOUNT": "Total Amount"})
            .sort_values("Total Amount", ascending=False)
        )
        total_df["Total Amount"] = total_df["Total Amount"].map(lambda v: f"${v:,.2f}")
        st.dataframe(total_df, use_container_width=True, height=180)

        grouped = (
            to_upsert.groupby(["WORK_ID", "FISCAL_YEAR", "MONTH"], as_index=False)["AMOUNT"]
            .sum()
        )
        user_email = st.session_state.get("auth_user", {}).get("email")
        rows = [
            {
                "work_id": r["WORK_ID"],
                "fiscal_year": int(r["FISCAL_YEAR"]),
                "month": int(r["MONTH"]),
                "amount": float(r["AMOUNT"] or 0),
                "source": apptio_file.name if apptio_file else None,
                "loaded_by": user_email,
            }
            for _, r in grouped.iterrows()
        ]
        # Enriched lines for taxonomy mapping (best-effort; still keep legacy totals upsert above).
        line_cols = ["WORK_ID", "FISCAL_YEAR", "MONTH", "AMOUNT", "LEDGER_ACCOUNT_L3_DESC", "PRODUCT_ID", "PRODUCT_NAME"]
        line_src = to_upsert.copy()
        for c in ["LEDGER_ACCOUNT_L3_DESC", "PRODUCT_ID", "PRODUCT_NAME"]:
            if c not in line_src.columns:
                line_src[c] = None
        line_src = line_src[line_cols].copy()
        line_src["AMOUNT"] = pd.to_numeric(line_src.get("AMOUNT"), errors="coerce").fillna(0.0)
        line_grouped = (
            line_src.groupby(["WORK_ID", "FISCAL_YEAR", "MONTH", "PRODUCT_ID", "LEDGER_ACCOUNT_L3_DESC"], as_index=False)
            .agg({"AMOUNT": "sum", "PRODUCT_NAME": "first"})
        )
        rows_lines = [
            {
                "work_id": r.get("WORK_ID"),
                "fiscal_year": int(r.get("FISCAL_YEAR")),
                "month": int(r.get("MONTH")),
                "amount": float(r.get("AMOUNT") or 0.0),
                "ledger_account_l3_desc": r.get("LEDGER_ACCOUNT_L3_DESC"),
                "product_id": r.get("PRODUCT_ID"),
                "product_name": r.get("PRODUCT_NAME"),
            }
            for _, r in line_grouped.iterrows()
        ]
        if st.button("Save Apptio actuals", icon=":material/save:", key="btn_save_apptio_actuals", type="primary"):
            try:
                legacy_saved = upsert_apptio_actuals(rows)
                try:
                    lines_saved = upsert_apptio_actuals_lines(
                        rows_lines,
                        updated_by=user_email,
                        source=apptio_file.name if apptio_file else None,
                    )
                except Exception:
                    # Legacy totals must remain saveable even if enriched table/mapping is not ready.
                    lines_saved = 0
                toast_success(f"Saved Apptio actuals: legacy_saved={legacy_saved}, lines_saved={lines_saved}.")
                try:
                    yrs_saved = sorted({int(y) for y in grouped.get('FISCAL_YEAR').dropna().astype(int).tolist()}) if grouped is not None and not grouped.empty else []
                except Exception:
                    yrs_saved = []
                for fy in (yrs_saved or []):
                    try:
                        stats = fetch_df(
                            """
                            SELECT
                              COUNT(*) AS ROWS_N,
                              COUNT(DISTINCT NULLIF(LTRIM(RTRIM(PRODUCT_ID)), '')) AS PRODUCT_IDS_N
                            FROM APPTIO_ACTUALS_LINES
                            WHERE TRY_CONVERT(INT, FISCAL_YEAR) = ?
                            """,
                            (int(fy),),
                        )
                        if stats is not None and not stats.empty:
                            st.caption(
                                f"APPTIO_ACTUALS_LINES now contains {int(stats.iloc[0]['ROWS_N'])} rows and "
                                f"{int(stats.iloc[0]['PRODUCT_IDS_N'])} distinct PRODUCT_IDs for FY{int(fy)}."
                            )
                    except Exception:
                        pass
                try:
                    _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                except Exception:
                    pass
            except Exception as e:
                toast_error(f"Could not save Apptio actuals: {e}")

# =========================# =========================================================
# 📦 Bulk Load (ONE sheet): Programs, Vendors, Applications, Applications, Teams & Invoices
# =========================================================
def _render_legacy_bulk_load_one_sheet() -> None:
    st.subheader("Bulk Load (one sheet): Programs, Vendors, Applications, Applications, Teams & Invoices")
    st.caption("Upload a single sheet that contains the columns for all sections. Map once, preview auto‑generates, then MERGE in a safe order with UUIDs.")
    # Downloadable Excel templates: Contracts-only (recommended) or Full (contracts + invoices)
    tmpl_mode = st.radio(
        "Template type",
        ["Contracts only (recommended)", "Core entities only (Programs/Vendors/Teams/Groups/Apps)"],
        index=0,
        horizontal=True,
        key="bulk_tmpl_mode"
    )
    if tmpl_mode.startswith("Contracts only"):
        tmpl = generate_contracts_template()
        fname = "next_contracts_template.xlsx"
        help_text = "Use this template to load Programs, Vendors, Groups, Applications, Teams, and Contracts. Planned invoices will be created per FY from contracts."
    else:
        tmpl = generate_core_entities_template()
        fname = "next_core_entities_template.xlsx"
        help_text = "Use this template to load Programs, Vendors, Teams (with Product Owner), Applications, and Application Instances."
    st.download_button(
        label="⬇️ Download Excel Template",
        data=tmpl,
        file_name=fname,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        help=help_text,
    )

    upl_one = st.file_uploader("Upload workbook (XLSX preferred; CSV allowed)", type=["xlsx","xlsm","xlsb","xls","csv"], key="one_workbook")
    sheet = None
    if upl_one:
        upl_one.seek(0); raw = upl_one.read(); upl_one.seek(0)
        sheets = _list_excel_sheets(raw)
        if sheets:
            sheet = st.selectbox("Worksheet (one sheet for everything)", sheets, index=0, key="one_sheet_select")
        else:
            st.info("No sheet list available (CSV or detection failed). I will parse the first/only sheet.")

        diag: Dict[str, Any] = {}
        try:
            df_src = _read_file_any(upl_one, sheet, diag)
            # De-duplicate column labels for display logic
            seen: Dict[str, int] = {}
            newcols: List[str] = []
            for col_lbl in df_src.columns:
                key_lbl = col_lbl
                if key_lbl in seen:
                    seen[key_lbl] += 1
                    key_lbl = f"{col_lbl}__{seen[col_lbl]}"
                else:
                    seen[key_lbl] = 0
                newcols.append(key_lbl)
            df_src.columns = newcols

            st.session_state["one_sheet_df"] = df_src
            st.success(f"Parsed {len(df_src)} rows from {upl_one.name}.")
        except Exception as e:
            with st.expander("Debug - diagnostics", expanded=True):
                st.write("**Why it failed**")
                st.exception(e)
                st.write("**Parse attempts**")
                st.json(diag)
            st.error("Could not parse the workbook. Try a clean XLSX/CSV with headers.")
            st.stop()

    # -------------------------
    # Column mapping (single source)
    # -------------------------
    if st.session_state["one_sheet_df"] is not None:
        df_src = st.session_state["one_sheet_df"]
        cols = df_src.columns.tolist()
        # Determine current template mode to tailor mapping/previews
        is_contracts_only = str(st.session_state.get("bulk_tmpl_mode", "")).startswith("Contracts only")

        def _default_pick(name: str) -> Optional[str]:
            lower = {c.lower(): c for c in cols if isinstance(c, str)}
            if name.lower() in lower:
                return lower[name.lower()]
            matches = [c for c in cols if isinstance(c, str) and name.lower() in c.lower()]
            return matches[0] if matches else None

        # Initialize mapping once
        if not st.session_state["colmap"]:
            st.session_state["colmap"] = {
                # Programs
                "PROGRAMNAME": _default_pick("PROGRAMNAME") or _default_pick("PROGRAM"),
                # Vendors
                "VENDORNAME": _default_pick("VENDORNAME") or _default_pick("VENDOR"),
                # Applications
                "GROUPNAME": _default_pick("APP GROUP NAME") or _default_pick("GROUPNAME") or _default_pick("GROUP"),
                # Applications
                "APPNAME": _default_pick("APPNAME") or _default_pick("APPLICATION"),
                # Teams
                "TEAMNAME": _default_pick("TEAMNAME") or _default_pick("TEAM"),
                # Contracts (required)
                "CONTRACT_START_FY": _default_pick("CONTRACT_START_FY") or _default_pick("START_FY"),
                "CONTRACT_END_FY": _default_pick("CONTRACT_END_FY") or _default_pick("END_FY"),
                "CONTRACT_RENEWAL_MONTH": _default_pick("CONTRACT_RENEWAL_MONTH") or _default_pick("CONTRACT_RENEWAL") or _default_pick("RENEWAL_MONTH"),
                "CONTRACT_ANNUAL_AMOUNT": _default_pick("CONTRACT_ANNUAL_AMOUNT") or _default_pick("ANNUAL_AMOUNT") or _default_pick("AMOUNT"),
                "CONTRACT_ESCALATION_PCT": _default_pick("CONTRACT_ESCALATION_PCT") or _default_pick("ESCALATION_PCT"),
                "CONTRACT_STATUS": _default_pick("CONTRACT_STATUS") or _default_pick("STATUS"),
                "CONTRACT_AGREEMENT_NUMBER": _default_pick("CONTRACT_AGREEMENT_NUMBER") or _default_pick("AGREEMENT_NUMBER"),
                # Invoices (required)
                "AMOUNT": _default_pick("AMOUNT"),
                "FISCAL_YEAR": _default_pick("FISCAL_YEAR") or _default_pick("YEAR"),
                "RENEWAL_MONTH": _default_pick("RENEWAL_MONTH") or _default_pick("RENEWAL MONTH"),
                # Invoices (optional)
                "INVOICE_STATUS": _default_pick("INVOICE_STATUS") or _default_pick("STATUS"),
                "CONTRACT_ACTIVE": _default_pick("CONTRACT_ACTIVE"),
                "SERIAL_NUMBER": _default_pick("SERIAL_NUMBER"),
                "WORK_ORDER": _default_pick("WORK_ORDER"),
                "COMPANY_CODE": _default_pick("COMPANY_CODE"),
                "COST_CENTER": _default_pick("COST_CENTER"),
                "PRODUCT_OWNER": _default_pick("PRODUCT_OWNER"),
                "NOTES": _default_pick("NOTES"),
            }

        cm: Dict[str, Optional[str]] = st.session_state["colmap"]

        st.markdown("### Column Mapping")
        # Group pickers into collapsible sections
        with st.expander("Core Entities (Programs, Vendors, Groups, Apps, Teams)", expanded=True):
            ce1, ce2 = st.columns(2)
            with ce1:
                st.markdown("**Program**")
                cm["PROGRAMNAME"] = st.selectbox("PROGRAMNAME (required)", options=["(none)"] + cols,
                                                 index=(cols.index(cm["PROGRAMNAME"]) + 1) if cm.get("PROGRAMNAME") in cols else 0,
                                                 key="map_programname")
                st.markdown("**Vendor**")
                cm["VENDORNAME"] = st.selectbox("VENDORNAME (required)", options=["(none)"] + cols,
                                                index=(cols.index(cm["VENDORNAME"]) + 1) if cm.get("VENDORNAME") in cols else 0,
                                                key="map_vendorname")
                st.markdown("**Application**")
                cm["GROUPNAME"] = st.selectbox("APP GROUP NAME (required)", options=["(none)"] + cols,
                                               index=(cols.index(cm["GROUPNAME"]) + 1) if cm.get("GROUPNAME") in cols else 0,
                                               key="map_groupname")
            with ce2:
                st.markdown("**Applications**")
                cm["APPNAME"] = st.selectbox("APPNAME (required)", options=["(none)"] + cols,
                                             index=(cols.index(cm["APPNAME"]) + 1) if cm.get("APPNAME") in cols else 0,
                                             key="map_appname")
                st.markdown("**Teams**")
                cm["TEAMNAME"] = st.selectbox("TEAMNAME (required)", options=["(none)"] + cols,
                                              index=(cols.index(cm["TEAMNAME"]) + 1) if cm.get("TEAMNAME") in cols else 0,
                                              key="map_teamname")

        with st.expander("Contracts Mapping", expanded=True):
            st.markdown("**Contracts (required)**")
            for key, label in [
                ("CONTRACT_START_FY", "CONTRACT_START_FY"),
                ("CONTRACT_END_FY", "CONTRACT_END_FY"),
                ("CONTRACT_ANNUAL_AMOUNT", "CONTRACT_ANNUAL_AMOUNT"),
                ("CONTRACT_RENEWAL_DATE", "CONTRACT_RENEWAL_DATE"),
            ]:
                cm[key] = st.selectbox(label, options=["(none)"] + cols,
                                       index=(cols.index(cm[key]) + 1) if cm.get(key) in cols else 0,
                                       key=f"map_contract_{key}")
            st.markdown("**Contracts (optional)**")
            for key, label in [
                ("CONTRACT_ESCALATION_PCT", "CONTRACT_ESCALATION_PCT"),
                ("CONTRACT_STATUS", "CONTRACT_STATUS"),
                ("CONTRACT_AGREEMENT_NUMBER", "CONTRACT_AGREEMENT_NUMBER"),
                ("CONTRACT_COMPANY_CODE", "CONTRACT_COMPANY_CODE"),
                ("CONTRACT_COST_CENTER", "CONTRACT_COST_CENTER"),
                ("CONTRACT_SERVICE_TYPE", "CONTRACT_SERVICE_TYPE"),
                ("INVOICE_RENEWAL_DATE", "INVOICE_RENEWAL_DATE"),
                ("CONTRACT_TOTAL_COST", "CONTRACT_TOTAL_COST"),
            ]:
                cm[key] = st.selectbox(label, options=["(none)"] + cols,
                                       index=(cols.index(cm[key]) + 1) if cm.get(key) in cols else 0,
                                       key=f"map_contract_opt_{key}")

        if not is_contracts_only:
            with st.expander("Invoices Mapping", expanded=False):
                st.markdown("**Invoices (required)**")
                cm["AMOUNT"] = st.selectbox("AMOUNT", options=["(none)"] + cols,
                                            index=(cols.index(cm["AMOUNT"]) + 1) if cm.get("AMOUNT") in cols else 0,
                                            key="map_inv_amt")
                cm["FISCAL_YEAR"] = st.selectbox("FISCAL_YEAR (Year)", options=["(none)"] + cols,
                                                 index=(cols.index(cm["FISCAL_YEAR"]) + 1) if cm.get("FISCAL_YEAR") in cols else 0,
                                                 key="map_inv_fy")
                cm["RENEWAL_MONTH"] = st.selectbox("RENEWAL_MONTH (1-12)", options=["(none)"] + cols,
                                                   index=(cols.index(cm["RENEWAL_MONTH"]) + 1) if cm.get("RENEWAL_MONTH") in cols else 0,
                                                   key="map_inv_rmonth")

                st.markdown("**Invoices (optional)**")
                for opt_key, label in [("INVOICE_STATUS", "INVOICE_STATUS"),("CONTRACT_ACTIVE", "CONTRACT_ACTIVE"),("SERIAL_NUMBER", "SERIAL_NUMBER"),("WORK_ORDER", "WORK_ORDER"),("COMPANY_CODE", "COMPANY_CODE"),("COST_CENTER", "COST_CENTER"),("PRODUCT_OWNER", "PRODUCT_OWNER"),("NOTES", "NOTES")]:
                    current = cm.get(opt_key)
                    cm[opt_key] = st.selectbox(label, options=["(none)"] + cols,
                                               index=(cols.index(current) + 1) if current in cols else 0,
                                               key=f"map_opt_{opt_key}")
        else:
            st.caption("Invoices are auto-generated from Contracts (no invoice fields to map).")

        st.markdown("---")

        # -------------------------
        # Build auto-previews from one sheet using the mapping
        # -------------------------
        def _col(name: str) -> Optional[str]:
            val = cm.get(name)
            return val if val and val != "(none)" and val in df_src.columns else None

        def _safe_num(series: pd.Series) -> pd.Series:
            return pd.to_numeric(series.astype(str).str.replace(",", ".", regex=False), errors="coerce")

        # Programs preview
        dfP = pd.DataFrame()
        if _col("PROGRAMNAME"):
            tmp = df_src[[_col("PROGRAMNAME")]].rename(columns={_col("PROGRAMNAME"): "PROGRAMNAME"})
            tmp["PROGRAMNAME"] = tmp["PROGRAMNAME"].astype(str).str.strip()
            tmp = tmp[~tmp["PROGRAMNAME"].apply(_blank_or_nan)]
            dfP = tmp.drop_duplicates(subset=["PROGRAMNAME"]).reset_index(drop=True)

        # Vendors preview
        dfV = pd.DataFrame()
        if _col("VENDORNAME"):
            tmp = df_src[[_col("VENDORNAME")]].rename(columns={_col("VENDORNAME"): "VENDORNAME"})
            tmp["VENDORNAME"] = tmp["VENDORNAME"].astype(str).str.strip()
            tmp = tmp[~tmp["VENDORNAME"].apply(_blank_or_nan)]
            dfV = tmp.drop_duplicates(subset=["VENDORNAME"]).reset_index(drop=True)

        # Groups preview  (TEAM × GROUP scoped + majority vendor)
        def _mode_nonblank(series: pd.Series) -> Optional[str]:
            s = series.dropna().astype(str).str.strip()
            s = s[s != ""]
            if s.empty:
                return None
            return s.value_counts().idxmax()

        dfG = pd.DataFrame()
        if _col("GROUPNAME"):
            tmp = df_src[[_col("GROUPNAME")]].rename(columns={_col("GROUPNAME"): "GROUPNAME"})
            tmp["GROUPNAME"] = tmp["GROUPNAME"].astype(str).str.strip()

            if _col("TEAMNAME"):
                tmp["TEAMNAME"] = df_src[_col("TEAMNAME")].astype(str).str.strip()
            else:
                tmp["TEAMNAME"] = None

            if _col("PROGRAMNAME"):
                tmp["PROGRAMNAME"] = df_src[_col("PROGRAMNAME")].astype(str).str.strip()

            if _col("VENDORNAME"):
                tmp["DEFAULT_VENDORNAME"] = df_src[_col("VENDORNAME")].astype(str).str.strip()

            # Drop blank group names; then aggregate by TEAMNAME × GROUPNAME
            tmp = tmp[~tmp["GROUPNAME"].apply(_blank_or_nan)].copy()

            # Choose vendor by majority occurrence within each TEAM×GROUP
            grp_keys = ["TEAMNAME", "GROUPNAME"]
            agg = {
                "PROGRAMNAME": "first",
                "DEFAULT_VENDORNAME": _mode_nonblank,
            }
            # Only aggregate columns that exist
            agg = {k: v for k, v in agg.items() if k in tmp.columns}
            dfG = (
                tmp.groupby([k for k in grp_keys if k in tmp.columns], dropna=False)
                .agg(agg)
                .reset_index()
            )


        # Applications preview
        dfA = pd.DataFrame()
        if _col("APPNAME"):
            tmp = df_src[[_col("APPNAME")]].rename(columns={_col("APPNAME"): "APPNAME"})
            tmp["APPNAME"] = tmp["APPNAME"].astype(str).str.strip()
            tmp = tmp[~tmp["APPNAME"].apply(_blank_or_nan)]
            if _col("GROUPNAME"):
                tmp["GROUPNAME"] = df_src[_col("GROUPNAME")].astype(str).str.strip()
            dfA = tmp.drop_duplicates(subset=["APPNAME","GROUPNAME"] if "GROUPNAME" in tmp.columns else ["APPNAME"]).reset_index(drop=True)

        # Teams preview
        dfT = pd.DataFrame()
        if _col("TEAMNAME"):
            tmp = df_src[[_col("TEAMNAME")]].rename(columns={_col("TEAMNAME"): "TEAMNAME"})
            tmp["TEAMNAME"] = tmp["TEAMNAME"].astype(str).str.strip()
            tmp = tmp[~tmp["TEAMNAME"].apply(_blank_or_nan)]
            if _col("PROGRAMNAME"):
                tmp["PROGRAMNAME"] = df_src[_col("PROGRAMNAME")].astype(str).str.strip()
            dfT = tmp.drop_duplicates(subset=["TEAMNAME"]).reset_index(drop=True)

        # -------------------------
        # Invoices preview (row-level requireds + validity)
        # -------------------------
        # Separate requireds for invoices and contracts
        INVOICE_REQUIRED_KEYS = [
            "PROGRAMNAME",      # Program Name
            "TEAMNAME",         # Team Name
            "GROUPNAME",        # Application
            "APPNAME",          # Application Instance
            "VENDORNAME",       # Vendor Name
            "FISCAL_YEAR",      # Year
            "RENEWAL_MONTH",    # Renewal Month
        ]
        CONTRACT_REQUIRED_KEYS = [
            "CONTRACT_START_FY","CONTRACT_END_FY","CONTRACT_ANNUAL_AMOUNT"
        ]

        def _has_required_mapping() -> Dict[str, bool]:
            return {key: bool(_col(key)) for key in REQUIRED_KEYS}

        def _row_is_missing(series_val: Any) -> bool:
            if series_val is None:
                return True
            s = str(series_val).strip()
            return s == "" or s.lower() == "nan"

        dfI = pd.DataFrame()
        def _has_invoice_mapping() -> Dict[str, bool]:
            return {key: bool(_col(key)) for key in INVOICE_REQUIRED_KEYS}
        invoice_diag = {"column_mapping": _has_invoice_mapping(), "row_counts": {}}
        missing_invoice_reqs = [key for key, ok in invoice_diag["column_mapping"].items() if not ok]

        if not missing_invoice_reqs:
            out = pd.DataFrame()
            out["AMOUNT"] = _safe_num(df_src[_col("AMOUNT")]) if _col("AMOUNT") else None
            out["FISCAL_YEAR"] = pd.to_numeric(df_src[_col("FISCAL_YEAR")], errors="coerce").astype("Int64")
            out["RENEWAL_MONTH"] = pd.to_numeric(df_src[_col("RENEWAL_MONTH")], errors="coerce").astype("Int64")

            # required name fields
            out["PROGRAMNAME"] = df_src[_col("PROGRAMNAME")].astype(str).str.strip()
            out["TEAMNAME"] = df_src[_col("TEAMNAME")].astype(str).str.strip()
            out["GROUPNAME"] = df_src[_col("GROUPNAME")].astype(str).str.strip()
            out["APPNAME"] = df_src[_col("APPNAME")].astype(str).str.strip()
            out["VENDORNAME"] = df_src[_col("VENDORNAME")].astype(str).str.strip()

            # optional extras
            for opt_name in ["INVOICE_STATUS","CONTRACT_ACTIVE","SERIAL_NUMBER","WORK_ORDER","COMPANY_CODE","COST_CENTER","PRODUCT_OWNER","NOTES"]:
                out[opt_name] = df_src[_col(opt_name)] if _col(opt_name) else None
            # contracts columns
            out["CONTRACT_START_FY"] = pd.to_numeric(df_src[_col("CONTRACT_START_FY")], errors="coerce").astype("Int64")
            out["CONTRACT_END_FY"] = pd.to_numeric(df_src[_col("CONTRACT_END_FY")], errors="coerce").astype("Int64")
            out["CONTRACT_ANNUAL_AMOUNT"] = _safe_num(df_src[_col("CONTRACT_ANNUAL_AMOUNT")])
            out["CONTRACT_ESCALATION_PCT"] = _safe_num(df_src[_col("CONTRACT_ESCALATION_PCT")]) if _col("CONTRACT_ESCALATION_PCT") else 0.0
            out["CONTRACT_STATUS"] = df_src[_col("CONTRACT_STATUS")].astype(str).str.strip() if _col("CONTRACT_STATUS") else "Active"
            out["CONTRACT_AGREEMENT_NUMBER"] = df_src[_col("CONTRACT_AGREEMENT_NUMBER")].astype(str).str.strip() if _col("CONTRACT_AGREEMENT_NUMBER") else None
            out["CONTRACT_COMPANY_CODE"] = df_src[_col("CONTRACT_COMPANY_CODE")].astype(str).str.strip() if _col("CONTRACT_COMPANY_CODE") else None
            out["CONTRACT_COST_CENTER"] = df_src[_col("CONTRACT_COST_CENTER")].astype(str).str.strip() if _col("CONTRACT_COST_CENTER") else None
            out["CONTRACT_SERVICE_TYPE"] = df_src[_col("CONTRACT_SERVICE_TYPE")].astype(str).str.strip() if _col("CONTRACT_SERVICE_TYPE") else None
            # Parse contract renewal date (assume month/day; attach fiscal year later per invoice)
            if _col("CONTRACT_RENEWAL_DATE"):
                try:
                    out["CONTRACT_RENEWAL_DATE"] = pd.to_datetime(df_src[_col("CONTRACT_RENEWAL_DATE")], errors="coerce").dt.date
                except Exception:
                    out["CONTRACT_RENEWAL_DATE"] = None
            else:
                out["CONTRACT_RENEWAL_DATE"] = None
            # Invoice renewal date (REQUIRED) — accept MM/DD/YYYY or MM/DD
            def _parse_md_or_mdy(val):
                try:
                    s = str(val).strip()
                except Exception:
                    return None
                if not s:
                    return None
                import re
                m = re.match(r"^\s*(\d{1,2})[\/-](\d{1,2})(?:[\/-](\d{2,4}))?\s*$", s)
                if m:
                    mm, dd, yy = m.group(1), m.group(2), m.group(3)
                    try:
                        mm_i, dd_i = int(mm), int(dd)
                        yy_i = int(yy) if yy else 2000  # placeholder year; only month/day are used downstream
                        return date(yy_i, mm_i, dd_i)
                    except Exception:
                        return None
                try:
                    return pd.to_datetime(s, errors="coerce").date()
                except Exception:
                    return None

            if _col("INVOICE_RENEWAL_DATE"):
                out["INVOICE_RENEWAL_DATE"] = df_src[_col("INVOICE_RENEWAL_DATE")].apply(_parse_md_or_mdy)
            else:
                out["INVOICE_RENEWAL_DATE"] = None
            # Optional total contract cost
            out["CONTRACT_TOTAL_COST"] = _safe_num(df_src[_col("CONTRACT_TOTAL_COST")]) if _col("CONTRACT_TOTAL_COST") else None

            # Derived date for preview — use INVOICE_RENEWAL_DATE day when available
            def _mk_date(row: pd.Series) -> Optional[date]:
                fy = row.get("FISCAL_YEAR")
                m  = row.get("RENEWAL_MONTH")
                try:
                    if pd.isna(fy) or pd.isna(m):
                        return None
                    mm = int(m); yy = int(fy)
                    dd = 1
                    try:
                        d_src = row.get("INVOICE_RENEWAL_DATE") or row.get("CONTRACT_RENEWAL_DATE")
                        if pd.notna(d_src):
                            d_parsed = pd.to_datetime(d_src, errors="coerce")
                            if pd.notna(d_parsed):
                                dd = int(d_parsed.day)
                    except Exception:
                        pass
                    if 1 <= mm <= 12:
                        return date(yy, mm, dd)
                    return None
                except Exception:
                    return None
            out["RENEWALDATE"] = out.apply(_mk_date, axis=1)

            # Row-level reasons
            def _row_reasons(r: pd.Series) -> List[str]:
                reasons: List[str] = []
                if _row_is_missing(r["PROGRAMNAME"]): reasons.append("missing PROGRAMNAME")
                if _row_is_missing(r["TEAMNAME"]): reasons.append("missing TEAMNAME")
                if _row_is_missing(r["GROUPNAME"]): reasons.append("missing GROUPNAME")
                if _row_is_missing(r["APPNAME"]): reasons.append("missing APPNAME")
                if _row_is_missing(r["VENDORNAME"]): reasons.append("missing VENDORNAME")
                if pd.isna(r["FISCAL_YEAR"]): reasons.append("missing FISCAL_YEAR")
                if pd.isna(r["RENEWAL_MONTH"]): reasons.append("missing RENEWAL_MONTH")
                # contract requireds
                if pd.isna(r["CONTRACT_START_FY"]): reasons.append("missing CONTRACT_START_FY")
                if pd.isna(r["CONTRACT_END_FY"]): reasons.append("missing CONTRACT_END_FY")
                if pd.isna(r["CONTRACT_ANNUAL_AMOUNT"]): reasons.append("missing CONTRACT_ANNUAL_AMOUNT")
                if r.get("INVOICE_RENEWAL_DATE") is None: reasons.append("missing INVOICE_RENEWAL_DATE")
                try:
                    mm = int(r["RENEWAL_MONTH"]) if pd.notna(r["RENEWAL_MONTH"]) else None
                    if mm is not None and (mm < 1 or mm > 12):
                        reasons.append("RENEWAL_MONTH out of range")
                except Exception:
                    reasons.append("RENEWAL_MONTH invalid")
                return reasons

            out["_REASONS"] = out.apply(_row_reasons, axis=1)
            out["_IS_VALID"] = out["_REASONS"].apply(lambda lst: len(lst) == 0)

            dfI = out

        # Store previews
        # Contracts preview (unique App×Team)
        dfK = pd.DataFrame()
        try:
            # Ensure contract required mapping exists
            contract_map_ok = all(_col(k) for k in CONTRACT_REQUIRED_KEYS)
            if contract_map_ok and _col("TEAMNAME") and _col("APPNAME"):
                dfK = pd.DataFrame({
                    "TEAMNAME": df_src[_col("TEAMNAME")].astype(str).str.strip(),
                    "APPNAME": df_src[_col("APPNAME")].astype(str).str.strip(),
                    "CONTRACT_START_FY": pd.to_numeric(df_src[_col("CONTRACT_START_FY")], errors="coerce").astype("Int64"),
                    "CONTRACT_END_FY": pd.to_numeric(df_src[_col("CONTRACT_END_FY")], errors="coerce").astype("Int64"),
                    "CONTRACT_ANNUAL_AMOUNT": _safe_num(df_src[_col("CONTRACT_ANNUAL_AMOUNT")]),
                    "CONTRACT_ESCALATION_PCT": _safe_num(df_src[_col("CONTRACT_ESCALATION_PCT")]) if _col("CONTRACT_ESCALATION_PCT") else 0.0,
                    "CONTRACT_STATUS": df_src[_col("CONTRACT_STATUS")].astype(str).str.strip() if _col("CONTRACT_STATUS") else "Active",
                    "CONTRACT_AGREEMENT_NUMBER": df_src[_col("CONTRACT_AGREEMENT_NUMBER")].astype(str).str.strip() if _col("CONTRACT_AGREEMENT_NUMBER") else None,
                    "CONTRACT_COMPANY_CODE": df_src[_col("CONTRACT_COMPANY_CODE")].astype(str).str.strip() if _col("CONTRACT_COMPANY_CODE") else None,
                    "CONTRACT_COST_CENTER": df_src[_col("CONTRACT_COST_CENTER")].astype(str).str.strip() if _col("CONTRACT_COST_CENTER") else None,
                    "CONTRACT_SERVICE_TYPE": df_src[_col("CONTRACT_SERVICE_TYPE")].astype(str).str.strip() if _col("CONTRACT_SERVICE_TYPE") else None,
                    "CONTRACT_TOTAL_COST": _safe_num(df_src[_col("CONTRACT_TOTAL_COST")]) if _col("CONTRACT_TOTAL_COST") else None,
                })
                if _col("GROUPNAME"):
                    dfK["GROUPNAME"] = df_src[_col("GROUPNAME")].astype(str).str.strip()
                # Derive CONTRACT_RENEWAL_MONTH from dates if mapping not provided
                try:
                    if _col("CONTRACT_RENEWAL_MONTH"):
                        dfK["CONTRACT_RENEWAL_MONTH"] = pd.to_numeric(df_src[_col("CONTRACT_RENEWAL_MONTH")], errors="coerce").astype("Int64")
                    else:
                        inv_m = pd.Series(dtype="Int64")
                        con_m = pd.Series(dtype="Int64")
                        if _col("INVOICE_RENEWAL_DATE"):
                            inv_m = df_src[_col("INVOICE_RENEWAL_DATE")].apply(_parse_md_or_mdy)
                            inv_m = pd.to_datetime(inv_m, errors="coerce").dt.month.astype("Int64")
                        if _col("CONTRACT_RENEWAL_DATE"):
                            con_m = pd.to_datetime(df_src[_col("CONTRACT_RENEWAL_DATE")], errors="coerce").dt.month.astype("Int64")
                        # Reindex to match dfK length if needed
                        if inv_m.shape[0] != dfK.shape[0]:
                            inv_m = inv_m.reindex(dfK.index)
                        if con_m.shape[0] != dfK.shape[0]:
                            con_m = con_m.reindex(dfK.index)
                        dfK["CONTRACT_RENEWAL_MONTH"] = inv_m.combine_first(con_m)
                except Exception:
                    pass
                # Renewal date
                if _col("CONTRACT_RENEWAL_DATE"):
                    try:
                        dfK["CONTRACT_RENEWAL_DATE"] = pd.to_datetime(df_src[_col("CONTRACT_RENEWAL_DATE")], errors="coerce").dt.date
                    except Exception:
                        dfK["CONTRACT_RENEWAL_DATE"] = None
                # Also include invoice renewal date for preview and later upsert_contract use
                if _col("INVOICE_RENEWAL_DATE"):
                    try:
                        dfK["INVOICE_RENEWAL_DATE"] = pd.to_datetime(df_src[_col("INVOICE_RENEWAL_DATE")], errors="coerce").dt.date
                    except Exception:
                        dfK["INVOICE_RENEWAL_DATE"] = None
                dfK = dfK[~dfK["TEAMNAME"].apply(_blank_or_nan) & ~dfK["APPNAME"].apply(_blank_or_nan)]
                # Prefer later rows on duplicates so user-entered rows override template examples
                dfK = dfK.drop_duplicates(subset=["TEAMNAME","APPNAME"], keep="last").reset_index(drop=True)
        except Exception:
            dfK = pd.DataFrame()

        st.session_state["previews"] = {
            "Programs": dfP, "Vendors": dfV, "Groups": dfG, "Apps": dfA, "Teams": dfT, "Contracts": dfK, "Invoices": dfI
        }

        # ----- Show previews + preflight -----
        st.markdown("#### Previews")
        p1, p2 = st.columns(2)
        with p1:
            st.markdown("**Programs (unique)**")
            st.dataframe(dfP.head(400), use_container_width=True, height=220)
            st.markdown("**Vendors (unique)**")
            st.dataframe(dfV.head(400), use_container_width=True, height=220)
            st.markdown("**Applications (unique)**")
            st.dataframe(dfG.head(400), use_container_width=True, height=240)
        with p2:
            st.markdown("**Applications (unique)**")
            st.dataframe(dfA.head(400), use_container_width=True, height=240)
            st.markdown("**Teams (unique)**")
            st.dataframe(dfT.head(400), use_container_width=True, height=220)
        st.markdown("**Contracts (App × Team, unique)**")
        st.dataframe(dfK.head(400), use_container_width=True, height=220)

        if not is_contracts_only:
            st.markdown("**Invoices — Row‑level Preflight**")
            if missing_invoice_reqs:
                st.error(
                    "Missing column mappings for required field(s): "
                    + ", ".join(missing_invoice_reqs)
                    + ". Map these columns to proceed."
                )
                df_valid_preview = pd.DataFrame()
                df_invalid_preview = pd.DataFrame()
            else:
                c_valid, c_invalid = st.columns(2)
                with c_valid:
                    st.caption("✅ Will be imported (valid rows)")
                    df_valid_preview = dfI[dfI["_IS_VALID"]].copy() if ("_IS_VALID" in dfI.columns) else pd.DataFrame()
                    show_cols_valid = ["PROGRAMNAME","TEAMNAME","GROUPNAME","APPNAME","VENDORNAME","FISCAL_YEAR","RENEWAL_MONTH","INVOICE_RENEWAL_DATE","RENEWALDATE","AMOUNT"]
                    show_cols_valid = [c for c in show_cols_valid if c in df_valid_preview.columns]
                    st.dataframe(df_valid_preview[show_cols_valid].head(400), use_container_width=True, height=260)
                with c_invalid:
                    st.caption("❌ Will be skipped (invalid rows) — with reasons")
                    df_invalid_preview = dfI[~dfI["_IS_VALID"]].copy() if ("_IS_VALID" in dfI.columns) else pd.DataFrame()
                    if not df_invalid_preview.empty and "_REASONS" in df_invalid_preview.columns:
                        df_invalid_preview["_REASONS_STR"] = df_invalid_preview["_REASONS"].apply(lambda xs: "; ".join(xs))
                    show_cols_invalid = ["PROGRAMNAME","TEAMNAME","GROUPNAME","APPNAME","VENDORNAME","FISCAL_YEAR","RENEWAL_MONTH","INVOICE_RENEWAL_DATE","AMOUNT","AMOUNT_NEXT_YEAR","_REASONS_STR"]
                    show_cols_invalid = [c for c in show_cols_invalid if c in df_invalid_preview.columns]
                    st.dataframe(df_invalid_preview[show_cols_invalid].head(400), use_container_width=True, height=260)
        else:
            st.markdown("**Contracts — Preflight**")
            # Check required mappings for contracts
            contract_required_map = ["TEAMNAME","APPNAME","CONTRACT_START_FY","CONTRACT_END_FY","CONTRACT_ANNUAL_AMOUNT"]
            contract_missing_map = [k for k in contract_required_map if not _col(k)]
            if contract_missing_map:
                st.error("Missing column mappings for required contract field(s): " + ", ".join(contract_missing_map))
            # Validate dfK rows
            def _contract_reasons(r: pd.Series) -> list:
                reasons = []
                if _blank_or_nan(r.get("TEAMNAME")): reasons.append("missing TEAMNAME")
                if _blank_or_nan(r.get("APPNAME")): reasons.append("missing APPNAME")
                if pd.isna(r.get("CONTRACT_START_FY")): reasons.append("missing CONTRACT_START_FY")
                if pd.isna(r.get("CONTRACT_END_FY")): reasons.append("missing CONTRACT_END_FY")
                if pd.isna(r.get("CONTRACT_ANNUAL_AMOUNT")): reasons.append("missing CONTRACT_ANNUAL_AMOUNT")
                if _blank_or_nan(r.get("CONTRACT_RENEWAL_DATE")): reasons.append("missing CONTRACT_RENEWAL_DATE")
                return reasons
            dfK_preview = dfK.copy() if isinstance(dfK, pd.DataFrame) else pd.DataFrame()
            if not dfK_preview.empty:
                dfK_preview["_REASONS"] = dfK_preview.apply(_contract_reasons, axis=1)
                dfK_valid = dfK_preview[dfK_preview["_REASONS"].apply(len) == 0]
                dfK_invalid = dfK_preview[dfK_preview["_REASONS"].apply(len) > 0]
            else:
                dfK_valid = pd.DataFrame(); dfK_invalid = pd.DataFrame()
            cc1, cc2 = st.columns(2)
            with cc1:
                st.caption("✅ Contracts to import (valid, unique Team×App)")
                show_cols = [
                    "TEAMNAME","APPNAME",
                    "CONTRACT_RENEWAL_DATE","INVOICE_RENEWAL_DATE","CONTRACT_RENEWAL_MONTH",
                    "CONTRACT_START_FY","CONTRACT_END_FY","CONTRACT_ANNUAL_AMOUNT"
                ]
                show_cols = [c for c in show_cols if c in dfK_valid.columns]
                st.dataframe(dfK_valid[show_cols].head(400), use_container_width=True, height=260)
            with cc2:
                st.caption("❌ Skipped contracts (invalid) — with reasons")
                if not dfK_invalid.empty:
                    dfK_invalid["_REASONS_STR"] = dfK_invalid["_REASONS"].apply(lambda xs: "; ".join(xs))
                show_cols2 = [
                    "TEAMNAME","APPNAME",
                    "CONTRACT_RENEWAL_DATE","INVOICE_RENEWAL_DATE","CONTRACT_RENEWAL_MONTH",
                    "CONTRACT_START_FY","CONTRACT_END_FY","CONTRACT_ANNUAL_AMOUNT",
                    "_REASONS_STR"
                ]
                show_cols2 = [c for c in show_cols2 if c in dfK_invalid.columns]
                st.dataframe(dfK_invalid[show_cols2].head(400), use_container_width=True, height=260)

        with st.expander("Preflight Summary", expanded=True):
            if not is_contracts_only:
                if missing_invoice_reqs:
                    st.warning("Map all required invoice columns to enable import.")
                    total_rows = int(dfI.shape[0])
                    st.write(f"- Total invoice rows detected: **{total_rows}**")
                else:
                    total_rows = int(dfI.shape[0])
                    valid_rows_n = int(df_valid_preview.shape[0])
                    invalid_rows_n = int(df_invalid_preview.shape[0])
                    st.write(
                        f"- Total invoice rows detected: **{total_rows}**  \n"
                        f"- Valid rows (will import): **{valid_rows_n}**  \n"
                        f"- Invalid rows (skipped): **{invalid_rows_n}**"
                    )
                    if invalid_rows_n > 0 and "_REASONS" in df_invalid_preview.columns:
                        agg: Dict[str, int] = {}
                        for reasons in df_invalid_preview["_REASONS"].tolist():
                            for reason in reasons:
                                agg[reason] = agg.get(reason, 0) + 1
                        if agg:
                            st.write("**Top issues:**")
                            for reason, n in sorted(agg.items(), key=lambda x: -x[1]):
                                st.write(f"- {reason}: **{n}**")
            else:
                total_contracts = int(dfK.shape[0]) if isinstance(dfK, pd.DataFrame) else 0
                valid_n = int(dfK_valid.shape[0]) if 'dfK_valid' in locals() else 0
                invalid_n = int(dfK_invalid.shape[0]) if 'dfK_invalid' in locals() else 0
                st.write(
                    f"- Total contracts detected (unique Team×App): **{total_contracts}**  \n"
                    f"- Valid contracts (will import): **{valid_n}**  \n"
                    f"- Invalid contracts (skipped): **{invalid_n}**"
                )
                if invalid_n > 0 and not dfK_invalid.empty and "_REASONS" in dfK_invalid.columns:
                    agg2: Dict[str, int] = {}
                    for reasons in dfK_invalid["_REASONS"].tolist():
                        for reason in reasons:
                            agg2[reason] = agg2.get(reason, 0) + 1
                    if agg2:
                        st.write("**Top issues:**")
                        for reason, n in sorted(agg2.items(), key=lambda x: -x[1]):
                            st.write(f"- {reason}: **{n}**")

            # Duplicate checks (block import if duplicates)
            has_dupes_contracts = False
            has_dupes_invoices = False
            try:
                dfK = st.session_state["previews"].get("Contracts")
                if isinstance(dfK, pd.DataFrame) and not dfK.empty:
                    dupk = dfK[dfK.duplicated(subset=["TEAMNAME","APPNAME"], keep=False)].copy()
                    if not dupk.empty:
                        has_dupes_contracts = True
                        st.error(f"Duplicate Contracts detected (same TEAMNAME+APPNAME): {len(dupk)} row(s).")
                        st.dataframe(dupk, use_container_width=True, height=180)
            except Exception:
                pass

        # -------------------------
        # Contracts → Planned invoices (preview)
        # -------------------------
        try:
            dfK_preview_src = st.session_state.get("previews", {}).get("Contracts")
            if isinstance(dfK_preview_src, pd.DataFrame) and not dfK_preview_src.empty:
                with st.expander("Contracts → Planned invoices (preview)", expanded=False):
                    # Use only rows that have required fields
                    dfKp = dfK_preview_src.copy()
                    for col in ("TEAMNAME","APPNAME","CONTRACT_START_FY","CONTRACT_END_FY"):
                        if col not in dfKp.columns:
                            dfKp[col] = None
                    dfKp = dfKp[
                        (~dfKp["TEAMNAME"].apply(_blank_or_nan)) &
                        (~dfKp["APPNAME"].apply(_blank_or_nan)) &
                        (pd.notna(dfKp["CONTRACT_START_FY"])) &
                        (pd.notna(dfKp["CONTRACT_END_FY"]))
                    ].reset_index(drop=True)
                    if dfKp.empty:
                        st.info("No valid contract rows to preview.")
                    else:
                        # Build maps from DB
                        teams_now = list_teams(); tmap_now = {}
                        if teams_now is not None and not teams_now.empty:
                            for _, r0 in teams_now.iterrows():
                                tmap_now[str(r0.TEAMNAME).strip().upper()] = r0.TEAMID
                        apps_now = list_applications(); amap_now = {}
                        if apps_now is not None and not apps_now.empty:
                            for _, r0 in apps_now.iterrows():
                                amap_now[str(r0.APPLICATIONNAME).strip().upper()] = r0.APPLICATIONID

                        preview_rows: List[Dict[str, Any]] = []
                        candidate_pairs: List[Tuple[str,str]] = []
                        # First resolve pairs and compute FY spans
                        for _, r in dfKp.iterrows():
                            tname = str(r.get("TEAMNAME") or "").strip()
                            aname = str(r.get("APPNAME") or "").strip()
                            gname = str(r.get("GROUPNAME") or "").strip() if "GROUPNAME" in dfKp.columns else ""
                            sfy = int(pd.to_numeric(r.get("CONTRACT_START_FY"), errors="coerce")) if pd.notna(r.get("CONTRACT_START_FY")) else None
                            efy = int(pd.to_numeric(r.get("CONTRACT_END_FY"), errors="coerce")) if pd.notna(r.get("CONTRACT_END_FY")) else None
                            fy_span = max(0, (efy - sfy + 1)) if (sfy is not None and efy is not None) else 0
                            # Resolve canonical app name like in import
                            final_appname = aname
                            if gname:
                                prefix = f"{gname} - "
                                if aname.upper().startswith(prefix.upper()):
                                    final_appname = aname
                                elif aname.upper() == gname.upper():
                                    final_appname = f"{gname} - Instance"
                                else:
                                    final_appname = f"{gname} - {aname}"
                            tid = tmap_now.get(tname.upper())
                            aid = amap_now.get(final_appname.upper()) or amap_now.get(aname.upper())
                            row: Dict[str, Any] = {
                                "TEAMNAME": tname,
                                "APPNAME": aname,
                                "CANONICAL_APPNAME": final_appname,
                                "START_FY": sfy,
                                "END_FY": efy,
                                "FY_SPAN": fy_span,
                            }
                            if tid and aid and fy_span > 0:
                                row["_TID"] = tid; row["_AID"] = aid
                                candidate_pairs.append((tid, aid))
                            else:
                                reason = []
                                if not tid: reason.append("Team not found")
                                if not aid: reason.append("Application not found")
                                if fy_span <= 0: reason.append("Invalid FY span")
                                row["UNMATCHED_REASON"] = "; ".join(reason)
                            preview_rows.append(row)

                        # Fetch existing recurring invoice FYs for these pairs
                        existing_map: set[Tuple[str,str,int]] = set()
                        a_ids = sorted(set(a for _, a in candidate_pairs))
                        t_ids = sorted(set(t for t, _ in candidate_pairs))
                        if a_ids and t_ids:
                            try:
                                placeholders_a = ",".join(["%s"] * len(a_ids))
                                placeholders_t = ",".join(["%s"] * len(t_ids))
                                sql_exist = f"""
                                    SELECT APPLICATIONID, TEAMID, FISCAL_YEAR
                                    FROM INVOICES
                                    WHERE COALESCE(INVOICE_TYPE,'Recurring Invoice') = 'Recurring Invoice'
                                      AND APPLICATIONID IN ({placeholders_a})
                                      AND TEAMID IN ({placeholders_t})
                                """
                                df_exist = fetch_df(sql_exist, tuple(list(a_ids) + list(t_ids)))
                                if df_exist is not None and not df_exist.empty:
                                    for _, re in df_exist.iterrows():
                                        try:
                                            existing_map.add((str(re["TEAMID"]), str(re["APPLICATIONID"]), int(re["FISCAL_YEAR"])) )
                                        except Exception:
                                            pass
                            except Exception:
                                pass

                        # Finalize counts
                        for row in preview_rows:
                            tid = row.get("_TID"); aid = row.get("_AID")
                            sfy = row.get("START_FY"); efy = row.get("END_FY")
                            if tid and aid and isinstance(sfy, int) and isinstance(efy, int) and efy >= sfy:
                                yrs = list(range(int(sfy), int(efy) + 1))
                                exists_n = sum(1 for y in yrs if (str(tid), str(aid), int(y)) in existing_map)
                                row["EXISTING_PLANNED"] = exists_n
                                row["WILL_CREATE"] = max(0, len(yrs) - exists_n)
                            else:
                                row.setdefault("EXISTING_PLANNED", None)
                                row.setdefault("WILL_CREATE", None)

                        out_df = pd.DataFrame(preview_rows)
                        # Tidy output
                        show_cols = [
                            "TEAMNAME","CANONICAL_APPNAME","START_FY","END_FY","FY_SPAN","EXISTING_PLANNED","WILL_CREATE","UNMATCHED_REASON"
                        ]
                        show_cols = [c for c in show_cols if c in out_df.columns]
                        st.dataframe(out_df[show_cols], use_container_width=True, height=260)
                        try:
                            total_create = int(out_df["WILL_CREATE"].fillna(0).sum())
                            st.caption(f"Estimated new planned invoices to create: {total_create}")
                        except Exception:
                            pass
        except Exception:
            pass
            try:
                if isinstance(df_valid_preview, pd.DataFrame) and not df_valid_preview.empty:
                    dupi = df_valid_preview[df_valid_preview.duplicated(subset=["TEAMNAME","APPNAME","FISCAL_YEAR"], keep=False)].copy()
                    if not dupi.empty:
                        has_dupes_invoices = True
                        st.error(f"Duplicate Invoices detected (same TEAMNAME+APPNAME+FISCAL_YEAR): {len(dupi)} row(s).")
                        st.dataframe(dupi, use_container_width=True, height=180)
            except Exception:
                pass

        # -------------------------
        # Import ALL in safe order
        # -------------------------
        prog_widget = st.progress(0, text="Ready.")
        status_txt = st.empty()

        def _tick(pct: int, msg: str, start_time: float) -> None:
            elapsed = time.time() - start_time
            prog_widget.progress(pct, text=f"{msg}  ⏱ {elapsed:,.1f}s")
            status_txt.caption(f"⏱ Elapsed: **{elapsed:,.1f} s**")

        # Enable import if either valid invoices exist OR contracts are present, and no duplicates
        # Ensure df_valid_preview is defined in all modes
        if 'df_valid_preview' not in locals():
            df_valid_preview = pd.DataFrame()
        disabled_btn = (df_valid_preview.empty and (dfK is None or dfK.empty))
        try:
            if isinstance(dfK, pd.DataFrame) and not dfK.empty and dfK.duplicated(subset=["TEAMNAME","APPNAME"]).any():
                disabled_btn = True
        except Exception:
            pass
        try:
            if isinstance(df_valid_preview, pd.DataFrame) and not df_valid_preview.empty and df_valid_preview.duplicated(subset=["TEAMNAME","APPNAME","FISCAL_YEAR"]).any():
                disabled_btn = True
        except Exception:
            pass

        if st.button("Import ALL", icon=":material/download:", type="primary", use_container_width=True, disabled=disabled_btn):
            start_time = time.time()
            import_batch_id = str(uuid.uuid4())
            try:
                pre = {
                    "vendors": _table_count("VENDORS"),
                    "programs": _table_count("PROGRAMS"),
                    "teams": _table_count("TEAMS"),
                    "groups": _table_count("APPLICATION_GROUPS"),
                    "apps": _table_count("APPLICATIONS"),
                    "invoices": _table_count("INVOICES"),
                }

                # Refresh previews (safe)
                dfP = _get_preview("Programs")
                dfV = _get_preview("Vendors")
                dfG = _get_preview("Groups")
                dfA = _get_preview("Apps")
                dfT = _get_preview("Teams")
                dfI = _get_preview("Invoices")
                dfK = _get_preview("Contracts")

                # ---------- Vendors ----------
                _tick(5, "Importing Vendors...", start_time)
                existing_vendors = list_vendors()
                vmap: Dict[str, str] = {}
                if existing_vendors is not None and not existing_vendors.empty:
                    for _, r0 in existing_vendors.iterrows():
                        vmap[str(r0.VENDORNAME).strip().upper()] = r0.VENDORID
                inserted_v = updated_v = 0
                if dfV is not None and not dfV.empty:
                    for _, row in dfV.iterrows():
                        vname = str(row["VENDORNAME"]).strip()
                        if _blank_or_nan(vname):
                            continue
                        existed = vname.upper() in vmap
                        vid = vmap.get(vname.upper()) or str(uuid.uuid4())
                        upsert_vendor(vid, vname)
                        vmap[vname.upper()] = vid
                        inserted_v += 0 if existed else 1
                        updated_v  += 1 if existed else 0

                # ---------- Programs ----------
                _tick(12, "Importing Programs...", start_time)
                existing_programs = list_programs()
                pmap: Dict[str, str] = {}
                if existing_programs is not None and not existing_programs.empty:
                    for _, r0 in existing_programs.iterrows():
                        pmap[str(r0.PROGRAMNAME).strip().upper()] = r0.PROGRAMID
                inserted_p = updated_p = 0
                if dfP is not None and not dfP.empty:
                    for _, row in dfP.iterrows():
                        pname = str(row["PROGRAMNAME"]).strip()
                        if _blank_or_nan(pname):
                            continue
                        existed = pname.upper() in pmap
                        pid = pmap.get(pname.upper()) or str(uuid.uuid4())
                        upsert_program(pid, pname, owner=None, fte=None)
                        pmap[pname.upper()] = pid
                        inserted_p += 0 if existed else 1
                        updated_p  += 1 if existed else 0

                # ---------- Teams ----------
                _tick(22, "Importing Teams...", start_time)
                existing_teams = list_teams()
                tmap: Dict[str, str] = {}
                if existing_teams is not None and not existing_teams.empty:
                    for _, r0 in existing_teams.iterrows():
                        tmap[str(r0.TEAMNAME).strip().upper()] = r0.TEAMID
                inserted_t = updated_t = 0
                if dfT is not None and not dfT.empty:
                    for _, row in dfT.iterrows():
                        tname = str(row["TEAMNAME"]).strip()
                        if _blank_or_nan(tname):
                            continue
                        pid = None
                        if "PROGRAMNAME" in row and not _blank_or_nan(row["PROGRAMNAME"]):
                            pid = pmap.get(str(row["PROGRAMNAME"]).strip().upper())
                        existed = tname.upper() in tmap
                        tid = tmap.get(tname.upper()) or str(uuid.uuid4())
                        upsert_team(
                            team_id=tid,
                            name=tname,
                            program_id=pid,
                            team_fte=None, delivery_team_fte=None,
                            contractor_c_fte=None, contractor_cs_fte=None
                        )
                        tmap[tname.upper()] = tid
                        inserted_t += 0 if existed else 1
                        updated_t  += 1 if existed else 0

                # ---------- Groups ----------
                _tick(35, "Importing Applications...", start_time)

                # Build TEAM map (already built above, but refresh safely)
                existing_teams2 = list_teams()
                tmap: Dict[str, str] = {}
                if existing_teams2 is not None and not existing_teams2.empty:
                    for _, r0 in existing_teams2.iterrows():
                        tmap[str(r0.TEAMNAME).strip().upper()] = r0.TEAMID

                # Build VENDOR map
                existing_vendors2 = list_vendors()
                vmap: Dict[str, str] = {}
                if existing_vendors2 is not None and not existing_vendors2.empty:
                    for _, r0 in existing_vendors2.iterrows():
                        vmap[str(r0.VENDORNAME).strip().upper()] = r0.VENDORID

                # Build a composite GROUP map from DB: (TEAMID, UPPER(GROUPNAME)) -> GROUPID
                gpair_map: Dict[Tuple[str, str], str] = {}
                cur_groups_all = fetch_df("""
                    SELECT g.GROUPID, g.GROUPNAME, g.TEAMID
                    FROM APPLICATION_GROUPS g
                """)
                if cur_groups_all is not None and not cur_groups_all.empty:
                    for _, r in cur_groups_all.iterrows():
                        gid = str(r["GROUPID"]).strip()
                        gname = str(r["GROUPNAME"]).strip().upper()
                        tid = str(r["TEAMID"]).strip() if pd.notna(r["TEAMID"]) else ""
                        if gid and gname and tid:
                            gpair_map[(tid, gname)] = gid

                inserted_g = updated_g = 0
                if dfG is not None and not dfG.empty:
                    for _, row in dfG.iterrows():
                        gname = str(row["GROUPNAME"]).strip()
                        tname = str(row["TEAMNAME"]).strip() if "TEAMNAME" in row else ""
                        if _blank_or_nan(gname) or _blank_or_nan(tname):
                            continue

                        tid = tmap.get(tname.upper())
                        if not tid:
                            # cannot place group without a team
                            continue

                        default_vendor_id = None
                        if "DEFAULT_VENDORNAME" in row and not _blank_or_nan(row["DEFAULT_VENDORNAME"]):
                            default_vendor_id = vmap.get(str(row["DEFAULT_VENDORNAME"]).strip().upper())

                        key = (tid, gname.upper())
                        existed = key in gpair_map
                        gid = gpair_map.get(key) or str(uuid.uuid4())

                        # Upsert by (TEAMID, GROUPNAME). Your upsert helper likely merges on GROUPID.
                        upsert_application_group(
                            gid,
                            gname,
                            team_id=tid,
                            default_vendor_id=default_vendor_id,
                            owner=None
                        )

                        gpair_map[key] = gid
                        inserted_g += 0 if existed else 1
                        updated_g  += 1 if existed else 0


                    # ---------- Applications ----------
                    _tick(48, "Importing Application Instances...", start_time)

                    # Build APPLICATION map by name (still global, since app names look canonical)
                    existing_apps = list_applications()
                    amap: Dict[str, str] = {}
                    if existing_apps is not None and not existing_apps.empty:
                        for _, r0 in existing_apps.iterrows():
                            amap[str(r0.APPLICATIONNAME).strip().upper()] = r0.APPLICATIONID

                    inserted_a = updated_a = 0
                    if dfA is not None and not dfA.empty:
                        # refresh composite group map (TEAMID × GROUPNAME)
                        cur_groups = fetch_df("SELECT GROUPID, GROUPNAME, TEAMID FROM APPLICATION_GROUPS")
                        gpair_map: Dict[Tuple[str,str], str] = {}
                        gname_map: Dict[str, str] = {}
                        if cur_groups is not None and not cur_groups.empty:
                            for _, r in cur_groups.iterrows():
                                gid = str(r["GROUPID"]).strip()
                                gname = str(r["GROUPNAME"]).strip().upper()
                                tid = str(r["TEAMID"]).strip()
                                if gid and gname and tid:
                                    gpair_map[(tid, gname)] = gid
                                    # Also map by name only (GROUPNAME is globally unique in schema)
                                    if gname not in gname_map:
                                        gname_map[gname] = gid

                        # Need team map to resolve TEAMNAME → TEAMID
                        teams_now = list_teams()
                        tmap_now: Dict[str, str] = {}
                        if teams_now is not None and not teams_now.empty:
                            for _, r0 in teams_now.iterrows():
                                tmap_now[str(r0.TEAMNAME).strip().upper()] = r0.TEAMID

                        for _, row in dfA.iterrows():
                            aname = str(row["APPNAME"]).strip()
                            gname = str(row["GROUPNAME"]).strip() if "GROUPNAME" in row else ""
                            tname = str(row["TEAMNAME"]).strip() if "TEAMNAME" in row else ""
                            if _blank_or_nan(aname):
                                continue

                            # If we have both Team & Group, resolve to the correct group's ID
                            gid = None
                            if not _blank_or_nan(gname) and not _blank_or_nan(tname):
                                tid = tmap_now.get(tname.upper())
                                if tid:
                                    gid = gpair_map.get((tid, gname.upper()))
                            # Fallback: resolve by group name only (unique globally)
                            if gid is None and not _blank_or_nan(gname):
                                gid = gname_map.get(gname.upper())

                            # Canonicalize application name to "Group - Suffix" when group provided
                            final_name = aname
                            if not _blank_or_nan(gname):
                                prefix = f"{gname} - "
                                if aname.upper().startswith(prefix.upper()):
                                    final_name = aname
                                elif aname.upper() == gname.upper():
                                    final_name = f"{gname} - Instance"
                                else:
                                    final_name = f"{gname} - {aname}"

                            existing_id_by_raw = amap.get(aname.upper())
                            existed = final_name.upper() in amap or (existing_id_by_raw is not None)

                            # If an app exists named exactly as the Group (bad data), rename it to canonical instead of creating a duplicate
                            if (not _blank_or_nan(gname)) and (aname.upper() == gname.upper()) and existing_id_by_raw and gid:
                                try:
                                    execute(
                                        "UPDATE APPLICATIONS SET APPLICATIONNAME=%s, GROUPID = COALESCE(GROUPID, %s) WHERE APPLICATIONID=%s",
                                        (final_name, gid, existing_id_by_raw),
                                    )
                                    # update local map
                                    amap.pop(aname.upper(), None)
                                    amap[final_name.upper()] = existing_id_by_raw
                                    updated_a += 1
                                    continue
                                except Exception:
                                    # fall through to normal upsert if rename fails
                                    pass

                            aid = amap.get(final_name.upper()) or existing_id_by_raw or str(uuid.uuid4())

                            upsert_application_instance(
                                application_id=aid,
                                group_id=(gid or None),
                                application_name=final_name,
                                add_info=None,
                                vendor_id=None
                            )
                            amap[final_name.upper()] = aid
                            inserted_a += 0 if existed else 1
                            updated_a  += 1 if existed else 0


                        # ---------- Contracts ----------
                        _tick(58, "Importing Contracts...", start_time)

                        ensure_contracts_table()
                        # Refresh maps for IDs
                        cur_teams3 = list_teams(); tmap3 = {}
                        if cur_teams3 is not None and not cur_teams3.empty:
                            for _, r0 in cur_teams3.iterrows():
                                tmap3[str(r0.TEAMNAME).strip().upper()] = r0.TEAMID

                        cur_apps3 = list_applications(); amap3 = {}
                        if cur_apps3 is not None and not cur_apps3.empty:
                            for _, r0 in cur_apps3.iterrows():
                                amap3[str(r0.APPLICATIONNAME).strip().upper()] = r0.APPLICATIONID

                        created_c = updated_c = 0
                        created_invoices_total = 0
                        created_contracts_detail: List[Dict[str, Any]] = []
                        if dfK is not None and not dfK.empty:
                            from db import sync_contract_invoices
                            unmatched_contracts: List[Dict[str, Any]] = []
                            failed_contracts: List[Dict[str, Any]] = []
                            for _, row in dfK.iterrows():
                                tname = str(row.get("TEAMNAME") or "").strip()
                                aname = str(row.get("APPNAME") or "").strip()
                                gname = str(row.get("GROUPNAME") or "").strip() if "GROUPNAME" in row else ""
                                if _blank_or_nan(tname) or _blank_or_nan(aname):
                                    continue
                                tid = tmap3.get(tname.upper())
                                # Resolve ApplicationID by canonical name "Group - Suffix" where possible
                                final_appname = aname
                                if gname:
                                    prefix = f"{gname} - "
                                    if aname.upper().startswith(prefix.upper()):
                                        final_appname = aname
                                    elif aname.upper() == gname.upper():
                                        final_appname = f"{gname} - Instance"
                                    else:
                                        final_appname = f"{gname} - {aname}"
                                aid = amap3.get(final_appname.upper()) or amap3.get(aname.upper())
                                if not tid or not aid:
                                    reason = ", ".join([
                                        r for r in [
                                            ("Team not found" if not tid else None),
                                            ("Application not found" if not aid else None)
                                        ] if r
                                    ]) or "Unresolved"
                                    unmatched_contracts.append({
                                        "TEAMNAME": tname,
                                        "APPNAME": aname,
                                        "CANONICAL_APPNAME": final_appname,
                                        "GROUPNAME": gname or None,
                                        "REASON": reason,
                                    })
                                    continue
                                # Safe parse contract fields
                                def _to_int_default(val, default):
                                    try:
                                        if val is None or (hasattr(pd, 'isna') and pd.isna(val)):
                                            return default
                                        s = str(val).strip()
                                        if s == "":
                                            return default
                                        v = int(float(s))
                                        return v
                                    except Exception:
                                        return default
                                def _to_float_default(val, default):
                                    try:
                                        if val is None or (hasattr(pd, 'isna') and pd.isna(val)):
                                            return default
                                        s = str(val).strip()
                                        if s == "":
                                            return default
                                        return float(s)
                                    except Exception:
                                        return default

                                start_fy_val = _to_int_default(row.get("CONTRACT_START_FY"), None)
                                end_fy_val   = _to_int_default(row.get("CONTRACT_END_FY"), None)
                                annual_val   = _to_float_default(row.get("CONTRACT_ANNUAL_AMOUNT"), None)
                                esc_val      = _to_float_default(row.get("CONTRACT_ESCALATION_PCT"), 0.0)
                                rmon_val     = _to_int_default(row.get("CONTRACT_RENEWAL_MONTH"), 1)
                                if rmon_val is None or rmon_val < 1 or rmon_val > 12:
                                    rmon_val = 1

                                if start_fy_val is None or end_fy_val is None or annual_val is None:
                                    failed_contracts.append({
                                        "TEAMNAME": tname,
                                        "APPNAME": final_appname,
                                        "REASON": "Missing/invalid START_FY, END_FY, or ANNUAL_AMOUNT",
                                        "START_FY": start_fy_val,
                                        "END_FY": end_fy_val,
                                        "ANNUAL_AMOUNT": annual_val,
                                    })
                                    continue
                                if end_fy_val < start_fy_val:
                                    failed_contracts.append({
                                        "TEAMNAME": tname,
                                        "APPNAME": final_appname,
                                        "REASON": "END_FY < START_FY",
                                        "START_FY": start_fy_val,
                                        "END_FY": end_fy_val,
                                    })
                                    continue

                                try:
                                    import uuid as _uuid
                                    cid = str(_uuid.uuid5(_uuid.NAMESPACE_URL, f"contract:{aid}:{tid}"))
                                    # Guard: prevent overlapping contracts for the same Application Instance across any Team
                                    try:
                                        overlap_b = fetch_df(
                                            """
                                            SELECT TOP 1 c.CONTRACT_ID, t.TEAMNAME, c.START_FY, c.END_FY
                                            FROM CONTRACTS c
                                            LEFT JOIN TEAMS t ON t.TEAMID = c.TEAMID
                                            WHERE c.APPLICATIONID = %s
                                              AND NOT (c.END_FY < %s OR c.START_FY > %s)
                                              AND NOT (c.TEAMID = %s)
                                            ORDER BY c.CREATED_AT DESC
                                            """,
                                            (aid, int(start_fy_val), int(end_fy_val), tid)
                                        )
                                    except Exception:
                                        overlap_b = None
                                    if overlap_b is not None and not overlap_b.empty:
                                        failed_contracts.append({
                                            "TEAMNAME": tname,
                                            "APPNAME": final_appname,
                                            "REASON": f"Overlapping contract exists for this Application Instance (Team: {str(overlap_b.iloc[0].get('TEAMNAME') or overlap_b.iloc[0].get('TEAMID') or '(unknown team)')}). Skipped.",
                                        })
                                        continue
                                    upsert_contract(
                                        contract_id=cid,
                                        application_id=aid,
                                        team_id=tid,
                                        start_fy=int(start_fy_val),
                                        end_fy=int(end_fy_val),
                                        renewal_month=int(rmon_val),
                                        annual_amount=float(annual_val),
                                        escalation_pct=float(esc_val),
                                        status=(str(row.get("CONTRACT_STATUS") or "Active")),
                                        agreement_number=(str(row.get("CONTRACT_AGREEMENT_NUMBER")).strip() or None),
                                        company_code=(str(row.get("CONTRACT_COMPANY_CODE")).strip() or None),
                                        cost_center=(str(row.get("CONTRACT_COST_CENTER")).strip() or None),
                                        service_type=(str(row.get("CONTRACT_SERVICE_TYPE")).strip() or None),
                                        contract_renewal_date=(str(row.get("CONTRACT_RENEWAL_DATE")) if pd.notna(row.get("CONTRACT_RENEWAL_DATE")) else None),
                                        invoice_renewal_date=(str(row.get("INVOICE_RENEWAL_DATE")) if pd.notna(row.get("INVOICE_RENEWAL_DATE")) else None),
                                        total_contract_cost=(float(row.get("CONTRACT_TOTAL_COST")) if pd.notna(row.get("CONTRACT_TOTAL_COST")) else None),
                                    )
                                    updated_c += 1
                                    try:
                                        n_created = int(sync_contract_invoices(cid) or 0)
                                        created_invoices_total += n_created
                                        created_contracts_detail.append({
                                            "TEAMNAME": tname,
                                            "APPNAME": final_appname,
                                            "CONTRACT_ID": cid,
                                            "START_FY": int(start_fy_val),
                                            "END_FY": int(end_fy_val),
                                            "PLANNED_INVOICES": n_created,
                                        })
                                    except Exception:
                                        pass
                                except Exception as e:
                                    failed_contracts.append({
                                        "TEAMNAME": tname,
                                        "APPNAME": final_appname,
                                        "REASON": f"Upsert failed: {str(e)[:120]}",
                                    })

                        # Optional: show a compact summary of contracts processed and planned invoices created
                        if created_contracts_detail:
                            try:
                                sum_df = pd.DataFrame(created_contracts_detail)
                                st.success(f"Contracts processed: {len(created_contracts_detail)}  ·  Planned invoices created/updated: {created_invoices_total}")
                                with st.expander("Contracts results (per Team × App)", expanded=False):
                                    # Show key columns only to keep it tidy
                                    show_cols = [
                                        "TEAMNAME", "APPNAME", "START_FY", "END_FY", "PLANNED_INVOICES", "CONTRACT_ID"
                                    ]
                                    show_cols = [c for c in show_cols if c in sum_df.columns]
                                    st.dataframe(sum_df[show_cols], use_container_width=True, height=220)
                                try:
                                    # Clear caches so other tabs reflect newly imported contracts/invoices immediately
                                    _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                                except Exception:
                                    pass
                            except Exception:
                                pass
                        if 'unmatched_contracts' in locals() and unmatched_contracts:
                            try:
                                um_df = pd.DataFrame(unmatched_contracts)
                                with st.expander("Contracts not matched (missing Team or Application)", expanded=True):
                                    st.warning("Some contract rows could not be resolved to an existing Team and/or Application. Fix names and reimport.")
                                    st.dataframe(um_df, use_container_width=True, height=220)
                            except Exception:
                                pass
                        if 'failed_contracts' in locals() and failed_contracts:
                            try:
                                fc_df = pd.DataFrame(failed_contracts)
                                with st.expander("Contracts failed to import (invalid values)", expanded=True):
                                    st.warning("Some contract rows had invalid/missing values (START_FY/END_FY/ANNUAL_AMOUNT or END_FY < START_FY). Fix and reimport.")
                                    st.dataframe(fc_df, use_container_width=True, height=220)
                            except Exception:
                                pass
                        # ---------- Invoices ----------
                        _tick(62, "Importing Invoices...", start_time)
                        created_i = 0
                        created_vendors = 0
                        created_apps = 0

                        # refresh all maps just before invoicing
                        cur_programs = list_programs();  pmap = {}
                        if cur_programs is not None and not cur_programs.empty:
                            for _, r0 in cur_programs.iterrows():
                                pmap[str(r0.PROGRAMNAME).strip().upper()] = r0.PROGRAMID

                        cur_teams = list_teams(); tmap = {}
                        if cur_teams is not None and not cur_teams.empty:
                            for _, r0 in cur_teams.iterrows():
                                tmap[str(r0.TEAMNAME).strip().upper()] = r0.TEAMID

                        cur_groups2 = list_application_groups(); gmap = {}
                        if cur_groups2 is not None and not cur_groups2.empty:
                            for _, r0 in cur_groups2.iterrows():
                                gmap[str(r0.GROUPNAME).strip().upper()] = r0.GROUPID

                        cur_vendors = list_vendors(); vmap = {}
                        if cur_vendors is not None and not cur_vendors.empty:
                            for _, r0 in cur_vendors.iterrows():
                                vmap[str(r0.VENDORNAME).strip().upper()] = r0.VENDORID

                        cur_apps2 = list_applications(); amap = {}
                        if cur_apps2 is not None and not cur_apps2.empty:
                            for _, r0 in cur_apps2.iterrows():
                                amap[str(r0.APPLICATIONNAME).strip().upper()] = r0.APPLICATIONID

                        required_present = (
                            dfI is not None and not dfI.empty and
                            all(c in dfI.columns for c in
                                ["PROGRAMNAME","TEAMNAME","GROUPNAME","APPNAME","VENDORNAME","FISCAL_YEAR","RENEWAL_MONTH","AMOUNT"])
                        )

                        if required_present and (df_valid_preview is not None) and (not df_valid_preview.empty):
                            for i, row in enumerate(df_valid_preview.itertuples(index=False), start=1):
                                pname = str(getattr(row, "PROGRAMNAME", "") or "").strip()
                                tname = str(getattr(row, "TEAMNAME", "") or "").strip()
                                gname = str(getattr(row, "GROUPNAME", "") or "").strip()
                                aname = str(getattr(row, "APPNAME", "") or "").strip()
                                vname = str(getattr(row, "VENDORNAME", "") or "").strip()
                                fy    = getattr(row, "FISCAL_YEAR", None)
                                rmonth= getattr(row, "RENEWAL_MONTH", None)
                                amt   = getattr(row, "AMOUNT", None)

                                if any(_blank_or_nan(x) for x in [pname, tname, gname, aname, vname]):
                                    continue
                                if any(pd.isna(x) for x in [fy, rmonth]):
                                    continue

                                vendorid_at_booking = vmap.get(vname.upper())
                                if not vendorid_at_booking and vname:
                                    # Auto-create missing vendor on the fly
                                    try:
                                        new_vid = str(uuid.uuid4())
                                        upsert_vendor(new_vid, vname)
                                        vmap[vname.upper()] = new_vid
                                        vendorid_at_booking = new_vid
                                        created_vendors += 1
                                    except Exception:
                                        continue

                                teamid = tmap.get(tname.upper())
                                if not teamid:
                                    continue

                                groupid = gmap.get(gname.upper()) if gname else None

                                # Resolve application by canonical name "Group - Suffix" when group is provided
                                final_appname = aname
                                if gname:
                                    prefix = f"{gname} - "
                                    if aname.upper().startswith(prefix.upper()):
                                        final_appname = aname
                                    elif aname.upper() == gname.upper():
                                        final_appname = f"{gname} - Instance"
                                    else:
                                        final_appname = f"{gname} - {aname}"

                                appid = amap.get(final_appname.upper()) if final_appname else None
                                if (appid is None) and groupid and final_appname:
                                    # Create missing application in target group
                                    try:
                                        new_aid = str(uuid.uuid4())
                                        upsert_application_instance(
                                            application_id=new_aid,
                                            group_id=groupid,
                                            application_name=final_appname,
                                            add_info=None,
                                            vendor_id=vendorid_at_booking,
                                        )
                                        amap[final_appname.upper()] = new_aid
                                        appid = new_aid
                                        created_apps += 1
                                    except Exception:
                                        pass

                                try:
                                    d_renew = date(int(fy), int(rmonth), 1)
                                except Exception:
                                    continue

                                inv_id = str(uuid.uuid4())
                                # Derive contract-active and default codes from CONTRACTS (App×Team)
                                contract_active_val = None
                                company_code_val = (str(getattr(row, "COMPANY_CODE", None)) if pd.notna(getattr(row, "COMPANY_CODE", None)) else None)
                                cost_center_val = (str(getattr(row, "COST_CENTER", None)) if pd.notna(getattr(row, "COST_CENTER", None)) else None)
                                service_type_val = None
                                try:
                                    if appid and teamid:
                                        cdfx = fetch_df(
                                            "SELECT TOP 1 UPPER(ISNULL(STATUS,'Active')) AS ST, COMPANY_CODE, COST_CENTER, SERVICE_TYPE FROM CONTRACTS WHERE APPLICATIONID=%s AND TEAMID=%s",
                                            (str(appid), str(teamid))
                                        )
                                        if cdfx is not None and not cdfx.empty:
                                            contract_active_val = (str(cdfx.iloc[0]["ST"]) == "ACTIVE")
                                            if not company_code_val and pd.notna(cdfx.iloc[0].get("COMPANY_CODE")):
                                                company_code_val = str(cdfx.iloc[0].get("COMPANY_CODE"))
                                            if not cost_center_val and pd.notna(cdfx.iloc[0].get("COST_CENTER")):
                                                cost_center_val = str(cdfx.iloc[0].get("COST_CENTER"))
                                            if pd.notna(cdfx.iloc[0].get("SERVICE_TYPE")):
                                                service_type_val = str(cdfx.iloc[0].get("SERVICE_TYPE"))
                                except Exception:
                                    pass

                                # derive invoice status from sheet (optional)
                                inv_status = None
                                try:
                                    inv_status = str(getattr(row, "INVOICE_STATUS", "") or "").strip()
                                except Exception:
                                    inv_status = None
                                upsert_invoice(
                                    invoice_id=inv_id,
                                    application_id=str(appid or ""),
                                    team_id=str(teamid or ""),
                                    renewal_date=d_renew,
                                    amount=float(amt) if not pd.isna(amt) else None,
                                    status=(inv_status if inv_status in ("Planned","Completed") else "Planned"),
                                    fiscal_year=int(fy) if not pd.isna(fy) else None,
                                    product_owner=(str(getattr(row, "PRODUCT_OWNER", None)) if pd.notna(getattr(row, "PRODUCT_OWNER", None)) else None),
                                    amount_next_year=None,
                                    contract_active=contract_active_val,
                                    company_code=company_code_val,
                                    cost_center=cost_center_val,
                                    serial_number=(str(getattr(row, "SERIAL_NUMBER", None)) if pd.notna(getattr(row, "SERIAL_NUMBER", None)) else None),
                                    work_order=(str(getattr(row, "WORK_ORDER", None)) if pd.notna(getattr(row, "WORK_ORDER", None)) else None),
                                    agreement_number=None,
                                    contract_due=None,   # now comes from CONTRACTS
                                    service_type=service_type_val,
                                    notes=(str(getattr(row, "NOTES", None)) if pd.notna(getattr(row, "NOTES", None)) else None),
                                    group_id=str(groupid or ""),
                                    programid_at_booking=None,
                                    vendorid_at_booking=str(vendorid_at_booking or ""),
                                    groupid_at_booking=str(groupid or ""),
                                    rollover_batch_id=import_batch_id,
                                    rolled_over_from_year=None,
                                    invoice_type="Recurring Invoice",
                                )

                                if i % 50 == 0:
                                    _tick(62 + min(30, int(i / max(len(df_valid_preview), 1) * 30)), f"Importing Invoices... {i} processed", start_time)
                                created_i += 1

                        _tick(96, "Finalizing...", start_time)
                        post = {
                            "vendors": _table_count("VENDORS"),
                            "programs": _table_count("PROGRAMS"),
                            "teams": _table_count("TEAMS"),
                            "groups": _table_count("APPLICATION_GROUPS"),
                            "apps": _table_count("APPLICATIONS"),
                            "invoices": _table_count("INVOICES"),
                        }

                        _tick(100, "Done.", start_time)
                        st.success(
                            "Import complete ✅\n\n"
                            f"- Vendors: **{post['vendors'] - pre['vendors']} new** (total: {post['vendors']})\n"
                            f"- Programs: **{post['programs'] - pre['programs']} new** (total: {post['programs']})\n"
                            f"- Teams: **{post['teams'] - pre['teams']} new** (total: {post['teams']})\n"
                            f"- Applications: **{post['groups'] - pre['groups']} new** (total: {post['groups']})\n"
                            f"- Applications: **{post['apps'] - pre['apps']} new** (total: {post['apps']})\n"
                            f"- Invoices: **{post['invoices'] - pre['invoices']} new** (total: {post['invoices']})\n"
                            f"  (Created vendors: {created_vendors}, created apps: {created_apps})\n\n"
                            f"**Batch ID:** `{import_batch_id}`"
                        )
                        if created_invoices_total:
                            st.caption(f"Planned recurring invoices created from contracts: {created_invoices_total}")
                        # Fallback: if contracts were uploaded but no planned invoices were created yet,
                        # try syncing again outside of the group/apps gating (covers contracts-only imports)
                        try:
                            need_fallback = (('dfK' in locals()) and isinstance(dfK, pd.DataFrame) and not dfK.empty and int(created_invoices_total or 0) == 0)
                        except Exception:
                            need_fallback = False
                        if need_fallback:
                            try:
                                ensure_contracts_table()
                                # Refresh maps for IDs
                                cur_teams_fb = list_teams(); tmap_fb = {}
                                if cur_teams_fb is not None and not cur_teams_fb.empty:
                                    for _, r0 in cur_teams_fb.iterrows():
                                        tmap_fb[str(r0.TEAMNAME).strip().upper()] = r0.TEAMID

                                cur_apps_fb = list_applications(); amap_fb = {}
                                if cur_apps_fb is not None and not cur_apps_fb.empty:
                                    for _, r0 in cur_apps_fb.iterrows():
                                        amap_fb[str(r0.APPLICATIONNAME).strip().upper()] = r0.APPLICATIONID

                                from db import sync_contract_invoices as _sync
                                fb_created = 0
                                for _, row in dfK.iterrows():
                                    tname = str(row.get("TEAMNAME") or "").strip()
                                    aname = str(row.get("APPNAME") or "").strip()
                                    gname = str(row.get("GROUPNAME") or "").strip() if "GROUPNAME" in row else ""
                                    if _blank_or_nan(tname) or _blank_or_nan(aname):
                                        continue
                                    tid = tmap_fb.get(tname.upper())
                                    final_appname = aname
                                    if gname:
                                        prefix = f"{gname} - "
                                        if aname.upper().startswith(prefix.upper()):
                                            final_appname = aname
                                        elif aname.upper() == gname.upper():
                                            final_appname = f"{gname} - Instance"
                                        else:
                                            final_appname = f"{gname} - {aname}"
                                    aid = amap_fb.get(final_appname.upper()) or amap_fb.get(aname.upper())
                                    if not tid or not aid:
                                        continue
                                    import uuid as _uuid
                                    cid = str(_uuid.uuid5(_uuid.NAMESPACE_URL, f"contract:{aid}:{tid}"))
                                    # Upsert contract (idempotent) then sync planned invoices
                                    try:
                                        # Safe parse fallback
                                        def _to_int_default(val, default):
                                            try:
                                                if val is None or (hasattr(pd, 'isna') and pd.isna(val)):
                                                    return default
                                                s = str(val).strip()
                                                if s == "":
                                                    return default
                                                return int(float(s))
                                            except Exception:
                                                return default
                                        def _to_float_default(val, default):
                                            try:
                                                if val is None or (hasattr(pd, 'isna') and pd.isna(val)):
                                                    return default
                                                s = str(val).strip()
                                                if s == "":
                                                    return default
                                                return float(s)
                                            except Exception:
                                                return default

                                        sfy = _to_int_default(row.get("CONTRACT_START_FY"), None)
                                        efy = _to_int_default(row.get("CONTRACT_END_FY"), None)
                                        amt = _to_float_default(row.get("CONTRACT_ANNUAL_AMOUNT"), None)
                                        esc = _to_float_default(row.get("CONTRACT_ESCALATION_PCT"), 0.0)
                                        rmon = _to_int_default(row.get("CONTRACT_RENEWAL_MONTH"), 1)
                                        if rmon is None or rmon < 1 or rmon > 12:
                                            rmon = 1
                                        if sfy is None or efy is None or amt is None or (efy < sfy):
                                            raise ValueError("Invalid START_FY/END_FY/ANNUAL_AMOUNT")
                                        upsert_contract(
                                            contract_id=cid,
                                            application_id=aid,
                                            team_id=tid,
                                            start_fy=int(sfy),
                                            end_fy=int(efy),
                                            renewal_month=int(rmon),
                                            annual_amount=float(amt),
                                            escalation_pct=float(esc),
                                            status=(str(row.get("CONTRACT_STATUS") or "Active")),
                                            agreement_number=(str(row.get("CONTRACT_AGREEMENT_NUMBER")).strip() or None),
                                            company_code=(str(row.get("CONTRACT_COMPANY_CODE")).strip() or None),
                                            cost_center=(str(row.get("CONTRACT_COST_CENTER")).strip() or None),
                                            service_type=(str(row.get("CONTRACT_SERVICE_TYPE")).strip() or None),
                                            contract_renewal_date=(str(row.get("CONTRACT_RENEWAL_DATE")) if pd.notna(row.get("CONTRACT_RENEWAL_DATE")) else None),
                                            invoice_renewal_date=(str(row.get("INVOICE_RENEWAL_DATE")) if pd.notna(row.get("INVOICE_RENEWAL_DATE")) else None),
                                            total_contract_cost=(float(row.get("CONTRACT_TOTAL_COST")) if pd.notna(row.get("CONTRACT_TOTAL_COST")) else None),
                                        )
                                    except Exception:
                                        pass
                                    try:
                                        fb_created += int(_sync(cid) or 0)
                                    except Exception:
                                        pass
                                if fb_created > 0:
                                    st.success(f"Created planned recurring invoices from contracts (fallback): {fb_created}")
                                    try:
                                        _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                                    except Exception:
                                        pass
                                else:
                                    st.info("No planned invoices created from contracts. Check Start/End FY and Team/App names.")
                            except Exception as e:
                                st.warning(f"Fallback planned invoice sync failed: {e}")

            except Exception as e:
                        st.error(f"Import failed: {e}")
                        st.exception(e)

        with st.expander("Sanity: Groups with ambiguous vendors per TEAM×GROUP", expanded=False):
            check = fetch_df("""
                WITH base AS (
                  SELECT t.TEAMNAME, g.GROUPNAME, g.DEFAULT_VENDORID, v.VENDORNAME
                  FROM APPLICATION_GROUPS g
                  JOIN TEAMS t ON t.TEAMID = g.TEAMID
                  LEFT JOIN VENDORS v ON v.VENDORID = g.DEFAULT_VENDORID
                ), counts AS (
                  SELECT TEAMNAME, GROUPNAME, COUNT(DISTINCT DEFAULT_VENDORID) AS VENDOR_IDS
                  FROM base
                  GROUP BY TEAMNAME, GROUPNAME
                ), names AS (
                  SELECT TEAMNAME, GROUPNAME, STRING_AGG(VENDORNAME, ', ') WITHIN GROUP (ORDER BY VENDORNAME) AS VENDORS
                  FROM (
                    SELECT DISTINCT TEAMNAME, GROUPNAME, VENDORNAME FROM base
                  ) d
                  GROUP BY TEAMNAME, GROUPNAME
                )
                SELECT c.TEAMNAME, c.GROUPNAME, c.VENDOR_IDS, n.VENDORS
                FROM counts c
                JOIN names n ON n.TEAMNAME = c.TEAMNAME AND n.GROUPNAME = c.GROUPNAME
                WHERE c.VENDOR_IDS > 1
                ORDER BY c.TEAMNAME, c.GROUPNAME
            """)
            if check is None or check.empty:
                st.success("No ambiguous vendors detected across TEAM×GROUP.")
            else:
                st.warning("Found TEAM×GROUP rows with multiple vendor IDs (shouldn’t happen):")
                st.dataframe(check, use_container_width=True, height=260)

        # --------------------------------------------
        # 🛠 One-click fix: set default vendor per TEAM×GROUP
        # --------------------------------------------
        with st.expander("One-click fix: default vendor (TEAM × GROUP) from majority invoices", expanded=False):
            st.caption(
                "For each Team × Group, pick the vendor that appears most often in INVOICES.VENDORID_AT_BOOKING "
                "(ties → most recent invoice wins; final tie → lowest VendorID). "
                "Only updates groups where the proposed vendor differs from the current default."
            )

            # Preview the proposed changes
            try:
                preview_sql = """
                WITH counts AS (
                  SELECT
                    ISNULL(i.TEAMID, g.TEAMID) AS TID,
                    ISNULL(i.GROUPID_AT_BOOKING, i.GROUPID) AS GID,
                    i.VENDORID_AT_BOOKING AS VID,
                    COUNT(*) AS CNT,
                    MAX(ISNULL(i.RENEWALDATE, CAST('1900-01-01' AS DATE))) AS LATEST_DATE
                  FROM INVOICES i
                  LEFT JOIN APPLICATION_GROUPS g ON g.GROUPID = i.GROUPID
                  WHERE i.VENDORID_AT_BOOKING IS NOT NULL AND LTRIM(RTRIM(i.VENDORID_AT_BOOKING)) <> ''
                  GROUP BY ISNULL(i.TEAMID, g.TEAMID), ISNULL(i.GROUPID_AT_BOOKING, i.GROUPID), i.VENDORID_AT_BOOKING
                ), ranked AS (
                  SELECT *, ROW_NUMBER() OVER (PARTITION BY TID, GID ORDER BY CNT DESC, LATEST_DATE DESC, VID) AS rn
                  FROM counts
                ), best AS (
                  SELECT * FROM ranked WHERE rn = 1
                )
                SELECT
                  t.TEAMNAME,
                  g.GROUPID,
                  g.GROUPNAME,
                  v_cur.VENDORNAME AS CURRENT_VENDOR,
                  v_new.VENDORNAME AS PROPOSED_VENDOR,
                  b.CNT            AS EVIDENCE_INVOICES
                FROM best b
                JOIN APPLICATION_GROUPS g ON g.GROUPID = b.GID AND g.TEAMID = b.TID
                JOIN TEAMS t ON t.TEAMID = g.TEAMID
                LEFT JOIN VENDORS v_cur ON v_cur.VENDORID = g.DEFAULT_VENDORID
                LEFT JOIN VENDORS v_new ON v_new.VENDORID = b.VID
                WHERE ISNULL(g.DEFAULT_VENDORID, '_') <> ISNULL(b.VID, '_')
                ORDER BY t.TEAMNAME, g.GROUPNAME
                """
                fix_preview = fetch_df(preview_sql)
            except Exception as e:
                fix_preview = pd.DataFrame()
                st.error(f"Could not build preview: {e}")

            if fix_preview is None or fix_preview.empty:
                st.success("No changes needed — current defaults already match the majority vendors per TEAM×GROUP.")
            else:
                st.info(f"Proposed updates: **{len(fix_preview)}** group(s) will change their default vendor.")
                st.dataframe(
                    fix_preview[["TEAMNAME","GROUPID","GROUPNAME","CURRENT_VENDOR","PROPOSED_VENDOR","EVIDENCE_INVOICES"]],
                    use_container_width=True,
                    height=320,
                )

                if st.button("✅ Apply fix (update default vendors)", type="primary", key="btn_apply_vendor_fix"):
                    try:
                        update_sql = """
                        WITH counts AS (
                          SELECT
                            ISNULL(i.TEAMID, g.TEAMID) AS TID,
                            ISNULL(i.GROUPID_AT_BOOKING, i.GROUPID) AS GID,
                            i.VENDORID_AT_BOOKING AS VID,
                            COUNT(*) AS CNT,
                            MAX(ISNULL(i.RENEWALDATE, CAST('1900-01-01' AS DATE))) AS LATEST_DATE
                          FROM INVOICES i
                          LEFT JOIN APPLICATION_GROUPS g ON g.GROUPID = i.GROUPID
                          WHERE i.VENDORID_AT_BOOKING IS NOT NULL AND LTRIM(RTRIM(i.VENDORID_AT_BOOKING)) <> ''
                          GROUP BY ISNULL(i.TEAMID, g.TEAMID), ISNULL(i.GROUPID_AT_BOOKING, i.GROUPID), i.VENDORID_AT_BOOKING
                        ), ranked AS (
                          SELECT *, ROW_NUMBER() OVER (PARTITION BY TID, GID ORDER BY CNT DESC, LATEST_DATE DESC, VID) AS rn
                          FROM counts
                        ), best AS (
                          SELECT * FROM ranked WHERE rn = 1
                        )
                        UPDATE g
                        SET DEFAULT_VENDORID = b.VID
                        FROM APPLICATION_GROUPS g
                        JOIN best b ON g.GROUPID = b.GID AND g.TEAMID = b.TID
                        WHERE ISNULL(g.DEFAULT_VENDORID, '_') <> ISNULL(b.VID, '_')
                        """
                        execute(update_sql)
                        st.success("Default vendors updated from invoice majorities. ✅")
                        st.rerun()
                    except Exception as e:
                        st.error(f"Vendor fix failed: {e}")

        # --------------------------------------------
        # 🩹 Backfill app → group links from name prefix
        # --------------------------------------------
        with st.expander("Backfill Application → Group links from name prefix ('Group - Suffix')", expanded=False):
            st.caption("For application instances with empty GROUPID, infer the group by matching the name prefix before ' - ' against existing Applications.")
            if st.button("Backfill GROUPID from Application Name", key="btn_backfill_groupid"):
                try:
                    # Preview how many would be updated
                    preview = fetch_df(
                        """
                        WITH parsed AS (
                          SELECT a.APPLICATIONID,
                                 CASE WHEN CHARINDEX(' - ', a.APPLICATIONNAME) > 0
                                      THEN LEFT(a.APPLICATIONNAME, CHARINDEX(' - ', a.APPLICATIONNAME) - 1)
                                      ELSE NULL END AS GNAME
                          FROM APPLICATIONS a
                          WHERE a.GROUPID IS NULL OR LTRIM(RTRIM(a.GROUPID)) = ''
                        )
                        SELECT COUNT(*) AS N
                        FROM parsed p
                        JOIN APPLICATION_GROUPS g ON UPPER(g.GROUPNAME) = UPPER(p.GNAME)
                        """
                    )
                    will_update = int(preview.iloc[0]["N"]) if preview is not None and not preview.empty else 0

                    execute(
                        """
                        WITH parsed AS (
                          SELECT a.APPLICATIONID,
                                 CASE WHEN CHARINDEX(' - ', a.APPLICATIONNAME) > 0
                                      THEN LEFT(a.APPLICATIONNAME, CHARINDEX(' - ', a.APPLICATIONNAME) - 1)
                                      ELSE NULL END AS GNAME
                          FROM APPLICATIONS a
                          WHERE a.GROUPID IS NULL OR LTRIM(RTRIM(a.GROUPID)) = ''
                        ), matched AS (
                          SELECT p.APPLICATIONID, g.GROUPID
                          FROM parsed p
                          JOIN APPLICATION_GROUPS g ON UPPER(g.GROUPNAME) = UPPER(p.GNAME)
                        )
                        UPDATE a
                          SET a.GROUPID = m.GROUPID
                        FROM APPLICATIONS a
                        JOIN matched m ON m.APPLICATIONID = a.APPLICATIONID;
                        """
                    )
                    st.success(f"Linked {will_update} application(s) to groups by name prefix.")
                    st.rerun()
                except Exception as e:
                    st.error(f"Backfill failed: {e}")

        # --------------------------------------------
        # 🔁 Rollback imported invoices by Batch ID
        # --------------------------------------------
        with st.expander("Rollback imported invoices by Batch ID", expanded=False):
            st.caption("This will **DELETE** all invoices with the selected `ROLLOVER_BATCH_ID`. It does not touch vendors/programs/teams/apps/groups.")

            def list_invoice_batches(limit: int = 50) -> pd.DataFrame:
                try:
                    return fetch_df(f"""
                        SELECT TOP {int(limit)}
                          ROLLOVER_BATCH_ID AS BATCH_ID,
                          COUNT(*) AS N
                        FROM INVOICES
                        WHERE ROLLOVER_BATCH_ID IS NOT NULL
                        GROUP BY ROLLOVER_BATCH_ID
                        ORDER BY MAX(ISNULL(RENEWALDATE, CAST('1900-01-01' AS DATE))) DESC, BATCH_ID DESC
                    """)
                except Exception:
                    return pd.DataFrame(columns=["BATCH_ID","N"])

            batches_df = list_invoice_batches(limit=100)
            if batches_df is None or batches_df.empty:
                st.info("No batches found yet.")
            else:
                st.dataframe(batches_df, use_container_width=True, height=220)

                batch_choices = ["(type a batch id)"] + batches_df["BATCH_ID"].astype(str).tolist()
                picked = st.selectbox("Pick a recent Batch ID", options=batch_choices, index=0, key="rollback_pick")
                typed = st.text_input("...or paste a Batch ID exactly", value="" if picked == "(type a batch id)" else picked, key="rollback_typed").strip()

                col_prev, col_del = st.columns([1,1])

                with col_prev:
                    if st.button("Preview rows in this batch", key="btn_preview_batch", use_container_width=True, disabled=(typed == "")):
                        try:
                            prev = fetch_df("""
                                SELECT TOP 500
                                  INVOICEID, TEAMID, APPLICATIONID, GROUPID, VENDORID_AT_BOOKING,
                                  FISCAL_YEAR, RENEWALDATE, AMOUNT, AMOUNT_NEXT_YEAR, STATUS, INVOICE_TYPE,
                                  ROLLOVER_BATCH_ID
                                FROM INVOICES
                                WHERE ROLLOVER_BATCH_ID = %s
                                ORDER BY FISCAL_YEAR DESC,
                                         CASE WHEN RENEWALDATE IS NULL THEN 1 ELSE 0 END ASC,
                                         RENEWALDATE DESC,
                                         INVOICEID
                            """, (typed,))
                            if prev is None or prev.empty:
                                st.warning("No invoices found for that Batch ID.")
                            else:
                                st.success(f"Found {len(prev)} row(s) (showing up to 500).")
                                st.dataframe(prev, use_container_width=True, height=320)
                        except Exception as e:
                            st.error(f"Preview failed: {e}")

                with col_del:
                    danger = st.checkbox("I understand this **permanently deletes** invoices in this batch.", value=False, key="confirm_delete_batch")
                    if st.button("Rollback (Delete Invoices in Batch)", icon=":material/delete:", type="secondary", use_container_width=True,
                                 disabled=(typed == "" or not danger), key="btn_delete_batch"):
                        try:
                            cnt = fetch_df("SELECT COUNT(*) AS N FROM INVOICES WHERE ROLLOVER_BATCH_ID = %s", (typed,))
                            n = int(cnt.iloc[0]["N"]) if cnt is not None and not cnt.empty else 0

                            execute("DELETE FROM INVOICES WHERE ROLLOVER_BATCH_ID = %s", (typed,))
                            st.success(f"Deleted **{n}** invoice(s) for batch `{typed}`.")
                        except Exception as e:
                            st.error(f"Rollback failed: {e}")

    else:
        st.info("Upload your workbook above to proceed.")


# =========================# =========================================================
# 📦 Bulk Load (v2 – MSSQL): Multi-sheet importer
# =========================================================
with _tab_loader:
    st.subheader("Data Loader (v2 – MSSQL)")
    st.caption("Templates + importer for Master Data and Finance. MSSQL only. Idempotent upserts. Transactional (rollback on error). Rates are managed in-app on the Rates page.")
    st.caption("Flow: 1) Download template  2) Upload workbook  3) Run validation  4) Plan changes  5) Import")

    with st.expander("Checklist (must-have vs optional)", expanded=False):
        st.markdown(bulkload_v2.get_checklist_markdown())

    col_dl, col_scope = st.columns([1, 1])
    with col_dl:
        st.download_button(
            "⬇️ Download Master Data template",
            data=bulkload_v2.build_master_template_bytes(),
            file_name="next_data_loader_master_template.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            help="Programs/Teams/Vendors/Applications/Applications (+ optional Group-Team Links).",
            key="bulk_v2_download_master",
        )
        st.download_button(
            "⬇️ Download Finance template",
            data=bulkload_v2.build_finance_template_bytes(),
            file_name="next_data_loader_finance_template.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            help="Contracts/Invoices/Program Addl Costs (assumes Master Data already exists).",
            key="bulk_v2_download_finance",
        )
        st.caption("Templates include an `Examples` sheet and Excel dropdowns (data validation). For dropdowns, open in Microsoft Excel Desktop.")
    with col_scope:
        scope_label = st.radio(
            "Loader type",
            ["Master Data", "Finance"],
            index=0,
            horizontal=True,
            key="bulk_v2_scope",
        )
        bulk_mode = {"Master Data": "master", "Finance": "finance"}[scope_label]

    upl = st.file_uploader(
        "Upload completed workbook",
        type=["xlsx", "xlsm", "xlsb", "xls", "csv"],
        key="bulk_v2_uploader",
    )

    if upl:
        dfs = bulkload_v2.read_workbook(upl)
        if st.session_state.get("bulk_v2_last_mode") != bulk_mode:
            st.session_state["bulk_v2_last_mode"] = bulk_mode
            st.session_state.pop("bulk_v2_validation_done", None)
            st.session_state.pop("bulk_v2_plan", None)
            st.session_state.pop("bulk_v2_blocking_df", None)
            st.session_state.pop("bulk_v2_warnings_df", None)
        detected = pd.DataFrame(
            [{"sheet": k, "rows": int(len(v.index)) if isinstance(v, pd.DataFrame) else 0, "cols": int(len(v.columns)) if isinstance(v, pd.DataFrame) else 0} for k, v in dfs.items()]
        ).sort_values(["sheet"])
        st.markdown("**Detected sheets**")
        st.dataframe(detected, use_container_width=True, height=220)

        c1, c2, c3 = st.columns([1, 1, 1])
        with c1:
            do_validate = st.button("Run validation", type="primary", key="bulk_v2_btn_validate", use_container_width=True)
        with c2:
            do_plan = st.button("Plan changes", key="bulk_v2_btn_plan", use_container_width=True)
        with c3:
            do_import = st.button("Import", key="bulk_v2_btn_import", use_container_width=True)

        if do_validate or ("bulk_v2_validation_done" not in st.session_state):
            blocking_df, warnings_df = bulkload_v2.validate_workbook(dfs, bulk_mode)
            st.session_state["bulk_v2_blocking_df"] = blocking_df
            st.session_state["bulk_v2_warnings_df"] = warnings_df
            st.session_state["bulk_v2_validation_done"] = True

        blocking_df = st.session_state.get("bulk_v2_blocking_df", pd.DataFrame())
        warnings_df = st.session_state.get("bulk_v2_warnings_df", pd.DataFrame())

        if blocking_df is not None and not blocking_df.empty:
            st.error(f"Blocking errors: {len(blocking_df)}")
            st.dataframe(blocking_df, use_container_width=True, height=260)
        else:
            st.success("No blocking validation errors.")

        if warnings_df is not None and not warnings_df.empty:
            st.warning(f"Warnings: {len(warnings_df)}")
            st.dataframe(warnings_df, use_container_width=True, height=220)

        rebuild_all_views = st.checkbox(
            "Rebuild ALL views after import (slow)",
            value=False,
            key="bulk_v2_rebuild_all_views",
            help="Defaults to rebuilding only the analytics views after import. Enable this to run a full rebuild of all view materializations.",
        )

        report_df = pd.concat(
            [
                (blocking_df.assign(severity="BLOCKING") if blocking_df is not None and not blocking_df.empty else pd.DataFrame(columns=["sheet", "row", "column", "error", "suggested_fix", "severity"])),
                (warnings_df.assign(severity="WARNING") if warnings_df is not None and not warnings_df.empty else pd.DataFrame(columns=["sheet", "row", "column", "error", "suggested_fix", "severity"])),
            ],
            ignore_index=True,
        )
        if report_df is not None and not report_df.empty:
            st.download_button(
                "⬇️ Download validation report (CSV)",
                data=report_df.to_csv(index=False).encode("utf-8"),
                file_name="bulkload_v2_validation_report.csv",
                mime="text/csv",
                key="bulk_v2_dl_report_csv",
            )
            bio = io.BytesIO()
            with pd.ExcelWriter(bio, engine="openpyxl") as w:
                report_df.to_excel(w, index=False, sheet_name="Validation")
            st.download_button(
                "⬇️ Download validation report (XLSX)",
                data=bio.getvalue(),
                file_name="bulkload_v2_validation_report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="bulk_v2_dl_report_xlsx",
            )

        if do_plan and (blocking_df is None or blocking_df.empty):
            st.session_state["bulk_v2_plan"] = bulkload_v2.plan_changes(dfs, None, mode=bulk_mode)

        plan = st.session_state.get("bulk_v2_plan")
        if plan and isinstance(plan, dict):
            st.markdown("**Preview (rows to insert / update / skipped)**")
            sheets_plan = plan.get("sheets") or {}
            summary_rows = []
            for sname, splan in sheets_plan.items():
                summary_rows.append(
                    {"sheet": sname, "rows": splan.get("rows"), "insert": splan.get("insert"), "update": splan.get("update"), "skip": splan.get("skip")}
                )
            st.dataframe(pd.DataFrame(summary_rows), use_container_width=True, height=320)

            with st.expander("Sample previews", expanded=False):
                for sname, splan in sheets_plan.items():
                    for kind in ("sample_insert", "sample_update", "sample_skip"):
                        samp = splan.get(kind)
                        if isinstance(samp, pd.DataFrame) and not samp.empty:
                            st.markdown(f"**{sname} – {kind.replace('sample_', '').replace('_', ' ')}**")
                            st.dataframe(samp, use_container_width=True, height=180)

        if do_import and plan and (blocking_df is None or blocking_df.empty):
            user_email = st.session_state.get("auth_user", {}).get("email")
            pb = st.progress(0.0)
            status = st.empty()

            def _cb(step: str, i: int, total: int) -> None:
                pct = 0.0 if total <= 0 else float(i) / float(total)
                pb.progress(min(1.0, max(0.0, pct)))
                status.info(f"{step} ({i}/{total})")

            result = bulkload_v2.apply_import(
                plan,
                None,
                progress_cb=_cb,
                user_email=user_email,
                source_filename=getattr(upl, "name", None),
                rebuild_all_views=rebuild_all_views,
            )
            pb.progress(1.0)

            if result.get("success"):
                toast_success("Bulk Load v2 import completed.")
                st.dataframe(pd.DataFrame(result.get("steps") or []), use_container_width=True, height=260)
                with st.expander("Import log", expanded=False):
                    st.code("\n".join(result.get("log") or []))
                try:
                    _settings_post_write_refresh("settings_cache_clear_legacy", bump_version=True)
                except Exception:
                    pass
                st.session_state["bulk_v2_validation_done"] = False
                st.session_state["bulk_v2_plan"] = None
            else:
                toast_error(f"Bulk Load v2 import failed: {result.get('error')}")
                with st.expander("Import log", expanded=True):
                    st.code("\n".join(result.get("log") or []))
    else:
        st.info("Upload your workbook above to proceed.")

    st.markdown("---")
    with st.expander("Legacy bulk load (deprecated)", expanded=False):
        st.warning("Deprecated: use Bulk Load (v2 – MSSQL) above for new imports.")
        enable_legacy = st.toggle("Enable legacy loader", value=False, key="bulk_v2_show_legacy")
        if enable_legacy:
            _render_legacy_bulk_load_one_sheet()
